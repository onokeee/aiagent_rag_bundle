"""brain.py — AIと分析。core.py から呼ばれる。

中身は元のファイル順のまま4つのかたまり:

  advanced / analysis   … 統計・予測・SQLでは書けない集計（純粋な計算）
  models                … 使えるモデルの一覧・既定・文脈量（llm と表裏）
  usage / exports / excel / charts / figures / docx_report / pptx_report / business
                        … 集計の出力（グラフ・Excel・PowerPoint・Word）
  mailer / custom_tools / tools/* / rag/* / llm
                        … メール下書き、AIに配る道具一式、文書検索、エージェント本体

core.py との関係は一方通行:

  core.py  … 土台（db / catalog / importer / jobs / chats …）と画面（Flask）
  brain.py … その上で動く「考える側」

core.py がデータ層を定義したあとで import brain する。そのため、
ここで `import db` や `catalog.load_meta()` と書けば今までどおり解決する
（core.py が自分自身を元のモジュール名で sys.modules に登録している）。

逆に core.py からは `llm.` `tools.` `models.` … で呼べる。
このファイルが自分自身をそれらの名前で登録するため。
"""
from __future__ import annotations

import sys as _sys

# 元のモジュール名でも import できるようにする（呼び出し側を変えないため）。
# ここに挙げた名前の中身は、このファイルにある。
for _alias in ("advanced", "analysis", "models", "usage", "exports", "excel",
               "charts", "figures", "docx_report", "pptx_report", "business",
               "mailer", "custom_tools", "llm", "tools", "rag", "results"):
    _sys.modules[_alias] = _sys.modules[__name__]
del _alias

results = _sys.modules[__name__]   # 統合前の書き方をそのまま使えるようにする
_results = _sys.modules[__name__]  # 同上

# 読み込みの順を間違えたときは、ここで分かる形で止める。
# brain.py は core.py から読み込まれる前提（core が土台を定義して、
# 自分自身を db / catalog … の名前で登録したあとに import brain する）。
if "catalog" not in _sys.modules:
    raise ImportError(
        "brain.py は core.py から読み込まれる前提です。"
        "単体では、ここで使う db / catalog / sqlusage といった名前を解決できません。"
        "アプリの起動は python core.py、"
        "コードから使うときは import core を先に行ってください。")

# core.py 側にあるもの。core がデータ層を定義したあとに import brain するので、
# この時点でどれも出来上がっている。
import sqlusage                                        # noqa: E402
from catalog import (db_text_cached, load_meta,        # noqa: E402
                     parse_endpoint_cols, profile_db)

# ==========================================================================
# ===== 元 advanced.py（統計・予測・業務分析。行番号は inspect で動的に取得される）
# ==========================================================================
"""高度な分析。回帰・仮説検定・予測・シミュレーションなど。

analysis.py が「SQLでは書けない集計」を担うのに対し、こちらは
統計モデルを当てる側を担当する。入出力の約束は analysis.py と同じで、
SELECT結果 (columns, rows) を受け取り、表示用の表と所見テキストを返す。

戻り値の共通形:
    {"title": 見出し, "tables": [{"name":..., "columns":[...], "rows":[...]}, ...],
     "notes": [所見の文字列, ...], "meta": {機械可読な値}}

所見(notes)は日本語で書く。数字だけ返してもLLMが読み違えるので、
「有意差あり/なし」「あてはまりが弱い」といった判断まで言語化してここで持たせる。
"""

import math
import warnings

import numpy as np
import pandas as pd
from scipy import stats

# _clean / _df / _out / _to_numeric / numeric_columns はこのファイルの末尾（元 analysis.py）にある

warnings.filterwarnings("ignore", category=FutureWarning)
warnings.filterwarnings("ignore", category=UserWarning, module="statsmodels")

# 既定の有意水準。業務レポートで使う 5% を採用する。
ALPHA = 0.05

TEST_METHODS = {
    "ttest_1samp": "1標本t検定（平均が基準値と違うか）",
    "ttest_ind": "2標本t検定（2群の平均差・Welch）",
    "ttest_rel": "対応のあるt検定（同じ対象の前後比較）",
    "mannwhitney": "Mann-WhitneyのU検定（2群・順位）",
    "wilcoxon": "Wilcoxonの符号順位検定（対応あり・順位）",
    "anova": "一元配置分散分析（3群以上の平均差）",
    "kruskal": "Kruskal-Wallis検定（3群以上・順位）",
    "chi2": "カイ二乗検定（独立性・クロス集計）",
    "chi2_goodness": "カイ二乗適合度検定（比率が想定通りか）",
    "proportion": "比率の検定（2群の割合の差）",
    "normality": "正規性の検定（Shapiro-Wilk）",
    "levene": "等分散性の検定（Levene）",
    "correlation": "無相関の検定（相関が偶然か）",
}
REGRESSION_METHODS = {
    "ols": "重回帰（最小二乗法）",
    "logistic": "ロジスティック回帰（0/1の予測）",
    "poisson": "ポアソン回帰（件数の予測）",
}
FORECAST_METHODS = {
    "naive": "直近値をそのまま延長",
    "drift": "直線の傾きで延長",
    "moving_average": "移動平均で延長",
    "linear": "線形トレンド（回帰）",
    "holt": "指数平滑（トレンドあり）",
    "holt_winters": "指数平滑（トレンド＋季節性）",
    "arima": "ARIMA",
}
DISTRIBUTIONS = {
    "norm": "正規分布", "lognorm": "対数正規分布", "expon": "指数分布",
    "gamma": "ガンマ分布", "uniform": "一様分布", "triang": "三角分布",
}
OUTLIER_METHODS_EXT = {
    "iqr": "四分位範囲（箱ひげ図の外側）",
    "zscore": "標準偏差からの距離",
    "modified_zscore": "中央値からの距離（MAD・外れ値に強い）",
    "percentile": "上下の指定パーセンタイル",
    "mahalanobis": "複数列をまとめて見る（マハラノビス距離）",
}


class AnalysisError(Exception):
    """分析できないときの理由（そのまま画面とLLMに見せる）。"""


def _advanced_num(df: pd.DataFrame, cols: list, need: int = 1) -> pd.DataFrame:
    """指定列を数値にして、欠損行を落とす。"""
    missing = [c for c in cols if c not in df.columns]
    if missing:
        raise AnalysisError(f"列が見つかりません: {', '.join(missing)}"
                            f"（SQLの結果にある列: {', '.join(map(str, df.columns))}）")
    out = _to_numeric(df.copy(), cols)[cols].dropna()
    if len(out) < need:
        raise AnalysisError(f"数値として使える行が {len(out)} 行しかありません"
                            f"（{need}行以上必要）。列の型か抽出条件を見直してください。")
    return out


def _p_note(p: float, what: str) -> str:
    if p is None or (isinstance(p, float) and math.isnan(p)):
        return f"{what}: p値を計算できませんでした。"
    if p < ALPHA:
        return (f"{what}: p値 {p:.4g} < {ALPHA} なので、"
                "偶然とは考えにくい差（有意差あり）です。")
    return (f"{what}: p値 {p:.4g} ≧ {ALPHA} なので、"
            "この結果からは差があるとは言えません（有意差なし）。")


def _effect_note(name: str, value: float) -> str:
    """効果量の目安。p値だけ見て「差がある」と早合点しないための添え物。"""
    a = abs(value)
    if name == "cohen_d":
        size = "小さい" if a < 0.5 else ("中くらい" if a < 0.8 else "大きい")
        return f"効果量 Cohen's d = {value:.3f}（差の大きさは{size}）"
    if name == "eta_squared":
        size = "小さい" if a < 0.06 else ("中くらい" if a < 0.14 else "大きい")
        return f"効果量 η² = {value:.3f}（群による説明力は{size}）"
    if name == "cramers_v":
        size = "弱い" if a < 0.2 else ("中くらい" if a < 0.4 else "強い")
        return f"効果量 Cramer's V = {value:.3f}（関連は{size}）"
    return f"効果量 {name} = {value:.3f}"


def _advanced_table(name: str, columns: list, rows: list) -> dict:
    return {"name": name, "columns": [str(c) for c in columns],
            "rows": [[_clean(v) for v in r] for r in rows]}


def _df_table(name: str, df: pd.DataFrame, index_label: str | None = None) -> dict:
    d = df.reset_index() if index_label else df
    if index_label:
        d = d.rename(columns={d.columns[0]: index_label})
    cols, rows = _out(d)
    return _advanced_table(name, cols, rows)


# =============================================================================
# 仮説検定
# =============================================================================

def hypothesis_test(columns: list, rows: list, method: str, *,
                    value_col: str | None = None, group_col: str | None = None,
                    value_col2: str | None = None, popmean: float = 0.0,
                    expected: list | None = None, alternative: str = "two-sided",
                    alpha: float = ALPHA) -> dict:
    """統計的仮説検定。method は TEST_METHODS のキー。"""
    if method not in TEST_METHODS:
        raise AnalysisError(f"未対応の検定です: {method}。"
                            f"使えるのは {', '.join(TEST_METHODS)} です。")
    df = _df(columns, rows)
    label = TEST_METHODS[method]
    tables, notes, meta = [], [], {"method": method, "alpha": alpha}

    def groups_of(vcol, gcol):
        d = df[[gcol, vcol]].copy()
        d[vcol] = pd.to_numeric(d[vcol], errors="coerce")
        d = d.dropna()
        gs = [(str(k), g[vcol].to_numpy()) for k, g in d.groupby(gcol)]
        gs = [(k, v) for k, v in gs if len(v) >= 2]
        if len(gs) < 2:
            raise AnalysisError(f"比較できる群が {len(gs)} しかありません"
                                f"（各群2件以上・2群以上が必要）。")
        return gs

    if method == "ttest_1samp":
        x = _advanced_num(df, [value_col], 2)[value_col].to_numpy()
        st, p = stats.ttest_1samp(x, popmean, alternative=alternative)
        d = (x.mean() - popmean) / (x.std(ddof=1) or np.nan)
        tables.append(_advanced_table("検定結果", ["項目", "値"], [
            ["件数", len(x)], ["平均", x.mean()], ["基準値", popmean],
            ["差", x.mean() - popmean], ["t値", st], ["p値", p]]))
        notes += [_p_note(p, f"{value_col} の平均と基準値 {popmean} の差"),
                  _effect_note("cohen_d", d)]
        meta.update(statistic=float(st), p_value=float(p), effect=float(d))

    elif method in ("ttest_ind", "mannwhitney", "levene"):
        gs = groups_of(value_col, group_col)
        if len(gs) > 2:
            raise AnalysisError(f"この検定は2群までです（いまは {len(gs)}群）。"
                                "3群以上なら anova か kruskal を使ってください。")
        (n1, a), (n2, b) = gs
        if method == "ttest_ind":
            st, p = stats.ttest_ind(a, b, equal_var=False, alternative=alternative)
            sd = math.sqrt(((len(a) - 1) * a.var(ddof=1) + (len(b) - 1) * b.var(ddof=1))
                           / max(len(a) + len(b) - 2, 1))
            eff = (a.mean() - b.mean()) / sd if sd else float("nan")
            notes.append(_effect_note("cohen_d", eff))
        elif method == "mannwhitney":
            st, p = stats.mannwhitneyu(a, b, alternative=alternative)
            eff = 2 * st / (len(a) * len(b)) - 1        # rank-biserial
            notes.append(f"効果量 rank-biserial = {eff:.3f}")
        else:
            st, p = stats.levene(a, b)
            eff = float("nan")
        tables.append(_advanced_table("群ごとの要約", ["群", "件数", "平均", "中央値", "標準偏差"], [
            [n1, len(a), a.mean(), np.median(a), a.std(ddof=1)],
            [n2, len(b), b.mean(), np.median(b), b.std(ddof=1)]]))
        tables.append(_advanced_table("検定結果", ["項目", "値"],
                             [["統計量", st], ["p値", p]]))
        notes.insert(0, _p_note(p, f"{n1} と {n2} の{'ばらつき' if method == 'levene' else value_col}の差"))
        meta.update(statistic=float(st), p_value=float(p), effect=float(eff))

    elif method in ("ttest_rel", "wilcoxon"):
        pair = _advanced_num(df, [value_col, value_col2], 2)
        a, b = pair[value_col].to_numpy(), pair[value_col2].to_numpy()
        if method == "ttest_rel":
            st, p = stats.ttest_rel(a, b, alternative=alternative)
            diff = a - b
            eff = diff.mean() / (diff.std(ddof=1) or np.nan)
            notes.append(_effect_note("cohen_d", eff))
        else:
            st, p = stats.wilcoxon(a, b, alternative=alternative)
            eff = float("nan")
        tables.append(_advanced_table("対応するデータ", ["項目", value_col, value_col2, "差"], [
            ["件数", len(a), len(b), len(a)],
            ["平均", a.mean(), b.mean(), (a - b).mean()],
            ["中央値", np.median(a), np.median(b), np.median(a - b)]]))
        tables.append(_advanced_table("検定結果", ["項目", "値"], [["統計量", st], ["p値", p]]))
        notes.insert(0, _p_note(p, f"{value_col} と {value_col2} の差"))
        meta.update(statistic=float(st), p_value=float(p), effect=float(eff))

    elif method in ("anova", "kruskal"):
        gs = groups_of(value_col, group_col)
        arrays = [v for _, v in gs]
        if method == "anova":
            st, p = stats.f_oneway(*arrays)
            grand = np.concatenate(arrays)
            ss_b = sum(len(v) * (v.mean() - grand.mean()) ** 2 for v in arrays)
            ss_t = ((grand - grand.mean()) ** 2).sum()
            eff = ss_b / ss_t if ss_t else float("nan")
            notes.append(_effect_note("eta_squared", eff))
        else:
            st, p = stats.kruskal(*arrays)
            eff = float("nan")
        tables.append(_advanced_table("群ごとの要約", ["群", "件数", "平均", "中央値", "標準偏差"],
                             [[k, len(v), v.mean(), np.median(v), v.std(ddof=1)]
                              for k, v in gs]))
        tables.append(_advanced_table("検定結果", ["項目", "値"],
                             [["統計量", st], ["p値", p], ["群の数", len(gs)]]))
        notes.insert(0, _p_note(p, f"{len(gs)}群の {value_col} の差"))
        if p < alpha and len(gs) > 2:
            pairs = []
            for i in range(len(gs)):
                for j in range(i + 1, len(gs)):
                    _, pp = stats.ttest_ind(gs[i][1], gs[j][1], equal_var=False)
                    # Bonferroni: 比較回数ぶん厳しくする
                    n_comp = len(gs) * (len(gs) - 1) / 2
                    pairs.append([gs[i][0], gs[j][0], pp, min(pp * n_comp, 1.0),
                                  "有意" if pp * n_comp < alpha else ""])
            tables.append(_advanced_table("どの組み合わせに差があるか（Bonferroni補正）",
                                 ["群1", "群2", "p値", "補正後p値", "判定"], pairs))
            notes.append("どの群どうしに差があるかは「補正後p値」が0.05未満の行を見てください。")
        meta.update(statistic=float(st), p_value=float(p), effect=float(eff))

    elif method == "chi2":
        ct = pd.crosstab(df[group_col], df[value_col])
        st, p, dof, exp = stats.chi2_contingency(ct)
        n = ct.to_numpy().sum()
        v = math.sqrt(st / (n * (min(ct.shape) - 1))) if min(ct.shape) > 1 else float("nan")
        tables.append(_df_table("クロス集計（実測）", ct, index_label=str(group_col)))
        tables.append(_df_table("期待度数（関連が無い場合）",
                                pd.DataFrame(exp, index=ct.index, columns=ct.columns).round(2),
                                index_label=str(group_col)))
        tables.append(_advanced_table("検定結果", ["項目", "値"],
                             [["カイ二乗値", st], ["p値", p], ["自由度", dof]]))
        notes += [_p_note(p, f"{group_col} と {value_col} の関連"), _effect_note("cramers_v", v)]
        if (exp < 5).mean() > 0.2:
            notes.append("※ 期待度数が5未満のマスが2割を超えています。"
                         "カテゴリをまとめるか、件数を増やした方が結果は安定します。")
        meta.update(statistic=float(st), p_value=float(p), effect=float(v), dof=int(dof))

    elif method == "chi2_goodness":
        counts = df[value_col].value_counts().sort_index()
        exp = np.array(expected, dtype=float) if expected else np.full(len(counts),
                                                                      counts.sum() / len(counts))
        exp = exp * counts.sum() / exp.sum()
        st, p = stats.chisquare(counts.to_numpy(), exp)
        tables.append(_advanced_table("実測と期待", ["区分", "実測", "期待"],
                             [[str(k), int(v), round(e, 2)]
                              for k, v, e in zip(counts.index, counts, exp)]))
        tables.append(_advanced_table("検定結果", ["項目", "値"], [["カイ二乗値", st], ["p値", p]]))
        notes.append(_p_note(p, f"{value_col} の分布と想定の違い"))
        meta.update(statistic=float(st), p_value=float(p))

    elif method == "proportion":
        ct = pd.crosstab(df[group_col], df[value_col])
        if ct.shape != (2, 2):
            raise AnalysisError(f"比率の検定は2群×2値のときに使えます"
                                f"（いまは {ct.shape[0]}群×{ct.shape[1]}値）。")
        st, p, dof, _ = stats.chi2_contingency(ct, correction=True)
        rates = ct.iloc[:, 1] / ct.sum(axis=1)
        tables.append(_advanced_table("群ごとの比率", ["群", "件数", f"{ct.columns[1]}の数", "比率"],
                             [[str(i), int(ct.loc[i].sum()), int(ct.loc[i].iloc[1]),
                               round(float(rates[i]), 4)] for i in ct.index]))
        tables.append(_advanced_table("検定結果", ["項目", "値"], [["カイ二乗値", st], ["p値", p]]))
        notes += [_p_note(p, "2群の比率の差"),
                  f"比率の差 = {abs(rates.iloc[0] - rates.iloc[1]):.4f}"]
        meta.update(statistic=float(st), p_value=float(p))

    elif method == "normality":
        x = _advanced_num(df, [value_col], 3)[value_col].to_numpy()
        if len(x) > 5000:
            x = np.random.default_rng(0).choice(x, 5000, replace=False)
            notes.append("※ 件数が多いので5000件を無作為抽出して検定しました。")
        st, p = stats.shapiro(x)
        tables.append(_advanced_table("検定結果", ["項目", "値"], [
            ["件数", len(x)], ["歪度", stats.skew(x)], ["尖度", stats.kurtosis(x)],
            ["W統計量", st], ["p値", p]]))
        notes.append(f"正規性: p値 {p:.4g} "
                     + ("< 0.05 なので正規分布とは言いにくいです"
                        "（順位を使う検定 mannwhitney / kruskal が無難）。"
                        if p < alpha else "≧ 0.05 なので正規分布として扱って差し支えありません。"))
        meta.update(statistic=float(st), p_value=float(p))

    elif method == "correlation":
        pair = _advanced_num(df, [value_col, value_col2], 3)
        r, p = stats.pearsonr(pair[value_col], pair[value_col2])
        rho, prho = stats.spearmanr(pair[value_col], pair[value_col2])
        tables.append(_advanced_table("検定結果", ["項目", "値"], [
            ["件数", len(pair)], ["ピアソンr", r], ["p値", p],
            ["スピアマンρ", rho], ["p値(ρ)", prho]]))
        notes += [_p_note(p, f"{value_col} と {value_col2} の相関"),
                  f"相関係数 r = {r:.3f}（"
                  + ("ほぼ無相関" if abs(r) < 0.2 else
                     "弱い相関" if abs(r) < 0.4 else
                     "中程度の相関" if abs(r) < 0.7 else "強い相関") + "）",
                  "相関は因果ではありません。第三の要因が両方を動かしている可能性を検討してください。"]
        meta.update(statistic=float(r), p_value=float(p))

    return {"title": f"{label}", "tables": tables, "notes": notes, "meta": meta}


# =============================================================================
# 回帰
# =============================================================================

def _holdout_score(sm, method: str, X: pd.DataFrame, y: pd.Series, cols) -> dict | None:
    """行の3割を隠して学習し、隠した側での成績を返す。

    学習に使ったデータで測ったR²は必ず良く出る。「予測に使えるのか」を
    聞かれたときに、その数字を答えると嘘になるので、別データで測り直す。
    """
    rng = np.random.default_rng(0)
    idx = rng.permutation(len(X))
    cut = int(len(X) * 0.7)
    tr, te = idx[:cut], idx[cut:]
    if len(te) < 5 or cut <= X.shape[1] + 1:
        return None
    Xtr = sm.add_constant(X.iloc[tr], has_constant="add")
    Xte = sm.add_constant(X.iloc[te], has_constant="add").reindex(columns=cols, fill_value=0.0)
    ytr, yte = y.iloc[tr], y.iloc[te]
    try:
        if method == "ols":
            m = sm.OLS(ytr, Xtr).fit()
            pred = np.asarray(m.predict(Xte), dtype=float)
            act = yte.to_numpy(dtype=float)
            ss_res = float(np.sum((act - pred) ** 2))
            ss_tot = float(np.sum((act - act.mean()) ** 2))
            r2 = 1 - ss_res / ss_tot if ss_tot else float("nan")
            mae = float(np.mean(np.abs(act - pred)))
            rows = [["検証データのR²", round(r2, 3)], ["平均絶対誤差(MAE)", round(mae, 4)],
                    ["学習に使った件数", len(tr)], ["検証に使った件数", len(te)]]
            note = (f"学習に使っていないデータでのR² = {r2:.3f}"
                    + ("。学習時とほぼ同じで、予測にも使えます。" if r2 > 0.5 else
                       "。学習時より大きく落ちる場合、この式は手元のデータに"
                       "合わせすぎていて、将来の予測には向きません。"))
        elif method == "logistic":
            m = sm.Logit(ytr, Xtr).fit(disp=0)
            p = np.asarray(m.predict(Xte), dtype=float)
            act = yte.to_numpy(dtype=float)
            acc = float(np.mean((p >= 0.5).astype(float) == act))
            base = float(max(act.mean(), 1 - act.mean()))
            rows = [["正解率", round(acc, 3)], ["全部多い方に賭けた場合", round(base, 3)],
                    ["検証に使った件数", len(te)]]
            note = (f"学習に使っていないデータでの正解率 = {acc:.1%}"
                    f"（何も考えず多い方に賭けると {base:.1%}）。"
                    + ("上回っているので、判別に意味があります。" if acc > base + 0.02 else
                       "差が無いため、この説明変数では判別できていません。"))
        else:
            return None
    except Exception:
        return None
    return {"rows": rows, "note": note}


def regression(columns: list, rows: list, target: str, features: list,
               method: str = "ols", predict: list | None = None,
               alpha: float = ALPHA) -> dict:
    """回帰分析。係数・有意性・あてはまり・診断をまとめて返す。"""
    import statsmodels.api as sm

    if method not in REGRESSION_METHODS:
        raise AnalysisError(f"未対応の手法です: {method}。"
                            f"使えるのは {', '.join(REGRESSION_METHODS)} です。")
    df = _df(columns, rows)
    features = [f for f in (features or []) if f != target]
    if not features:
        raise AnalysisError("説明変数を1つ以上指定してください。")

    # 文字列の列はダミー変数にする（部署などをそのまま説明変数にできるように）
    use = df[[target] + features].copy()
    use[target] = pd.to_numeric(use[target], errors="coerce")
    num_feats, cat_feats = [], []
    for f in features:
        s = pd.to_numeric(use[f], errors="coerce")
        if s.notna().mean() >= 0.8:
            use[f] = s
            num_feats.append(f)
        else:
            cat_feats.append(f)
    use = use.dropna(subset=[target] + num_feats)
    X = use[num_feats].copy()
    for f in cat_feats:
        d = pd.get_dummies(use[f].astype(str), prefix=f, drop_first=True, dtype=float)
        X = pd.concat([X, d], axis=1)
    X = X.dropna()
    y = use.loc[X.index, target]
    if len(X) <= X.shape[1] + 1:
        raise AnalysisError(f"データが {len(X)} 行しかなく、説明変数 {X.shape[1]} 個に対して"
                            "足りません。行を増やすか説明変数を減らしてください。")

    Xc = sm.add_constant(X, has_constant="add")
    tables, notes = [], []
    if method == "ols":
        model = sm.OLS(y, Xc).fit()
        fit_rows = [["決定係数 R²", model.rsquared], ["自由度調整済みR²", model.rsquared_adj],
                    ["F値", model.fvalue], ["モデルのp値", model.f_pvalue],
                    ["AIC", model.aic], ["件数", int(model.nobs)]]
    elif method == "logistic":
        uniq = set(pd.unique(y.dropna()))
        if not uniq <= {0, 1}:
            raise AnalysisError(f"ロジスティック回帰の目的変数は0か1にしてください"
                                f"（いまの値: {sorted(uniq)[:5]}）。"
                                "SQL側で CASE WHEN ... THEN 1 ELSE 0 END にしてください。")
        model = sm.Logit(y, Xc).fit(disp=0)
        fit_rows = [["疑似R²(McFadden)", model.prsquared], ["対数尤度", model.llf],
                    ["AIC", model.aic], ["件数", int(model.nobs)]]
    else:
        model = sm.GLM(y, Xc, family=sm.families.Poisson()).fit()
        fit_rows = [["対数尤度", model.llf], ["AIC", model.aic], ["件数", int(model.nobs)]]

    coef = pd.DataFrame({
        "変数": model.params.index,
        "係数": model.params.to_numpy(),
        "標準誤差": model.bse.to_numpy(),
        "p値": model.pvalues.to_numpy(),
    })
    ci = model.conf_int()
    coef["95%下限"], coef["95%上限"] = ci.iloc[:, 0].to_numpy(), ci.iloc[:, 1].to_numpy()
    coef["判定"] = np.where(coef["p値"] < alpha, "有意", "")
    if method == "logistic":
        coef["オッズ比"] = np.exp(coef["係数"])
    elif method == "poisson":
        coef["倍率"] = np.exp(coef["係数"])
    cols, rws = _out(coef)
    tables.append(_advanced_table("係数", cols, rws))
    tables.append(_advanced_table("あてはまり", ["項目", "値"], fit_rows))

    sig = coef[(coef["p値"] < alpha) & (coef["変数"] != "const")]
    if len(sig):
        top = sig.reindex(sig["係数"].abs().sort_values(ascending=False).index)
        notes.append("効いている変数: " + "、".join(
            f"{r['変数']}（係数 {r['係数']:.4g}）" for _, r in top.head(5).iterrows()))
    else:
        notes.append("p値が0.05を下回る説明変数はありませんでした。"
                     "この説明変数の組み合わせでは目的変数を説明できていません。")
    if method == "ols":
        notes.append(f"あてはまり: R² = {model.rsquared:.3f}（目的変数のばらつきの"
                     f"{model.rsquared * 100:.1f}%を説明）。"
                     + ("説明力は弱いので、変数の追加を検討してください。"
                        if model.rsquared < 0.3 else ""))
        # 多重共線性。説明変数どうしが強く相関していると係数の解釈ができない
        if X.shape[1] > 1:
            from statsmodels.stats.outliers_influence import variance_inflation_factor
            vifs = []
            arr = Xc.to_numpy(dtype=float)
            for i, name in enumerate(Xc.columns):
                if name == "const":
                    continue
                try:
                    vifs.append([name, variance_inflation_factor(arr, i)])
                except Exception:
                    pass
            if vifs:
                tables.append(_advanced_table("多重共線性(VIF)", ["変数", "VIF"], vifs))
                bad = [n for n, v in vifs if v and v > 10]
                if bad:
                    notes.append(f"※ VIFが10を超える変数（{', '.join(bad)}）があります。"
                                 "説明変数どうしが似すぎていて、係数の意味を読み違えます。"
                                 "どちらかを外してください。")
        resid = model.resid
        dw = float(sm.stats.durbin_watson(resid))
        notes.append(f"残差の自己相関 Durbin-Watson = {dw:.2f}"
                     + ("（2から離れており、時系列の並びが残差に残っています）"
                        if dw < 1.5 or dw > 2.5 else "（おおむね問題なし）"))

    # 未知のデータでも通用するか。学習に使っていない行で確かめる。
    # R²は当てはめた本人のデータで測ると必ず良く出るので、予測用途では過大評価になる。
    if len(X) >= 30:
        hold = _holdout_score(sm, method, X, y, Xc.columns)
        if hold:
            tables.append(_advanced_table("検証（学習に使っていないデータでの成績）",
                                 ["項目", "値"], hold["rows"]))
            notes.append(hold["note"])
    notes.append("回帰は相関の構造を示すもので、因果を証明するものではありません。")

    meta = {"method": method, "n": len(X),
            "coefficients": {str(k): _clean(v) for k, v in model.params.items()},
            "r2": _clean(getattr(model, "rsquared", None)),
            "formula": f"{target} ~ " + " + ".join(map(str, X.columns))}

    if predict:
        rowsp = []
        for case in predict:
            vec = {"const": 1.0}
            for c in X.columns:
                vec[c] = float(case.get(c, 0) or 0)
            xs = pd.DataFrame([[vec.get(c, 0.0) for c in Xc.columns]], columns=Xc.columns)
            yhat = float(model.predict(xs)[0])
            rowsp.append([", ".join(f"{k}={v}" for k, v in case.items()), yhat])
        tables.append(_advanced_table("予測", ["入力", "予測値"], rowsp))

    return {"title": REGRESSION_METHODS[method], "tables": tables,
            "notes": notes, "meta": meta}


# =============================================================================
# 予測（時系列）
# =============================================================================

def lag_correlation(columns: list, rows: list, target: str, features: list | None,
                    max_lag: int = 6, method: str = "pearson") -> dict:
    """時差相関。「今月の広告費は翌月の売上に効くのか」を見る。

    行は時点の昇順に並んでいる前提（SQLで ORDER BY してもらう）。
    lag=k は「説明側をk期ずらして、後の target と突き合わせる」の意味。
    """
    df = _df(columns, rows)
    if target not in df.columns:
        raise AnalysisError(f"target の列 '{target}' がありません。")
    feats = [c for c in (features or df.columns) if c != target and c in df.columns]
    feats = [c for c in feats if pd.to_numeric(df[c], errors="coerce").notna().mean() >= 0.8]
    if not feats:
        raise AnalysisError("比べる数値列がありません。columns で指定してください。")
    y = pd.to_numeric(df[target], errors="coerce")
    max_lag = max(1, min(int(max_lag or 6), max(1, len(df) // 3)))

    out, best_lines = [], []
    for f in feats:
        x = pd.to_numeric(df[f], errors="coerce")
        row, best = [f], (0.0, 0)
        for k in range(0, max_lag + 1):
            pair = pd.DataFrame({"x": x.shift(k), "y": y}).dropna()
            r = (float(pair["x"].corr(pair["y"], method=method))
                 if len(pair) >= 3 else float("nan"))
            row.append(None if math.isnan(r) else round(r, 3))
            if not math.isnan(r) and abs(r) > abs(best[0]):
                best = (r, k)
        out.append(row)
        if best[1] > 0:
            best_lines.append(
                f"{f} は {best[1]}期ずらしたときが最も強く（相関 {best[0]:+.2f}）、"
                f"同時点（{row[1] if row[1] is not None else '—'}）より強い。"
                f"{f}の効果が{best[1]}期あとに出ている可能性があります。")
        else:
            best_lines.append(f"{f} は同時点が最も強い（相関 {best[0]:+.2f}）。ずれは見られません。")

    notes = ["lag=k は「説明側をk期ずらして、あとの " + target + " と比べた」という意味です。"]
    notes += best_lines
    notes.append("時差があるからといって原因とは限りません。両方が同じ季節要因で"
                 "動いているだけのこともあります。")
    return {"title": f"{target} との時差相関（最大{max_lag}期）",
            "tables": [_advanced_table("時差ごとの相関",
                              ["列"] + [f"lag={k}" for k in range(0, max_lag + 1)], out)],
            "notes": notes, "meta": {"max_lag": max_lag}}


def partial_correlation(columns: list, rows: list, features: list | None,
                        control: list, method: str = "pearson") -> dict:
    """偏相関。指定した列の影響を取り除いてから相関を見る。

    「気温を除いても売上と来店数は連動しているか」のように、
    第3の変数のせいで相関して見えているだけ、を切り分けるためのもの。
    """
    df = _df(columns, rows)
    ctrl = [c for c in (control or []) if c in df.columns]
    if not ctrl:
        raise AnalysisError("control（影響を取り除きたい列）を1つ以上指定してください。")
    feats = [c for c in (features or df.columns) if c in df.columns and c not in ctrl]
    feats = [c for c in feats if pd.to_numeric(df[c], errors="coerce").notna().mean() >= 0.8]
    if len(feats) < 2:
        raise AnalysisError("偏相関には数値列が2つ以上必要です。")

    d = _advanced_num(df, feats + ctrl, 3)
    # 各列から control 成分を回帰で抜き、その残差どうしの相関を見る
    resid = {}
    c_mat = np.column_stack([np.ones(len(d))] + [d[c].to_numpy(float) for c in ctrl])
    for f in feats:
        yv = d[f].to_numpy(float)
        beta, *_ = np.linalg.lstsq(c_mat, yv, rcond=None)
        resid[f] = yv - c_mat @ beta
    rd = pd.DataFrame(resid)
    pc = rd.corr(method=method).round(3)
    raw = d[feats].corr(method=method).round(3)

    table = pc.reset_index().rename(columns={"index": "列"})
    cols, rws = _out(table)

    notes = [f"{'、'.join(ctrl)} の影響を取り除いた相関です。"]
    for i, a in enumerate(feats):
        for b in feats[i + 1:]:
            r0, r1 = float(raw.loc[a, b]), float(pc.loc[a, b])
            if abs(r0) >= 0.3 and abs(r1) < abs(r0) * 0.5:
                notes.append(f"{a} と {b}: 見かけ {r0:+.2f} → 取り除くと {r1:+.2f}。"
                             f"この関係の多くは {'、'.join(ctrl)} で説明できます。")
            elif abs(r1) >= 0.3:
                notes.append(f"{a} と {b}: 見かけ {r0:+.2f} → 取り除いても {r1:+.2f}。"
                             "取り除いた変数だけでは説明できない関係が残っています。")
    if len(notes) == 1:
        notes.append("目立った変化はありませんでした。")
    return {"title": f"偏相関（{'、'.join(ctrl)} の影響を除く）",
            "tables": [_advanced_table("偏相関行列", cols, rws)], "notes": notes,
            "meta": {"control": ctrl}}


def detect_anomalies(columns: list, rows: list, time_col: str, value_col: str,
                     window: int = 7, threshold: float = 3.0,
                     season_length: int | None = None,
                     changepoints: bool = True) -> dict:
    """時系列の「いつもと違う時点」と「いつから変わったか」を出す。

    静的な外れ値（outliers）は、右肩上がりのデータだと最近の値を全部
    外れ値と言ってしまう。ここでは直前の期間と比べるので、水準が変わっても
    「その時点として不自然か」を見られる。
    """
    df = _df(columns, rows)
    for c in (time_col, value_col):
        if c not in df.columns:
            raise AnalysisError(f"列が見つかりません: {c}")
    d = df[[time_col, value_col]].copy()
    d[value_col] = pd.to_numeric(d[value_col], errors="coerce")
    d = d.dropna().sort_values(time_col).reset_index(drop=True)
    if len(d) < 8:
        raise AnalysisError(f"異常検知には最低8点必要です（いま {len(d)} 点）。")

    y = d[value_col]
    window = max(3, min(int(window or 7), max(3, len(d) // 3)))
    base = y

    notes = []
    if season_length and len(d) >= season_length * 2:
        # 曜日・月の周期そのものは異常ではない。季節ぶんを引いてから見る
        from statsmodels.tsa.seasonal import seasonal_decompose
        try:
            dec = seasonal_decompose(y, period=int(season_length),
                                     model="additive", extrapolate_trend="freq")
            base = y - dec.seasonal
            notes.append(f"周期{season_length}の季節変動を取り除いてから判定しました。")
        except Exception:
            pass

    # 中央値と MAD（中央絶対偏差）。平均と標準偏差だと異常値自身に引っ張られる
    med = base.rolling(window, center=True, min_periods=3).median()
    mad = (base - med).abs().rolling(window, center=True, min_periods=3).median()
    scale = mad * 1.4826                       # MADを標準偏差の尺度に合わせる係数
    scale = scale.replace(0, np.nan).fillna(base.std(ddof=0) or 1.0)
    score = (base - med) / scale

    d["移動中央値"] = med.round(4)
    d["乖離スコア"] = score.round(2)
    d["判定"] = np.where(score.abs() >= threshold,
                         np.where(score > 0, "高い", "低い"), "")
    hit = d[d["判定"] != ""].copy()
    hit = hit.reindex(hit["乖離スコア"].abs().sort_values(ascending=False).index)

    cols, rws = _out(hit[[time_col, value_col, "移動中央値", "乖離スコア", "判定"]].head(100))
    tables = [_advanced_table("いつもと違う時点", cols, rws)]

    notes.insert(0, f"前後{window}期の中央値から、ばらつきの{threshold}倍以上離れた時点を"
                    f"拾いました。{len(hit)}件（全{len(d)}期中 {len(hit) / len(d) * 100:.1f}%）。")
    if len(hit):
        worst = hit.iloc[0]
        notes.append(f"最も外れているのは {worst[time_col]}（{value_col} = "
                     f"{worst[value_col]:,.4g}、通常は {worst['移動中央値']:,.4g} 前後）。")
    else:
        notes.append("目立って外れた時点はありませんでした。")

    cps = []
    if changepoints and len(d) >= 12:
        cps = _changepoints(base.to_numpy(dtype=float))
        if cps:
            crows = []
            for i in cps:
                before, after = base[:i].mean(), base[i:].mean()
                crows.append([str(d[time_col].iloc[i]), round(float(before), 4),
                              round(float(after), 4),
                              f"{(after - before) / before * 100:+.1f}%" if before else "—"])
            tables.append(_advanced_table("水準が変わった時点", ["時点", "その前の平均", "その後の平均", "変化"],
                                 crows))
            notes.append("「水準が変わった時点」は、そこを境に平均が段差になっている所です。"
                         "施策や仕様変更の日付と突き合わせてみてください。")
        else:
            notes.append("平均が段差になるような変化点は見つかりませんでした。")

    return {"title": f"{value_col} の異常検知",
            "tables": tables, "notes": notes,
            "meta": {"anomalies": len(hit), "window": window,
                     "threshold": threshold, "changepoints": len(cps)}}


def _changepoints(y: np.ndarray, max_cuts: int = 3, min_seg: int = 4) -> list:
    """平均が切り替わった位置を探す（二分割を繰り返す素朴な方法）。

    残差平方和がいちばん減る切れ目を選び、減り方が小さくなったらやめる。
    ライブラリを増やさずに「いつから変わったか」に答えるための最小限の実装。
    """
    def best_cut(a: int, b: int):
        seg = y[a:b]
        if len(seg) < min_seg * 2:
            return None
        total = float(((seg - seg.mean()) ** 2).sum())
        best, gain = None, 0.0
        for i in range(min_seg, len(seg) - min_seg):
            left, right = seg[:i], seg[i:]
            after = float(((left - left.mean()) ** 2).sum()
                          + ((right - right.mean()) ** 2).sum())
            if total - after > gain:
                best, gain = a + i, total - after
        # 全体のばらつきの1割も説明できない切れ目は、ただの揺らぎとみなす
        return best if best is not None and gain > total * 0.10 else None

    cuts, segments = [], [(0, len(y))]
    while len(cuts) < max_cuts and segments:
        found = []
        for a, b in segments:
            c = best_cut(a, b)
            if c is not None:
                found.append((c, a, b))
        if not found:
            break
        c, a, b = found[0]
        cuts.append(c)
        segments = [s for s in segments if s != (a, b)] + [(a, c), (c, b)]
    return sorted(cuts)


def survival_analysis(columns: list, rows: list, duration_col: str,
                      event_col: str | None = None, group_col: str | None = None,
                      fit_weibull: bool = True) -> dict:
    """生存時間分析。「どれだけ持つか」「いつ辞めるか」を扱う。

    設備の故障間隔（MTBF）、社員の在籍期間、顧客の継続期間に同じ道具が使える。
    event_col は「そのできごとが起きたか（1）／まだ起きていないか（0）」。
    まだ起きていない分（打ち切り）を捨てて平均を取ると、必ず短く見積もる。
    """
    df = _df(columns, rows)
    if duration_col not in df.columns:
        raise AnalysisError(f"列が見つかりません: {duration_col}")
    d = df.copy()
    d[duration_col] = pd.to_numeric(d[duration_col], errors="coerce")
    d = d[d[duration_col].notna() & (d[duration_col] >= 0)]
    if len(d) < 5:
        raise AnalysisError(f"分析には最低5件必要です（いま {len(d)} 件）。")
    if event_col and event_col in d.columns:
        ev = pd.to_numeric(d[event_col], errors="coerce").fillna(0)
        d["_event"] = (ev > 0).astype(int)
    else:
        d["_event"] = 1                     # 指定が無ければ全件で起きたものとして扱う

    def km(sub: pd.DataFrame):
        """カプラン・マイヤー法。打ち切りを含めても偏らない継続率の出し方。"""
        t = np.sort(sub.loc[sub["_event"] == 1, duration_col].unique())
        surv, out = 1.0, []
        for ti in t:
            at_risk = int((sub[duration_col] >= ti).sum())
            died = int(((sub[duration_col] == ti) & (sub["_event"] == 1)).sum())
            if at_risk:
                surv *= (1 - died / at_risk)
            out.append([float(ti), at_risk, died, round(surv, 4)])
        return out

    def median_surv(curve):
        for ti, _, _, s in curve:
            if s <= 0.5:
                return ti
        return None

    tables, notes = [], []
    groups = [(None, d)] if not (group_col and group_col in d.columns) \
        else list(d.groupby(group_col))
    summary = []
    for name, sub in groups:
        curve = km(sub)
        label = "全体" if name is None else str(name)
        if name is None:
            cols = ["時間", "対象数", "発生数", "継続率"]
            tables.append(_advanced_table("継続率（カプラン・マイヤー）", cols, curve))
        med = median_surv(curve)
        events = int(sub["_event"].sum())
        mtbf = float(sub.loc[sub["_event"] == 1, duration_col].mean()) if events else None
        summary.append([label, len(sub), events,
                        round(mtbf, 3) if mtbf is not None else None,
                        med if med is not None else "未到達"])
    tables.insert(0, _advanced_table("要約", [group_col or "対象", "件数", "発生件数",
                                    "平均(発生ぶんのみ)", "継続率が50%になる時間"], summary))

    censored = int((d["_event"] == 0).sum())
    if censored:
        notes.append(f"まだ起きていない（打ち切り）が {censored} 件あります。"
                     "これを捨てて平均すると短く見積もるため、継続率は"
                     "カプラン・マイヤー法で計算しています。")
    for row in summary:
        notes.append(f"{row[0]}: {row[1]}件中 {row[2]}件で発生。"
                     f"継続率が50%を切るのは {row[4]}。")

    if fit_weibull and int(d["_event"].sum()) >= 8:
        # 形状パラメータは「時間とともに壊れやすくなるか」を表す
        try:
            obs = d.loc[d["_event"] == 1, duration_col]
            obs = obs[obs > 0]
            shape, loc, scale = stats.weibull_min.fit(obs, floc=0)
            # 対象は設備とは限らない（在籍期間・契約の継続期間にも使う）ので、
            # 読み方も打ち手も特定の業種の言葉にしない
            trend = ("時間とともに起きやすくなる（経年で発生率が上がる型）" if shape > 1.2 else
                     "時間とともに起きにくくなる（初期に集中する型）" if shape < 0.8 else
                     "時間によらず一定の確率で起きる（偶発型）")
            tables.append(_advanced_table("Weibull分布の当てはめ",
                                 ["項目", "値"],
                                 [["形状パラメータ m", round(float(shape), 3)],
                                  ["尺度パラメータ η", round(float(scale), 3)],
                                  ["読み方", trend]]))
            notes.append(f"Weibullの形状パラメータ = {shape:.2f} → {trend}。"
                         + ("経過時間を基準にした事前の手当てが効きます。" if shape > 1.2 else
                            "初期段階での対応が効きます。" if shape < 0.8 else
                            "一定周期での対応をしても発生率は下がりません。"))
        except Exception:
            pass

    return {"title": "生存時間分析", "tables": tables, "notes": notes,
            "meta": {"n": len(d), "events": int(d["_event"].sum()),
                     "censored": censored}}


def _best_arima_order(y: np.ndarray) -> tuple:
    """ARIMAの次数を情報量規準(AIC)で選ぶ。

    以前は (1,1,1) 固定だった。データによっては当てはまらないのに
    「ARIMAで予測した」とだけ言うことになるので、素直な範囲を総当たりする。
    """
    import warnings

    from statsmodels.tsa.api import ARIMA

    best, best_aic = (1, 1, 1), np.inf
    s = pd.Series(y)
    # 総当たりの途中で収束しない組み合わせは当然出る。警告はログを埋めるだけなので黙らせる
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        for p in range(0, 3):
            for dd in range(0, 2):
                for q in range(0, 3):
                    if p == 0 and q == 0:
                        continue
                    try:
                        aic = float(ARIMA(s, order=(p, dd, q)).fit().aic)
                    except Exception:
                        continue
                    if np.isfinite(aic) and aic < best_aic:
                        best, best_aic = (p, dd, q), aic
    return best


def _backtest_note(columns: list, rows: list, time_col: str, value_col: str,
                   method: str, season_length: int | None, window: int,
                   exog: dict | None) -> tuple:
    """原点をずらしながら何度も予測して、誤差率(MAPE)の平均を出す。

    末尾を1回だけ隠す方式だと、たまたま当たった／外れただけの数字になる。
    予測できる範囲で3回まで試し、ばらつきも見えるようにする。
    """
    n = len(rows)
    hold = max(1, min(3, n // 6))
    folds, errs, maes, skipped = min(3, max(1, (n - 8) // hold)) if n >= 10 else 1, [], [], 0
    scale = float(np.nanmean(np.abs(pd.to_numeric(
        [r[list(columns).index(value_col)] for r in rows], errors="coerce"))))
    for f in range(folds):
        cut = n - hold * (f + 1)
        if cut < 6:
            break
        try:
            back = forecast(columns, rows[:cut], time_col, value_col,
                            periods=hold, method=method,
                            season_length=season_length, window=window,
                            exog=exog, _backtest=False)
            got = np.array([r[1] for r in back["tables"][1]["rows"]], dtype=float)
        except Exception:
            continue
        act = np.array([pd.to_numeric(r[list(columns).index(value_col)], errors="coerce")
                        for r in rows[cut:cut + hold]], dtype=float)
        if len(act) != len(got) or np.isnan(act).any():
            continue
        maes.append(float(np.mean(np.abs(act - got))))
        # 実績が0に近い期を割り算に含めると誤差率が跳ね上がる。数だけ覚えて除く
        small = np.abs(act) < max(scale * 0.01, 1e-9)
        skipped += int(small.sum())
        with np.errstate(divide="ignore", invalid="ignore"):
            e = np.abs((act - got) / np.where(small, np.nan, act))
        e = e[~np.isnan(e)]
        if len(e):
            errs.append(float(np.mean(e)) * 100)
    return {"mape": float(np.mean(errs)) if errs else None,
            "sd": float(np.std(errs)) if len(errs) > 1 else None,
            "mae": float(np.mean(maes)) if maes else None,
            "scale": scale, "folds": len(maes), "hold": hold, "skipped": skipped}


def forecast(columns: list, rows: list, time_col: str, value_col: str,
             periods: int = 6, method: str = "auto",
             season_length: int | None = None, window: int = 3,
             exog: dict | None = None, _backtest: bool = True) -> dict:
    """将来の値を予測する。時系列は time_col の昇順に並べ替えて使う。

    exog を渡すと説明変数つきで予測する（例: 広告費を来期こう置いたら売上はどうなるか）。
      {"columns": ["広告費"], "future": [[120], [130], ...]}
    """
    df = _df(columns, rows)
    if time_col not in df.columns or value_col not in df.columns:
        raise AnalysisError(f"列が見つかりません（{time_col} / {value_col}）。")
    keep = [time_col, value_col] + list((exog or {}).get("columns") or [])
    missing = [c for c in keep if c not in df.columns]
    if missing:
        raise AnalysisError(f"列が見つかりません: {', '.join(missing)}")
    d = df[keep].copy()
    for c in keep[1:]:
        d[c] = pd.to_numeric(d[c], errors="coerce")
    d = d.dropna().sort_values(time_col)
    if len(d) < 4:
        raise AnalysisError(f"予測には最低4点必要です（いま {len(d)} 点）。")
    y = d[value_col].to_numpy(dtype=float)
    labels = [str(v) for v in d[time_col]]
    periods = max(1, min(int(periods or 6), 120))

    ex_cols = list((exog or {}).get("columns") or [])
    ex_future = (exog or {}).get("future") or []
    if ex_cols:
        # 説明変数つきは ARIMA(X) でしか扱えないので、方法を寄せる
        if len(ex_future) != periods:
            raise AnalysisError(
                f"説明変数つきの予測には、予測する期数と同じ数だけ将来の値が要ります"
                f"（periods={periods} に対して future は {len(ex_future)} 件）。")
        method = "arima"

    if method == "auto":
        if season_length and len(y) >= season_length * 2:
            method = "holt_winters"
        elif len(y) >= 8:
            method = "holt"
        else:
            method = "linear"
    if method not in FORECAST_METHODS:
        raise AnalysisError(f"未対応の予測方法です: {method}。"
                            f"使えるのは {', '.join(FORECAST_METHODS)} です。")

    notes, lower, upper = [], None, None
    idx = np.arange(len(y), dtype=float)
    future_idx = np.arange(len(y), len(y) + periods, dtype=float)

    if method == "naive":
        pred = np.full(periods, y[-1])
    elif method == "drift":
        slope = (y[-1] - y[0]) / max(len(y) - 1, 1)
        pred = y[-1] + slope * np.arange(1, periods + 1)
    elif method == "moving_average":
        w = max(2, min(int(window or 3), len(y)))
        pred = np.full(periods, y[-w:].mean())
    elif method == "linear":
        sl, ic, r, p, se = stats.linregress(idx, y)
        pred = ic + sl * future_idx
        resid_sd = np.std(y - (ic + sl * idx), ddof=2) if len(y) > 2 else 0.0
        lower, upper = pred - 1.96 * resid_sd, pred + 1.96 * resid_sd
        notes.append(f"傾き = {sl:.4g}/期（{'増加' if sl > 0 else '減少'}傾向）、"
                     f"決定係数 R² = {r ** 2:.3f}、傾きのp値 = {p:.4g}")
    else:
        from statsmodels.tsa.api import ARIMA, ExponentialSmoothing
        s = pd.Series(y)
        try:
            if method == "holt":
                fit = ExponentialSmoothing(s, trend="add").fit()
            elif method == "holt_winters":
                if not season_length or len(y) < season_length * 2:
                    raise AnalysisError(
                        f"季節性ありの予測には、季節の長さ（season_length）の2周期ぶん"
                        f"以上のデータが必要です（いま {len(y)}点 / 季節 {season_length}）。")
                fit = ExponentialSmoothing(s, trend="add", seasonal="add",
                                           seasonal_periods=int(season_length)).fit()
            else:
                # 列名を持たせたまま渡す。そうしないと係数が x1 になり、どの変数か分からなくなる
                ex_hist = (d[ex_cols].astype(float).reset_index(drop=True)
                           if ex_cols else None)
                order = _best_arima_order(y)
                fit = ARIMA(s, order=order, exog=ex_hist).fit()
                notes.append(f"ARIMAの次数は当てはまりの良さ(AIC)で選びました: {order}")
                if ex_cols:
                    coefs = ", ".join(
                        f"{c}={float(fit.params.get(c, float('nan'))):+.4g}" for c in ex_cols)
                    notes.append(f"説明変数の係数（1単位増えたときの{value_col}への効き）: {coefs}")
            if ex_cols:
                fut = pd.DataFrame(list(ex_future), columns=ex_cols).astype(float)
                pred = np.asarray(fit.forecast(periods, exog=fut), dtype=float)
            else:
                pred = np.asarray(fit.forecast(periods), dtype=float)
            resid_sd = float(np.std(fit.resid, ddof=1)) if len(fit.resid) > 1 else 0.0
            lower, upper = pred - 1.96 * resid_sd, pred + 1.96 * resid_sd
        except AnalysisError:
            raise
        except Exception as e:
            raise AnalysisError(f"{FORECAST_METHODS[method]}の当てはめに失敗しました: {e}。"
                                "linear など簡単な方法を試してください。") from e

    fut_labels = [f"+{i}期" for i in range(1, periods + 1)]
    frows = []
    for i, v in enumerate(pred):
        frows.append([fut_labels[i], v,
                      lower[i] if lower is not None else None,
                      upper[i] if upper is not None else None])
    tables = [
        _advanced_table("実績", [time_col, value_col],
               [[labels[i], y[i]] for i in range(len(y))]),
        _advanced_table("予測", ["期", "予測値", "下限(95%)", "上限(95%)"], frows),
    ]
    # 当てはまりの目安。原点をずらして何度も試し、平均とばらつきで示す
    if _backtest and len(y) >= 8 and not ex_cols:
        bt = _backtest_note(columns, rows, time_col, value_col,
                            method, season_length, window, exog)
        if bt["mae"] is not None:
            line = f"過去データで{bt['folds']}回試した検証（毎回{bt['hold']}期先まで予測）: "
            line += f"平均が外れた幅 = {bt['mae']:,.4g}"
            if bt["scale"]:
                line += f"（{value_col}の平均 {bt['scale']:,.4g} に対して {bt['mae'] / bt['scale'] * 100:.0f}%）"
            if bt["mape"] is not None:
                line += f"、誤差率(MAPE) = {bt['mape']:.1f}%"
                if bt["sd"] is not None:
                    line += f"（回ごとのばらつき ±{bt['sd']:.1f}ポイント）"
                line += ("。かなり当たります。" if bt["mape"] < 10 else
                         "。実用的な精度です。" if bt["mape"] < 20 else
                         "。外れやすいので参考程度に見てください。")
            notes.append(line)
            if bt["skipped"]:
                notes.append(f"※ 実績が0に近い期が {bt['skipped']} 回あり、"
                             "誤差率の計算からは外しています（割り算が跳ね上がるため）。"
                             "そういう期がある場合は、誤差率より「外れた幅」で見てください。")
    notes.append(f"予測方法: {FORECAST_METHODS[method]}。"
                 "将来は過去の延長でしか計算していません。"
                 "施策や外部要因の変化は反映されないため、幅（上限・下限）も併せて見てください。")

    return {"title": f"{value_col} の予測（{periods}期先まで）", "tables": tables,
            "notes": notes,
            "meta": {"method": method, "periods": periods,
                     "history": [_clean(v) for v in y],
                     "labels": labels,
                     "forecast": [_clean(v) for v in pred],
                     "lower": [_clean(v) for v in (lower if lower is not None else [])],
                     "upper": [_clean(v) for v in (upper if upper is not None else [])],
                     "future_labels": fut_labels}}


def timeseries(columns: list, rows: list, time_col: str, value_col: str,
               window: int = 3, season_length: int | None = None) -> dict:
    """時系列の見方をまとめる（移動平均・前期比・季節分解・自己相関）。"""
    df = _df(columns, rows)
    d = df[[time_col, value_col]].copy()
    d[value_col] = pd.to_numeric(d[value_col], errors="coerce")
    d = d.dropna().sort_values(time_col).reset_index(drop=True)
    if len(d) < 3:
        raise AnalysisError(f"時系列の分析には3点以上必要です（いま {len(d)} 点）。")
    s = d[value_col]
    out = pd.DataFrame({
        time_col: d[time_col].astype(str),
        value_col: s,
        f"移動平均({window})": s.rolling(int(window or 3), min_periods=1).mean().round(4),
        "前期比": (s.pct_change() * 100).round(2),
        "累計": s.cumsum(),
    })
    if season_length and len(s) > season_length:
        out[f"前年同期比({season_length}期前)"] = (
            (s / s.shift(int(season_length)) - 1) * 100).round(2)
    cols, rws = _out(out)
    tables = [_advanced_table("推移", cols, rws)]
    notes = []

    sl, ic, r, p, _ = stats.linregress(np.arange(len(s)), s.to_numpy())
    notes.append(f"トレンド: 1期あたり {sl:+.4g}（p値 {p:.4g}）"
                 + ("、統計的に意味のある傾きです。" if p < ALPHA
                    else "、傾きは誤差の範囲です。"))
    notes.append(f"直近値 {s.iloc[-1]:,.4g} / 平均 {s.mean():,.4g} / "
                 f"最大 {s.max():,.4g} / 最小 {s.min():,.4g}")

    if season_length and len(s) >= season_length * 2:
        from statsmodels.tsa.seasonal import seasonal_decompose
        try:
            dec = seasonal_decompose(s, model="additive", period=int(season_length))
            comp = pd.DataFrame({
                time_col: d[time_col].astype(str),
                "実績": s, "トレンド": dec.trend.round(4),
                "季節": dec.seasonal.round(4), "残差": dec.resid.round(4)})
            c2, r2 = _out(comp)
            tables.append(_advanced_table("季節分解", c2, r2))
            amp = float(dec.seasonal.max() - dec.seasonal.min())
            notes.append(f"季節変動の振れ幅 = {amp:,.4g}"
                         f"（平均の {amp / (s.mean() or 1) * 100:.1f}%）")
        except Exception as e:
            notes.append(f"季節分解はできませんでした: {e}")

    if len(s) > 4:
        lags = min(6, len(s) - 2)
        ac = [[k, float(s.autocorr(lag=k))] for k in range(1, lags + 1)]
        tables.append(_advanced_table("自己相関（何期前と似ているか）", ["ラグ", "自己相関"], ac))
        best = max(ac, key=lambda x: abs(x[1]) if not math.isnan(x[1]) else 0)
        if abs(best[1]) > 0.5:
            notes.append(f"{best[0]}期前との相関が {best[1]:.2f} と高く、"
                         f"{best[0]}期の周期がありそうです。")

    return {"title": f"{value_col} の時系列分析", "tables": tables, "notes": notes,
            "meta": {"slope": _clean(sl), "p_value": _clean(p),
                     "labels": [str(v) for v in d[time_col]],
                     "values": [_clean(v) for v in s]}}


# =============================================================================
# 外れ値（analysis.outliers の拡張）
# =============================================================================

def outliers_ext(columns: list, rows: list, target: str | list, method: str = "iqr",
                 threshold: float | None = None) -> dict:
    """外れ値の抽出。1列でも複数列（マハラノビス距離）でも扱える。"""
    if method not in OUTLIER_METHODS_EXT:
        raise AnalysisError(f"未対応の方法です: {method}。"
                            f"使えるのは {', '.join(OUTLIER_METHODS_EXT)} です。")
    df = _df(columns, rows)
    targets = [target] if isinstance(target, str) else list(target)
    notes, meta = [], {"method": method}

    if method == "mahalanobis":
        d = _advanced_num(df, targets, len(targets) + 1)
        arr = d.to_numpy(dtype=float)
        cov = np.cov(arr, rowvar=False)
        try:
            inv = np.linalg.pinv(np.atleast_2d(cov))
        except np.linalg.LinAlgError as e:
            raise AnalysisError(f"分散共分散行列を逆行列にできません: {e}") from e
        center = arr.mean(axis=0)
        md = np.sqrt(np.einsum("ij,jk,ik->i", arr - center, inv, arr - center))
        cut = float(threshold) if threshold else math.sqrt(stats.chi2.ppf(0.975, len(targets)))
        flag = md > cut
        res = df.loc[d.index].copy()
        res["距離"] = md.round(4)
        res = res[flag].sort_values("距離", ascending=False)
        notes.append(f"{len(targets)}列をまとめて見て、距離が {cut:.2f} を超えた"
                     f"{int(flag.sum())}件を外れ値としました"
                     f"（全 {len(d)}件中 {flag.mean() * 100:.1f}%）。")
    else:
        col = targets[0]
        s = _advanced_num(df, [col], 3)[col]
        if method == "iqr":
            thr = float(threshold) if threshold is not None else 1.5
            q1, q3 = s.quantile(0.25), s.quantile(0.75)
            iqr = q3 - q1
            lo, hi = q1 - thr * iqr, q3 + thr * iqr
        elif method == "zscore":
            thr = float(threshold) if threshold is not None else 3.0
            m, sd = s.mean(), s.std(ddof=1)
            lo, hi = m - thr * sd, m + thr * sd
        elif method == "modified_zscore":
            thr = float(threshold) if threshold is not None else 3.5
            med = s.median()
            mad = (s - med).abs().median()
            scale = 1.4826 * mad
            if not scale:
                raise AnalysisError("中央絶対偏差が0のため判定できません"
                                    "（同じ値ばかりです）。")
            lo, hi = med - thr * scale, med + thr * scale
        else:                                      # percentile
            thr = float(threshold) if threshold is not None else 1.0
            lo, hi = s.quantile(thr / 100), s.quantile(1 - thr / 100)
        flag = (s < lo) | (s > hi)
        res = df.loc[s.index][flag.to_numpy()].copy()
        res = res.assign(**{f"{col}_逸脱": np.where(s[flag] < lo, "下振れ", "上振れ")})
        notes.append(f"{col} の正常範囲を {lo:,.4g} 〜 {hi:,.4g} と見て、"
                     f"外れた {int(flag.sum())}件を抽出しました"
                     f"（全 {len(s)}件中 {flag.mean() * 100:.1f}%）。")
        meta.update(lower=_clean(lo), upper=_clean(hi))

    if len(res) > 500:
        notes.append(f"※ {len(res)}件のうち上位500件だけ表示します。")
        res = res.head(500)
    cols, rws = _out(res)
    meta["count"] = len(rws)
    notes.append("外れ値＝誤りとは限りません。実際に起きた特異な事象か、"
                 "入力ミスかを、元データに当たって確かめてください。")
    return {"title": f"外れ値（{OUTLIER_METHODS_EXT[method]}）",
            "tables": [_advanced_table("外れ値", cols, rws)], "notes": notes, "meta": meta}


# =============================================================================
# 分布
# =============================================================================

def distribution(columns: list, rows: list, target: str, bins: int = 20,
                 fit: list | None = None, group_col: str | None = None) -> dict:
    """分布の形を見る。ヒストグラムの度数表と、当てはまる分布の推定。"""
    df = _df(columns, rows)
    s = _advanced_num(df, [target], 3)[target]
    counts, edges = np.histogram(s.to_numpy(dtype=float), bins=int(bins or 20))
    hist = [[f"{edges[i]:,.4g} 〜 {edges[i + 1]:,.4g}", int(counts[i]),
             round(float(counts[i]) / len(s) * 100, 2)] for i in range(len(counts))]
    tables = [_advanced_table("度数分布", ["区間", "件数", "割合(%)"], hist)]
    q = s.quantile([0, .05, .25, .5, .75, .95, 1])
    tables.append(_advanced_table("要約", ["項目", "値"], [
        ["件数", len(s)], ["平均", s.mean()], ["標準偏差", s.std(ddof=1)],
        ["最小", q.iloc[0]], ["5%", q.iloc[1]], ["25%", q.iloc[2]], ["中央値", q.iloc[3]],
        ["75%", q.iloc[4]], ["95%", q.iloc[5]], ["最大", q.iloc[6]],
        ["歪度", stats.skew(s)], ["尖度", stats.kurtosis(s)]]))
    notes = []
    sk = float(stats.skew(s))
    notes.append(f"歪み: {sk:+.2f}"
                 + ("（右に長い裾。平均より中央値を見た方が実感に合います）" if sk > 0.5
                    else "（左に長い裾）" if sk < -0.5 else "（左右対称に近い）"))

    if group_col and group_col in df.columns:
        summary = df.loc[s.index].assign(**{target: s}).groupby(group_col)[target].agg(
            ["count", "mean", "median", "std", "min", "max"]).round(4)
        summary.columns = ["件数", "平均", "中央値", "標準偏差", "最小", "最大"]
        tables.append(_df_table("グループ別", summary, index_label=str(group_col)))

    fit = fit or ["norm", "lognorm"]
    frows = []
    for name in fit:
        if name not in DISTRIBUTIONS:
            continue
        dist = getattr(stats, name)
        try:
            data = s.to_numpy(dtype=float)
            if name in ("lognorm", "expon", "gamma") and (data <= 0).any():
                frows.append([DISTRIBUTIONS[name], "―", "―", "0以下の値があるため当てはめ不可"])
                continue
            params = dist.fit(data)
            ks, p = stats.kstest(data, name, args=params)
            frows.append([DISTRIBUTIONS[name], round(float(ks), 4), round(float(p), 4),
                          "当てはまる" if p >= ALPHA else "当てはまらない"])
        except Exception as e:
            frows.append([DISTRIBUTIONS[name], "―", "―", f"失敗: {e}"])
    if frows:
        tables.append(_advanced_table("分布の当てはめ（KS検定）",
                             ["分布", "KS統計量", "p値", "判定"], frows))
        ok = [r[0] for r in frows if r[3] == "当てはまる"]
        notes.append("当てはまる分布: " + ("、".join(ok) if ok else
                     "候補なし（実データ特有の形なので、シミュレーションでは"
                     "実データからの再抽出(empirical)を使ってください）"))

    return {"title": f"{target} の分布", "tables": tables, "notes": notes,
            "meta": {"bins": [_clean(e) for e in edges],
                     "counts": [int(c) for c in counts],
                     "mean": _clean(s.mean()), "std": _clean(s.std(ddof=1))}}


# =============================================================================
# シミュレーション
# =============================================================================

def _sample(rng, spec: dict, n: int, data: pd.DataFrame | None):
    """1つの入力変数のサンプルを作る。"""
    kind = (spec.get("dist") or "normal").lower()
    if kind == "empirical":
        col = spec.get("column")
        if data is None or col not in data.columns:
            raise AnalysisError(f"empirical には実データの列が要ります（{col} が見つかりません）。")
        pool = pd.to_numeric(data[col], errors="coerce").dropna().to_numpy()
        if not len(pool):
            raise AnalysisError(f"{col} に数値がありません。")
        return rng.choice(pool, size=n, replace=True)
    if kind in ("normal", "norm"):
        return rng.normal(float(spec.get("mean", 0)), abs(float(spec.get("std", 1))), n)
    if kind in ("uniform", "unif"):
        return rng.uniform(float(spec.get("min", 0)), float(spec.get("max", 1)), n)
    if kind in ("triangular", "triang"):
        lo, mode, hi = (float(spec.get("min", 0)), float(spec.get("mode", 0.5)),
                        float(spec.get("max", 1)))
        return rng.triangular(lo, min(max(mode, lo), hi), hi, n)
    if kind == "lognormal":
        return rng.lognormal(float(spec.get("mean", 0)), abs(float(spec.get("std", 1))), n)
    if kind == "poisson":
        return rng.poisson(float(spec.get("lam", 1)), n).astype(float)
    if kind == "binomial":
        return rng.binomial(int(spec.get("n", 1)), float(spec.get("p", 0.5)), n).astype(float)
    if kind in ("fixed", "const"):
        return np.full(n, float(spec.get("value", 0)))
    raise AnalysisError(f"未対応の分布です: {kind}。"
                        "normal / uniform / triangular / lognormal / poisson / "
                        "binomial / empirical / fixed から選んでください。")


_ALLOWED_FORMULA = set("0123456789.+-*/()<>=, _abcdefghijklmnopqrstuvwxyz"
                       "ABCDEFGHIJKLMNOPQRSTUVWXYZ")


def _check_formula(formula: str) -> str:
    """eval に渡す前に式を検証する。

    env から __builtins__ を外すだけでは `().__class__` 経由の抜け道が残るので、
    使える文字そのものを絞る。変数名に日本語を使うため、ASCII以外の文字は
    そのまま通し（記号は Unicode の記号類だけ弾く）、`__` は禁止する。
    """
    f = (formula or "").strip()
    if not f:
        raise AnalysisError("計算式を指定してください。")
    if "__" in f:
        raise AnalysisError("計算式に `__` は使えません。")
    bad = {ch for ch in f if ch.isascii() and ch not in _ALLOWED_FORMULA}
    if bad:
        raise AnalysisError(
            f"計算式に使えない文字が入っています: {' '.join(sorted(bad))}。"
            "使えるのは 数字・変数名・+ - * / ( ) と比較記号だけです。")
    return f


def monte_carlo(formula: str, variables: dict, trials: int = 10000,
                columns: list | None = None, rows: list | None = None,
                seed: int = 0, targets: list | None = None) -> dict:
    """モンテカルロ・シミュレーション。

    formula は変数名を使った式（例: "(単価 - 原価) * 数量 - 固定費"）。
    variables は {変数名: {"dist": "normal", "mean": 100, "std": 10}} の形。
    """
    formula = _check_formula(formula)
    if not variables:
        raise AnalysisError("入力変数を1つ以上指定してください。")
    trials = max(100, min(int(trials or 10000), 200000))
    data = _df(columns, rows) if columns else None
    rng = np.random.default_rng(int(seed or 0))

    samples, spec_rows = {}, []
    for name, spec in variables.items():
        samples[name] = _sample(rng, spec or {}, trials, data)
        spec_rows.append([name, (spec or {}).get("dist", "normal"),
                          ", ".join(f"{k}={v}" for k, v in (spec or {}).items()
                                    if k != "dist")])

    # 式は numpy 配列に対して評価する。名前は変数名だけに限定する。
    env = {"__builtins__": {}, "np": np, "min": np.minimum, "max": np.maximum,
           "abs": np.abs, "where": np.where, "sqrt": np.sqrt, "log": np.log,
           "exp": np.exp}
    env.update(samples)
    try:
        result = np.asarray(eval(formula, env), dtype=float)  # noqa: S307
    except Exception as e:
        raise AnalysisError(f"計算式を評価できませんでした: {e}。"
                            f"使える変数: {', '.join(variables)}") from e
    if result.ndim == 0:
        result = np.full(trials, float(result))
    result = result[np.isfinite(result)]
    if not len(result):
        raise AnalysisError("計算結果がすべて無限大かNaNになりました。式を見直してください。")

    qs = [1, 5, 10, 25, 50, 75, 90, 95, 99]
    pct = np.percentile(result, qs)
    tables = [
        _advanced_table("入力の前提", ["変数", "分布", "パラメータ"], spec_rows),
        _advanced_table("結果の要約", ["項目", "値"], [
            ["試行回数", len(result)], ["平均", result.mean()],
            ["標準偏差", result.std(ddof=1)], ["最小", result.min()],
            ["最大", result.max()], ["中央値", np.median(result)]]),
        _advanced_table("パーセンタイル", ["区分", "値"],
               [[f"P{q}", v] for q, v in zip(qs, pct)]),
    ]
    counts, edges = np.histogram(result, bins=20)
    tables.append(_advanced_table("結果の分布", ["区間", "件数"],
                         [[f"{edges[i]:,.4g} 〜 {edges[i + 1]:,.4g}", int(counts[i])]
                          for i in range(len(counts))]))

    notes = [f"平均 {result.mean():,.4g}、"
             f"9割方 {pct[1]:,.4g} 〜 {pct[7]:,.4g} の範囲に収まります（P5〜P95）。"]
    p_neg = float((result < 0).mean())
    if p_neg > 0:
        notes.append(f"マイナスになる確率 = {p_neg * 100:.1f}%")
    for t in (targets or []):
        try:
            th = float(t)
        except (TypeError, ValueError):
            continue
        notes.append(f"{th:,.4g} を超える確率 = {float((result > th).mean()) * 100:.1f}%")

    # 感度分析。どの入力が結果を動かしているか
    sens = []
    for name, arr in samples.items():
        if np.std(arr) == 0:
            continue
        r = float(np.corrcoef(arr[:len(result)], result)[0, 1])
        sens.append([name, r, abs(r)])
    if sens:
        sens.sort(key=lambda x: -x[2])
        tables.append(_advanced_table("感度（結果との相関）", ["変数", "相関", "影響の大きさ"], sens))
        notes.append(f"結果を最も左右するのは「{sens[0][0]}」です（相関 {sens[0][1]:.2f}）。"
                     "精度を上げたいなら、まずこの変数の見積もりを詰めてください。")
    notes.append("この結果は入力の前提に完全に依存します。前提が外れれば結論も変わります。")

    return {"title": "モンテカルロ・シミュレーション", "tables": tables, "notes": notes,
            "meta": {"trials": len(result), "mean": _clean(result.mean()),
                     "percentiles": {f"P{q}": _clean(v) for q, v in zip(qs, pct)},
                     "bins": [_clean(e) for e in edges],
                     "counts": [int(c) for c in counts]}}


def scenario(formula: str, scenarios: dict, base: dict | None = None) -> dict:
    """シナリオ比較（楽観・標準・悲観など）と感度の確認。"""
    formula = _check_formula(formula)
    if not scenarios:
        raise AnalysisError("シナリオを1つ以上指定してください。")
    env_base = dict(base or {})
    names, values, rows_out = [], [], []
    keys = sorted({k for s in scenarios.values() for k in s} | set(env_base))
    for label, over in scenarios.items():
        env = {"__builtins__": {}, "min": min, "max": max, "abs": abs}
        env.update(env_base)
        env.update(over or {})
        try:
            v = float(eval(formula, env))  # noqa: S307
        except Exception as e:
            raise AnalysisError(f"シナリオ「{label}」の計算に失敗しました: {e}") from e
        names.append(label)
        values.append(v)
        rows_out.append([label] + [env.get(k) for k in keys] + [v])

    tables = [_advanced_table("シナリオ比較", ["シナリオ"] + keys + ["結果"], rows_out)]
    best, worst = max(zip(names, values), key=lambda x: x[1]), min(zip(names, values), key=lambda x: x[1])
    notes = [f"最も良いのは「{best[0]}」で {best[1]:,.4g}、"
             f"最も悪いのは「{worst[0]}」で {worst[1]:,.4g}。"
             f"振れ幅は {best[1] - worst[1]:,.4g} です。"]

    # トルネード（各変数を単独で±10%動かしたときの影響）
    if env_base:
        tor = []
        try:
            b = float(eval(formula, {"__builtins__": {}, **env_base}))  # noqa: S307
            for k, v in env_base.items():
                if not isinstance(v, (int, float)):
                    continue
                lo_env, hi_env = dict(env_base), dict(env_base)
                lo_env[k], hi_env[k] = v * 0.9, v * 1.1
                lo = float(eval(formula, {"__builtins__": {}, **lo_env}))   # noqa: S307
                hi = float(eval(formula, {"__builtins__": {}, **hi_env}))   # noqa: S307
                tor.append([k, lo, hi, abs(hi - lo)])
            tor.sort(key=lambda x: -x[3])
            if tor:
                tables.append(_advanced_table("感度（各変数を±10%動かしたとき）",
                                     ["変数", "-10%のとき", "+10%のとき", "影響幅"], tor))
                notes.append(f"基準値 {b:,.4g}。影響が大きい順に "
                             + "、".join(x[0] for x in tor[:3]) + " です。")
        except Exception:
            pass
    return {"title": "シナリオ分析", "tables": tables, "notes": notes,
            "meta": {"scenarios": {n: _clean(v) for n, v in zip(names, values)}}}


def bootstrap(columns: list, rows: list, target: str, statistic: str = "mean",
              trials: int = 5000, group_col: str | None = None,
              seed: int = 0) -> dict:
    """ブートストラップ法で統計量の信頼区間を出す。分布の仮定が要らない。"""
    funcs = {"mean": ("平均", np.mean), "median": ("中央値", np.median),
             "std": ("標準偏差", lambda a: np.std(a, ddof=1)),
             "sum": ("合計", np.sum), "p90": ("90パーセンタイル",
                                             lambda a: np.percentile(a, 90))}
    if statistic not in funcs:
        raise AnalysisError(f"未対応の統計量です: {statistic}。"
                            f"使えるのは {', '.join(funcs)} です。")
    label, fn = funcs[statistic]
    df = _df(columns, rows)
    trials = max(200, min(int(trials or 5000), 50000))
    rng = np.random.default_rng(int(seed or 0))

    def ci(arr):
        boots = np.array([fn(rng.choice(arr, size=len(arr), replace=True))
                          for _ in range(trials)])
        return fn(arr), np.percentile(boots, 2.5), np.percentile(boots, 97.5)

    out_rows = []
    if group_col and group_col in df.columns:
        d = df[[group_col, target]].copy()
        d[target] = pd.to_numeric(d[target], errors="coerce")
        for k, g in d.dropna().groupby(group_col):
            if len(g) < 3:
                continue
            v, lo, hi = ci(g[target].to_numpy(dtype=float))
            out_rows.append([str(k), len(g), v, lo, hi])
        cols = ["群", "件数", label, "95%下限", "95%上限"]
    else:
        s = _advanced_num(df, [target], 3)[target].to_numpy(dtype=float)
        v, lo, hi = ci(s)
        out_rows.append([len(s), v, lo, hi])
        cols = ["件数", label, "95%下限", "95%上限"]

    notes = [f"{trials:,}回の再抽出から求めた信頼区間です。"
             "分布の形を仮定しないので、正規分布でないデータにも使えます。"]
    if len(out_rows) > 1:
        los = [r[3] for r in out_rows]
        his = [r[4] for r in out_rows]
        overlap = not (max(los) > min(his))
        notes.append("群どうしの信頼区間が"
                     + ("重なっていません。実質的な差がありそうです。" if not overlap
                        else "重なっています。群による差は断定できません。"))
    return {"title": f"{target} の{label}と信頼区間（ブートストラップ）",
            "tables": [_advanced_table("推定", cols, out_rows)], "notes": notes,
            "meta": {"statistic": statistic, "trials": trials}}


# =============================================================================
# クラスタリング / ABC分析
# =============================================================================

def _silhouette(x: np.ndarray, labels: np.ndarray, rng) -> float:
    """シルエット係数（-1〜1。大きいほど分かれ方が良い）。

    scikit-learn は入れていないので自前で出す。全点の総当たり距離は
    件数の2乗で効いてくるため、多いときは標本を抜いて評価する。
    """
    from scipy.spatial.distance import cdist

    n = len(x)
    if n > 800:                       # 総当たりが重くなる手前で標本に切り替える
        pick = rng.choice(n, 800, replace=False)
        x, labels = x[pick], labels[pick]
        n = len(x)
    uniq = np.unique(labels)
    if len(uniq) < 2:
        return -1.0
    dist = cdist(x, x)
    scores = []
    for i in range(n):
        same = labels == labels[i]
        same[i] = False
        if not same.any():
            continue                  # 1点だけのクラスタは評価に使えない
        a = dist[i][same].mean()
        b = min(dist[i][labels == u].mean() for u in uniq if u != labels[i])
        if max(a, b) > 0:
            scores.append((b - a) / max(a, b))
    return float(np.mean(scores)) if scores else -1.0


def _cluster_naming(flat: pd.DataFrame, features: list) -> list:
    """各クラスタが「何のグループか」を、全体平均とのズレから言葉にする。

    平均値の表だけ渡されても人は読み解けない。「単価は高いが頻度は低い層」
    のように、目立つ特徴だけを拾って一言にする。
    """
    names = []
    for _, row in flat.iterrows():
        marks = []
        for f in features:
            col = flat[f"{f}の平均"]
            base, sd = col.mean(), col.std(ddof=0)
            if sd and abs(row[f"{f}の平均"] - base) >= sd * 0.8:
                marks.append(f"{f}が{'高い' if row[f'{f}の平均'] > base else '低い'}")
        names.append("・".join(marks[:3]) if marks else "平均的")
    return names


def clustering(columns: list, rows: list, features: list, k: int | str = 3,
               label_col: str | None = None, seed: int = 0,
               categorical: list | None = None) -> dict:
    """k-meansでグループ分けする（顧客や店舗のセグメント分け）。

    k に "auto" を渡すと、シルエット係数が最も高い分割数を自分で選ぶ。
    categorical に区分の列を渡すと、0/1に開いてから一緒に分ける
    （地域や会員区分のように、数値でないが効いている属性を捨てない）。
    """
    from scipy.cluster.vq import kmeans2, whiten

    df = _df(columns, rows)
    auto = str(k).lower() == "auto"
    want = 3 if auto else int(k or 3)
    d = _advanced_num(df, list(features), max(want, 3))

    parts = [d.to_numpy(dtype=float)]
    used_cat = []
    for c in (categorical or []):
        if c not in df.columns:
            raise AnalysisError(f"列が見つかりません: {c}")
        dummies = pd.get_dummies(df.loc[d.index, c].astype(str), prefix=str(c))
        if 1 < dummies.shape[1] <= 20:      # 種類が多すぎる列は分割の役に立たない
            parts.append(dummies.to_numpy(dtype=float))
            used_cat.append(c)
    arr = np.hstack(parts)
    scaled = np.nan_to_num(whiten(arr))

    rng = np.random.default_rng(int(seed or 0))
    notes = []
    if auto:
        # 2〜8グループを試して、最も素直に分かれる数を採る
        upper = max(2, min(8, len(d) // 2))
        scored = []
        for cand in range(2, upper + 1):
            try:
                _, lb = kmeans2(scaled, cand, minit="++", seed=int(seed or 0))
                if len(np.unique(lb)) < cand:
                    continue                # 空のクラスタが出た分割は採らない
                scored.append((_silhouette(scaled, lb, rng), cand))
            except Exception:
                continue
        if not scored:
            raise AnalysisError("グループ分けできませんでした。件数か列を見直してください。")
        best, k = max(scored)
        notes.append("分割数はシルエット係数（分かれ方の良さ。1に近いほど良い）で選びました: "
                     + "、".join(f"{c}群={s:.2f}" for s, c in sorted(scored, key=lambda t: t[1]))
                     + f" → {k}群を採用")
        if best < 0.25:
            notes.append("どの分割数でも係数が低く、はっきりした群には分かれていません。"
                         "この結果は「たまたまの区切り」に近いので、扱いは慎重に。")
    else:
        k = max(2, min(want, max(2, len(d) // 2)))

    centroid, labels = kmeans2(scaled, k, minit="++", seed=int(seed or 0))

    res = df.loc[d.index].copy()
    res["クラスタ"] = [f"C{i + 1}" for i in labels]
    summary = res.groupby("クラスタ")[list(features)].agg(["count", "mean"])
    flat = pd.DataFrame({"クラスタ": summary.index})
    flat["件数"] = summary[(features[0], "count")].to_numpy()
    for f in features:
        flat[f"{f}の平均"] = summary[(f, "mean")].round(4).to_numpy()
    flat.insert(1, "特徴", _cluster_naming(flat, list(features)))
    cols, rws = _out(flat)
    tables = [_advanced_table("クラスタの特徴", cols, rws)]

    show = res
    if label_col and label_col in res.columns:
        show = res[[label_col] + list(features) + ["クラスタ"]]
    if len(show) > 500:
        show = show.head(500)
    c2, r2 = _out(show)
    tables.append(_advanced_table("割り当て（先頭500件）", c2, r2))

    sizes = flat["件数"].tolist()
    notes.append(f"{k}グループに分けました（件数: {', '.join(map(str, sizes))}）。")
    notes.append("各列は尺度が違うと結果が歪むため、標準化してから分けています。")
    if used_cat:
        notes.append(f"区分の列も0/1に開いて使いました: {', '.join(used_cat)}")
    for _, row in flat.iterrows():
        notes.append(f"{row['クラスタ']}（{row['件数']}件）: {row['特徴']}")
    return {"title": f"クラスタ分析（{k}グループ）", "tables": tables, "notes": notes,
            "meta": {"k": int(k), "sizes": [int(s) for s in sizes]}}


def abc_analysis(columns: list, rows: list, label_col: str, value_col: str,
                 thresholds: list | None = None) -> dict:
    """ABC分析（パレート）。売上の8割を占める品目を切り出す。"""
    df = _df(columns, rows)
    if label_col not in df.columns:
        raise AnalysisError(f"列が見つかりません: {label_col}")
    d = df[[label_col, value_col]].copy()
    d[value_col] = pd.to_numeric(d[value_col], errors="coerce")
    d = d.dropna().groupby(label_col, as_index=False)[value_col].sum()
    if d.empty:
        raise AnalysisError("集計できる行がありません。")
    d = d.sort_values(value_col, ascending=False).reset_index(drop=True)
    total = d[value_col].sum()
    d["構成比(%)"] = (d[value_col] / total * 100).round(2)
    d["累計構成比(%)"] = d["構成比(%)"].cumsum().round(2)
    cuts = thresholds or [70, 90]
    d["区分"] = np.where(d["累計構成比(%)"] <= cuts[0], "A",
                         np.where(d["累計構成比(%)"] <= cuts[1], "B", "C"))
    cols, rws = _out(d)
    # 見出しは渡された列名から作る。「金額」と決め打つと、停止時間や工数を渡したときに
    # 単位の違う見出しのままExcel・PowerPointへ出て行ってしまう
    cnt = "件数" if value_col != "件数" else "行数"      # 値の列が「件数」でも衝突しないように
    summary = d.groupby("区分").agg(**{cnt: (label_col, "count"),
                                      value_col: (value_col, "sum")}).reset_index()
    summary[f"{value_col}構成比(%)"] = (summary[value_col] / total * 100).round(2)
    summary[f"{cnt}構成比(%)"] = (summary[cnt] / len(d) * 100).round(2)
    c2, r2 = _out(summary)

    a = d[d["区分"] == "A"]
    notes = [f"全 {len(d)} 件のうち、上位 {len(a)} 件"
             f"（{len(a) / len(d) * 100:.1f}%）で {value_col} の "
             f"{a['構成比(%)'].sum():.1f}% を占めます。"]
    if len(a):
        notes.append("A区分: " + "、".join(map(str, a[label_col].head(10))))
    return {"title": f"ABC分析（{value_col}）",
            "tables": [_advanced_table("区分の要約", c2, r2), _advanced_table("明細", cols, rws)],
            "notes": notes,
            "meta": {"labels": [str(v) for v in d[label_col]],
                     "values": [_clean(v) for v in d[value_col]],
                     "cumulative": [_clean(v) for v in d["累計構成比(%)"]]}}


# ==========================================================================
# ===== 元 analysis.py
# SQLiteだけでは書けない集計・統計を pandas で行う。
#
# このアプリのSQLite(3.32)には次が無い:
#   STDDEV / VARIANCE / MEDIAN / CORR / PERCENTILE / SQRT / POWER / PIVOT構文
# そのため「相関」「中央値」「ばらつき」「クロス集計」はSQLでは実質書けない。
# ここではSELECT結果(columns, rows)を受け取り、同じ形(columns, rows)で返す。
# ==========================================================================
import math

import numpy as np
import pandas as pd

AGG_FUNCS = {
    "sum": ("合計", "sum"),
    "mean": ("平均", "mean"),
    "count": ("件数", "count"),
    "median": ("中央値", "median"),
    "min": ("最小", "min"),
    "max": ("最大", "max"),
    "std": ("標準偏差", "std"),
    "nunique": ("種類数", "nunique"),
}
CORR_METHODS = ("pearson", "spearman")
MARGIN_NAME = "合計"


def _df(columns: list, rows: list) -> pd.DataFrame:
    return pd.DataFrame([list(r) for r in rows], columns=list(columns))


def _to_numeric(df: pd.DataFrame, cols) -> pd.DataFrame:
    for c in cols:
        if c in df.columns:
            df[c] = pd.to_numeric(df[c], errors="coerce")
    return df


def numeric_columns(columns: list, rows: list) -> list:
    """数値として扱える列を推定する（8割以上が数値なら数値列とみなす）。"""
    df = _df(columns, rows)
    out = []
    for c in df.columns:
        s = pd.to_numeric(df[c], errors="coerce")
        if len(s) and s.notna().mean() >= 0.8:
            out.append(c)
    return out


def _clean(v):
    """JSONに載せられる形へ（NaN/Infとnumpy型を素の値にする）。"""
    if isinstance(v, (np.integer,)):
        return int(v)
    if isinstance(v, (np.floating, float)):
        f = float(v)
        return None if (math.isnan(f) or math.isinf(f)) else round(f, 6)
    if isinstance(v, (np.bool_,)):
        return bool(v)
    return v


def _out(df: pd.DataFrame):
    """DataFrame を (columns, rows) に戻す。"""
    cols = [str(c) for c in df.columns]
    rows = [tuple(_clean(v) for v in r) for r in df.itertuples(index=False, name=None)]
    return cols, rows


# --- クロス集計 -----------------------------------------------------------------

def _sorted_by(pt: pd.DataFrame, rank_by: str) -> pd.DataFrame:
    """大きい順に並べ替える。rank_by が列名ならその列、それ以外は行の合計で。

    合計（margins）の行は必ずいちばん大きいので、そのまま並べると先頭に来て
    「1位は合計」になってしまう。合計はどの分類でもないので、末尾に置く。
    """
    target = None
    for c in pt.columns:
        if str(c) == str(rank_by):
            target = c
            break
    key = pt[target] if target is not None else pt.sum(axis=1, numeric_only=True)
    order = list(key.sort_values(ascending=False).index)
    total = [i for i in order if _is_margin_row(i)]
    if total:
        order = [i for i in order if i not in total] + total
    return pt.loc[order]


def _is_margin_row(idx) -> bool:
    """その行が合計（margins）の行か。多段の見出しにも効くようにする。"""
    if isinstance(idx, tuple):
        return any(str(x) == MARGIN_NAME for x in idx)
    return str(idx) == MARGIN_NAME


def _as_percent(pt: pd.DataFrame, mode: str) -> pd.DataFrame:
    """実数を構成比(%)に置き換える。合計が0の行や列は0のままにする。"""
    num = pt.select_dtypes("number")
    if mode == "row":
        denom = num.sum(axis=1).replace(0, np.nan)
        out = num.div(denom, axis=0)
    elif mode == "column":
        denom = num.sum(axis=0).replace(0, np.nan)
        out = num.div(denom, axis=1)
    else:
        total = num.to_numpy().sum()
        out = num / (total if total else np.nan)
    pt = pt.copy()
    pt[num.columns] = (out * 100).round(1).fillna(0.0)
    return pt


#: 構成比の取り方。クロス集計は「実数で見たい」より「割合で見たい」ことが多い。
PERCENT_MODES = {
    "row": "行内の構成比（行ごとに合計100%）",
    "column": "列内の構成比（列ごとに合計100%）",
    "total": "全体に対する構成比（表全体で100%）",
}


def pivot(columns: list, rows: list, index: list, cols: str | None, values: str,
          aggfunc: str = "sum", fill_value=0, margins: bool = False,
          percent: str | None = None, rank_by: str | None = None):
    """クロス集計表を作る。SQLiteにPIVOT構文が無いのでここで行う。

    index   : 行にする列（複数可）
    cols    : 列に展開する列（省略可。省略時は index ごとの集計表になる）
    values  : 集計する値の列
    percent : row / column / total を指定すると実数を構成比(%)に置き換える
    rank_by : 指定した列（または合計）の大きい順に並べ、順位の列を先頭に足す
    """
    if not index:
        raise ValueError("index（行にする列）を1つ以上指定してください。")
    if not values:
        raise ValueError("values（集計する値の列）を指定してください。")
    if aggfunc not in AGG_FUNCS:
        raise ValueError(f"aggfunc は {', '.join(AGG_FUNCS)} のいずれかです。")
    if percent and percent not in PERCENT_MODES:
        raise ValueError(f"percent は {', '.join(PERCENT_MODES)} のいずれかです。")

    df = _df(columns, rows)
    missing = [c for c in list(index) + ([cols] if cols else []) + [values]
               if c not in df.columns]
    if missing:
        raise ValueError(f"指定列が結果にありません: {missing} / 利用可能: {list(df.columns)}")

    if aggfunc not in ("count", "nunique"):
        _to_numeric(df, [values])

    pt = pd.pivot_table(
        df, index=list(index), columns=cols, values=values,
        aggfunc=AGG_FUNCS[aggfunc][1],
        fill_value=fill_value, margins=margins and not percent,
        margins_name=MARGIN_NAME, dropna=False, observed=False,
    )
    if isinstance(pt, pd.Series):
        pt = pt.to_frame(name=values)

    # 並べ替えは構成比にする前に行う（%にすると行内の大小が消えることがある）
    if rank_by:
        pt = _sorted_by(pt, rank_by)
    if percent:
        pt = _as_percent(pt, percent)

    pt = pt.reset_index()
    if rank_by:
        # 合計の行には順位を振らない（どの分類でもないので、1位にすると嘘になる）
        marks, n = [], 0
        for _i, row in pt.iterrows():
            if any(str(v) == MARGIN_NAME for v in row.values):
                marks.append("")
            else:
                n += 1
                marks.append(n)
        pt.insert(0, "順位", marks)
    # 列がMultiIndex（valuesとcolsの2段）になる場合があるので平坦化する
    flat = []
    for c in pt.columns:
        if isinstance(c, tuple):
            parts = [str(p) for p in c if str(p) != ""]
            flat.append(" / ".join(parts) if parts else values)
        else:
            flat.append(str(c))
    pt.columns = flat
    return _out(pt)


# --- 基本統計量 -----------------------------------------------------------------

_DESC_LABELS = {"count": "件数", "mean": "平均", "std": "標準偏差", "min": "最小",
                "25%": "25%", "50%": "中央値", "75%": "75%", "max": "最大"}


def _advanced_describe(columns: list, rows: list, targets: list | None = None,
             group_by: str | None = None):
    """基本統計量（件数/平均/標準偏差/最小/四分位/中央値/最大）。"""
    df = _df(columns, rows)
    targets = list(targets or []) or numeric_columns(columns, rows)
    targets = [c for c in targets if c in df.columns and c != group_by]
    if not targets:
        raise ValueError("数値として集計できる列がありません。columns で対象列を指定してください。")
    _to_numeric(df, targets)

    if group_by:
        if group_by not in df.columns:
            raise ValueError(f"group_by の列 '{group_by}' が結果にありません。")
        out = []
        for key, g in df.groupby(group_by, dropna=False):
            d = g[targets].describe().T.reset_index().rename(columns={"index": "列"})
            d.insert(0, group_by, key)
            out.append(d)
        res = pd.concat(out, ignore_index=True)
    else:
        res = df[targets].describe().T.reset_index().rename(columns={"index": "列"})

    res = res.rename(columns=_DESC_LABELS)
    for c in res.columns:
        if c not in ("列", group_by):
            res[c] = pd.to_numeric(res[c], errors="coerce").round(3)
    return _out(res)


# --- 相関 -----------------------------------------------------------------------

def correlation(columns: list, rows: list, targets: list | None = None,
                method: str = "pearson"):
    """数値列どうしの相関行列。1列目が列名なので、そのままヒートマップにできる。"""
    if method not in CORR_METHODS:
        raise ValueError(f"method は {', '.join(CORR_METHODS)} のいずれかです。")
    df = _df(columns, rows)
    targets = list(targets or []) or numeric_columns(columns, rows)
    targets = [c for c in targets if c in df.columns]
    if len(targets) < 2:
        raise ValueError("相関には数値列が2つ以上必要です。columns で対象列を指定してください。")
    _to_numeric(df, targets)
    corr = df[targets].corr(method=method).round(3).reset_index()
    corr = corr.rename(columns={"index": "列"})
    return _out(corr)


def correlation_pairs(columns: list, rows: list, method: str = "pearson"):
    """相関の強い組み合わせを、強さ順のリストで返す（LLMへの説明用）。"""
    cols, mrows = correlation(columns, rows, None, method)
    names = cols[1:]
    pairs = []
    for i, r in enumerate(mrows):
        for j, v in enumerate(r[1:]):
            if j > i and v is not None:
                pairs.append({"a": names[i], "b": names[j], "corr": v})
    pairs.sort(key=lambda p: abs(p["corr"]), reverse=True)
    return pairs



# ==========================================================================
# ===== 元 models.py
# 使うモデルの選択。
#
# 2段構えになっている。
#   管理者 … 「モデル設定」画面で、選ばせる候補・既定・画像対応の判定を決める
#   利用者 … チャット画面のプルダウンで、その候補から自分の1つを選ぶ
#
# 利用者が選んだモデルは prefs.py（ユーザーごとのファイル）に残るので、
# ログアウトしても次に入ったときは同じモデルのまま。
#
# 候補の決まり方は 管理者の設定 > env の OPENAI_MODELS > APIの /models の順。
# 「そのモデルは画像を送れるか」も、ここで一元的に判断する。
# ==========================================================================
import threading
import time

import yaml

import config
import prefs

_models_lock = threading.Lock()
_models_cache: dict = {"at": 0.0, "models": []}
_CACHE_SEC = 300


# =============================================================================
# 管理者が決める設定（data/model_settings.yaml）
#
# env を初期値として、このファイルの内容で上書きする。
# メール設定と同じ考え方で、env は「まだ画面で決めていないときの値」。
# =============================================================================

ADMIN_KEYS = ("models", "default", "vision", "context_overrides", "api_key",
              "chat_url", "models_url")

#: カタログのインライン上限として認める範囲。
#: 下限は「1DBぶんの詳細（実測で平均5.3K字）が入る」ことを目安にした。
#: 上限は、いちばん広いモデルでも文脈を食い尽くさないところで止める。
INLINE_LIMIT_MIN = 4_000
INLINE_LIMIT_MAX = 400_000


def _read_admin() -> dict:
    p = config.MODEL_SETTINGS_FILE
    if not p.exists():
        return {}
    try:
        data = yaml.safe_load(p.read_text(encoding="utf-8")) or {}
    except Exception as e:
        print(f"[models] 設定を読めませんでした: {p} ({e})")
        return {}
    return {k: v for k, v in data.items() if k in ADMIN_KEYS} \
        if isinstance(data, dict) else {}


def _write_admin(data: dict) -> None:
    p = config.MODEL_SETTINGS_FILE
    p.parent.mkdir(parents=True, exist_ok=True)
    with _models_lock:
        p.write_text(yaml.safe_dump({k: data[k] for k in ADMIN_KEYS if k in data},
                                    allow_unicode=True, sort_keys=False),
                     encoding="utf-8")


def _vision_keys() -> list[str]:
    ov = _read_admin().get("vision")
    keys = ov if isinstance(ov, list) else None
    return [str(k).strip().lower() for k in (keys or config.OPENAI_VISION_MODELS)
            if str(k).strip()]


def default_model() -> str:
    """未選択のユーザーが使うモデル。"""
    return str(_read_admin().get("default") or config.OPENAI_MODEL or "").strip()


def llm_api_key() -> str:
    """LLM接続に実際に使うAPIキー。「モデル設定」画面で保存した値 > env の順。"""
    return str(_read_admin().get("api_key") or "").strip() or config.OPENAI_API_KEY


def llm_api_key_source() -> str:
    """キーの出所（"screen"=画面 / "env" / ""=未設定）。キー本体は決して返さない。"""
    if str(_read_admin().get("api_key") or "").strip():
        return "screen"
    return "env" if config.OPENAI_API_KEY else ""


# --- 接続先URL（フルパス2本。画面保存 > env の OPENAI_BASE_URL から導出）------------

def _env_chat_url() -> str:
    base = (config.OPENAI_BASE_URL or "").strip().rstrip("/")
    return base + "/chat/completions" if base else ""


def _env_models_url() -> str:
    base = (config.OPENAI_BASE_URL or "").strip().rstrip("/")
    return base + "/models" if base else ""


def llm_chat_url() -> str:
    """チャット（AI呼び出し）のフルURL。「モデル設定」画面で保存した値 > env の順。"""
    return str(_read_admin().get("chat_url") or "").strip() or _env_chat_url()


def llm_models_url() -> str:
    """モデル一覧取得のフルURL。「モデル設定」画面で保存した値 > env の順。"""
    return str(_read_admin().get("models_url") or "").strip() or _env_models_url()


def is_vision(model: str) -> bool:
    """画像を送れるモデルか。名前に手がかりが含まれるかで判断する。

    モデル名は環境によって違うので、完全一致ではなく部分一致にしている
    （「モデル設定」画面、または env の OPENAI_VISION_MODELS で調整できる）。
    """
    low = str(model or "").lower()
    return any(key in low for key in _vision_keys())


def prompt_inline_limit() -> int:
    """カタログをそのまま入れる量の天井（文字）。

    実効値はモデルの文脈から自動で決まる（inline_limit_for）。この天井は
    「文脈が100万トークンあるモデルでも、カタログに割く量はここまで」という
    安全弁で、画面から変えるものではない。
    """
    return INLINE_LIMIT_MAX


#: モデルの文脈のうち、カタログに使ってよい割合。
#: 残り半分はツール定義・会話の履歴・SQL結果・回答のために空けておく。
CATALOG_CONTEXT_RATIO = 0.5


def context_overrides() -> dict:
    """管理者が「モデル設定」画面で登録した文脈量 {モデル名(小文字): トークン}。

    公式の表（config.MODEL_CONTEXT_WINDOWS）に無いモデル — ゲートウェイ独自の名前や
    他社モデル — の実力を教えるための口。表より優先する。
    """
    raw = _read_admin().get("context_overrides") or {}
    out = {}
    for k, v in (raw.items() if isinstance(raw, dict) else []):
        try:
            n = int(v)
        except (TypeError, ValueError):
            continue
        if n > 0 and str(k).strip():
            out[str(k).strip().lower()] = n
    return out


def inline_limit_for(model: str | None = None) -> int:
    """カタログをそのまま入れる上限（文字）。選択中モデルの文脈量から自動で決める。

        上限 = 文脈(トークン) × 0.5 ÷ 0.55(日本語1文字あたりの概算トークン)

    文脈の半分をカタログに、残りをツール定義・履歴・回答に使う配分。
    文脈量は 管理者の登録 > 公式の表 > 既定値 の順で決まる（context_window）。
    天井（INLINE_LIMIT_MAX）と床（INLINE_LIMIT_MIN）で丸める。
    """
    if not model:
        return prompt_inline_limit()
    import llm                     # 循環importを避ける（llm側もmodelsを遅延importしている）
    context, _ = context_window(model)
    capacity = int(context * CATALOG_CONTEXT_RATIO / llm.TOKENS_PER_CHAR_TEXT)
    return max(INLINE_LIMIT_MIN, min(prompt_inline_limit(), capacity))


def catalog_total_chars() -> int:
    """全DBの詳細カタログの合計文字数（キャッシュ済みテキストを測るだけ）。"""
    import catalog as catalog_mod
    import db as db_mod
    return catalog_mod.inline_length(
        [{"path": str(f), "alias": db_mod.alias_for(f), "tables": None}
         for f in db_mod.list_db_files()])


def context_window(model: str) -> tuple:
    """そのモデルが一度に読める量（トークン）と、それが確かな値かどうか。

    戻り値: (トークン数, 分かっているモデルか)
    名前は環境によって違うので、前方一致の長い方から当てる。
    """
    low = str(model or "").lower()
    # 管理者の登録（完全一致 → 部分一致）> 公式の表（長い名前から部分一致）> 既定値（推定）
    ov = context_overrides()
    if low in ov:
        return ov[low], True
    for key in sorted(ov, key=len, reverse=True):
        if key and key in low:
            return ov[key], True
    for key in sorted(config.MODEL_CONTEXT_WINDOWS, key=len, reverse=True):
        if key in low:
            return config.MODEL_CONTEXT_WINDOWS[key], True
    return config.MODEL_CONTEXT_DEFAULT, False


def _from_api() -> list[str]:
    """APIに聞ける環境なら、使えるモデルの一覧を取ってくる。"""
    import llm
    if not llm.is_configured():
        return []
    now = time.time()
    with _models_lock:
        if _models_cache["models"] and now - _models_cache["at"] < _CACHE_SEC:
            return list(_models_cache["models"])
    try:
        got = sorted(m.id for m in llm.models_client().models.list().data)
    except Exception as e:
        print(f"[models] 一覧を取得できませんでした: {e}")
        got = []
    with _models_lock:
        _models_cache["at"], _models_cache["models"] = now, got
    return list(got)


def source() -> str:
    """候補がどこから来ているか。"admin" | "env" | "default"

    画面に出す文言を、実態とずれないようにするためのもの。
    """
    if [str(m).strip() for m in (_read_admin().get("models") or []) if str(m).strip()]:
        return "admin"
    return "env" if config.OPENAI_MODELS else "default"


def available(refresh: bool = False) -> list[str]:
    """チャット画面のプルダウンに出す候補。

    決まり方は 管理者の設定 > env の OPENAI_MODELS > 既定＋利用中のモデル。

    APIが返す一覧はここでは使わない。以前は最後の手段として使っていたが、
    それだと何も設定していないときに babbage-002 のような使えないモデルまで
    100件以上並び、「モデル設定」で絞ったつもりが効いていないように見えた。

    何も決めていないときは既定だけにしたいところだが、それだと以前の一覧から
    選んでいた人が黙って別のモデルに変わってしまう。決まるまでの間は、
    すでに誰かが選んでいるモデルも残す（画面で候補を決めれば、そちらが優先）。
    """
    if refresh:
        with _models_lock:
            _models_cache["at"] = 0.0
    admin = [str(m).strip() for m in (_read_admin().get("models") or []) if str(m).strip()]
    names = admin or list(config.OPENAI_MODELS) or sorted(users_by_model())
    # 既定のモデルは必ず候補に入れる（一覧に出てこないAPIもあるため）
    d = default_model()
    if d and d not in names:
        names.insert(0, d)
    return names


def users_by_model() -> dict:
    """いま誰がどのモデルを選んでいるか。{モデル名: [ユーザー名, ...]}

    候補から外すとその人は既定に戻る。外す前に影響が見えるようにするため、
    利用者のフォルダを読んで集める（読むだけで、何も書き換えない）。
    """
    out: dict = {}
    root = config.USER_META_DIR
    if not root.exists():
        return out
    for d in sorted(root.iterdir()):
        p = d / "prefs.yaml"
        if not p.is_file():
            continue
        try:
            data = yaml.safe_load(p.read_text(encoding="utf-8")) or {}
        except Exception:
            continue
        m = str((data or {}).get("model") or "").strip()
        if m:
            out.setdefault(m, []).append(d.name)
    return out


def model_catalog(refresh: bool = False) -> list[str]:
    """管理者が候補を選ぶときに見せる「選べる全部」。

    available() は絞り込んだ後の一覧なので、管理画面ではこちらを使う。
    """
    if refresh:
        with _models_lock:
            _models_cache["at"] = 0.0
    return sorted(set(list(config.OPENAI_MODELS) + _from_api()))


def current(user=None) -> str:
    """そのユーザーがいま使うモデル。

    選んでいても、管理者が候補から外していれば既定に戻す。
    外したモデルを使い続けられると、絞り込んだ意味がなくなるため。
    """
    if user:
        chosen = prefs.get_model(user)
        if chosen and chosen in available():
            return chosen
    return default_model()


def choose(user, model: str) -> str:
    """モデルを選ぶ。管理者が決めた候補の中からだけ。"""
    model = str(model or "").strip()
    if not model:
        raise ValueError("モデル名が空です。")
    allowed = available()
    if model not in allowed:
        raise ValueError(f"{model} は選べません。"
                         f"選べるのは {'、'.join(allowed) or '（候補なし）'} です。")
    prefs.set_model(user, model)
    who = getattr(user, "username", None) or user
    print(f"[models] {who} のモデルを {model} にしました")
    return model


def _scope_note(total: int, limit: int) -> str:
    """カタログがそのモデルに収まらないときの、画面向けの説明文。

    安心（使えるデータは変わらない）を先に、数字は根拠として後ろに置く。
    数字が先だと、初見の利用者は大きさだけを見て「何かが壊れている」と読む。
    """
    detail = (f"（カタログ全体 約{total:,}字 ＞ このモデルが一度に読める量 "
              f"約{limit:,}字のため）")
    if config.SCOPE_MODE == "all":
        return ("このモデルにはカタログ全体が収まらず、詳細（列名・コード値）が"
                "渡らない要約モードになります。文脈の大きいモデルを選ぶか、"
                "env の SCOPE_MODE を見直してください" + detail + "。")
    return ("使えるデータは変わりません。カタログが大きいので、質問ごとに関係する"
            "データだけへ自動で絞り、詳細を保って渡します" + detail + "。")


def status(user: str | None = None, refresh: bool = False) -> dict:
    cur = current(user)
    names = available(refresh)
    total = catalog_total_chars()
    limit = inline_limit_for(cur)
    fits = total <= limit
    return {
        "current": cur,
        "models": [{"id": m, "vision": is_vision(m),
                    "catalog_fits": total <= inline_limit_for(m)} for m in names],
        "vision": is_vision(cur),
        "from_env": bool(config.OPENAI_MODELS),
        "image_max_mb": config.IMAGE_MAX_MB,
        "image_max_count": config.IMAGE_MAX_COUNT,
        # 選択中モデルにカタログ全体が収まるか（チャット画面の警告表示に使う）
        "scope": {"mode": config.SCOPE_MODE, "catalog_chars": total,
                  "limit_chars": limit, "fits": fits,
                  "note": "" if fits else _scope_note(total, limit)},
    }


# --- 管理画面向け -----------------------------------------------------------

def admin_status(refresh: bool = False, scope: list[dict] | None = None) -> dict:
    """「モデル設定」画面に渡す内容。

    scope を渡すと、そのデータ範囲で「文脈をどれだけ使うか」も一緒に返す
    （上限を決めるのに、いまの実測値が要るため）。
    """
    ov = _read_admin()
    chosen = [str(m).strip() for m in (ov.get("models") or []) if str(m).strip()]
    names = chosen or list(config.OPENAI_MODELS)
    out = {
        "models": names,
        "default": default_model(),
        "vision": _vision_keys(),
        "catalog": model_catalog(refresh),
        "source": source(),
        "effective": available(),
        "in_use": users_by_model(),
        "from_env": not chosen,
        "env_models": list(config.OPENAI_MODELS),
        "env_default": config.OPENAI_MODEL,
        "settings_file": str(config.MODEL_SETTINGS_FILE),
        "llm_ready": _llm_ready(),
        # APIキーは値を返さない。設定済みかどうかと出所だけ（画面表示用）
        "api_key_set": bool(llm_api_key()),
        "api_key_source": llm_api_key_source(),
        # 接続先URL（フルパス）。URLは秘密ではないので値も返す
        "chat_url": llm_chat_url(),
        "models_url": llm_models_url(),
        "chat_url_source": ("screen" if str(_read_admin().get("chat_url") or "").strip()
                            else "env"),
        "models_url_source": ("screen" if str(_read_admin().get("models_url") or "").strip()
                              else "env"),
        "env_chat_url": _env_chat_url(),
        "env_models_url": _env_models_url(),
        "context_overrides": context_overrides(),
        "env_context_default": config.MODEL_CONTEXT_DEFAULT,
        "catalog_chars": catalog_total_chars(),
        # 候補ごとの文脈量とカタログ上限。出所も返す（登録 / 公式の表 / 推定）
        "contexts": [_context_row(m) for m in names],
    }
    if scope is not None:
        import llm
        base = llm.budget(scope, model=default_model(), admin=True)
        out["budget"] = base
        # 候補それぞれで、いまのカタログがどれだけ文脈を食うか
        out["per_model"] = []
        for m in (names or [default_model()]):
            ctx, known = context_window(m)
            out["per_model"].append({
                "id": m, "context": ctx, "context_known": known,
                "now_pct": round(base["now_tokens"] / ctx * 100, 1) if ctx else 0.0,
                "at_limit_pct": round(base["at_limit_tokens"] / ctx * 100, 1) if ctx else 0.0,
            })
    return out


def _context_row(model: str) -> dict:
    """モデル設定画面の1行ぶん。文脈量がどこから来た値かも添える。"""
    low = str(model or "").lower()
    ov = context_overrides()
    if low in ov or any(k in low for k in ov):
        source = "override"
    elif any(k in low for k in config.MODEL_CONTEXT_WINDOWS):
        source = "table"
    else:
        source = "default"
    ctx, _ = context_window(model)
    limit = inline_limit_for(model)
    return {"id": model, "context": ctx, "source": source, "limit_chars": limit,
            "fits": catalog_total_chars() <= limit}


def _llm_ready() -> bool:
    import llm
    return llm.is_configured()


def save_admin(data: dict, user: str | None = None) -> dict:
    """「モデル設定」画面からの保存。"""
    models = [str(m).strip() for m in (data.get("models") or []) if str(m).strip()]
    if not models:
        raise ValueError("選択できるモデルを1つ以上残してください。")
    if len(models) != len(set(models)):
        raise ValueError("同じモデルが重複しています。")
    for m in models:
        if len(m) > 120:
            raise ValueError(f"モデル名が長すぎます: {m[:40]}…")

    default = str(data.get("default") or "").strip() or models[0]
    if default not in models:
        raise ValueError(f"既定のモデル {default} が候補に入っていません。")

    vision = [str(v).strip().lower() for v in (data.get("vision") or []) if str(v).strip()]

    # 文脈量の登録 {モデル名: トークン}。表に無いモデルの実力を教える口
    overrides = {}
    for k, v in (data.get("context_overrides") or {}).items():
        name = str(k).strip().lower()
        if not name:
            continue
        try:
            n = int(str(v).replace(",", "").replace("_", ""))
        except (TypeError, ValueError):
            raise ValueError(f"「{k}」の文脈量は数字（トークン数）で指定してください。") from None
        if n < 1_000 or n > 10_000_000:
            raise ValueError(f"「{k}」の文脈量 {n:,} は範囲外です（1,000〜10,000,000）。")
        overrides[name] = n

    # APIキー。値が来たときだけ更新する。応答にもログにもキーの値は出さない
    key_in = data.get("api_key")
    key_new = str(key_in).strip() if isinstance(key_in, str) else ""
    key_clear = bool(data.get("api_key_clear"))
    if key_new:
        if any(c.isspace() for c in key_new):
            raise ValueError("APIキーに空白や改行が入っています。コピーし直してください。")
        if not (8 <= len(key_new) <= 500):
            raise ValueError("APIキーの長さが不自然です。値を確かめてください。")

    keep = _read_admin()  # 保存済みのAPIキーを巻き添えで消さないため、上書きで重ねる
    keep.update({"models": models, "default": default, "vision": vision,
                 "context_overrides": overrides})

    # 接続先URL（フルパス2本）。空欄で保存すると env の値に戻る。
    # env と同じ値なら上書きとして持たない（envを変えたとき追随できるように）。
    url_changed = False
    for field, suffix, envval, label in (
            ("chat_url", "/chat/completions", _env_chat_url(), "チャット"),
            ("models_url", "/models", _env_models_url(), "モデル一覧")):
        if field not in data:
            continue
        u = str(data.get(field) or "").strip().rstrip("/")
        if u:
            if any(c.isspace() for c in u):
                raise ValueError(f"{label}のURLに空白が入っています。")
            if not (u.startswith("http://") or u.startswith("https://")):
                raise ValueError(f"{label}のURLは http:// か https:// で始めてください。")
            if not u.endswith(suffix):
                raise ValueError(
                    f"{label}のURLは {suffix} で終わるフルパスで入力してください"
                    f"（例: https://api.openai.com/v1{suffix}）。")
        old = str(keep.get(field) or "").strip()
        new = "" if u == envval else u
        if new:
            keep[field] = new
        else:
            keep.pop(field, None)
        if old != new:
            url_changed = True

    if key_clear:
        keep.pop("api_key", None)
    elif key_new:
        keep["api_key"] = key_new
    _write_admin(keep)
    if key_clear or key_new or url_changed:
        reset_llm_client()  # 次のAI呼び出しから新しい接続先・キーを使う（再起動不要）
    print(f"[models] モデル設定を更新しました（{user or '不明'}）: "
          f"候補{len(models)}件 / 既定={default} / 画像判定={len(vision)}件 / "
          f"文脈量の登録={len(overrides)}件"
          + (" / APIキーを更新" if key_new else "")
          + (" / APIキーをenvに戻した" if key_clear else "")
          + (" / 接続先URLを変更" if url_changed else ""))
    return admin_status()


# ==========================================================================
# ===== 元 usage.py
# このアプリ自身が、誰にどう使われているかを数える。
#
# sqlusage.py が「どの結合が通ったか」だけを見るのに対し、こちらは利用そのものを見る。
# 見たいのは利用者数ではなく、次の3つ。
#
#   伸びているか   … 使われ続けているのか、最初の週だけだったのか
#   何に使われるか … よく呼ばれる機能と、まったく呼ばれない機能
#   どこで転ぶか   … 失敗した質問。これがカタログを直す入口になる
#
# 失敗の中身は分けて数える。「列が無い」はカタログ不足で人間が直せるが、
# 「LLM呼び出しに失敗」は設定や回線の問題で、カタログをいくら直しても減らない。
# 混ぜて「エラー率5%」と出すと、直せないものを直そうとして時間を溶かす。
#
# 材料は data/users/<ユーザー>/chats/*.json（会話の実体）と、
# data/import_history.jsonl（取り込みの記録）。どちらも読むだけで書き換えない。
#
# 戻り値の形は advanced.py / business.py と同じ {"title", "tables", "notes", "meta"}。
# 画面もLLMも同じ入れ物で受け取れる。
#
# 時系列は発言ごとの時刻（表示物の at）で数える。この仕組みを入れる前の
# 古い会話には at が無いので、会話の開始時刻で代用し、その旨を所見に明示する。
# ==========================================================================
import json
import re
from collections import Counter
from datetime import datetime, timedelta
from pathlib import Path

import config
import history

#: 質問文・エラー文をそのまま並べるときの上限。多すぎると読まれない。
MAX_LIST = 40

METHODS = {
    "summary": "全体像（期間・利用者・会話数・失敗率）",
    "users": "利用者ごとの利用量",
    "trend": "日ごと・曜日・時間帯の推移",
    "tools": "呼ばれた機能の回数",
    "databases": "実際に使われたDB",
    "errors": "失敗の内訳と、直し方の当たり",
    "questions": "実際に聞かれた質問",
}

#: 失敗の分類。上から順に当てる（先に当たったものを採る）。
#: 「誰が直せるか」で分ける。カタログ担当・管理者・利用者では打ち手が違う。
_ERROR_KINDS = (
    ("カタログ不足（列・テーブルの取り違え）",
     r"no such (table|column)|列名が違います|テーブルが見つかりません",
     "describe_table で確認できる情報が足りていない。"
     "カタログの列説明・コード値・結合定義を書き足すと減る。"),
    ("SQLの誤り（構文・集計）",
     r"SQL実行エラー|syntax error|ambiguous|misuse of aggregate",
     "例文（Q&SQL）を足すと、AIが型を真似るので減る。"),
    ("分析に足りるデータが無い",
     r"行しかなく|0行でした|データが0行|足りません",
     "抽出条件が狭すぎる。期間を広げるか、分析の指定（説明変数など）を減らす。"),
    ("ツールの引数不足",
     r"(には|は).{0,30}(必要|指定してください)|引数|列が結果にありません|指定列",
     "ツールの説明文を具体的にすると、AIの指定ミスが減る。"),
    ("LLM・API側の問題",
     r"LLM呼び出しに失敗|Error code:|timeout|接続",
     "カタログでは直らない。モデル設定・APIキー・回線を確認する。"),
    ("実行時間切れ",
     r"時間がかかりすぎ|タイムアウト|interrupted",
     "対象データを絞るか、集計済みのユーザー定義ツールを用意する。"),
)

_WEEKDAYS = ("月", "火", "水", "木", "金", "土", "日")


def _usage_dt(value) -> datetime | None:
    try:
        return datetime.fromisoformat(str(value))
    except (TypeError, ValueError):
        return None


def _usage_table(name: str, columns: list, rows: list) -> dict:
    return {"name": name, "columns": columns, "rows": [tuple(r) for r in rows]}


def _usage_out(title: str, tables: list, notes: list, meta: dict | None = None) -> dict:
    return {"title": title, "tables": tables, "notes": notes, "meta": meta or {}}


def classify_error(message: str) -> tuple[str, str]:
    """エラー文を「誰が直せるか」で分類する。戻り値: (分類, 打ち手)"""
    for name, pattern, fix in _ERROR_KINDS:
        if re.search(pattern, message, re.IGNORECASE):
            return name, fix
    return "その他", "内容を読んで個別に判断する。"


# =============================================================================
# 材料集め
# =============================================================================

def _asked(log: list, fallback: datetime | None) -> list[dict]:
    """1発言ぶんの記録。質問の時刻と、答え終わった時刻を組にする。

    表示物には at が入っている（質問は積んだ時刻、応答は保存した時刻）。
    この2つの差が、その質問で待たされた時間になる。
    at を持たない古い会話は、会話の開始時刻で代用する。
    """
    out: list[dict] = []
    for item in log:
        at = _usage_dt(item.get("at")) or fallback
        if item.get("role") == "user" and item.get("kind") == "text":
            out.append({"at": at, "text": str(item.get("content") or "").strip(),
                        "done": None, "failed": False, "errors": [],
                        "exact": bool(item.get("at"))})
        elif out:
            if at and out[-1]["at"] and at >= out[-1]["at"]:
                out[-1]["done"] = at
            if item.get("kind") == "error":
                out[-1]["failed"] = True
                out[-1]["errors"].append(str(item.get("message") or ""))
    return out


def collect(days: int | None = None, user: str | None = None) -> list[dict]:
    """会話ファイルを1会話1レコードに畳む。

    days を指定すると、その日数より前に始まった会話は捨てる。
    捨てるのは開始時刻で判定するので、古い会話を今日まで続けていた場合も対象外。
    """
    root = Path(config.USER_META_DIR)
    if not root.exists():
        return []
    limit = datetime.now() - timedelta(days=days) if days else None

    out = []
    for f in root.glob("*/chats/*.json"):
        if f.name == "index.json":
            continue
        who = f.parent.parent.name
        if user and who.lower() != user.lower():
            continue
        try:
            data = json.loads(f.read_text(encoding="utf-8"))
        except Exception:
            continue                      # 壊れた1本のために全体を止めない
        created = _usage_dt(data.get("created_at"))
        if limit and created and created < limit:
            continue

        log = data.get("render_log") or []
        asked = _asked(log, created)
        # 最初の質問より前に出た失敗は、どの質問のものとも言えない
        first_q = next((i for i, x in enumerate(log)
                        if x.get("role") == "user" and x.get("kind") == "text"), len(log))
        orphans = [str(x.get("message") or "") for x in log[:first_q]
                   if x.get("kind") == "error"]
        kinds = Counter(str(i.get("kind") or "") for i in log)
        tools = [(tc.get("function") or {}).get("name")
                 for m in (data.get("messages") or [])
                 for tc in (m.get("tool_calls") or [])]
        out.append({
            "user": who,
            "id": data.get("id") or f.stem,
            "title": str(data.get("title") or "").strip(),
            "created": created,
            "updated": _usage_dt(data.get("updated_at")),
            "asked": asked,
            "questions": [a["text"] for a in asked if a["text"]],
            "tools": [t for t in tools if t],
            "errors": [str(i.get("message") or "") for i in log
                       if i.get("kind") == "error"],
            "orphan_errors": orphans,
            "dbs": list(data.get("db_names") or []),
            "sqls": kinds.get("sql", 0),
            "charts": kinds.get("chart", 0),
            "tables": kinds.get("table", 0),
            "files": kinds.get("file", 0),
            "reports": kinds.get("report", 0),
        })
    out.sort(key=lambda r: r["created"] or datetime.min)
    return out


def _period_note(records: list[dict]) -> str:
    days = [r["created"] for r in records if r["created"]]
    if not days:
        return "期間: 不明"
    return f"期間: {min(days):%Y-%m-%d} 〜 {max(days):%Y-%m-%d}"


def _empty(days: int | None) -> dict:
    span = f"直近{days}日には" if days else ""
    return _usage_out("利用状況", [], [f"{span}会話の記録がありませんでした。"
                                "まだ誰も使っていないか、data/users/ が空です。"])


# =============================================================================
# 分析
# =============================================================================

def summary(records: list[dict], days: int | None = None) -> dict:
    """全体像。まずこれを見て、気になった軸を他のメソッドで掘る。"""
    if not records:
        return _empty(days)

    n_chats = len(records)
    n_turns = sum(len(r["questions"]) for r in records)
    users = {r["user"] for r in records}
    errs = [e for r in records for e in r["errors"]]
    chats_with_err = sum(1 for r in records if r["errors"])
    active_days = {r["created"].date() for r in records if r["created"]}

    rows = [
        ("会話", f"{n_chats:,} 件"),
        ("質問", f"{n_turns:,} 回"),
        ("利用者", f"{len(users)} 人"),
        ("使われた日", f"{len(active_days)} 日"),
        ("1会話あたりの質問", f"{n_turns / n_chats:.1f} 回"),
        ("失敗を含む会話", f"{chats_with_err} 件（{chats_with_err / n_chats * 100:.0f}%）"),
        ("作った表・グラフ", f"表 {sum(r['tables'] for r in records):,} / "
                            f"グラフ {sum(r['charts'] for r in records):,}"),
        ("出したファイル", f"{sum(r['files'] for r in records):,} 件"),
    ]

    notes = [_period_note(records)]
    # 1回で終わった会話が多いなら、続けて聞ける場になっていない可能性がある
    one_shot = sum(1 for r in records if len(r["questions"]) <= 1)
    if n_chats >= 5:
        notes.append(
            f"1問だけで終わった会話が {one_shot}/{n_chats} 件"
            f"（{one_shot / n_chats * 100:.0f}%）。"
            + ("会話を続けて掘り下げる使い方が根づいています。"
               if one_shot / n_chats < 0.5 else
               "多くが単発です。最初の答えで満足したか、続きを諦めたかのどちらかなので、"
               "失敗の内訳（errors）も合わせて見てください。"))
    if errs:
        kinds = Counter(classify_error(e)[0] for e in errs)
        top, n = kinds.most_common(1)[0]
        notes.append(f"失敗 {len(errs)} 件のうち最も多いのは「{top}」{n} 件。"
                     "内訳は errors で確認できます。")
    else:
        notes.append("記録された失敗はありません。")

    heavy = Counter(r["user"] for r in records).most_common(1)
    if heavy and len(users) > 1:
        who, cnt = heavy[0]
        notes.append(f"最も使っているのは {who}（{cnt} 件 / 全体の "
                     f"{cnt / n_chats * 100:.0f}%）。")

    return _usage_out("利用状況の全体像", [_usage_table("全体", ["項目", "値"], rows)], notes,
                {"chats": n_chats, "turns": n_turns, "users": len(users),
                 "errors": len(errs), "active_days": len(active_days)})


def by_user(records: list[dict]) -> dict:
    """利用者ごと。誰が使っていて、誰が離れたかを見る。"""
    if not records:
        return _empty(None)

    per: dict[str, list] = {}
    for r in records:
        per.setdefault(r["user"], []).append(r)

    rows = []
    for who, rs in sorted(per.items(), key=lambda kv: -len(kv[1])):
        last = max((r["updated"] or r["created"] for r in rs
                    if (r["updated"] or r["created"])), default=None)
        turns = sum(len(r["questions"]) for r in rs)
        errs = sum(len(r["errors"]) for r in rs)
        rows.append((who, len(rs), turns, round(turns / len(rs), 1), errs,
                     f"{last:%Y-%m-%d}" if last else "—"))
    notes = [_period_note(records),
             "「最終利用」が古い人は、使えなかったのか、必要が無かったのかを直接聞くのが早いです。"]

    today = datetime.now()
    stale = [r[0] for r in rows
             if r[5] != "—" and (today - datetime.fromisoformat(r[5])).days >= 14]
    if stale:
        notes.append(f"2週間以上使っていないのは {len(stale)} 人（{'、'.join(stale[:5])}"
                     f"{' ほか' if len(stale) > 5 else ''}）。")
    return _usage_out("利用者ごとの利用量",
                [_usage_table("利用者別", ["利用者", "会話", "質問", "1会話あたり", "失敗", "最終利用"],
                        rows)],
                notes, {"users": len(rows)})


def trend(records: list[dict]) -> dict:
    """日ごと・曜日・時間帯。定着したのか、一度きりだったのかを見る。

    数えるのは会話ではなく質問1件ずつ。1本の会話で何度も聞いていれば、
    その回数だけ数える（実際にどれだけ使われたかは、そちらの方が近い）。
    """
    if not records:
        return _empty(None)

    asked = [a for r in records for a in r["asked"] if a["at"]]
    if not asked:
        return _empty(None)

    daily = Counter(a["at"].date() for a in asked)
    dow = Counter(a["at"].weekday() for a in asked)
    hour = Counter(a["at"].hour for a in asked)

    day_rows = [(str(d), n) for d, n in sorted(daily.items())]
    dow_rows = [(_WEEKDAYS[i], dow.get(i, 0)) for i in range(7)]
    hour_rows = [(f"{h:02d}時", hour.get(h, 0)) for h in range(24) if hour.get(h)]

    notes = [_period_note(records), f"質問 {len(asked)} 件を、聞かれた時刻で数えています。"]
    rough = sum(1 for a in asked if not a["exact"])
    if rough:
        notes.append(f"うち {rough} 件は時刻を持たない古い会話で、"
                     "会話の開始時刻で代用しています。")

    # 応答にかかった時間。待たされているなら、対象データの絞り込みや
    # ユーザー定義ツールの用意で短くできる。
    waits = [(a["done"] - a["at"]).total_seconds() for a in asked
             if a["done"] and a["at"] and (a["done"] - a["at"]).total_seconds() >= 0]
    if waits:
        waits.sort()
        mid = waits[len(waits) // 2]
        slow = sum(1 for w in waits if w >= 30)
        notes.append(f"回答までの時間は中央値 {mid:.0f} 秒、最長 {waits[-1]:.0f} 秒。"
                     + (f"30秒以上待った質問が {slow} 件あります。" if slow else ""))
    if len(daily) >= 2:
        days_sorted = sorted(daily)
        span = (days_sorted[-1] - days_sorted[0]).days + 1
        notes.append(f"{span} 日のうち {len(daily)} 日に利用がありました"
                     f"（{len(daily) / span * 100:.0f}%）。")
        half = len(days_sorted) // 2
        first = sum(daily[d] for d in days_sorted[:half])
        last = sum(daily[d] for d in days_sorted[half:])
        if first:
            notes.append(f"前半 {first} 件 → 後半 {last} 件（{(last - first) / first * 100:+.0f}%）。"
                         + ("使われ方が伸びています。" if last > first else
                            "落ちています。失敗の内訳（errors）と合わせて見てください。"))
    if hour_rows:
        peak = max(hour_rows, key=lambda t: t[1])
        notes.append(f"最も使われる時間帯は {peak[0]}（{peak[1]} 件）。")

    return _usage_out("利用の推移",
                [_usage_table("日ごと", ["日付", "質問"], day_rows),
                 _usage_table("曜日", ["曜日", "質問"], dow_rows),
                 _usage_table("時間帯", ["時間", "質問"], hour_rows)],
                notes, {"active_days": len(daily), "questions": len(asked),
                        "median_wait_sec": round(mid) if waits else None})


def by_tool(records: list[dict]) -> dict:
    """呼ばれた機能。使われていない機能は、説明文が悪いか、要らないかのどちらか。"""
    if not records:
        return _empty(None)

    calls = Counter(t for r in records for t in r["tools"])
    total = sum(calls.values())
    rows = [(name, n, f"{n / total * 100:.1f}%") for name, n in calls.most_common()]

    notes = [_period_note(records)]
    if total:
        notes.append(f"ツール呼び出しは合計 {total:,} 回。"
                     f"種類は {len(calls)} 種です。")
        top = calls.most_common(3)
        notes.append("よく使われるのは " +
                     "、".join(f"{n}（{c}回）" for n, c in top) + "。")
    try:
        import tools as _tools
        unused = sorted(set(_tools._HANDLERS) - set(calls))
        if unused:
            notes.append(f"一度も呼ばれていない組み込みツールが {len(unused)} 種あります"
                         f"（{'、'.join(unused[:8])}"
                         f"{' ほか' if len(unused) > 8 else ''}）。"
                         "要らないなら「ツール」タブで無効にすると、AIの選択肢が減って"
                         "呼び分けが安定します。使ってほしいなら説明文を具体的に書き直します。")
    except Exception:
        pass
    return _usage_out("呼ばれた機能", [_usage_table("ツール別", ["ツール", "回数", "割合"], rows)],
                notes, {"total_calls": total, "kinds": len(calls)})


def by_database(records: list[dict]) -> dict:
    """どのテーブルが使われたか。

    かつては「どのDBを使ったか」を数えていたが、このアプリのDBは常に1つで、
    「1会話で使ったDBの数」は 0 か 1 にしかならない（＝データを使ったかどうか
    しか言っていない）。読む人の役に立つのは、実際にSQLが触った
    テーブルとまとまりなので、そちらを数える。
    """
    if not records:
        return _empty(None)

    # 実行されたSQLからテーブル名を拾う（ER図の利用状況と同じ数え方）。
    sqls, _ = sqlusage.collect_sqls()
    known = set()
    for f in db.list_db_files():
        known |= set(catalog.profile_db(f)["tables"])
    per_table: Counter = Counter()
    for sql in sqls:
        low = str(sql).lower()
        for t in known:
            if re.search(r'(?<![\w."])' + re.escape(t.lower()) + r'(?![\w])', low):
                per_table[t] += 1

    total = len(records)
    # 「データを使った」＝SQLが実行された会話。r["tables"] は画面に出した表の数
    # （文書の検索結果や、AIが文字列から作った表も含む）なので使わない
    used = sum(1 for r in records if r.get("sqls"))
    rows = [(t, n) for t, n in per_table.most_common(MAX_LIST)]

    # まとまり（表名の接頭辞）でまとめた眺め。どの領域が使われているかが分かる
    per_group: Counter = Counter()
    for t, n in per_table.items():
        per_group[t.split("__", 1)[0] if "__" in t else "（まとまり無し）"] += n
    group_rows = [(g, n) for g, n in per_group.most_common()]

    tables = []
    if group_rows:
        tables.append(_usage_table("まとまり別", ["まとまり", "SQLで使われた回数"], group_rows))
    if rows:
        tables.append(_usage_table("テーブル別", ["テーブル", "SQLで使われた回数"], rows))
    tables.append(_usage_table(
        "データを使ったか", ["区分", "会話"],
        [("データを使った会話", used), ("データを使わなかった会話", total - used)]))

    hit = sum(1 for sql in sqls
              if any(re.search(r'(?<![\w."])' + re.escape(t.lower()) + r'(?![\w])',
                               str(sql).lower()) for t in known))
    notes = [_period_note(records),
             "実行されたSQLに出てきたテーブルを数えています"
             "（1つのSQLに複数のテーブルが出れば、それぞれ1回ずつ）。"]
    if sqls and hit < len(sqls):
        notes.append(
            f"記録に残るSQL {len(sqls)} 本のうち、いまあるテーブルを使っているのは "
            f"{hit} 本です。残りは、既に削除されたテーブルを使っていたか、"
            "テーブルを読まないSQL（AIが文字列だけで表を組み立てたもの）です。")
    if total:
        notes.append(f"データを使った会話は {used}/{total} 件"
                     f"（{used / total * 100:.0f}%）。"
                     + ("残りは文書の検索や、ツールを使わない質問です。"
                        if used < total else ""))
    if not per_table:
        notes.append("SQLが実行された記録がまだありません。")
    else:
        cold = sorted(known - set(per_table))
        if cold:
            notes.append(f"残っている会話の中では一度も使われていないテーブルが "
                         f"{len(cold)}/{len(known)} 個あります"
                         f"（例: {'、'.join(cold[:5])}）。"
                         "取り込んだばかりの表もここに入るので、"
                         "しばらく使われないままなら説明や例文を足す目安にしてください。")
    return _usage_out("使われたデータ", tables, notes,
                      {"tables": len(per_table)})


def errors(records: list[dict]) -> dict:
    """失敗の内訳。カタログを直して減るものと、そうでないものを分ける。"""
    if not records:
        return _empty(None)

    # 失敗は「どの質問で起きたか」まで対応づける。時刻と質問文が揃っていないと、
    # カタログの何を直せばよいかを後から辿れない。
    items = [(r["user"], a["at"], e, a["text"])
             for r in records for a in r["asked"] for e in a["errors"]]
    # 質問より前に出た失敗（会話を開いた直後など）。数を合わせるために拾っておく。
    items += [(r["user"], r["created"], e, "")
              for r in records for e in r["orphan_errors"]]
    items.sort(key=lambda t: t[1] or datetime.min)
    if not items:
        return _usage_out("失敗の内訳", [],
                    [_period_note(records),
                     "記録された失敗はありません。"], {"errors": 0})

    kinds: Counter = Counter()
    fixes: dict[str, str] = {}
    for _, _, msg, _ in items:
        kind, fix = classify_error(msg)
        kinds[kind] += 1
        fixes.setdefault(kind, fix)
    kind_rows = [(name, n, f"{n / len(items) * 100:.0f}%", fixes.get(name, ""))
                 for name, n in kinds.most_common()]

    detail_rows = [(f"{dt:%m-%d %H:%M}" if dt else "—", who, (q or "")[:40],
                    msg.splitlines()[0][:80])
                   for who, dt, msg, q in items[-MAX_LIST:]]

    notes = [_period_note(records),
             f"失敗 {len(items)} 件。分類は「誰が直せるか」で分けています。"]
    # 「カタログ画面で直せるもの」だけを足す。打ち手はいちばん多い分類のものを出す
    # （合計だけ言われても、何から手を付ければよいか分からないため）。
    fixable = {k: n for k, n in kinds.items()
               if k.startswith(("カタログ", "SQL", "ツール"))}
    # 下の集計で必ず使うので、1件も無いときのために先に置く。
    # 置かないと「AI側の問題だけ」「文書検索の失敗だけ」の期間で画面ごと落ちる
    catalog_side = 0
    if fixable:
        catalog_side = sum(fixable.values())
        top = max(fixable.items(), key=lambda kv: kv[1])
        notes.append(f"うち {catalog_side} 件（{catalog_side / len(items) * 100:.0f}%）は"
                     "カタログ画面での手当てで減らせます。"
                     f"いちばん多いのは「{top[0]}」{top[1]} 件で、{fixes.get(top[0], '')}")
    outside = kinds.get("LLM・API側の問題", 0)
    if outside:
        notes.append(f"{outside} 件はモデル・API側の問題で、カタログを直しても減りません。")
    return _usage_out("失敗の内訳",
                [_usage_table("分類", ["分類", "件数", "割合", "打ち手"], kind_rows),
                 _usage_table(f"直近の失敗（最大{MAX_LIST}件）",
                        ["日付", "利用者", "質問", "内容"], detail_rows)],
                notes, {"errors": len(items), "catalog_fixable": catalog_side})


def questions(records: list[dict]) -> dict:
    """実際に聞かれた質問。例文とカタログに反映するための材料。"""
    if not records:
        return _empty(None)

    asked = [(a["at"], r["user"], a["text"], a["failed"])
             for r in records for a in r["asked"] if a["text"]]
    if not asked:
        return _usage_out("聞かれた質問", [], [_period_note(records), "質問の記録がありません。"])
    asked.sort(key=lambda t: t[0] or datetime.min)

    rows = [(f"{dt:%m-%d %H:%M}" if dt else "—", who, q[:60], "×" if bad else "")
            for dt, who, q, bad in asked[-MAX_LIST:]]

    # かつてここで「よく出る語」を数えていたが、区切りが記号だけだったので
    # 文がまるごと1語になり（例:「8月の装置別MTTRを教えて」）、傾向が読めなかった。
    # 形態素解析を入れてまで出す価値のある表ではないと判断して廃止した。
    notes = [_period_note(records),
             f"質問 {len(asked)} 件を記録しています。",
             "うまく答えられた質問は、チャットの⭐から例文としてカタログに登録できます。"
             "例文が増えるほど、同じ聞き方への精度が上がります。"]
    failed = sum(1 for *_, bad in asked if bad)
    if failed:
        notes.append(f"失敗を含む会話の質問が {failed} 件（× 印）。"
                     "この質問文がそのまま、カタログに足りない語彙の一覧になります。")
    return _usage_out("聞かれた質問",
                [_usage_table(f"直近の質問（最大{MAX_LIST}件）",
                        ["日付", "利用者", "質問", "失敗"], rows)],
                notes, {"questions": len(asked)})


def imports(days: int | None = None) -> dict:
    """取り込みの実績。チャットとは別系統なので、まとめてここから見えるようにする。"""
    recs = history.recent_import_records(limit=2000)
    if not recs:
        return _usage_out("取り込みの実績", [], ["取り込みの記録がありません。"], {"runs": 0})

    limit = datetime.now() - timedelta(days=days) if days else None
    picked = []
    for r in recs:
        at = _usage_dt(r.get("at") or r.get("started"))
        if limit and at and at < limit:
            continue
        picked.append((at, r))
    if not picked:
        return _usage_out("取り込みの実績", [], [f"直近{days}日に取り込みの記録がありません。"],
                    {"runs": 0})

    ok = sum(1 for _, r in picked if r.get("ok"))
    kinds = Counter(history.IMPORT_RECORD_KINDS.get(str(r.get("kind")), str(r.get("kind")))
                    for _, r in picked)
    rows = [(k, n) for k, n in kinds.most_common()]
    tables = Counter(f"{r.get('db_file')} / {r.get('table')}" for _, r in picked)
    tbl_rows = [(name, n) for name, n in tables.most_common(20)]

    notes = [f"取り込み {len(picked)} 回。成功 {ok} 件 / 失敗 {len(picked) - ok} 件"
             f"（成功率 {ok / len(picked) * 100:.0f}%）。"]
    fails = [r for _, r in picked if not r.get("ok")]
    if fails:
        notes.append("直近の失敗: " + str(fails[-1].get("message", ""))[:100])
    return _usage_out("取り込みの実績",
                [_usage_table("実行のしかた別", ["種別", "回数"], rows),
                 _usage_table("テーブル別", ["DB / テーブル", "回数"], tbl_rows)],
                notes, {"runs": len(picked), "ok": ok})


#: メソッド名 -> 実処理。records を取らない imports だけ形が違う。
_METHOD_FUNCS = {
    "summary": summary,
    "users": by_user,
    "trend": trend,
    "tools": by_tool,
    "databases": by_database,
    "errors": errors,
    "questions": questions,
}


def analyze(method: str = "summary", days: int | None = None,
            user: str | None = None) -> dict:
    """入口。method はこのモジュールの METHODS のいずれか。"""
    method = (method or "summary").strip().lower()
    if method == "imports":
        return imports(days)
    fn = _METHOD_FUNCS.get(method)
    if fn is None:
        raise ValueError(f"method は {'、'.join([*METHODS, 'imports'])} "
                         f"のいずれかです（受け取った値: {method}）")
    records = collect(days=days, user=user)
    res = fn(records, days) if fn is summary else fn(records)
    if user:
        res["notes"] = [f"対象: {user} のみ", *res.get("notes", [])]
    return res


# ==========================================================================
# ===== 元 exports.py
# ダウンロードさせるファイル（CSV / テキスト / ZIP）の組み立てと、自動保存用HTML。
#
# xlsx の組み立ては excel.py。いずれもディスクには書かず、メモリ上のバイト列を返す。
# ==========================================================================
import csv
import datetime as _dt
import io
import re
import zipfile

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
CSV_MIME = "text/csv"
TEXT_MIME = "text/plain"
MD_MIME = "text/markdown"
ZIP_MIME = "application/zip"

# CSVの文字コード。Excelでそのまま開ける utf-8-sig を既定にする。
ENCODINGS = {
    "utf-8-sig": "UTF-8（BOM付き／Excelで文字化けしない・推奨）",
    "utf-8": "UTF-8（BOMなし）",
    "cp932": "Shift_JIS（cp932／古いWindows向け）",
}
DEFAULT_ENCODING = "utf-8-sig"
EXPORT_DELIMITERS = {"comma": ",", "tab": "\t", "semicolon": ";"}


def safe_filename(name: str | None, ext: str, default: str = "export") -> str:
    """ダウンロード用のファイル名を整える（末尾に日時、指定の拡張子）。"""
    s = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", str(name or "")).strip().strip(".")
    s = re.sub(r"\.(xlsx|csv|txt|md|zip)$", "", s, flags=re.IGNORECASE) or default
    stamp = _dt.datetime.now().strftime("%Y%m%d_%H%M")
    return f"{s[:80]}_{stamp}.{ext.lstrip('.')}"


def _encode(text: str, encoding: str) -> bytes:
    enc = encoding if encoding in ENCODINGS else DEFAULT_ENCODING
    # Shift_JIS に無い文字（絵文字や一部の漢字）で落ちないよう置換する
    return text.encode(enc, errors="replace" if enc == "cp932" else "strict")


def build_csv(columns: list, rows: list, encoding: str = DEFAULT_ENCODING,
              delimiter: str = "comma") -> bytes:
    """1つの結果セットを CSV のバイト列にする。"""
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=EXPORT_DELIMITERS.get(delimiter, ","),
                   lineterminator="\r\n", quoting=csv.QUOTE_MINIMAL)
    w.writerow([str(c) for c in columns])
    for r in rows:
        w.writerow(["" if v is None else v for v in r])
    return _encode(buf.getvalue(), encoding)


def build_zip(files: list[dict]) -> bytes:
    """[{"filename": str, "data": bytes}, ...] を1つのZIPにまとめる。"""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        used = set()
        for f in files:
            name = str(f.get("filename") or "file")
            base, i = name, 2
            while name.lower() in used:
                stem, _, ext = base.rpartition(".")
                name = f"{stem}_{i}.{ext}" if stem else f"{base}_{i}"
                i += 1
            used.add(name.lower())
            z.writestr(name, f["data"])
    return buf.getvalue()


def table_to_text(columns: list, rows: list, style: str = "markdown") -> str:
    """結果セットを本文に埋め込めるテキスト表にする。"""
    cols = [str(c) for c in columns]
    body = [["" if v is None else str(v) for v in r] for r in rows]
    if style == "markdown":
        out = ["| " + " | ".join(cols) + " |",
               "| " + " | ".join("---" for _ in cols) + " |"]
        out += ["| " + " | ".join(r) + " |" for r in body]
        return "\n".join(out)
    if style == "tsv":
        return "\n".join(["\t".join(cols)] + ["\t".join(r) for r in body])
    # 等幅（プレーンテキスト用に桁を揃える）
    widths = [max(len(cols[i]), *(len(r[i]) for r in body)) if body else len(cols[i])
              for i in range(len(cols))]
    line = "-+-".join("-" * w for w in widths)
    out = [" | ".join(c.ljust(widths[i]) for i, c in enumerate(cols)), line]
    out += [" | ".join(v.ljust(widths[i]) for i, v in enumerate(r)) for r in body]
    return "\n".join(out)


def build_text(body: str, encoding: str = DEFAULT_ENCODING) -> bytes:
    return _encode(str(body or ""), encoding)



# ==========================================================================
# ===== 元 excel.py
# SELECT結果から Excel ブック(.xlsx)を組み立てる。
#
# ファイルはディスクに書かず、メモリ上のバイト列として返す。
#
# グラフはExcelネイティブのグラフとして入れる（画像ではない）。
# 受け取った側が範囲や種類を変えられるうえ、画像化ライブラリ（Chrome等）が
# 要らないので、サーバの環境に左右されない。
# ==========================================================================
import datetime as _dt  # noqa: F401  （シート値の型判定で使用）
import io
import re


from openpyxl import Workbook
from openpyxl.chart import (AreaChart, BarChart, LineChart, PieChart, Reference,
                            ScatterChart, Series)
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# Excelのシート名に使えない文字と長さ制限
_BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")
_SHEET_NAME_MAX = 31
_MAX_WIDTH = 60          # 列幅の上限（文字数）
_HEADER_FILL = PatternFill("solid", fgColor="1F3B5C")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_BAND_FILL = PatternFill("solid", fgColor="F5F8FC")
_THIN = Side(style="thin", color="D5DBE2")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)

# グラフの種類 -> (openpyxlのクラス, 積み上げ方)
EXCEL_CHART_TYPES = {
    "bar": (BarChart, "col", None),
    "bar_stacked": (BarChart, "col", "stacked"),
    "bar_percent": (BarChart, "col", "percentStacked"),
    "hbar": (BarChart, "bar", None),
    "hbar_stacked": (BarChart, "bar", "stacked"),
    "line": (LineChart, None, None),
    "line_stacked": (LineChart, None, "stacked"),
    "area": (AreaChart, None, None),
    "area_stacked": (AreaChart, None, "stacked"),
    "pie": (PieChart, None, None),
    "scatter": (ScatterChart, None, None),
}
_SERIES_COLORS = ["1F4E79", "F4B183", "70AD47", "C55A11", "7F7F7F",
                  "2E75B6", "A9D18E", "FFD966", "9DC3E6", "BFBFBF"]


def safe_sheet_name(name: str, used: set) -> str:
    """Excelの制約に合わせてシート名を整え、重複を避ける。"""
    s = _BAD_SHEET_CHARS.sub("_", str(name or "Sheet")).strip() or "Sheet"
    s = s[:_SHEET_NAME_MAX]
    base, i = s, 2
    while s.lower() in used:
        suffix = f"_{i}"
        s = base[: _SHEET_NAME_MAX - len(suffix)] + suffix
        i += 1
    used.add(s.lower())
    return s


def _cell_value(v):
    """openpyxl が扱えない型は文字列に落とす。"""
    if v is None or isinstance(v, (int, float, bool, str, _dt.datetime, _dt.date, _dt.time)):
        return v
    if isinstance(v, bytes):
        return f"<BLOB {len(v)} bytes>"
    return str(v)


def _autosize(ws, columns: list, rows: list):
    """見出しと先頭200行から列幅を決める。"""
    for ci, col in enumerate(columns, start=1):
        width = len(str(col))
        for r in rows[:200]:
            v = r[ci - 1] if ci - 1 < len(r) else None
            if v is not None:
                width = max(width, len(str(v)))
        ws.column_dimensions[get_column_letter(ci)].width = min(width + 2, _MAX_WIDTH)


def _add_chart(ws, spec: dict, columns: list, rows: list, header_row: int):
    """シートのデータ範囲からExcelネイティブのグラフを作って貼る。

    spec: {"type": 種類, "category_column": 横軸の列名, "value_columns": [系列の列名],
           "title": 見出し, "y_title": .., "x_title": .., "anchor": "H2",
           "data_labels": bool, "width": cm, "height": cm}
    """
    kind = str(spec.get("type") or "bar").lower()
    if kind not in EXCEL_CHART_TYPES:
        raise ValueError(f"未対応のグラフ種類です: {kind}。"
                         f"使えるのは {', '.join(EXCEL_CHART_TYPES)} です。")
    if not rows:
        raise ValueError("グラフにできる行がありません。")

    cat = spec.get("category_column") or (columns[0] if columns else None)
    if cat not in columns:
        raise ValueError(f"横軸の列 '{cat}' がありません。ある列: {', '.join(map(str, columns))}")
    vals = spec.get("value_columns") or [c for c in columns if c != cat]
    missing = [v for v in vals if v not in columns]
    if missing:
        raise ValueError(f"系列の列 {', '.join(map(str, missing))} がありません。"
                         f"ある列: {', '.join(map(str, columns))}")
    if not vals:
        raise ValueError("系列にする数値列がありません。")

    cls, direction, grouping = EXCEL_CHART_TYPES[kind]
    chart = cls()
    chart.title = spec.get("title") or None
    chart.style = 2
    if direction:
        chart.type = direction
    if grouping:
        chart.grouping = grouping
        chart.overlap = 100
    last = header_row + len(rows)
    cat_ref = Reference(ws, min_col=columns.index(cat) + 1, min_row=header_row + 1,
                        max_row=last)

    if kind == "scatter":
        # 散布図は x も数値列。1列目を x、残りを y にする。
        x_ref = Reference(ws, min_col=columns.index(vals[0]) + 1,
                          min_row=header_row + 1, max_row=last)
        for v in vals[1:] or vals[:1]:
            y_ref = Reference(ws, min_col=columns.index(v) + 1, min_row=header_row,
                              max_row=last)
            s = Series(y_ref, x_ref, title_from_data=True)
            s.marker.symbol = "circle"
            s.graphicalProperties.line.noFill = True
            chart.series.append(s)
    else:
        for i, v in enumerate(vals):
            ref = Reference(ws, min_col=columns.index(v) + 1, min_row=header_row,
                            max_row=last)
            chart.add_data(ref, titles_from_data=True)
        chart.set_categories(cat_ref)
        for i, s in enumerate(chart.series):
            color = _SERIES_COLORS[i % len(_SERIES_COLORS)]
            try:
                if kind.startswith("line"):
                    s.graphicalProperties.line.solidFill = color
                    s.smooth = False
                else:
                    s.graphicalProperties.solidFill = color
                    s.graphicalProperties.line.solidFill = color
            except AttributeError:
                pass

    if spec.get("data_labels") or (kind == "pie" and spec.get("data_labels") is not False):
        chart.dataLabels = DataLabelList()
        chart.dataLabels.showVal = kind != "pie"
        chart.dataLabels.showPercent = kind == "pie"
    if kind not in ("pie",):
        chart.y_axis.title = spec.get("y_title") or None
        chart.x_axis.title = spec.get("x_title") or None
        chart.y_axis.numFmt = spec.get("number_format") or "#,##0"
        chart.x_axis.delete = False      # これが無いとExcelで軸が消えることがある
        chart.y_axis.delete = False
    chart.width = float(spec.get("width") or 20)     # cm
    chart.height = float(spec.get("height") or 10)
    chart.legend.position = "b"
    if len(vals) <= 1 and kind != "pie":
        chart.legend = None

    anchor = spec.get("anchor") or f"{get_column_letter(len(columns) + 2)}{header_row}"
    ws.add_chart(chart, anchor)
    return chart


def build_excel(sheets: list[dict], title: str | None = None) -> bytes:
    """[{"name", "columns", "rows", "note"?, "charts"?}, ...] から xlsx を作る。

    charts は同じシートのデータから作るグラフの指定（複数可）。
    """
    if not sheets:
        raise ValueError("シートが1つもありません。")
    wb = Workbook()
    wb.remove(wb.active)
    used: set = set()

    for sh in sheets:
        columns = list(sh.get("columns") or [])
        rows = list(sh.get("rows") or [])
        ws = wb.create_sheet(safe_sheet_name(sh.get("name"), used))

        start = 1
        note = str(sh.get("note") or "").strip()
        if note:
            ws.cell(row=1, column=1, value=note).font = Font(italic=True, color="666666")
            start = 3

        for ci, col in enumerate(columns, start=1):
            c = ws.cell(row=start, column=ci, value=str(col))
            c.font = _HEADER_FONT
            c.fill = _HEADER_FILL
            c.alignment = Alignment(vertical="center", horizontal="center")
            c.border = _BORDER
        for ri, row in enumerate(rows, start=start + 1):
            banded = (ri - start) % 2 == 0
            for ci in range(1, len(columns) + 1):
                cell = ws.cell(row=ri, column=ci,
                               value=_cell_value(row[ci - 1] if ci - 1 < len(row) else None))
                cell.border = _BORDER
                if banded:
                    cell.fill = _BAND_FILL
                if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                    cell.number_format = "#,##0.####"

        ws.freeze_panes = ws.cell(row=start + 1, column=1)
        if columns and rows:
            ws.auto_filter.ref = (f"A{start}:"
                                  f"{get_column_letter(len(columns))}{start + len(rows)}")
        _autosize(ws, columns, rows)

        charts = sh.get("charts")
        if isinstance(charts, dict):
            charts = [charts]
        for i, spec in enumerate(charts or []):
            spec = dict(spec or {})
            spec.setdefault("anchor",
                            f"{get_column_letter(len(columns) + 2)}"
                            f"{start + i * 21}")
            _add_chart(ws, spec, columns, rows, start)

    if title:
        wb.properties.title = str(title)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ==========================================================================
# ===== 元 charts.py
# SELECT結果からグラフを組み立てる。
#
# チャット画面とユーザー定義ツールの両方がここを通るので、
# 対応するグラフ種別を増やすときはこのファイルだけを直せばよい。
#
# 種別の追加手順:
#   1. CHART_SPECS に (説明, 必要な指定, 分類) を足す
#   2. _BUILDERS に組み立て関数を足す
# validate() と画面の説明文は CHART_SPECS から自動で作られる。
# ==========================================================================
import numpy as np
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
from plotly.subplots import make_subplots

# 列の指定をリストで受け取るもの（存在チェックの仕方が違う）
LIST_FIELDS = ("path", "dimensions")

# 種別 -> (日本語の説明, 必要な指定, 分類)
CHART_SPECS: dict[str, tuple[str, tuple, str]] = {
    # --- 比較 ---------------------------------------------------------------
    "bar":         ("棒。カテゴリ別の比較", ("x", "y"), "比較"),
    "hbar":        ("横棒。項目名が長いときや順位表に", ("x", "y"), "比較"),
    "stacked_bar": ("積み上げ棒。内訳つきの比較", ("x", "y"), "比較"),
    "percent_bar": ("100%積み上げ棒。構成比の比較", ("x", "y"), "比較"),
    "lollipop":    ("ロリポップ。棒より軽く順位を見せる", ("x", "y"), "比較"),
    "dumbbell":    ("ダンベル。2時点の差を1行で比べる", ("x", "y", "y2"), "比較"),
    "pareto":      ("パレート図。棒＋累積比率で重点を見つける", ("x", "y"), "比較"),
    "pyramid":     ("人口ピラミッド。左右に分けた横棒", ("x", "y", "color"), "比較"),
    "marimekko":   ("マリメッコ。幅も高さも意味を持つ積み上げ", ("x", "y", "size"), "比較"),
    "radar":       ("レーダー。複数指標のバランス", ("x", "y"), "比較"),
    "polar_bar":   ("極座標の棒。方位や時間帯の分布", ("x", "y"), "比較"),
    "bump":        ("バンプ。順位の入れ替わりを追う", ("x", "y", "color"), "比較"),
    # --- 推移 ---------------------------------------------------------------
    "line":        ("折れ線。時系列や推移", ("x", "y"), "推移"),
    "step":        ("階段。在庫や料金など段階的に変わる値", ("x", "y"), "推移"),
    "area":        ("面。積み上げの推移", ("x", "y"), "推移"),
    "area_percent": ("100%面。構成比の推移", ("x", "y"), "推移"),
    "range_area":  ("幅つき折れ線。予測の上下限や信頼区間", ("x", "y", "lower", "upper"), "推移"),
    "slope":       ("スロープ。2時点の順位・水準の変化", ("x", "y", "color"), "推移"),
    "candlestick": ("ローソク足。始値・高値・安値・終値", ("x", "open", "high", "low", "close"), "推移"),
    "ohlc":        ("OHLC。ローソク足の棒型", ("x", "open", "high", "low", "close"), "推移"),
    "gantt":       ("ガントチャート。作業や期間の並び", ("x", "start", "end"), "推移"),
    "calendar":    ("カレンダーヒートマップ。日ごとの多寡", ("x", "y"), "推移"),
    "control_chart": ("管理図。平均±3σを外れた点を見つける", ("x", "y"), "推移"),
    # --- 構成 ---------------------------------------------------------------
    "pie":         ("円。構成比", ("x", "y"), "構成"),
    "donut":       ("ドーナツ。構成比（中央に合計）", ("x", "y"), "構成"),
    "treemap":     ("ツリーマップ。階層つき構成比", ("path", "y"), "構成"),
    "sunburst":    ("サンバースト。階層つき構成比（円形）", ("path", "y"), "構成"),
    "icicle":      ("アイシクル。階層を短冊で並べる", ("path", "y"), "構成"),
    "funnel":      ("ファネル。段階ごとの減少", ("x", "y"), "構成"),
    "waterfall":   ("ウォーターフォール。増減の内訳", ("x", "y"), "構成"),
    "sankey":      ("サンキー。流れと量（どこからどこへ）", ("source", "target", "y"), "構成"),
    # --- 分布 ---------------------------------------------------------------
    "histogram":   ("ヒストグラム。1つの数値の分布", ("x",), "分布"),
    "density":     ("密度曲線。ヒストグラムをなめらかに", ("x",), "分布"),
    "ecdf":        ("累積分布。「〇〇以下が何%か」を読む", ("x",), "分布"),
    "box":         ("箱ひげ。カテゴリ別のばらつき", ("y",), "分布"),
    "violin":      ("バイオリン。分布の形まで見る", ("y",), "分布"),
    "strip":       ("ストリップ。個々の点を並べる", ("y",), "分布"),
    "ridgeline":   ("リッジライン。群ごとの分布を重ねる", ("x", "color"), "分布"),
    "qq":          ("Q-Qプロット。正規分布からのズレ", ("x",), "分布"),
    # --- 関係 ---------------------------------------------------------------
    "scatter":     ("散布。2つの数値の相関", ("x", "y"), "関係"),
    "bubble":      ("バブル。散布＋大きさで3指標", ("x", "y", "size"), "関係"),
    "histogram2d": ("2次元ヒストグラム。点が多すぎるときの散布図", ("x", "y"), "関係"),
    "contour":     ("等高線。2変数の密度", ("x", "y"), "関係"),
    "heatmap":     ("ヒートマップ。2軸の集計をマス目で", ("x", "y"), "関係"),
    "matrix":      ("行列ヒートマップ。集計済みのクロス表や相関行列をそのまま色で", (), "関係"),
    "scatter_matrix": ("散布図行列。数値列を総当たりで見る", ("dimensions",), "関係"),
    "parallel_coordinates": ("平行座標。多変量の傾向を線で追う", ("dimensions",), "関係"),
    "parallel_categories": ("平行カテゴリ。区分の組み合わせの多さ", ("dimensions",), "関係"),
    "scatter3d":   ("3D散布。3つの数値の関係", ("x", "y", "z"), "関係"),
    "surface":     ("3D曲面。集計済みのクロス表を立体で", (), "関係"),
    "network":     ("ネットワーク。つながりの図", ("source", "target"), "関係"),
    # --- 指標 ---------------------------------------------------------------
    "indicator":   ("数値の大写し。KPIを1つ見せる", ("value",), "指標"),
    "gauge":       ("ゲージ。目標に対する達成度", ("value",), "指標"),
    "bullet":      ("ブレット。実績と目標を並べる", ("value",), "指標"),
}

CHART_TYPES = tuple(CHART_SPECS)


def type_help(category: str | None = None) -> str:
    """LLMに見せる一覧。分類を指定するとその分だけ返す。"""
    items = [(k, v) for k, v in CHART_SPECS.items()
             if category is None or v[2] == category]
    return " / ".join(f"{k}={v[0]}" for k, v in items)


def types_in(category: str) -> list[str]:
    return [k for k, v in CHART_SPECS.items() if v[2] == category]


def required_fields(chart_type: str) -> tuple:
    spec = CHART_SPECS.get(chart_type)
    return spec[1] if spec else ("x", "y")


def validate(item: dict, columns: list) -> list[str]:
    """指定された列が結果に存在するか検証し、問題点を返す。"""
    ct = item.get("chart_type") or "bar"
    errs = []
    if ct not in CHART_SPECS:
        return [f"未対応のグラフ種別です: {ct} / 使えるのは {', '.join(CHART_TYPES)}"]
    for f in required_fields(ct):
        v = item.get(f)
        if f in LIST_FIELDS:
            cols = list(v or [])
            if not cols:
                errs.append(f"{ct} には {f}（列名のリスト）が必要です。")
            errs += [f"{f} の列 '{c}' が結果にありません。利用可能: {columns}"
                     for c in cols if c not in columns]
        elif not v:
            errs.append(f"{ct} には {f} の指定が必要です。")
        elif v not in columns:
            errs.append(f"指定列 '{v}' が結果にありません。利用可能: {columns}")
    # 任意指定も、指定されていれば存在チェック
    for f in ("color", "size", "text", "y2", "z", "lower", "upper", "target", "facet"):
        v = item.get(f)
        if not v:
            continue
        # target だけは列名ではなく目標値（数値）で来ることがある
        if f == "target" and isinstance(v, (int, float)) and not isinstance(v, bool):
            continue
        if v not in columns:
            errs.append(f"指定列 '{v}' が結果にありません。利用可能: {columns}")
    return errs


# =============================================================================
# 下ごしらえ
# =============================================================================

def _scale(name):
    """色スケール名を色のリストに直す。

    名前のまま渡すと、周辺分布つきのグラフで plotly が文字列を1文字ずつ
    色として読み、"Blues" が 'B' 扱いになって落ちる。
    """
    return getattr(px.colors.sequential, str(name or "Blues"), None) or "Blues"


def _numeric(df: pd.DataFrame, *cols):
    for c in cols:
        if c and c in df.columns:
            df[c] = pd.to_numeric(df[c], errors="coerce").fillna(df[c])
    return df


def _num_series(df: pd.DataFrame, col) -> pd.Series:
    """必ず数値のSeriesにする（欠損は落とさず NaN のまま）。"""
    return pd.to_numeric(df[col], errors="coerce")


class _Ctx:
    """組み立て関数に渡す、よく使う値の詰め合わせ。"""

    def __init__(self, item: dict):
        self.item = item
        self.df = pd.DataFrame(item["rows"], columns=item["columns"])
        self.x, self.y = item.get("x"), item.get("y")
        self.title = item.get("title", "")
        for f in ("color", "size", "text", "y2", "z", "lower", "upper",
                  "target", "facet", "source", "start", "end",
                  "open", "high", "low", "close", "value"):
            setattr(self, f, item.get(f) if item.get(f) in self.df.columns else None)
        # target は列名でなく数値で来ることもある（目標値）
        self.target_value = item.get("target")
        self.path = [c for c in (item.get("path") or []) if c in self.df.columns]
        self.dimensions = [c for c in (item.get("dimensions") or []) if c in self.df.columns]
        _numeric(self.df, self.y, self.size, self.y2, self.z,
                 self.lower, self.upper, self.open, self.high, self.low, self.close)

    def get(self, key, default=None):
        return self.item.get(key, default)


# =============================================================================
# 比較
# =============================================================================

def _bar(c, orientation=None, barmode=None):
    return px.bar(c.df, x=c.x, y=c.y, color=c.color, text=c.text, title=c.title,
                  barmode=barmode or c.get("barmode") or "group",
                  orientation=orientation or c.get("orientation") or "v",
                  facet_col=c.facet)


def _hbar(c):
    # 横棒は「値が大きいものを上」に。並べ替えないと読みにくい
    d = c.df.sort_values(c.y) if c.y in c.df.columns else c.df
    return px.bar(d, x=c.y, y=c.x, color=c.color, text=c.text, title=c.title,
                  orientation="h", barmode=c.get("barmode") or "group")


def _stacked_bar(c):
    return _bar(c, barmode="stack")


def _percent_bar(c):
    d = c.df.copy()
    total = d.groupby(c.x)[c.y].transform("sum")
    d["_割合"] = _num_series(d, c.y) / total.replace(0, np.nan) * 100
    fig = px.bar(d, x=c.x, y="_割合", color=c.color, title=c.title, barmode="stack",
                 text=d["_割合"].round(1).astype(str) + "%")
    fig.update_yaxes(title_text="構成比(%)", range=[0, 100])
    return fig


def _lollipop(c):
    d = c.df.sort_values(c.y)
    fig = go.Figure()
    for _, r in d.iterrows():
        fig.add_shape(type="line", x0=0, x1=r[c.y], y0=r[c.x], y1=r[c.x],
                      line=dict(color="#9DC3E6", width=2))
    fig.add_trace(go.Scatter(x=d[c.y], y=d[c.x].astype(str), mode="markers",
                             marker=dict(size=12, color="#1F4E79"), name=c.y))
    fig.update_layout(title=c.title, xaxis_title=c.y, yaxis_title=c.x)
    return fig


def _dumbbell(c):
    d = c.df
    fig = go.Figure()
    for _, r in d.iterrows():
        fig.add_shape(type="line", x0=r[c.y], x1=r[c.y2], y0=r[c.x], y1=r[c.x],
                      line=dict(color="#BFBFBF", width=3))
    fig.add_trace(go.Scatter(x=d[c.y], y=d[c.x].astype(str), mode="markers",
                             name=str(c.y), marker=dict(size=12, color="#9DC3E6")))
    fig.add_trace(go.Scatter(x=d[c.y2], y=d[c.x].astype(str), mode="markers",
                             name=str(c.y2), marker=dict(size=12, color="#1F4E79")))
    fig.update_layout(title=c.title, xaxis_title="値", yaxis_title=c.x)
    return fig


def _pareto(c):
    d = c.df.copy()
    d[c.y] = _num_series(d, c.y)
    d = d.dropna(subset=[c.y]).sort_values(c.y, ascending=False)
    total = d[c.y].sum() or 1
    d["_累積"] = d[c.y].cumsum() / total * 100
    fig = make_subplots(specs=[[{"secondary_y": True}]])
    fig.add_trace(go.Bar(x=d[c.x].astype(str), y=d[c.y], name=str(c.y),
                         marker_color="#2E75B6"), secondary_y=False)
    fig.add_trace(go.Scatter(x=d[c.x].astype(str), y=d["_累積"], name="累積構成比",
                             mode="lines+markers", line=dict(color="#C55A11")),
                  secondary_y=True)
    fig.add_hline(y=80, line_dash="dot", line_color="#C55A11", secondary_y=True,
                  annotation_text="80%")
    fig.update_yaxes(title_text=str(c.y), secondary_y=False)
    fig.update_yaxes(title_text="累積構成比(%)", range=[0, 105], secondary_y=True)
    fig.update_layout(title=c.title)
    return fig


def _pyramid(c):
    """人口ピラミッド。color の2種類を左右に振り分ける。"""
    d = c.df.copy()
    d[c.y] = _num_series(d, c.y)
    groups = list(pd.unique(d[c.color].dropna()))[:2]
    if len(groups) < 2:
        raise ValueError(f"人口ピラミッドには color 列に2種類の値が必要です"
                         f"（いま: {groups}）。")
    left, right = groups
    fig = go.Figure()
    dl, dr = d[d[c.color] == left], d[d[c.color] == right]
    fig.add_trace(go.Bar(y=dl[c.x].astype(str), x=-dl[c.y], name=str(left),
                         orientation="h", marker_color="#2E75B6"))
    fig.add_trace(go.Bar(y=dr[c.x].astype(str), x=dr[c.y], name=str(right),
                         orientation="h", marker_color="#F4B183"))
    fig.update_layout(title=c.title, barmode="overlay", bargap=0.1,
                      xaxis=dict(title=str(c.y),
                                 tickvals=None, ticktext=None))
    fig.update_xaxes(tickformat="~s")
    return fig


def _marimekko(c):
    """幅=size、高さ=y の積み上げ。x ごとの規模と内訳を同時に見せる。"""
    d = c.df.copy()
    d[c.y] = _num_series(d, c.y)
    d[c.size] = _num_series(d, c.size)
    widths = d.groupby(c.x, sort=False)[c.size].max()
    total_w = widths.sum() or 1
    fig = go.Figure()
    keys = list(widths.index)
    lefts, acc = {}, 0.0
    for k in keys:
        lefts[k] = acc
        acc += float(widths[k]) / total_w * 100
    groups = list(pd.unique(d[c.color].dropna())) if c.color else [None]
    for gi, g in enumerate(groups):
        sub = d if g is None else d[d[c.color] == g]
        xs, ys, ws = [], [], []
        for k in keys:
            row = sub[sub[c.x] == k]
            if row.empty:
                continue
            w = float(widths[k]) / total_w * 100
            xs.append(lefts[k] + w / 2)
            ws.append(w)
            ys.append(float(row[c.y].iloc[0]))
        fig.add_trace(go.Bar(x=xs, y=ys, width=ws, name=str(g) if g is not None else str(c.y),
                             marker_color=px.colors.qualitative.Set2[gi % 8]))
    fig.update_layout(title=c.title, barmode="stack", bargap=0,
                      xaxis_title=f"{c.x}（幅 = {c.size}）", yaxis_title=str(c.y))
    return fig


def _radar(c):
    fig = px.line_polar(c.df, r=c.y, theta=c.x, color=c.color, line_close=True,
                        title=c.title)
    fig.update_traces(fill="toself", opacity=0.5)
    return fig


def _polar_bar(c):
    return px.bar_polar(c.df, r=c.y, theta=c.x, color=c.color, title=c.title)


def _bump(c):
    """順位の推移。値が小さいほど上位なので、y軸を反転する。"""
    d = c.df.copy()
    d[c.y] = _num_series(d, c.y)
    fig = px.line(d, x=c.x, y=c.y, color=c.color, markers=True, title=c.title,
                  text=c.text)
    fig.update_traces(marker=dict(size=11))
    fig.update_yaxes(autorange="reversed", title_text=f"{c.y}（上が上位）",
                     dtick=1)
    return fig


# =============================================================================
# 推移
# =============================================================================

def _line(c):
    return px.line(c.df, x=c.x, y=c.y, color=c.color, text=c.text, title=c.title,
                   markers=True, facet_col=c.facet)


def _step(c):
    fig = px.line(c.df, x=c.x, y=c.y, color=c.color, title=c.title, markers=True)
    fig.update_traces(line_shape="hv")
    return fig


def _area(c):
    return px.area(c.df, x=c.x, y=c.y, color=c.color, title=c.title)


def _area_percent(c):
    d = c.df.copy()
    d[c.y] = _num_series(d, c.y)
    total = d.groupby(c.x)[c.y].transform("sum")
    d["_割合"] = d[c.y] / total.replace(0, np.nan) * 100
    fig = px.area(d, x=c.x, y="_割合", color=c.color, title=c.title)
    fig.update_yaxes(title_text="構成比(%)", range=[0, 100])
    return fig


def _range_area(c):
    d = c.df
    fig = go.Figure()
    fig.add_trace(go.Scatter(x=d[c.x], y=d[c.upper], mode="lines", name="上限",
                             line=dict(width=0), showlegend=False))
    fig.add_trace(go.Scatter(x=d[c.x], y=d[c.lower], mode="lines", name="幅（95%）",
                             line=dict(width=0), fill="tonexty",
                             fillcolor="rgba(46,117,182,.18)"))
    fig.add_trace(go.Scatter(x=d[c.x], y=d[c.y], mode="lines+markers", name=str(c.y),
                             line=dict(color="#1F4E79", width=2)))
    fig.update_layout(title=c.title, xaxis_title=str(c.x), yaxis_title=str(c.y))
    return fig


def _slope(c):
    d = c.df.copy()
    d[c.y] = _num_series(d, c.y)
    fig = px.line(d, x=c.x, y=c.y, color=c.color, markers=True, title=c.title)
    fig.update_traces(line=dict(width=2))
    # 端に系列名を出す（凡例を目で追わなくて済む）
    first = str(d[c.x].iloc[0])
    for name, g in d.groupby(c.color):
        head = g[g[c.x].astype(str) == first]
        if len(head):
            fig.add_annotation(x=head[c.x].iloc[0], y=head[c.y].iloc[0], text=str(name),
                               xanchor="right", showarrow=False, xshift=-6, font_size=11)
    fig.update_layout(showlegend=False)
    return fig


def _candlestick(c):
    return go.Figure(go.Candlestick(
        x=c.df[c.x], open=c.df[c.open], high=c.df[c.high],
        low=c.df[c.low], close=c.df[c.close])).update_layout(
            title=c.title, xaxis_rangeslider_visible=False)


def _ohlc(c):
    return go.Figure(go.Ohlc(
        x=c.df[c.x], open=c.df[c.open], high=c.df[c.high],
        low=c.df[c.low], close=c.df[c.close])).update_layout(
            title=c.title, xaxis_rangeslider_visible=False)


def _gantt(c):
    fig = px.timeline(c.df, x_start=c.start, x_end=c.end, y=c.x, color=c.color,
                      text=c.text, title=c.title)
    fig.update_yaxes(autorange="reversed")     # 上から順に並べる
    return fig


def _calendar(c):
    """日付ごとの値を、週×曜日のマス目にする。"""
    d = c.df.copy()
    d[c.x] = pd.to_datetime(d[c.x], errors="coerce")
    d[c.y] = _num_series(d, c.y)
    d = d.dropna(subset=[c.x])
    if d.empty:
        raise ValueError(f"{c.x} を日付として読めませんでした。")
    d["_日付"] = d[c.x].dt.normalize()
    d = d.groupby("_日付", as_index=False)[c.y].sum().rename(columns={c.y: "_値"})
    d["_週"] = d["_日付"].dt.isocalendar().week.astype(int)
    d["_年"] = d["_日付"].dt.isocalendar().year.astype(int)
    d["_通週"] = (d["_年"] - d["_年"].min()) * 53 + d["_週"]
    names = ["月", "火", "水", "木", "金", "土", "日"]
    d["_曜日"] = d["_日付"].dt.weekday
    pivot = d.pivot_table(index="_曜日", columns="_通週", values="_値", aggfunc="sum")
    pivot = pivot.reindex(range(7))
    labels = (d.groupby("_通週")["_日付"].min().dt.strftime("%m/%d")
              .reindex(pivot.columns).tolist())
    fig = px.imshow(pivot.to_numpy(), x=labels, y=names, aspect="auto",
                    color_continuous_scale=_scale(c.get("colorscale")),
                    title=c.title, labels=dict(color=str(c.y)))
    fig.update_xaxes(title_text="週（週初の日付）", side="top")
    return fig


def _control_chart(c):
    """管理図。平均と±3σを引き、外れた点を赤くする。"""
    d = c.df.copy()
    d[c.y] = _num_series(d, c.y)
    d = d.dropna(subset=[c.y])
    m, sd = d[c.y].mean(), d[c.y].std(ddof=1)
    ucl, lcl = m + 3 * sd, m - 3 * sd
    out = (d[c.y] > ucl) | (d[c.y] < lcl)
    fig = go.Figure()
    fig.add_trace(go.Scatter(x=d[c.x], y=d[c.y], mode="lines+markers", name=str(c.y),
                             line=dict(color="#2E75B6"),
                             marker=dict(size=8,
                                         color=np.where(out, "#B02A2A", "#2E75B6"))))
    for val, name, dash in ((m, "平均", "solid"), (ucl, "上方管理限界(+3σ)", "dash"),
                            (lcl, "下方管理限界(-3σ)", "dash")):
        fig.add_hline(y=val, line_dash=dash, line_color="#7F7F7F",
                      annotation_text=f"{name} {val:,.4g}", annotation_position="right")
    fig.update_layout(title=c.title or "管理図", xaxis_title=str(c.x),
                      yaxis_title=str(c.y))
    return fig


# =============================================================================
# 構成
# =============================================================================

def _pie(c):
    return px.pie(c.df, names=c.x, values=c.y, color=c.color, title=c.title,
                  hole=0.45 if c.get("chart_type") == "donut" else 0)


def _treemap(c):
    return px.treemap(c.df, path=c.path, values=c.y, color=c.color, title=c.title)


def _sunburst(c):
    return px.sunburst(c.df, path=c.path, values=c.y, color=c.color, title=c.title)


def _icicle(c):
    return px.icicle(c.df, path=c.path, values=c.y, color=c.color, title=c.title)


def _funnel(c):
    # plotly は x=値 / y=段階 なので入れ替える（x に段階、y に値を受け取る仕様）
    return px.funnel(c.df, x=c.y, y=c.x, color=c.color, title=c.title)


def _waterfall(c):
    d = c.df
    measure = ["relative"] * len(d)
    # 「合計」「total」で終わる行は合計として扱う
    for i, v in enumerate(d[c.x].astype(str)):
        if v.strip() in ("合計", "計", "total", "Total", "TOTAL"):
            measure[i] = "total"
    fig = go.Figure(go.Waterfall(
        x=d[c.x].astype(str), y=_num_series(d, c.y), measure=measure,
        text=d[c.y], textposition="outside"))
    fig.update_layout(title=c.title, waterfallgap=0.3)
    return fig


def _sankey(c):
    d = c.df.copy()
    d[c.y] = _num_series(d, c.y)
    labels = list(dict.fromkeys(d[c.source].astype(str).tolist()
                                + d[c.target].astype(str).tolist()))
    idx = {v: i for i, v in enumerate(labels)}
    fig = go.Figure(go.Sankey(
        node=dict(label=labels, pad=16, thickness=16,
                  line=dict(color="#BFBFBF", width=0.5)),
        link=dict(source=[idx[str(v)] for v in d[c.source]],
                  target=[idx[str(v)] for v in d[c.target]],
                  value=d[c.y].fillna(0).tolist())))
    fig.update_layout(title=c.title, font_size=12)
    return fig


# =============================================================================
# 分布
# =============================================================================

def _histogram(c):
    return px.histogram(c.df, x=c.x, color=c.color, title=c.title,
                        nbins=int(c.get("nbins")) if c.get("nbins") else None,
                        facet_col=c.facet, marginal=c.get("marginal"))


def _density(c):
    """ヒストグラム＋カーネル密度推定の曲線。"""
    from scipy import stats as sstats
    d = c.df.copy()
    d[c.x] = pd.to_numeric(d[c.x], errors="coerce")
    d = d.dropna(subset=[c.x])
    if len(d) < 3:
        raise ValueError("密度曲線には3行以上の数値が必要です。")
    fig = px.histogram(d, x=c.x, color=c.color, histnorm="probability density",
                       opacity=0.55, nbins=int(c.get("nbins") or 30), title=c.title)
    groups = [(None, d)] if not c.color else list(d.groupby(c.color))
    xs = np.linspace(d[c.x].min(), d[c.x].max(), 200)
    for name, g in groups:
        if g[c.x].nunique() < 2:
            continue
        try:
            kde = sstats.gaussian_kde(g[c.x].to_numpy())
        except np.linalg.LinAlgError:
            continue
        fig.add_trace(go.Scatter(x=xs, y=kde(xs), mode="lines",
                                 name=f"{name} 密度" if name is not None else "密度",
                                 line=dict(width=2)))
    fig.update_layout(bargap=0.02)
    return fig


def _ecdf(c):
    return px.ecdf(c.df, x=c.x, color=c.color, title=c.title, markers=False)


def _box(c):
    return px.box(c.df, x=c.x, y=c.y, color=c.color, title=c.title, points="outliers",
                  facet_col=c.facet)


def _violin(c):
    return px.violin(c.df, x=c.x, y=c.y, color=c.color, title=c.title, box=True,
                     points=False)


def _strip(c):
    return px.strip(c.df, x=c.x, y=c.y, color=c.color, title=c.title)


def _ridgeline(c):
    """群ごとの分布を少しずつずらして重ねる。"""
    d = c.df.copy()
    d[c.x] = pd.to_numeric(d[c.x], errors="coerce")
    d = d.dropna(subset=[c.x])
    fig = go.Figure()
    for name, g in d.groupby(c.color):
        fig.add_trace(go.Violin(x=g[c.x], name=str(name), side="positive",
                                width=2.2, points=False, meanline_visible=True,
                                orientation="h"))
    fig.update_layout(title=c.title, violingap=0, violinmode="overlay",
                      xaxis_title=str(c.x), showlegend=False)
    return fig


def _qq(c):
    """正規Q-Qプロット。点が直線に乗るほど正規分布に近い。"""
    from scipy import stats as sstats
    s = pd.to_numeric(c.df[c.x], errors="coerce").dropna().sort_values()
    if len(s) < 3:
        raise ValueError("Q-Qプロットには3行以上の数値が必要です。")
    theo = sstats.norm.ppf((np.arange(1, len(s) + 1) - 0.5) / len(s))
    theo = theo * s.std(ddof=1) + s.mean()
    fig = go.Figure()
    fig.add_trace(go.Scatter(x=theo, y=s, mode="markers", name="実測",
                             marker=dict(color="#2E75B6", size=7)))
    lo, hi = float(min(theo.min(), s.min())), float(max(theo.max(), s.max()))
    fig.add_trace(go.Scatter(x=[lo, hi], y=[lo, hi], mode="lines", name="正規分布の直線",
                             line=dict(color="#C55A11", dash="dash")))
    fig.update_layout(title=c.title or f"{c.x} のQ-Qプロット",
                      xaxis_title="正規分布ならこうなる", yaxis_title="実測")
    return fig


# =============================================================================
# 関係
# =============================================================================

def _scatter(c):
    return px.scatter(c.df, x=c.x, y=c.y, color=c.color, text=c.text, title=c.title,
                      facet_col=c.facet,
                      trendline="ols" if c.get("trendline") else None)


def _bubble(c):
    return px.scatter(c.df, x=c.x, y=c.y, color=c.color, size=c.size, text=c.text,
                      title=c.title, size_max=50)


def _histogram2d(c):
    return px.density_heatmap(c.df, x=c.x, y=c.y, title=c.title,
                              nbinsx=int(c.get("nbins") or 30),
                              nbinsy=int(c.get("nbins") or 30),
                              color_continuous_scale=_scale(c.get("colorscale")),
                              marginal_x="histogram", marginal_y="histogram")


def _contour(c):
    fig = px.density_contour(c.df, x=c.x, y=c.y, color=c.color, title=c.title)
    fig.update_traces(contours_coloring="fill", contours_showlabels=True)
    return fig


def _heatmap(c):
    if c.color:
        return px.density_heatmap(c.df, x=c.x, y=c.y, z=c.color, histfunc="sum",
                                  title=c.title, text_auto=True,
                                  color_continuous_scale=_scale(c.get("colorscale")))
    return px.density_heatmap(c.df, x=c.x, y=c.y, histfunc="count", title=c.title,
                              text_auto=True,
                              color_continuous_scale=_scale(c.get("colorscale")))


def _matrix(c):
    """集計済みの表をそのまま行列として塗る。"""
    label = c.x if c.x in c.df.columns else c.df.columns[0]
    m = c.df.set_index(label)
    m = m.apply(lambda s: pd.to_numeric(s, errors="coerce")).dropna(axis=1, how="all")
    fig = px.imshow(m, text_auto=True, aspect="auto", title=c.title,
                    color_continuous_scale=_scale(c.get("colorscale")))
    fig.update_xaxes(side="top")
    return fig


def _scatter_matrix(c):
    d = c.df.copy()
    for col in c.dimensions:
        d[col] = pd.to_numeric(d[col], errors="coerce")
    fig = px.scatter_matrix(d, dimensions=c.dimensions, color=c.color, title=c.title)
    fig.update_traces(diagonal_visible=False, showupperhalf=False,
                      marker=dict(size=4, opacity=0.6))
    return fig


def _parallel_coordinates(c):
    d = c.df.copy()
    for col in c.dimensions:
        d[col] = pd.to_numeric(d[col], errors="coerce")
    d = d.dropna(subset=c.dimensions)
    color = c.color if (c.color and pd.api.types.is_numeric_dtype(
        pd.to_numeric(d[c.color], errors="coerce"))) else None
    if color:
        d[color] = pd.to_numeric(d[color], errors="coerce")
    return px.parallel_coordinates(d, dimensions=c.dimensions, color=color,
                                   title=c.title,
                                   color_continuous_scale=_scale(c.get("colorscale")))


def _parallel_categories(c):
    return px.parallel_categories(c.df, dimensions=c.dimensions, title=c.title,
                                  color=(pd.to_numeric(c.df[c.color], errors="coerce")
                                         if c.color else None))


def _scatter3d(c):
    return px.scatter_3d(c.df, x=c.x, y=c.y, z=c.z, color=c.color, size=c.size,
                         text=c.text, title=c.title)


def _surface(c):
    """集計済みのクロス表を立体にする（1列目が行ラベル）。"""
    label = c.x if c.x in c.df.columns else c.df.columns[0]
    m = c.df.set_index(label)
    m = m.apply(lambda s: pd.to_numeric(s, errors="coerce")).dropna(axis=1, how="all")
    fig = go.Figure(go.Surface(z=m.to_numpy(), x=list(m.columns),
                               y=[str(i) for i in m.index],
                               colorscale=_scale(c.get("colorscale"))))
    fig.update_layout(title=c.title, scene=dict(
        xaxis_title="列", yaxis_title=str(label), zaxis_title="値"))
    return fig


def _network(c):
    """つながりの図。円周上にノードを並べ、関係を線で結ぶ。"""
    d = c.df
    nodes = list(dict.fromkeys(d[c.source].astype(str).tolist()
                               + d[c.target].astype(str).tolist()))
    n = len(nodes)
    if not n:
        raise ValueError("つながりが1件もありません。")
    ang = {v: 2 * np.pi * i / n for i, v in enumerate(nodes)}
    pos = {v: (np.cos(a), np.sin(a)) for v, a in ang.items()}
    weights = _num_series(d, c.y) if c.y in d.columns else pd.Series([1] * len(d))
    wmax = float(weights.max() or 1)
    edge_x, edge_y = [], []
    for (_, r), w in zip(d.iterrows(), weights):
        x0, y0 = pos[str(r[c.source])]
        x1, y1 = pos[str(r[c.target])]
        edge_x += [x0, x1, None]
        edge_y += [y0, y1, None]
    fig = go.Figure()
    fig.add_trace(go.Scatter(x=edge_x, y=edge_y, mode="lines", hoverinfo="skip",
                             line=dict(color="rgba(120,140,170,.45)", width=1.5),
                             showlegend=False))
    deg = pd.Series(d[c.source].astype(str).tolist()
                    + d[c.target].astype(str).tolist()).value_counts()
    fig.add_trace(go.Scatter(
        x=[pos[v][0] for v in nodes], y=[pos[v][1] for v in nodes],
        mode="markers+text", text=nodes, textposition="top center",
        marker=dict(size=[12 + 26 * deg.get(v, 1) / max(deg.max(), 1) for v in nodes],
                    color="#2E75B6"),
        hovertext=[f"{v}: {deg.get(v, 0)}件" for v in nodes], hoverinfo="text",
        showlegend=False))
    fig.update_layout(title=c.title, xaxis=dict(visible=False),
                      yaxis=dict(visible=False, scaleanchor="x"),
                      plot_bgcolor="rgba(0,0,0,0)")
    if wmax:
        fig.update_layout(margin=dict(l=20, r=20, t=50, b=20))
    return fig


# =============================================================================
# 指標
# =============================================================================

def _indicator_value(c) -> float:
    s = pd.to_numeric(c.df[c.value], errors="coerce").dropna()
    if s.empty:
        raise ValueError(f"{c.value} に数値がありません。")
    mode = (c.get("agg") or "sum").lower()
    return float({"sum": s.sum, "mean": s.mean, "max": s.max,
                  "min": s.min, "last": lambda: s.iloc[-1]}.get(mode, s.sum)())


def _target_of(c, value: float):
    if c.target and c.target in c.df.columns:
        t = pd.to_numeric(c.df[c.target], errors="coerce").dropna()
        return float(t.sum()) if len(t) else None
    try:
        return float(c.target_value) if c.target_value is not None else None
    except (TypeError, ValueError):
        return None


def _indicator(c):
    v = _indicator_value(c)
    t = _target_of(c, v)
    fig = go.Figure(go.Indicator(
        mode="number+delta" if t else "number", value=v,
        number=dict(valueformat=c.get("valueformat") or ",.4~f",
                    suffix=c.get("suffix") or ""),
        delta=dict(reference=t, relative=True, valueformat=".1%") if t else None,
        title=dict(text=c.title or str(c.value))))
    return fig


def _gauge(c):
    v = _indicator_value(c)
    t = _target_of(c, v)
    top = float(c.get("max") or (max(v, t or 0) * 1.25) or 1)
    fig = go.Figure(go.Indicator(
        mode="gauge+number+delta" if t else "gauge+number", value=v,
        number=dict(valueformat=c.get("valueformat") or ",.4~f",
                    suffix=c.get("suffix") or ""),
        delta=dict(reference=t, relative=True, valueformat=".1%") if t else None,
        title=dict(text=c.title or str(c.value)),
        gauge=dict(axis=dict(range=[0, top]), bar=dict(color="#2E75B6"),
                   steps=[dict(range=[0, top * 0.5], color="#F2F6FB"),
                          dict(range=[top * 0.5, top * 0.8], color="#D9E2F3")],
                   threshold=(dict(line=dict(color="#C55A11", width=3), value=t)
                              if t else None))))
    return fig


def _bullet(c):
    v = _indicator_value(c)
    t = _target_of(c, v)
    top = float(c.get("max") or (max(v, t or 0) * 1.25) or 1)
    fig = go.Figure(go.Indicator(
        mode="number+gauge+delta" if t else "number+gauge", value=v,
        delta=dict(reference=t) if t else None,
        number=dict(valueformat=c.get("valueformat") or ",.4~f"),
        title=dict(text=c.title or str(c.value)),
        gauge=dict(shape="bullet", axis=dict(range=[0, top]),
                   bar=dict(color="#1F4E79", thickness=0.6),
                   steps=[dict(range=[0, top * 0.6], color="#F2F6FB"),
                          dict(range=[top * 0.6, top * 0.85], color="#D9E2F3")],
                   threshold=(dict(line=dict(color="#C55A11", width=3), value=t)
                              if t else None))))
    fig.update_layout(height=190)
    return fig


# =============================================================================
# 組み立ての振り分け
# =============================================================================

_BUILDERS = {
    "bar": _bar, "hbar": _hbar, "stacked_bar": _stacked_bar,
    "percent_bar": _percent_bar, "lollipop": _lollipop, "dumbbell": _dumbbell,
    "pareto": _pareto, "pyramid": _pyramid, "marimekko": _marimekko,
    "radar": _radar, "polar_bar": _polar_bar, "bump": _bump,
    "line": _line, "step": _step, "area": _area, "area_percent": _area_percent,
    "range_area": _range_area, "slope": _slope, "candlestick": _candlestick,
    "ohlc": _ohlc, "gantt": _gantt, "calendar": _calendar,
    "control_chart": _control_chart,
    "pie": _pie, "donut": _pie, "treemap": _treemap, "sunburst": _sunburst,
    "icicle": _icicle, "funnel": _funnel, "waterfall": _waterfall, "sankey": _sankey,
    "histogram": _histogram, "density": _density, "ecdf": _ecdf, "box": _box,
    "violin": _violin, "strip": _strip, "ridgeline": _ridgeline, "qq": _qq,
    "scatter": _scatter, "bubble": _bubble, "histogram2d": _histogram2d,
    "contour": _contour, "heatmap": _heatmap, "matrix": _matrix,
    "scatter_matrix": _scatter_matrix,
    "parallel_coordinates": _parallel_coordinates,
    "parallel_categories": _parallel_categories,
    "scatter3d": _scatter3d, "surface": _surface, "network": _network,
    "indicator": _indicator, "gauge": _gauge, "bullet": _bullet,
}


def build_figure(item: dict):
    """render アイテム（kind="chart"）から plotly の figure を作る。"""
    ct = item.get("chart_type", "bar")
    builder = _BUILDERS.get(ct)
    if builder is None:
        raise ValueError(f"未対応のグラフ種別です: {ct} / "
                         f"使えるのは {', '.join(CHART_TYPES)}")
    fig = builder(_Ctx(item))
    fig.update_layout(margin=dict(l=55, r=20, t=50, b=50))
    return fig


def build_dual_figure(item: dict):
    """棒(左軸)+折れ線(右軸)の2軸グラフ。"""
    df = pd.DataFrame(item["rows"], columns=item["columns"])
    x = item["x"]
    bar_y = item.get("bar_y") or []
    line_y = item.get("line_y") or []
    _numeric(df, *bar_y, *line_y)

    fig = make_subplots(specs=[[{"secondary_y": True}]])
    for col in bar_y:
        fig.add_trace(go.Bar(x=df[x], y=df[col], name=col), secondary_y=False)
    for col in line_y:
        fig.add_trace(go.Scatter(x=df[x], y=df[col], name=col, mode="lines+markers"),
                      secondary_y=True)
    fig.update_layout(title=item.get("title", ""), barmode="group",
                      legend=dict(orientation="h", yanchor="bottom", y=1.02))
    fig.update_xaxes(title_text=x)
    fig.update_yaxes(title_text=item.get("left_title") or "（左軸）", secondary_y=False)
    fig.update_yaxes(title_text=item.get("right_title") or "（右軸）", secondary_y=True)
    return fig


# ==========================================================================
# ===== 元 figures.py
# グラフを画像（PNG）にする。Word や PowerPoint に貼るために使う。
#
# plotly の画像化は Chrome を裏で動かす（kaleido）。社内サーバに Chrome が
# 入っていないこともあるので、失敗したら None を返し、呼び出し側は
# 表や説明文だけで文書を作れるようにしてある。文書作成そのものは止めない。
#
# 画像は1文書内で何度も作るため、同じ図は使い回す（同じ処理を2回しない）。
# ==========================================================================
import hashlib
import threading

import config

_lock = threading.Lock()
_cache: dict[str, bytes] = {}
_MAX_CACHE = 40
# 一度失敗したら、その実行中は再挑戦しない（1枚あたり数秒待たされるため）
_broken: list[str] = []


def why_unavailable() -> str:
    return _broken[0] if _broken else ""


def render(fig, width: int | None = None, height: int | None = None,
           scale: float | None = None) -> bytes | None:
    """plotly の figure を PNG のバイト列にする。できなければ None。"""
    if _broken:
        return None
    w = int(width or config.REPORT_IMAGE_WIDTH)
    h = int(height or config.REPORT_IMAGE_HEIGHT)
    s = float(scale or config.REPORT_IMAGE_SCALE)
    try:
        key = hashlib.sha1(
            (fig.to_json() + f"|{w}x{h}@{s}").encode("utf-8")).hexdigest()
    except Exception:
        key = None
    if key:
        with _lock:
            hit = _cache.get(key)
        if hit is not None:
            return hit
    try:
        data = fig.to_image(format="png", width=w, height=h, scale=s)
    except Exception as e:
        msg = str(e).splitlines()[0][:200]
        _broken.append(f"グラフを画像にできませんでした（{type(e).__name__}: {msg}）。"
                       "文書には表と説明だけを入れます。"
                       "画像も入れたい場合は、サーバに Chrome/Chromium を用意して"
                       "kaleido が使える状態にしてください。")
        print(f"[figures] 画像化を無効にしました: {msg}")
        return None
    if key:
        with _lock:
            _cache[key] = data
            while len(_cache) > _MAX_CACHE:
                _cache.pop(next(iter(_cache)))
    return data


def for_print(fig, *, width=None, height=None):
    """紙・スライド向けに見た目を整えてから画像にする。

    画面はマウスで拡大できるが、紙とスライドはできない。
    文字を大きめに、余白を詰め、目盛りに桁区切りを入れる。
    """
    fig = _polish(fig)
    return render(fig, width=width, height=height)


def _polish(fig):
    """印刷向けの体裁に整える（元の figure は壊さない）。"""
    import copy
    fig = copy.deepcopy(fig)
    fig.update_layout(
        template="plotly_white",
        font=dict(family=config.REPORT_FONT_JA + ", sans-serif", size=15,
                  color="#1F1F1F"),
        title=dict(font=dict(size=17)),
        margin=dict(l=70, r=30, t=50, b=60),
        legend=dict(orientation="h", yanchor="bottom", y=1.02, x=0,
                    font=dict(size=13)),
        paper_bgcolor="white", plot_bgcolor="white",
        colorway=["#1F4E79", "#F4B183", "#70AD47", "#C55A11", "#7F7F7F",
                  "#2E75B6", "#A9D18E", "#FFD966", "#9DC3E6", "#BFBFBF"],
    )
    fig.update_xaxes(showgrid=False, linecolor="#BFBFBF", ticks="outside",
                     tickfont=dict(size=13))
    fig.update_yaxes(gridcolor="#E8E8E8", zerolinecolor="#BFBFBF",
                     tickfont=dict(size=13), tickformat=",")
    return fig


# ==========================================================================
# ===== 元 docx_report.py
# Wordレポートの生成。そのまま配布・回覧できる体裁で作る。
#
# 作りの方針:
#   - 表紙 → 目次 → 要約 → 本編 → 結論 → 付録 の順。報告書の型に合わせる。
#   - 図と表には通し番号とキャプションを付ける（「図3のとおり」と本文から呼べる）。
#   - 日本語フォントを明示的に当てる。指定しないと英字フォントが当たり、
#     開いた瞬間に体裁が崩れて見える。
#   - グラフは画像として貼る（Wordにネイティブのグラフが無いため）。
#     画像化できない環境では、同じ内容の表に自動で置き換える。
#
# 1セクション = 1つの dict:
#     {heading, body, bullets, table:{columns,rows}, image:bytes, caption,
#      note, callout, page_break}
# ==========================================================================
import io
import re
from datetime import datetime

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import config

DOCX_NAVY = RGBColor(0x1F, 0x3B, 0x5C)
DOCX_ACCENT = RGBColor(0x2E, 0x75, 0xB6)
DOCX_HILITE = RGBColor(0xC5, 0x5A, 0x11)
DOCX_INK = RGBColor(0x22, 0x26, 0x2B)
DOCX_MUTED = RGBColor(0x6B, 0x72, 0x80)
DOCX_BAND = "F5F8FC"
HEADER_BG = "1F3B5C"
CALLOUT_BG = "FDF3E7"

DOCX_MAX_TABLE_ROWS = 40


class DocxReportError(Exception):
    """レポートを作れない理由（そのまま画面に出す）。"""


# =============================================================================
# 体裁の下ごしらえ
# =============================================================================

def _jp_font(run, size=None, bold=None, color=None, name=None):
    """日本語フォントを当てる（東アジア用は XML で直接指定する）。"""
    font = name or config.REPORT_FONT_JA
    run.font.name = font
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(attr), font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color


def _shade(cell_or_par, hex_color: str):
    el = cell_or_par._tc if hasattr(cell_or_par, "_tc") else cell_or_par._p
    pr = el.get_or_add_tcPr() if hasattr(el, "get_or_add_tcPr") else \
        el.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    pr.append(shd)


def _border(par, *, size=6, color="C55A11", where="left"):
    pPr = par._p.get_or_add_pPr()
    borders = pPr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        pPr.append(borders)
    b = OxmlElement(f"w:{where}")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), str(size * 4))
    b.set(qn("w:space"), "8")
    b.set(qn("w:color"), color)
    borders.append(b)


def _field(par, instr: str):
    """Wordのフィールド（ページ番号や目次）を入れる。開いたときに計算される。"""
    r1 = par.add_run()._element
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    r1.append(fld)
    r2 = par.add_run()._element
    txt = OxmlElement("w:instrText")
    txt.set(qn("xml:space"), "preserve")
    txt.text = instr
    r2.append(txt)
    r3 = par.add_run()._element
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    r3.append(sep)
    par.add_run("　")                        # 未計算のときに出る仮の文字
    r5 = par.add_run()._element
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r5.append(end)


def _setup_styles(doc):
    """標準スタイルを日本語向けに整える。"""
    normal = doc.styles["Normal"]
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = DOCX_INK
    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(attr), config.REPORT_FONT_JA)
    normal.paragraph_format.line_spacing = 1.4
    normal.paragraph_format.space_after = Pt(6)

    for name, size, color, before in (("Heading 1", 16, DOCX_NAVY, 18),
                                      ("Heading 2", 13, DOCX_NAVY, 14),
                                      ("Heading 3", 11.5, DOCX_ACCENT, 10)):
        st = doc.styles[name]
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = color
        rPr = st.element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.insert(0, rFonts)
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rFonts.set(qn(attr), config.REPORT_FONT_JA)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_with_next = True


def _setup_page(doc, footer_text: str):
    sec = doc.sections[0]
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.2)
    sec.right_margin = Cm(2.2)

    p = sec.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if footer_text:
        r = p.add_run(footer_text + "　　")
        _jp_font(r, size=8.5, color=DOCX_MUTED)
    _field(p, "PAGE")
    r = p.add_run(" / ")
    _jp_font(r, size=8.5, color=DOCX_MUTED)
    _field(p, "NUMPAGES")
    for r in p.runs:
        _jp_font(r, size=8.5, color=DOCX_MUTED)


def _para(doc, text="", *, size=10.5, bold=False, color=DOCX_INK, align=None,
          space_after=6, style=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    for i, line in enumerate(str(text).split("\n")):
        if i:
            p.add_run().add_break()
        r = p.add_run(line)
        _jp_font(r, size=size, bold=bold, color=color)
    return p


def _docx_fmt(v) -> str:
    if v is None:
        return ""
    if isinstance(v, bool):
        return "はい" if v else "いいえ"
    if isinstance(v, float):
        if v == int(v) and abs(v) < 1e15:
            return f"{int(v):,}"
        return f"{v:,.2f}".rstrip("0").rstrip(".")
    if isinstance(v, int):
        return f"{v:,}"
    return str(v)


# =============================================================================
# 部品
# =============================================================================

def _cover(doc, args: dict):
    for _ in range(4):
        doc.add_paragraph()
    _para(doc, args.get("title", "レポート"), size=26, bold=True, color=DOCX_NAVY,
          align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    if args.get("subtitle"):
        _para(doc, args["subtitle"], size=13, color=DOCX_MUTED,
              align=WD_ALIGN_PARAGRAPH.CENTER, space_after=28)

    # 表紙の線
    line = doc.add_paragraph()
    line.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _border(line, size=3, color="1F3B5C", where="bottom")

    for _ in range(6):
        doc.add_paragraph()
    org = args.get("org") or config.REPORT_ORG
    for text in (args.get("date") or datetime.now().strftime("%Y年%m月%d日"),
                 org, args.get("author")):
        if text:
            _para(doc, text, size=11, color=DOCX_MUTED,
                  align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    doc.add_page_break()


def _toc(doc):
    _para(doc, "目次", size=15, bold=True, color=DOCX_NAVY, space_after=10)
    p = doc.add_paragraph()
    _field(p, r'TOC \o "1-2" \h \z \u')
    _para(doc, "※ 目次はWordで開いたあと、この部分を選んで F9 を押すと最新になります。",
          size=8.5, color=DOCX_MUTED, space_after=0)
    doc.add_page_break()


def _summary_box(doc, points: list):
    if not points:
        return
    _para(doc, "要点", size=12, bold=True, color=DOCX_HILITE, space_after=4)
    for s in points:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(str(s))
        _jp_font(r, size=11, bold=True, color=DOCX_INK)
    doc.add_paragraph()


def _docx_callout(doc, text: str, label="ポイント"):
    p = _para(doc, f"【{label}】{text}", size=10.5, color=DOCX_INK, space_after=10)
    _border(p, size=6, color="C55A11", where="left")
    _shade(p, CALLOUT_BG)
    p.paragraph_format.left_indent = Cm(0.4)
    p.paragraph_format.space_before = Pt(6)


def _table(doc, columns, rows, *, caption=None, number=None, note=None):
    limit = DOCX_MAX_TABLE_ROWS
    shown, cut = rows[:limit], max(0, len(rows) - limit)
    t = doc.add_table(rows=len(shown) + 1, cols=len(columns))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True

    for j, c in enumerate(columns):
        cell = t.cell(0, j)
        cell.text = ""
        _shade(cell, HEADER_BG)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(str(c))
        _jp_font(r, size=9.5, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    for i, row in enumerate(shown, start=1):
        for j in range(len(columns)):
            v = row[j] if j < len(row) else ""
            cell = t.cell(i, j)
            cell.text = ""
            if i % 2 == 0:
                _shade(cell, DOCX_BAND)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            if isinstance(v, (int, float)) and not isinstance(v, bool):
                p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r = p.add_run(_docx_fmt(v))
            _jp_font(r, size=9.5)

    tail = []
    if caption:
        cap = _para(doc, f"表{number} {caption}" if number else caption,
                    size=9, color=DOCX_MUTED, align=WD_ALIGN_PARAGRAPH.LEFT,
                    space_after=4)
        cap.paragraph_format.space_before = Pt(3)
    if cut:
        tail.append(f"全 {len(rows):,} 行のうち上位 {limit} 行を掲載")
    if note:
        tail.append(note)
    if tail:
        _para(doc, "　".join(tail), size=8.5, color=DOCX_MUTED, space_after=10)
    return t


def _image(doc, data: bytes, *, caption=None, number=None, width_cm=16.0):
    doc.add_picture(io.BytesIO(data), width=Cm(width_cm))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if caption:
        _para(doc, f"図{number} {caption}" if number else caption, size=9,
              color=DOCX_MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)


# =============================================================================
# 組み立て
# =============================================================================

def build_docx(sections: list[dict], *, title="レポート", subtitle="", summary=None,
          conclusion="", recommendations=None, caveats=None, footer="",
          org="", author="", toc=True, appendix=None) -> bytes:
    """セクションのリストから .docx のバイト列を作る。"""
    if not sections:
        raise DocxReportError("セクションが1つもありません。")

    doc = Document()
    _setup_styles(doc)
    _setup_page(doc, footer or config.REPORT_ORG or "")
    _cover(doc, {"title": title, "subtitle": subtitle, "org": org, "author": author})
    if toc:
        _toc(doc)

    if summary:
        doc.add_heading("要約", level=1)
        _summary_box(doc, summary)

    fig_no, tbl_no = 0, 0
    for i, s in enumerate(sections, 1):
        if not s.get("heading"):
            raise DocxReportError(f"{i}番目のセクションに heading がありません。")
        if s.get("page_break"):
            doc.add_page_break()
        doc.add_heading(s["heading"], level=int(s.get("level") or 1))
        if s.get("body"):
            _para(doc, s["body"])
        for b in (s.get("bullets") or []):
            text = b.get("text", "") if isinstance(b, dict) else str(b)
            level = int(b.get("level", 0)) if isinstance(b, dict) else 0
            p = doc.add_paragraph(style="List Bullet" if not level
                                  else "List Bullet 2")
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(text)
            _jp_font(r, size=10.5)
        if s.get("image"):
            fig_no += 1
            _image(doc, s["image"], caption=s.get("caption") or s["heading"],
                   number=fig_no)
        if s.get("table"):
            t = s["table"]
            if not t.get("columns"):
                raise DocxReportError(f"「{s['heading']}」の表に columns がありません。")
            tbl_no += 1
            _table(doc, t["columns"], t.get("rows") or [],
                   caption=s.get("table_caption") or s.get("caption") or s["heading"],
                   number=tbl_no, note=t.get("note"))
        if s.get("callout"):
            _docx_callout(doc, s["callout"])
        if s.get("note"):
            p = _para(doc, s["note"], size=10, color=DOCX_MUTED, space_after=10)
            _border(p, size=4, color="D5DBE2", where="left")
            p.paragraph_format.left_indent = Cm(0.4)

    if conclusion:
        doc.add_heading("結論", level=1)
        _para(doc, conclusion)
    if recommendations:
        doc.add_heading("推奨する打ち手", level=1)
        for i, r in enumerate(recommendations, 1):
            if isinstance(r, dict):
                text = r.get("text", "")
                extra = "　".join(x for x in
                                  (f"担当: {r['owner']}" if r.get("owner") else "",
                                   f"期限: {r['due']}" if r.get("due") else "") if x)
            else:
                text, extra = str(r), ""
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(f"{i}. ")
            _jp_font(run, size=10.5, bold=True, color=DOCX_HILITE)
            run = p.add_run(text)
            _jp_font(run, size=10.5, bold=True)
            if extra:
                run = p.add_run(f"（{extra}）")
                _jp_font(run, size=9.5, color=DOCX_MUTED)
    if caveats:
        doc.add_heading("前提・注意", level=1)
        for c in caveats:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(str(c))
            _jp_font(r, size=10, color=DOCX_MUTED)

    for extra in (appendix or []):
        doc.add_page_break()
        doc.add_heading(extra.get("heading", "付録"), level=1)
        if extra.get("body"):
            _para(doc, extra["body"])
        if extra.get("table"):
            tbl_no += 1
            _table(doc, extra["table"]["columns"], extra["table"].get("rows") or [],
                   caption=extra.get("caption"), number=tbl_no)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def docx_safe_filename(name: str | None, default: str = "report") -> str:
    base = re.sub(r'[\\/:*?"<>|]', "_", str(name or default)).strip() or default
    return base if base.lower().endswith(".docx") else base + ".docx"


def outline_docx(sections: list[dict]) -> list[str]:
    out = []
    for i, s in enumerate(sections, 1):
        bits = []
        if s.get("image"):
            bits.append("図")
        if s.get("table"):
            bits.append(f"表{len(s['table'].get('rows') or [])}行")
        if s.get("bullets"):
            bits.append(f"箇条書き{len(s['bullets'])}件")
        out.append(f"{i}. {s.get('heading', '')}"
                   + (f"（{'・'.join(bits)}）" if bits else ""))
    return out


# ==========================================================================
# ===== 元 pptx_report.py
# PowerPointレポートの生成。会議でそのまま映せる体裁で作る。
#
# 作りの方針:
#   - 1スライド1メッセージ。上部の「キーメッセージ」に結論を1行で書き、
#     図表はその根拠として下に置く。読み手は上の1行だけで用が足りる。
#   - 日本語フォントを明示的に指定する。指定しないと英字フォントが当たり、
#     開いた瞬間に「ちゃんとしていない資料」に見える。
#   - グラフはPowerPointネイティブ（編集可）を既定にし、
#     ネイティブで表現できない種類だけ画像として貼る。
#
# 1スライド = 1つの dict。kind で中身が決まる:
#     title    表紙
#     agenda   目次
#     section  中扉
#     message  文字だけ（結論・考察）
#     table    表
#     chart    グラフ
#     kpi      数字を大きく並べる
#     compare  2つ並べて比較
#     closing  まとめ／次のアクション
# ==========================================================================
import io
import re
from datetime import datetime

from pptx import Presentation
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.dml.color import RGBColor as _pptx_RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn as _pptx_qn
from pptx.util import Emu, Inches, Pt as _pptx_Pt

import config

SLIDE_W = Inches(13.333)          # 16:9
SLIDE_H = Inches(7.5)

PPTX_CHART_TYPES = {
    "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "bar_stacked": XL_CHART_TYPE.COLUMN_STACKED,
    "bar_percent": XL_CHART_TYPE.COLUMN_STACKED_100,
    "hbar": XL_CHART_TYPE.BAR_CLUSTERED,
    "hbar_stacked": XL_CHART_TYPE.BAR_STACKED,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "area": XL_CHART_TYPE.AREA,
    "area_stacked": XL_CHART_TYPE.AREA_STACKED,
    "pie": XL_CHART_TYPE.PIE,
    "doughnut": XL_CHART_TYPE.DOUGHNUT,
    "scatter": XL_CHART_TYPE.XY_SCATTER,
    "radar": XL_CHART_TYPE.RADAR_MARKERS,
}
SLIDE_KINDS = ("title", "agenda", "section", "message", "table", "chart",
               "kpi", "compare", "closing")

# 配色。1枚に何色も出さない。強調は1色だけ使う。
PPTX_NAVY = _pptx_RGBColor(0x1F, 0x3B, 0x5C)
PPTX_ACCENT = _pptx_RGBColor(0x2E, 0x75, 0xB6)
PPTX_HILITE = _pptx_RGBColor(0xC5, 0x5A, 0x11)
PPTX_INK = _pptx_RGBColor(0x22, 0x26, 0x2B)
PPTX_MUTED = _pptx_RGBColor(0x6B, 0x72, 0x80)
LINE = _pptx_RGBColor(0xD5, 0xDB, 0xE2)
PPTX_BAND = _pptx_RGBColor(0xF5, 0xF8, 0xFC)
WHITE = _pptx_RGBColor(0xFF, 0xFF, 0xFF)
GOOD = _pptx_RGBColor(0x1E, 0x7A, 0x3C)
BAD = _pptx_RGBColor(0xB0, 0x2A, 0x2A)
SERIES = ["1F4E79", "F4B183", "70AD47", "C55A11", "7F7F7F",
          "2E75B6", "A9D18E", "FFD966", "9DC3E6", "BFBFBF"]

PPTX_MAX_TABLE_ROWS = 12               # これを超えると字が小さくなって読めない
MAX_CATEGORIES = 24

# 余白（本文の左右端）
MARGIN = Inches(0.62)
BODY_W = SLIDE_W - MARGIN * 2


class PptxReportError(Exception):
    """レポートを作れない理由（そのまま画面に出す）。"""


# =============================================================================
# 文字まわり
# =============================================================================

def _jp(run):
    """日本語フォントを当てる。

    python-pptx は latin フォントしか設定しないので、
    日本語部分に別のフォントが当たってしまう。東アジア用を直接書く。
    """
    run.font.name = config.REPORT_FONT_JA
    rPr = run.font._element          # これ自体が rPr（文字の書式）
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(_pptx_qn(tag))
        if el is None:
            el = rPr.makeelement(_pptx_qn(tag), {})
            rPr.append(el)
        el.set("typeface", config.REPORT_FONT_JA)


def _text(frame, lines, *, size=18, bold=False, color=PPTX_INK, align=PP_ALIGN.LEFT,
          space_after=4, line_spacing=1.25):
    """text_frame に段落を流し込む。lines は文字列か (文字列, 上書き) の並び。"""
    frame.word_wrap = True
    items = lines if isinstance(lines, (list, tuple)) else [lines]
    first = True
    for item in items:
        opts = {}
        if isinstance(item, tuple):
            item, opts = item
        for line in str(item).split("\n"):
            p = frame.paragraphs[0] if first else frame.add_paragraph()
            first = False
            p.text = line
            p.alignment = opts.get("align", align)
            p.space_after = _pptx_Pt(opts.get("space_after", space_after))
            p.line_spacing = opts.get("line_spacing", line_spacing)
            if opts.get("level"):
                p.level = opts["level"]
            for run in p.runs:
                run.font.size = _pptx_Pt(opts.get("size", size))
                run.font.bold = opts.get("bold", bold)
                run.font.color.rgb = opts.get("color", color)
                _jp(run)


def _pptx_box(slide, left, top, width, height, lines, **kw):
    shape = slide.shapes.add_textbox(left, top, width, height)
    _text(shape.text_frame, lines, **kw)
    return shape


def _rect(slide, left, top, width, height, fill=None, line=None,
          shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, left, top, width, height)
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = _pptx_Pt(1)
    s.shadow.inherit = False
    return s


def _pptx_clean(v) -> str:
    return "" if v is None else str(v)


def _num(v):
    if v is None:
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[,\s¥$%]", "", str(v))
    try:
        return float(s)
    except ValueError:
        return None


def _pptx_fmt(v) -> str:
    if isinstance(v, bool):
        return "はい" if v else "いいえ"
    if isinstance(v, float):
        if v == int(v) and abs(v) < 1e15:
            return f"{int(v):,}"
        # 12.30 ではなく 12.3 と出す（末尾の0は読み手には意味が無い）
        return f"{v:,.2f}".rstrip("0").rstrip(".")
    if isinstance(v, int):
        return f"{v:,}"
    return _pptx_clean(v)


# =============================================================================
# 共通の枠（ヘッダ・キーメッセージ・フッタ）
# =============================================================================

def _pptx_blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _header(slide, title: str, message: str = "") -> Emu:
    """見出しと、その下のキーメッセージ帯。戻り値は本文を始めてよい上端。"""
    _pptx_box(slide, MARGIN, Inches(0.30), BODY_W, Inches(0.55), title,
         size=25, bold=True, color=PPTX_NAVY)
    _rect(slide, MARGIN, Inches(0.92), BODY_W, Emu(12700), fill=PPTX_NAVY)

    if not message:
        return Inches(1.18)
    # 結論を1行で。ここだけ読めば分かるようにする。
    bar = _rect(slide, MARGIN, Inches(1.06), BODY_W, Inches(0.62), fill=PPTX_BAND)
    bar.line.color.rgb = LINE
    bar.line.width = _pptx_Pt(0.75)
    tf = bar.text_frame
    tf.margin_left, tf.margin_right = Inches(0.16), Inches(0.16)
    tf.margin_top = tf.margin_bottom = Inches(0.04)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    _text(tf, message, size=15, bold=True, color=PPTX_NAVY, space_after=0)
    return Inches(1.86)


def _footer(slide, text: str, page: int | None = None):
    _rect(slide, MARGIN, SLIDE_H - Inches(0.52), BODY_W, Emu(9525), fill=LINE)
    if text:
        _pptx_box(slide, MARGIN, SLIDE_H - Inches(0.46), Inches(9), Inches(0.32),
             text, size=9.5, color=PPTX_MUTED)
    if page:
        _pptx_box(slide, SLIDE_W - MARGIN - Inches(0.8), SLIDE_H - Inches(0.46),
             Inches(0.8), Inches(0.32), str(page), size=9.5, color=PPTX_MUTED,
             align=PP_ALIGN.RIGHT)


def _notes(slide, text: str):
    if text:
        slide.notes_slide.notes_text_frame.text = str(text)


def _source(slide, top, text: str):
    """出典・条件。数字の資料には必ず要る。"""
    if text:
        _pptx_box(slide, MARGIN, top, BODY_W, Inches(0.3), f"出所: {text}",
             size=9.5, color=PPTX_MUTED)


# =============================================================================
# スライドの種類ごと
# =============================================================================

def _slide_title(prs, spec):
    slide = _pptx_blank(prs)
    _rect(slide, 0, 0, SLIDE_W, Inches(3.05), fill=PPTX_NAVY)
    _rect(slide, 0, Inches(3.05), SLIDE_W, Inches(0.06), fill=PPTX_HILITE)
    _pptx_box(slide, Inches(0.9), Inches(1.05), SLIDE_W - Inches(1.8), Inches(1.2),
         spec.get("title", "レポート"), size=40, bold=True, color=WHITE)
    if spec.get("subtitle"):
        _pptx_box(slide, Inches(0.9), Inches(2.25), SLIDE_W - Inches(1.8), Inches(0.5),
             spec["subtitle"], size=18, color=_pptx_RGBColor(0xC5, 0xD5, 0xE8))

    y = Inches(3.55)
    for line in (spec.get("lines") or [])[:4]:
        _rect(slide, Inches(0.9), y + Inches(0.10), Inches(0.09), Inches(0.22),
              fill=PPTX_HILITE)
        _pptx_box(slide, Inches(1.15), y, SLIDE_W - Inches(2.2), Inches(0.42), line,
             size=16, color=PPTX_INK)
        y += Inches(0.52)

    org = spec.get("org") or config.REPORT_ORG
    foot = " ／ ".join(x for x in (org, spec.get("author")) if x)
    _pptx_box(slide, Inches(0.9), SLIDE_H - Inches(1.05), Inches(8), Inches(0.34),
         spec.get("date") or datetime.now().strftime("%Y年%m月%d日"),
         size=13, color=PPTX_MUTED)
    if foot:
        _pptx_box(slide, Inches(0.9), SLIDE_H - Inches(0.72), Inches(8), Inches(0.34),
             foot, size=13, color=PPTX_MUTED)
    _notes(slide, spec.get("notes", ""))
    return slide


def _slide_agenda(prs, spec):
    slide = _pptx_blank(prs)
    top = _header(slide, spec.get("title") or "本日の内容", spec.get("message", ""))
    items = spec.get("items") or []
    y = top + Inches(0.18)
    step = min(Inches(0.72), (SLIDE_H - y - Inches(0.9)) / max(len(items), 1))
    for i, it in enumerate(items, 1):
        label = it.get("text") if isinstance(it, dict) else str(it)
        note = it.get("note", "") if isinstance(it, dict) else ""
        n = _rect(slide, MARGIN, y, Inches(0.44), Inches(0.44), fill=PPTX_NAVY,
                  shape=MSO_SHAPE.OVAL)
        tf = n.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        _text(tf, str(i), size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER,
              space_after=0)
        _pptx_box(slide, MARGIN + Inches(0.62), y + Inches(0.02), BODY_W - Inches(0.8),
             Inches(0.4), label, size=17, bold=True, color=PPTX_INK)
        if note:
            _pptx_box(slide, MARGIN + Inches(0.62), y + Inches(0.34),
                 BODY_W - Inches(0.8), Inches(0.3), note, size=12, color=PPTX_MUTED)
        y += step
    _notes(slide, spec.get("notes", ""))
    return slide


def _slide_section(prs, spec):
    slide = _pptx_blank(prs)
    _rect(slide, 0, 0, SLIDE_W, SLIDE_H, fill=PPTX_NAVY)
    _rect(slide, MARGIN, Inches(3.18), Inches(0.9), Inches(0.07), fill=PPTX_HILITE)
    _pptx_box(slide, MARGIN, Inches(3.42), SLIDE_W - MARGIN * 2, Inches(0.9),
         spec.get("title", ""), size=32, bold=True, color=WHITE)
    if spec.get("subtitle"):
        _pptx_box(slide, MARGIN, Inches(4.35), SLIDE_W - MARGIN * 2, Inches(0.5),
             spec["subtitle"], size=15, color=_pptx_RGBColor(0xC5, 0xD5, 0xE8))
    _notes(slide, spec.get("notes", ""))
    return slide


def _bullets(slide, left, top, width, height, items, *, size=16):
    """箇条書き。dict なら {text, level, strong} を見る。"""
    y = top
    for item in items:
        if isinstance(item, dict):
            text, level = item.get("text", ""), int(item.get("level", 0))
            strong = bool(item.get("strong"))
        else:
            text, level, strong = str(item), 0, False
        if y > top + height - Inches(0.3):
            break
        mark_x = left + Inches(0.26) * level
        if level == 0:
            _rect(slide, mark_x, y + Inches(0.13), Inches(0.10), Inches(0.10),
                  fill=PPTX_HILITE if strong else PPTX_ACCENT)
        else:
            _pptx_box(slide, mark_x, y - Inches(0.02), Inches(0.2), Inches(0.3), "－",
                 size=13, color=PPTX_MUTED)
        _pptx_box(slide, mark_x + Inches(0.24), y - Inches(0.05),
             width - Inches(0.26) * level - Inches(0.24), Inches(0.42), text,
             size=size - 1.5 * level, bold=strong,
             color=PPTX_HILITE if strong else PPTX_INK)
        # 行数ぶん送る（おおよそでよい。重ならなければ十分）
        per = max(1, int(width / Inches(0.135) / max(size - 1.5 * level, 1) * 1.9))
        lines = max(1, (len(text) + per - 1) // per)
        y += Inches(0.34) * lines + Inches(0.12)
    return y


def _slide_message(prs, spec):
    slide = _pptx_blank(prs)
    top = _header(slide, spec.get("title", ""), spec.get("message", ""))
    y = top + Inches(0.12)
    if spec.get("lead"):
        _pptx_box(slide, MARGIN, y, BODY_W, Inches(0.6), spec["lead"], size=17,
             color=PPTX_INK)
        y += Inches(0.72)
    items = spec.get("bullets") or []
    if items:
        y = _bullets(slide, MARGIN, y, BODY_W, SLIDE_H - y - Inches(1.2), items,
                     size=17)
    if spec.get("body"):
        _pptx_box(slide, MARGIN, y + Inches(0.1), BODY_W,
             SLIDE_H - y - Inches(1.1), spec["body"], size=14, color=PPTX_INK)
    if spec.get("callout"):
        _pptx_callout(slide, MARGIN, SLIDE_H - Inches(1.55), BODY_W, spec["callout"])
    _notes(slide, spec.get("notes", ""))
    return slide


def _pptx_callout(slide, left, top, width, text, label="ポイント"):
    """強調枠。1枚に1つだけ置く。"""
    box = _rect(slide, left, top, width, Inches(0.92), fill=_pptx_RGBColor(0xFD, 0xF3, 0xE7))
    box.line.color.rgb = PPTX_HILITE
    box.line.width = _pptx_Pt(1.25)
    _rect(slide, left, top, Inches(0.07), Inches(0.92), fill=PPTX_HILITE)
    _pptx_box(slide, left + Inches(0.22), top + Inches(0.09), Inches(2), Inches(0.26),
         label, size=10.5, bold=True, color=PPTX_HILITE)
    _pptx_box(slide, left + Inches(0.22), top + Inches(0.34), width - Inches(0.44),
         Inches(0.52), text, size=14, color=PPTX_INK)


def _slide_kpi(prs, spec):
    slide = _pptx_blank(prs)
    top = _header(slide, spec.get("title", ""), spec.get("message", ""))
    items = (spec.get("items") or [])[:4]
    if not items:
        raise PptxReportError("kpi スライドには items が必要です。")
    gap = Inches(0.32)
    width = (BODY_W - gap * (len(items) - 1)) / len(items)
    card_h = Inches(2.35)
    for i, it in enumerate(items):
        left = MARGIN + (width + gap) * i
        card = _rect(slide, left, top + Inches(0.25), width, card_h, fill=PPTX_BAND,
                     line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
        card.adjustments[0] = 0.04
        _rect(slide, left, top + Inches(0.25), width, Inches(0.07), fill=PPTX_ACCENT)
        _pptx_box(slide, left, top + Inches(0.48), width, Inches(0.34),
             _pptx_clean(it.get("label")), size=13, color=PPTX_MUTED, align=PP_ALIGN.CENTER)
        _pptx_box(slide, left, top + Inches(0.86), width, Inches(0.85),
             _pptx_fmt(it.get("value")) + _pptx_clean(it.get("unit")),
             size=34, bold=True, color=PPTX_NAVY, align=PP_ALIGN.CENTER)
        d = _num(it.get("delta"))
        if d is not None:
            mark = "▲" if d > 0 else ("▼" if d < 0 else "―")
            good = it.get("higher_is_better", True)
            col = PPTX_MUTED if d == 0 else (GOOD if (d > 0) == bool(good) else BAD)
            _pptx_box(slide, left, top + Inches(1.72), width, Inches(0.34),
                 f"{mark} {_pptx_fmt(abs(d))}{_pptx_clean(it.get('delta_unit'))}"
                 + (f"（{it['delta_label']}）" if it.get("delta_label") else ""),
                 size=13, bold=True, color=col, align=PP_ALIGN.CENTER)
        if it.get("note"):
            _pptx_box(slide, left, top + Inches(2.06), width, Inches(0.3),
                 it["note"], size=10.5, color=PPTX_MUTED, align=PP_ALIGN.CENTER)

    y = top + card_h + Inches(0.45)
    if spec.get("bullets"):
        y = _bullets(slide, MARGIN, y, BODY_W, SLIDE_H - y - Inches(1.0),
                     spec["bullets"], size=15)
    elif spec.get("comment"):
        _pptx_box(slide, MARGIN, y, BODY_W, Inches(1.0), spec["comment"], size=14)
    if spec.get("callout"):
        _pptx_callout(slide, MARGIN, SLIDE_H - Inches(1.5), BODY_W, spec["callout"])
    _source(slide, SLIDE_H - Inches(0.85), spec.get("source", ""))
    _notes(slide, spec.get("notes", ""))
    return slide


def _add_table(slide, left, top, width, height, columns, rows, *,
               font=11.5, highlight_rows=()):
    shape = slide.shapes.add_table(len(rows) + 1, len(columns), left, top,
                                   width, height)
    table = shape.table
    for j, c in enumerate(columns):
        cell = table.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = PPTX_NAVY
        cell.margin_left = cell.margin_right = Inches(0.07)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        _text(cell.text_frame, _pptx_clean(c), size=font, bold=True, color=WHITE,
              space_after=0, line_spacing=1.0)
    for i, row in enumerate(rows, start=1):
        for j in range(len(columns)):
            v = row[j] if j < len(row) else ""
            cell = table.cell(i, j)
            cell.margin_left = cell.margin_right = Inches(0.07)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = (
                _pptx_RGBColor(0xFD, 0xF3, 0xE7) if (i - 1) in highlight_rows
                else (PPTX_BAND if i % 2 == 0 else WHITE))
            _text(cell.text_frame, _pptx_fmt(v), size=font, space_after=0,
                  line_spacing=1.0,
                  align=PP_ALIGN.RIGHT if isinstance(v, (int, float))
                  and not isinstance(v, bool) else PP_ALIGN.LEFT,
                  bold=(i - 1) in highlight_rows)
    return table


def _slide_table(prs, spec):
    slide = _pptx_blank(prs)
    top = _header(slide, spec.get("title", ""), spec.get("message", ""))
    cols = [_pptx_clean(c) for c in (spec.get("columns") or [])]
    rows = spec.get("rows") or []
    if not cols:
        raise PptxReportError("table スライドには columns が必要です。")
    limit = int(spec.get("max_rows") or PPTX_MAX_TABLE_ROWS)
    shown, cut = rows[:limit], max(0, len(rows) - limit)

    comment = spec.get("comment")
    width = BODY_W if not comment else BODY_W - Inches(3.5)
    avail = SLIDE_H - top - Inches(1.05)
    height = min(Inches(0.36) * (len(shown) + 1), avail)
    _add_table(slide, MARGIN, top + Inches(0.1), width, height, cols, shown,
               font=11.5 if len(cols) <= 7 else 10,
               highlight_rows=set(spec.get("highlight_rows") or []))
    if comment:
        _pptx_box(slide, MARGIN + width + Inches(0.3), top + Inches(0.1),
             Inches(3.2), avail, comment, size=13.5)
    note = spec.get("source", "")
    if cut:
        note = (note + f"（全 {len(rows):,} 行のうち上位 {limit} 行）").strip()
    _source(slide, SLIDE_H - Inches(0.85), note)
    if spec.get("callout"):
        _pptx_callout(slide, MARGIN, SLIDE_H - Inches(1.5), BODY_W, spec["callout"])
    _notes(slide, spec.get("notes", ""))
    return slide


def _style_chart(chart, spec, series_count):
    chart.has_title = False
    chart.font.size = _pptx_Pt(12)
    chart.font.name = config.REPORT_FONT_JA
    kind = str(spec.get("chart") or "bar").lower()

    if series_count > 1 or kind in ("pie", "doughnut"):
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = _pptx_Pt(12)
    else:
        chart.has_legend = False

    try:
        for i, s in enumerate(chart.series):
            s.format.fill.solid()
            s.format.fill.fore_color.rgb = _pptx_RGBColor.from_string(
                SERIES[i % len(SERIES)])
            s.format.line.color.rgb = _pptx_RGBColor.from_string(SERIES[i % len(SERIES)])
    except (AttributeError, ValueError, NotImplementedError):
        pass

    # 数値軸は桁区切り。目盛り線は薄く。
    try:
        va = chart.value_axis
        va.has_major_gridlines = True
        va.major_gridlines.format.line.color.rgb = _pptx_RGBColor(0xE8, 0xE8, 0xE8)
        va.format.line.color.rgb = LINE
        va.tick_labels.number_format = spec.get("number_format") or "#,##0"
        va.tick_labels.number_format_is_linked = False
        va.tick_labels.font.size = _pptx_Pt(11.5)
    except (AttributeError, ValueError, NotImplementedError):
        pass
    try:
        ca = chart.category_axis
        ca.has_major_gridlines = False
        ca.format.line.color.rgb = LINE
        ca.tick_labels.font.size = _pptx_Pt(11.5)
    except (AttributeError, ValueError, NotImplementedError):
        pass

    # 値ラベル。多すぎると潰れるので、少ないときだけ。
    want = spec.get("data_labels")
    cats = len(spec.get("categories") or [])
    if want is None:
        want = kind in ("pie", "doughnut") or (cats and cats <= 8 and series_count <= 2)
    if want:
        try:
            plot = chart.plots[0]
            plot.has_data_labels = True
            dl = plot.data_labels
            dl.font.size = _pptx_Pt(11)
            dl.font.name = config.REPORT_FONT_JA
            if kind in ("pie", "doughnut"):
                dl.show_percentage = True
                dl.number_format = "0.0%"
                dl.number_format_is_linked = False
                dl.position = XL_LABEL_POSITION.OUTSIDE_END
            else:
                dl.number_format = spec.get("number_format") or "#,##0"
                dl.number_format_is_linked = False
                if kind in ("bar", "hbar"):
                    dl.position = XL_LABEL_POSITION.OUTSIDE_END
        except (AttributeError, ValueError, NotImplementedError):
            pass


def _slide_chart(prs, spec):
    slide = _pptx_blank(prs)
    top = _header(slide, spec.get("title", ""), spec.get("message", ""))
    series = spec.get("series") or []
    cats = [_pptx_clean(c) for c in (spec.get("categories") or [])]

    # 画像として渡された図（PowerPointで描けない種類）はそのまま貼る
    if spec.get("image"):
        avail_h = SLIDE_H - top - Inches(1.05)
        has_side = bool(spec.get("comment"))
        w = BODY_W if not has_side else BODY_W - Inches(3.5)
        slide.shapes.add_picture(io.BytesIO(spec["image"]), MARGIN,
                                 top + Inches(0.08), width=w)
        if has_side:
            _pptx_box(slide, MARGIN + w + Inches(0.3), top + Inches(0.1), Inches(3.2),
                 avail_h, spec["comment"], size=13.5)
        _source(slide, SLIDE_H - Inches(0.85), spec.get("source", ""))
        if spec.get("callout"):
            _pptx_callout(slide, MARGIN, SLIDE_H - Inches(1.5), BODY_W, spec["callout"])
        _notes(slide, spec.get("notes", ""))
        return slide

    if not series:
        raise PptxReportError("chart スライドには series か image が必要です。")
    kind = str(spec.get("chart") or "bar").lower()
    if kind not in PPTX_CHART_TYPES:
        raise PptxReportError(f"未対応のグラフ種類です: {kind}。"
                          f"使えるのは {', '.join(PPTX_CHART_TYPES)} です。")
    if len(cats) > MAX_CATEGORIES:
        cats = cats[:MAX_CATEGORIES]
        series = [{**s, "values": (s.get("values") or [])[:MAX_CATEGORIES]}
                  for s in series]
        spec = {**spec, "source": (spec.get("source", "")
                                   + f"（上位{MAX_CATEGORIES}件）").strip()}

    if kind == "scatter":
        data = XyChartData()
        for s in series:
            sd = data.add_series(_pptx_clean(s.get("name") or "系列"))
            for x, y in zip(s.get("x") or [], s.get("values") or s.get("y") or []):
                if _num(x) is not None and _num(y) is not None:
                    sd.add_data_point(_num(x), _num(y))
    else:
        data = CategoryChartData()
        data.categories = cats or [str(i + 1) for i in
                                   range(len(series[0].get("values") or []))]
        for s in series:
            data.add_series(_pptx_clean(s.get("name") or "系列"),
                            [_num(v) for v in (s.get("values") or [])],
                            number_format=spec.get("number_format") or "#,##0")

    has_side = bool(spec.get("comment"))
    width = BODY_W if not has_side else BODY_W - Inches(3.5)
    height = SLIDE_H - top - Inches(1.05)
    frame = slide.shapes.add_chart(PPTX_CHART_TYPES[kind], MARGIN, top + Inches(0.08),
                                   width, height, data)
    _style_chart(frame.chart, {**spec, "categories": cats}, len(series))

    if has_side:
        _pptx_box(slide, MARGIN + width + Inches(0.3), top + Inches(0.15), Inches(3.2),
             height - Inches(0.2), spec["comment"], size=13.5)
    _source(slide, SLIDE_H - Inches(0.85), spec.get("source", ""))
    if spec.get("callout"):
        _pptx_callout(slide, MARGIN, SLIDE_H - Inches(1.5), BODY_W, spec["callout"])
    _notes(slide, spec.get("notes", ""))
    return slide


def _slide_compare(prs, spec):
    """左右に並べて比べる（案A/案B、前年/今年 など）。"""
    slide = _pptx_blank(prs)
    top = _header(slide, spec.get("title", ""), spec.get("message", ""))
    panes = (spec.get("panes") or [])[:2]
    if len(panes) != 2:
        raise PptxReportError("compare スライドには panes を2つ指定してください。")
    gap = Inches(0.4)
    w = (BODY_W - gap) / 2
    h = SLIDE_H - top - Inches(1.05)
    for i, pane in enumerate(panes):
        left = MARGIN + (w + gap) * i
        head = _rect(slide, left, top + Inches(0.05), w, Inches(0.46),
                     fill=PPTX_NAVY if i == 0 else PPTX_ACCENT)
        head.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        _text(head.text_frame, _pptx_clean(pane.get("title")), size=15, bold=True,
              color=WHITE, align=PP_ALIGN.CENTER, space_after=0)
        y = top + Inches(0.62)
        if pane.get("value") is not None:
            _pptx_box(slide, left, y, w, Inches(0.8),
                 _pptx_fmt(pane["value"]) + _pptx_clean(pane.get("unit")),
                 size=30, bold=True, color=PPTX_NAVY, align=PP_ALIGN.CENTER)
            y += Inches(0.9)
        if pane.get("image"):
            slide.shapes.add_picture(io.BytesIO(pane["image"]), left, y, width=w)
            y += Inches(2.6)
        if pane.get("bullets"):
            _bullets(slide, left + Inches(0.05), y, w - Inches(0.1),
                     top + h - y, pane["bullets"], size=14)
    if spec.get("callout"):
        _pptx_callout(slide, MARGIN, SLIDE_H - Inches(1.5), BODY_W, spec["callout"])
    _source(slide, SLIDE_H - Inches(0.85), spec.get("source", ""))
    _notes(slide, spec.get("notes", ""))
    return slide


def _slide_closing(prs, spec):
    """まとめと次のアクション。担当と期限まで書けるようにする。"""
    slide = _pptx_blank(prs)
    top = _header(slide, spec.get("title") or "まとめと次のアクション",
                  spec.get("message", ""))
    y = top + Inches(0.1)
    if spec.get("summary"):
        _pptx_box(slide, MARGIN, y, BODY_W, Inches(0.3), "まとめ", size=12,
             bold=True, color=PPTX_MUTED)
        y = _bullets(slide, MARGIN, y + Inches(0.34), BODY_W, Inches(2.0),
                     spec["summary"], size=16) + Inches(0.15)

    actions = spec.get("actions") or []
    if actions:
        _pptx_box(slide, MARGIN, y, BODY_W, Inches(0.3), "次のアクション", size=12,
             bold=True, color=PPTX_MUTED)
        y += Inches(0.34)
        rows = []
        for a in actions:
            if isinstance(a, dict):
                rows.append([a.get("text", ""), a.get("owner", ""), a.get("due", "")])
            else:
                rows.append([str(a), "", ""])
        h = min(Inches(0.36) * (len(rows) + 1), SLIDE_H - y - Inches(0.8))
        _add_table(slide, MARGIN, y, BODY_W, h, ["やること", "担当", "期限"], rows,
                   font=12.5)
    _notes(slide, spec.get("notes", ""))
    return slide


_PPTX_BUILDERS = {"title": _slide_title, "agenda": _slide_agenda,
             "section": _slide_section, "message": _slide_message,
             "table": _slide_table, "chart": _slide_chart, "kpi": _slide_kpi,
             "compare": _slide_compare, "closing": _slide_closing,
             # 旧称
             "text": _slide_message}


# =============================================================================
# 組み立て
# =============================================================================

def build_pptx(slides: list[dict], title: str | None = None,
          subtitle: str | None = None, footer: str | None = None,
          agenda: bool = True) -> bytes:
    """スライド定義のリストから .pptx のバイト列を作る。"""
    if not slides:
        raise PptxReportError("スライドが1枚もありません。")
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, SLIDE_H

    specs = list(slides)
    if title and (specs[0].get("kind") or "").lower() != "title":
        specs.insert(0, {"kind": "title", "title": title, "subtitle": subtitle or ""})

    # 中扉があれば、その並びから目次を自動で作る
    if agenda and not any((s.get("kind") or "") == "agenda" for s in specs):
        sections = [s.get("title", "") for s in specs
                    if (s.get("kind") or "") == "section"]
        if len(sections) >= 2:
            at = 1 if (specs[0].get("kind") or "") == "title" else 0
            specs.insert(at, {"kind": "agenda", "title": "本日の内容",
                              "items": sections})

    page = 0
    for i, spec in enumerate(specs):
        kind = str(spec.get("kind") or "message").lower()
        if kind not in _PPTX_BUILDERS:
            raise PptxReportError(f"{i + 1}枚目: 未対応の種類です: {kind}。"
                              f"使えるのは {', '.join(SLIDE_KINDS)} です。")
        try:
            slide = _PPTX_BUILDERS[kind](prs, spec)
        except PptxReportError:
            raise
        except Exception as e:
            raise PptxReportError(f"{i + 1}枚目（{kind}）の作成に失敗しました: {e}") from e
        if kind not in ("title", "section"):
            page += 1
            _footer(slide, footer or config.REPORT_ORG or "", page)

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def pptx_safe_filename(name: str | None, default: str = "report") -> str:
    base = re.sub(r'[\\/:*?"<>|]', "_", str(name or default)).strip() or default
    return base if base.lower().endswith(".pptx") else base + ".pptx"


def outline_pptx(slides: list[dict]) -> list[str]:
    """何が入ったかの一覧（画面とLLMへの報告用）。"""
    labels = {"title": "表紙", "agenda": "目次", "section": "中扉",
              "message": "説明", "text": "説明", "table": "表", "chart": "グラフ",
              "kpi": "KPI", "compare": "比較", "closing": "まとめ"}
    out = []
    for i, s in enumerate(slides, 1):
        kind = str(s.get("kind") or "message").lower()
        extra = ""
        if kind == "chart":
            extra = (f"（{s.get('chart', 'bar')}・{len(s.get('series') or [])}系列）"
                     if not s.get("image") else "（画像）")
        elif kind == "table":
            extra = f"（{len(s.get('rows') or [])}行）"
        out.append(f"{i}. [{labels.get(kind, kind)}{extra}] {s.get('title', '')}")
    return out


# ==========================================================================
# ===== 元 business.py
# 業務でよく聞かれる分析。SQLでは書きにくく、pandas なら素直に書けるもの。
#
# advanced.py が統計の道具箱なのに対して、こちらは「現場の問い」に対応する。
#   期間比較    先月と比べてどうか。落ちた原因はどの区分か
#   ファネル    見積 → 受注 → 請求 → 入金 のどこで落ちているか
#   コホート    いつ始めた人が、どれだけ続いているか
#   併売        何と何が一緒に買われているか
#
# 戻り値の形は advanced.py と同じ {"title", "tables", "notes", "meta"}。
# 画面もLLMも同じ入れ物で受け取れるようにしてある。
# ==========================================================================
import numpy as np
import pandas as pd

# 表の作り方・数値の丸めは統計側と揃える（同じ見た目で出す）


def _business_numeric(df: pd.DataFrame, col: str) -> pd.Series:
    return pd.to_numeric(df[col], errors="coerce")


def _pct(a: float, b: float) -> float | None:
    """b に対する a の割合(%)。分母が0のときは出さない。"""
    return round(a / b * 100, 1) if b else None


def _delta_note(cur: float, prev: float, unit: str = "") -> str:
    diff = cur - prev
    rate = f"{diff / prev * 100:+.1f}%" if prev else "—"
    return f"{prev:,.4g}{unit} → {cur:,.4g}{unit}（{diff:+,.4g}{unit} / {rate}）"


# =============================================================================
# 期間比較（前月比・前年同月比）と寄与度分解
# =============================================================================

def compare_periods(columns: list, rows: list, period_col: str, value_col: str,
                    dimension_col: str | None = None,
                    current: str | None = None, previous: str | None = None,
                    qty_col: str | None = None, top: int = 15) -> dict:
    """2つの期間を比べ、差がどこから来たのかまで分解する。

    「先月と比べて売上が5%落ちた」で終わらせず、
    「どの区分が押し下げたのか」「数量が減ったのか単価が下がったのか」まで出す。
    period_col の値は文字列として比べるので、'2026-01' でも '2026年1月' でもよい。
    """
    df = _df(columns, rows)
    for c in (period_col, value_col):
        if c not in df.columns:
            raise AnalysisError(f"列が見つかりません: {c}"
                                f"（ある列: {', '.join(map(str, df.columns))}）")
    df[value_col] = _business_numeric(df, value_col)
    df = df.dropna(subset=[value_col])
    if df.empty:
        raise AnalysisError(f"{value_col} に数値がありません。")

    periods = [str(p) for p in sorted(df[period_col].astype(str).unique())]
    if len(periods) < 2 and not (current and previous):
        raise AnalysisError(
            f"比べるには期間が2つ以上必要です（いま {len(periods)} 個: {'、'.join(periods)}）。"
            "SQL側で2期間ぶんのデータを取ってください。")
    cur = str(current) if current else periods[-1]
    prev = str(previous) if previous else periods[-2]
    for p in (cur, prev):
        if p not in periods:
            raise AnalysisError(f"期間 '{p}' がデータにありません（ある期間: {'、'.join(periods)}）。")

    df["_p"] = df[period_col].astype(str)
    cur_df = df[df["_p"] == cur]
    prev_df = df[df["_p"] == prev]
    cur_total = float(cur_df[value_col].sum())
    prev_total = float(prev_df[value_col].sum())
    diff_total = cur_total - prev_total

    tables = [_advanced_table("全体", ["項目", prev, cur, "差分", "増減率(%)"],
                     [[value_col, round(prev_total, 4), round(cur_total, 4),
                       round(diff_total, 4), _pct(diff_total, abs(prev_total))]])]
    notes = [f"{value_col}: {_delta_note(cur_total, prev_total)}"]

    meta = {"current": cur, "previous": prev,
            "current_total": _clean(cur_total), "previous_total": _clean(prev_total)}

    if dimension_col and dimension_col in df.columns:
        a = prev_df.groupby(dimension_col)[value_col].sum()
        b = cur_df.groupby(dimension_col)[value_col].sum()
        seg = pd.DataFrame({prev: a, cur: b}).fillna(0.0)
        seg["差分"] = seg[cur] - seg[prev]
        seg["増減率(%)"] = np.where(seg[prev] != 0,
                                   (seg["差分"] / seg[prev].abs() * 100).round(1), np.nan)
        # 寄与度 = その区分の差分が、全体の差分のうち何割を占めるか
        seg["寄与度(%)"] = (seg["差分"] / abs(diff_total) * 100).round(1) if diff_total else np.nan
        seg = seg.sort_values("差分")
        show = pd.concat([seg.head(top), seg.tail(top)]).drop_duplicates()
        show = show.sort_values("差分", ascending=False).reset_index()
        cols, rws = _out(show.round(4))
        tables.append(_advanced_table(f"{dimension_col}別の内訳（増減の大きい順）", cols, rws))

        down = seg[seg["差分"] < 0].head(3)
        up = seg[seg["差分"] > 0].tail(3).iloc[::-1]
        if len(down):
            notes.append("押し下げた区分: " + "、".join(
                f"{i}（{r['差分']:+,.4g}"
                + (f" / 全体の変化の{abs(r['寄与度(%)']):.0f}%" if diff_total else "") + "）"
                for i, r in down.iterrows()))
        if len(up):
            notes.append("押し上げた区分: " + "、".join(
                f"{i}（{r['差分']:+,.4g}"
                + (f" / 全体の変化の{abs(r['寄与度(%)']):.0f}%" if diff_total else "") + "）"
                for i, r in up.iterrows()))
        # 新しく出てきた・消えた区分は、増減率だけ見ていると見落とす
        gone = [str(i) for i in seg.index[(seg[prev] > 0) & (seg[cur] == 0)]][:5]
        born = [str(i) for i in seg.index[(seg[prev] == 0) & (seg[cur] > 0)]][:5]
        if gone:
            notes.append(f"{cur} で無くなった{dimension_col}: {'、'.join(gone)}")
        if born:
            notes.append(f"{cur} で新たに出た{dimension_col}: {'、'.join(born)}")

    if qty_col and qty_col in df.columns:
        # 値の変化を「件数が動いたぶん」と「1件あたりが動いたぶん」に割る。
        # 「金額」「単価」と決め打つと、生産数量×工数のような組でも同じ見出しが出てしまう
        df[qty_col] = _business_numeric(df, qty_col)
        q0 = float(prev_df[qty_col].sum())
        q1 = float(cur_df[qty_col].sum())
        p0 = prev_total / q0 if q0 else 0.0
        p1 = cur_total / q1 if q1 else 0.0
        vol = (q1 - q0) * p0                      # 量の要因（1件あたりは前期のまま）
        price = (p1 - p0) * q1                    # 1件あたりの要因（量は当期）
        per = f"{value_col}/{qty_col}"
        tables.append(_advanced_table("増減の要因分解",
                             ["要因", value_col, "全体の変化に占める割合(%)"],
                             [[f"{qty_col} が変わったぶん", round(vol, 4), _pct(vol, abs(diff_total))],
                              [f"{per} が変わったぶん", round(price, 4), _pct(price, abs(diff_total))],
                              ["合計", round(vol + price, 4), None]]))
        notes.append(f"{qty_col} {q0:,.4g} → {q1:,.4g}、{per} の平均 {p0:,.4g} → {p1:,.4g}。"
                     f"変化の内訳は {qty_col} {vol:+,.4g}、{per} {price:+,.4g}。"
                     + (f"{qty_col} の影響が大きいです。" if abs(vol) > abs(price) else
                        f"{per} の影響が大きいです。"))
        meta.update({"volume_effect": _clean(vol), "price_effect": _clean(price)})

    notes.append(f"比較した期間: {prev} と {cur}。"
                 "期間の長さや営業日数が違うと単純比較はできません。"
                 "日数が違う場合は1日あたりに直して比べてください。")
    return {"title": f"{value_col} の期間比較（{prev} → {cur}）",
            "tables": tables, "notes": notes, "meta": meta}


# =============================================================================
# ファネル（段階ごとの通過と滞留）
# =============================================================================

def _passed(s: pd.Series) -> pd.Series:
    """その段階を通過したか。日付なら「入っていれば通過」、数値なら0より大。"""
    num = pd.to_numeric(s, errors="coerce")
    if num.notna().mean() >= 0.8:
        return num.fillna(0) > 0
    return s.notna() & (s.astype(str).str.strip() != "")


def funnel_analysis(columns: list, rows: list, steps: list,
                    labels: list | None = None, group_col: str | None = None,
                    date_steps: bool = True) -> dict:
    """段階ごとの通過数・転換率・離脱と、段階間の滞留日数を出す。

    1行 = 1案件。steps には段階を表す列を順に並べる
    （例: 見積日, 受注日, 請求日, 入金日）。値が入っていればその段階を通過した扱い。
    """
    df = _df(columns, rows)
    steps = [s for s in (steps or []) if s]
    if len(steps) < 2:
        raise AnalysisError("steps に段階の列を2つ以上、順番に並べて指定してください。")
    missing = [s for s in steps if s not in df.columns]
    if missing:
        raise AnalysisError(f"列が見つかりません: {', '.join(missing)}"
                            f"（ある列: {', '.join(map(str, df.columns))}）")
    names = list(labels or []) + steps[len(labels or []):]

    flags = pd.DataFrame({s: _passed(df[s]) for s in steps})
    total = len(df)
    counts = [int(flags[s].sum()) for s in steps]

    frows = []
    for i, s in enumerate(steps):
        prev = counts[i - 1] if i else counts[0]
        frows.append([names[i], counts[i],
                      _pct(counts[i], counts[0]),
                      _pct(counts[i], prev) if i else None,
                      (prev - counts[i]) if i else 0])
    tables = [_advanced_table("段階ごとの通過",
                     ["段階", "件数", "最初からの通過率(%)", "直前からの転換率(%)", "離脱数"],
                     frows)]

    notes = [f"対象 {total:,} 件。{names[0]} {counts[0]:,} 件から "
             f"{names[-1]} {counts[-1]:,} 件まで、"
             f"通過率は {_pct(counts[-1], counts[0])}% です。"]
    # いちばん漏れている段階を名指しする。ここが改善の的になる
    drops = [(counts[i - 1] - counts[i], i) for i in range(1, len(steps))]
    if drops:
        worst = max(drops)
        if worst[0] > 0:
            i = worst[1]
            notes.append(f"最も落ちているのは {names[i - 1]} → {names[i]} で、"
                         f"{worst[0]:,} 件（{100 - (_pct(counts[i], counts[i - 1]) or 0):.1f}%）が"
                         "先へ進んでいません。")

    # 段階間の滞留日数。日付として読めるときだけ出す
    if date_steps:
        lag_rows = []
        for i in range(1, len(steps)):
            a = pd.to_datetime(df[steps[i - 1]], errors="coerce")
            b = pd.to_datetime(df[steps[i]], errors="coerce")
            days = (b - a).dt.total_seconds() / 86400
            days = days[days.notna() & (days >= 0)]
            if len(days) >= 3:
                lag_rows.append([f"{names[i - 1]} → {names[i]}", len(days),
                                 round(float(days.mean()), 1),
                                 round(float(days.median()), 1),
                                 round(float(days.quantile(0.9)), 1)])
        if lag_rows:
            tables.append(_advanced_table("段階間の日数", ["区間", "件数", "平均", "中央値", "90%点"],
                                 lag_rows))
            slow = max(lag_rows, key=lambda r: r[3])
            notes.append(f"最も時間がかかるのは {slow[0]} で、中央値 {slow[3]} 日"
                         f"（1割は {slow[4]} 日以上）。")

    if group_col and group_col in df.columns:
        grows = []
        for name, sub in df.groupby(group_col):
            f = pd.DataFrame({s: _passed(sub[s]) for s in steps})
            first, last = int(f[steps[0]].sum()), int(f[steps[-1]].sum())
            grows.append([str(name), first, last, _pct(last, first)])
        grows.sort(key=lambda r: (r[3] is None, r[3]))
        tables.append(_advanced_table(f"{group_col}別の通過率",
                             [group_col, names[0], names[-1], "通過率(%)"], grows))
        if len(grows) >= 2 and grows[0][3] is not None and grows[-1][3] is not None:
            notes.append(f"通過率が最も低いのは {grows[0][0]}（{grows[0][3]}%）、"
                         f"最も高いのは {grows[-1][0]}（{grows[-1][3]}%）。")

    return {"title": "ファネル分析", "tables": tables, "notes": notes,
            "meta": {"steps": names, "counts": counts, "total": total}}


# =============================================================================
# コホート（いつ始めた人が、どれだけ続いているか）
# =============================================================================

def cohort_analysis(columns: list, rows: list, id_col: str, period_col: str,
                    value_col: str | None = None, max_periods: int = 12) -> dict:
    """初回の期でグループ分けし、その後どれだけ残っているかを見る。

    期の並びはデータに出てくる値の昇順で決める。'2026-01' でも '第1四半期' でも、
    並べたときに正しい順になっていれば動く。
    """
    df = _df(columns, rows)
    for c in (id_col, period_col):
        if c not in df.columns:
            raise AnalysisError(f"列が見つかりません: {c}"
                                f"（ある列: {', '.join(map(str, df.columns))}）")
    df = df[df[id_col].notna() & df[period_col].notna()].copy()
    if df.empty:
        raise AnalysisError("対象データがありません。")
    df["_p"] = df[period_col].astype(str)

    order = {p: i for i, p in enumerate(sorted(df["_p"].unique()))}
    if len(order) < 2:
        raise AnalysisError(f"期が1つしかありません（{list(order)}）。"
                            "複数の期にまたがるデータを取ってください。")
    df["_i"] = df["_p"].map(order)
    first = df.groupby(id_col)["_i"].min().rename("_c")
    df = df.join(first, on=id_col)
    df["経過"] = df["_i"] - df["_c"]
    max_periods = max(1, min(int(max_periods or 12), len(order)))
    df = df[df["経過"] < max_periods]

    rev = {i: p for p, i in order.items()}
    people = df.pivot_table(index="_c", columns="経過", values=id_col,
                            aggfunc="nunique", fill_value=0)
    size = people[0] if 0 in people.columns else people.max(axis=1)

    keep = people.div(size, axis=0).mul(100).round(1)
    keep.index = [f"{rev[i]}（{int(size[i])}件)" for i in keep.index]
    keep.columns = [f"+{c}期" for c in keep.columns]
    k = keep.reset_index().rename(columns={"index": "コホート", "_c": "コホート"})
    cols, rws = _out(k)
    tables = [_advanced_table("継続率(%)", cols, rws)]

    p = people.copy()
    p.index = [f"{rev[i]}" for i in p.index]
    p.columns = [f"+{c}期" for c in p.columns]
    c2, r2 = _out(p.reset_index().rename(columns={"index": "コホート", "_c": "コホート"}))
    tables.append(_advanced_table("人数", c2, r2))

    notes = []
    if len(keep.columns) > 1:
        avg = keep.iloc[:, 1].mean()
        notes.append(f"初回の次の期に残っているのは平均 {avg:.1f}% です。")
    if len(keep.columns) > 3:
        avg3 = keep.iloc[:, 3].mean()
        notes.append(f"3期あとに残っているのは平均 {avg3:.1f}%。"
                     + ("落ち方が急なので、初期の定着に手を打つ余地があります。"
                        if avg3 < 30 else "比較的よく定着しています。"))
    # 新しいコホートほど良くなっているか（施策の効果が出ているか）
    if len(keep) >= 3 and len(keep.columns) > 1:
        early, late = keep.iloc[0, 1], keep.iloc[-1, 1]
        if abs(early - late) >= 5:
            notes.append(f"最初のコホート {early:.1f}% に対し、直近は {late:.1f}%。"
                         + ("改善しています。" if late > early else
                            "悪化しています。対象の性質か初期の対応を確かめてください。"))

    if value_col and value_col in df.columns:
        df[value_col] = _business_numeric(df, value_col)
        amt = df.pivot_table(index="_c", columns="経過", values=value_col,
                             aggfunc="sum", fill_value=0).round(2)
        amt.index = [f"{rev[i]}" for i in amt.index]
        amt.columns = [f"+{c}期" for c in amt.columns]
        c3, r3 = _out(amt.reset_index().rename(columns={"index": "コホート", "_c": "コホート"}))
        tables.append(_advanced_table(f"{value_col}の合計", c3, r3))
        per = df.groupby("_c")[value_col].sum() / size
        notes.append("1件あたりの累計 " + value_col + ": " + "、".join(
            f"{rev[i]} {v:,.4g}" for i, v in per.items()))

    notes.append("直近のコホートは経過期間が短いぶん、右側のマスが空きます。"
                 "同じ経過期数どうし（縦ではなく列で）比べてください。")
    return {"title": "コホート分析（継続率）", "tables": tables, "notes": notes,
            "meta": {"cohorts": len(keep), "periods": len(keep.columns)}}


# =============================================================================
# 併売（何と何が一緒に買われているか）
# =============================================================================

def market_basket(columns: list, rows: list, transaction_col: str, item_col: str,
                  min_support: float = 1.0, top: int = 25,
                  max_items: int = 60) -> dict:
    """同じ伝票に一緒に入っている品目の組み合わせを見つける。

    リフト値は「たまたま一緒になる確率」に対して何倍かを表す。
    1.0 を大きく超える組み合わせが、置き場所や提案の手がかりになる。
    """
    df = _df(columns, rows)
    for c in (transaction_col, item_col):
        if c not in df.columns:
            raise AnalysisError(f"列が見つかりません: {c}"
                                f"（ある列: {', '.join(map(str, df.columns))}）")
    d = df[[transaction_col, item_col]].dropna().astype(str).drop_duplicates()
    n_tx = d[transaction_col].nunique()
    if n_tx < 10:
        raise AnalysisError(f"{transaction_col} が {n_tx} 件しかありません。10件以上必要です。")

    freq = d[item_col].value_counts()
    # 組み合わせの数は品目数の2乗で増える。よく出るものだけに絞って現実的な時間に収める
    keep = list(freq.head(max(2, int(max_items))).index)
    d = d[d[item_col].isin(keep)]
    cut = len(freq) - len(keep)

    baskets = d.groupby(transaction_col)[item_col].apply(set)
    baskets = baskets[baskets.map(len) >= 2]
    if baskets.empty:
        raise AnalysisError(f"2つ以上の{item_col}を含む{transaction_col}がありません。"
                            "1件1明細のデータになっていないか確認してください。")

    pair_count: dict = {}
    for items in baskets:
        picked = sorted(items)
        for i, a in enumerate(picked):
            for b in picked[i + 1:]:
                pair_count[(a, b)] = pair_count.get((a, b), 0) + 1

    out = []
    for (a, b), c in pair_count.items():
        support = c / n_tx * 100
        if support < float(min_support or 0):
            continue
        ca, cb = int(freq[a]), int(freq[b])
        conf_ab = c / ca * 100 if ca else 0.0
        conf_ba = c / cb * 100 if cb else 0.0
        lift = (c / n_tx) / ((ca / n_tx) * (cb / n_tx)) if ca and cb else 0.0
        out.append([a, b, c, round(support, 2), round(conf_ab, 1), round(conf_ba, 1),
                    round(lift, 2)])
    if not out:
        raise AnalysisError(f"支持度 {min_support}% 以上の組み合わせがありませんでした。"
                            "min_support を下げてください。")
    out.sort(key=lambda r: r[6], reverse=True)
    shown = out[: max(1, int(top))]

    # 「買った」と断定しない。同じ作業票で一緒に使われた部品なども対象になるため
    tables = [_advanced_table("よく一緒に出てくる組み合わせ",
                     ["対象A", "対象B", "同時件数", "支持度(%)",
                      "AならBも(%)", "BならAも(%)", "リフト"], shown)]
    tables.append(_advanced_table("よく出る対象", [item_col, "件数", "出現率(%)"],
                        [[i, int(c), round(c / n_tx * 100, 1)]
                         for i, c in freq.head(15).items()]))

    notes = [f"対象 {n_tx:,} 件の{transaction_col}、うち2つ以上を含むのは {len(baskets):,} 件です。"]
    if cut > 0:
        notes.append(f"{item_col} が多いため、出現の多い上位 {len(keep)} 件に絞って計算しました"
                     f"（{cut} 件を除外）。")
    best = shown[0]
    notes.append(f"最も結びつきが強いのは「{best[0]}」と「{best[1]}」で、リフト {best[6]}倍。"
                 f"「{best[0]}」を含む{transaction_col}の {best[4]}% に"
                 f"「{best[1]}」も含まれます。")
    notes.append("リフトは「たまたま一緒になる確率」に対する倍率です。1.0前後なら関係なし、"
                 "2.0を超えると強い結びつきと見ます。ただし件数が少ない組は偶然でも"
                 "大きな値になるので、同時件数も併せて見てください。")
    return {"title": "併売分析", "tables": tables, "notes": notes,
            "meta": {"transactions": int(n_tx), "pairs": len(out)}}


# ==========================================================================
# ===== 元 mailer.py
# メールの下書きと送信（SMTP）。
#
# 方針: 「作る」と「送る」を必ず分ける。
# LLMは compose（下書き作成）までしかできず、実際の送信は
# 画面でユーザーが本文と宛先を見て承認したときだけ実行される。
# 宛先を間違えた1通は取り消せないので、AIの判断だけでは外に出さない。
#
# 宛先はDBのテーブルから探す。人の情報がどのテーブルにあるかは
# DBごとに違うので、列名と実際の値（@を含むか等）から推測する。
# ==========================================================================
import mimetypes
import re
import smtplib
import ssl
import threading
from dataclasses import dataclass, field
from datetime import datetime
from email.header import Header
from email.message import EmailMessage
from email.utils import formataddr, formatdate, make_msgid, parseaddr

import yaml

import config
import db

# ざっくりだが実用上これで十分。厳密なRFC準拠より、明らかな入力ミスを弾く方が大事。
EMAIL_RE = re.compile(r"^[^@\s,;:<>\"]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}$")

# 宛先探しで手がかりにする列名（部分一致・大文字小文字は無視）
MAIL_HINTS = ("mail", "メール", "eメール", "address", "アドレス", "宛先", "email")
NAME_HINTS = ("name", "氏名", "名前", "担当", "社員名", "person", "user", "顧客名", "得意先")
DEPT_HINTS = ("部署", "部門", "所属", "課", "dept", "department", "division", "組織", "拠点")

_mailer_lock = threading.Lock()
_sent_log: list[dict] = []          # 直近の送信記録（画面表示用）
_MAX_LOG = 200


class MailError(Exception):
    """送信できない理由（そのまま画面に出す）。"""


# =============================================================================
# 設定
# =============================================================================

#: 管理者へ知らせる警告の種類（jobs.problems() の kind と対応）。
#: 画面のチェックボックスもこの順で並べる。
ALERT_KINDS = ("failed", "degraded", "overdue")
ALERT_KIND_LABELS = {
    "failed": "取り込みが失敗した",
    "degraded": "値の型がずれた",
    "overdue": "動いていない",
}


@dataclass
class SmtpSettings:
    host: str = ""
    port: int = 25
    security: str = "none"          # none / starttls / ssl
    user: str = ""
    password: str = ""
    sender: str = ""
    sender_name: str = ""
    timeout: int = 20
    allow_addresses: list = field(default_factory=list)
    senders: list = field(default_factory=list)      # 画面で選べる差出人の候補
    max_recipients: int = 20
    dry_run: bool = True
    alert_to: list = field(default_factory=list)     # 取り込みの警告を知らせる管理者
    # 警告メールを送るか。宛先を消さずに一時的に止められるようにしておく
    # （棚卸しや長期メンテのあいだだけ黙らせたい、という要望が普通にあるため）。
    alert_enabled: bool = True
    # どの警告を送るか。3種類は緊急度がまったく違うので、種類ごとに選べる。
    alert_kinds: list = field(default_factory=lambda: list(ALERT_KINDS))

    @property
    def configured(self) -> bool:
        return bool(self.host and self.sender)

    def problems(self) -> list[str]:
        out = []
        if not self.host:
            out.append("送信サーバのホスト名が未設定です。「メール設定」画面で登録してください。")
        if not self.sender:
            out.append("差出人アドレスが未設定です。「メール設定」画面で登録してください。")
        if not self.allow_addresses:
            out.append("送信できる宛先が1件も登録されていません。"
                       "登録するまでメールは送れません（「メール設定」画面で追加してください）。")
        elif not EMAIL_RE.match(self.sender):
            out.append(f"差出人アドレスの形式が正しくありません: {self.sender}")
        if self.security not in ("none", "starttls", "ssl"):
            out.append(f"暗号化は none / starttls / ssl のいずれかです: {self.security}")
        if self.user and not self.password:
            out.append("認証ユーザー名を設定したのに、パスワードが空です。")
        return out

    def allows(self, address: str) -> bool:
        """この宛先に送ってよいか。

        許可リストに載っているアドレスだけに送れる。**空のときは誰にも送れない。**
        「未設定なら全員に送れる」だと、設定を忘れたまま社外へ出てしまうため。
        """
        if not self.allow_addresses:
            return False
        return str(address).strip().lower() in [a.lower() for a in self.allow_addresses]

    @property
    def restricted(self) -> bool:
        """常に True。許可リスト方式なので、制限が外れることはない。"""
        return True


# =============================================================================
# 画面から変えられる設定（data/mail_settings.yaml）
#
# env の値を初期値として、このファイルの内容を上書きで重ねる。
# この設定ファイルは中身をそのまま画面に出すので、秘密は入れないこと。
# =============================================================================

# 送信サーバ（接続先）。暗号化と認証は社内リレー前提で画面に出さないので、
# 変えたい環境では env の SMTP_SECURITY / SMTP_USER / SMTP_PASSWORD を使う。
SERVER_KEYS = ("host", "port", "timeout")
# 差出人と宛先まわり
EDITABLE_KEYS = SERVER_KEYS + ("sender", "sender_name", "senders",
                               "allow_addresses", "max_recipients", "dry_run",
                               "alert_to", "alert_enabled", "alert_kinds",
                               "ok_domains")


def _read_overrides() -> dict:
    p = config.SMTP_SETTINGS_FILE
    if not p.exists():
        return {}
    try:
        data = yaml.safe_load(p.read_text(encoding="utf-8"))
    except Exception as e:
        print(f"[mailer] 設定を読めませんでした: {p} ({e})")
        return {}
    return {k: v for k, v in (data or {}).items() if k in EDITABLE_KEYS} \
        if isinstance(data, dict) else {}


def _write_overrides(data: dict) -> None:
    p = config.SMTP_SETTINGS_FILE
    p.parent.mkdir(parents=True, exist_ok=True)
    with _mailer_lock:
        p.write_text(yaml.safe_dump({k: data[k] for k in EDITABLE_KEYS if k in data},
                                    allow_unicode=True, sort_keys=False),
                     encoding="utf-8")


def settings() -> SmtpSettings:
    """env の値に、画面から保存した設定を重ねて返す。"""
    ov = _read_overrides()
    senders = [str(s).strip() for s in (ov.get("senders") or []) if str(s).strip()]
    sender = str(ov.get("sender") or config.SMTP_SENDER or "").strip()
    if not sender and senders:
        sender = senders[0]
    return SmtpSettings(
        host=str(ov.get("host", config.SMTP_HOST) or "").strip(),
        port=int(ov.get("port", config.SMTP_PORT) or 25),
        security=str(config.SMTP_SECURITY or "none").lower(),
        user=str(config.SMTP_USER or "").strip(),
        password=config.SMTP_PASSWORD,
        sender=sender,
        sender_name=str(ov.get("sender_name", config.SMTP_SENDER_NAME) or "").strip(),
        timeout=int(ov.get("timeout", config.SMTP_TIMEOUT) or 20),
        allow_addresses=[str(a).strip() for a in (ov.get("allow_addresses") or [])
                         if str(a).strip()],
        senders=senders,
        max_recipients=int(ov.get("max_recipients", config.SMTP_MAX_RECIPIENTS) or 20),
        dry_run=bool(ov["dry_run"]) if "dry_run" in ov else config.SMTP_DRY_RUN,
        alert_to=[str(a).strip() for a in (ov.get("alert_to") or []) if str(a).strip()],
        # 既定はON。以前の設定ファイルにキーが無い場合、今までどおり通知する
        alert_enabled=bool(ov["alert_enabled"]) if "alert_enabled" in ov else True,
        alert_kinds=([k for k in (ov.get("alert_kinds") or []) if k in ALERT_KINDS]
                     if "alert_kinds" in ov else list(ALERT_KINDS)),
    )


# --- 宛先に登録してよいドメイン（「メール設定」画面で変更。config.py は初期値）--------

def _parse_domains(value) -> list:
    """入力（"@a.co.jp; b.jp" のような文字列、またはリスト）を正規化して返す。"""
    if isinstance(value, str):
        parts = re.split(r"[;,、\s]+", value)
    else:
        parts = list(value or [])
    out = []
    for d in parts:
        d = str(d).strip().lstrip("@").rstrip(".").lower()
        if d and d not in out:
            out.append(d)
    return out


def send_ok_domains() -> list:
    """宛先に登録してよいドメイン。画面で保存した値 > config.py の初期値。"""
    ov = _read_overrides()
    if "ok_domains" in ov:
        return _parse_domains(ov.get("ok_domains"))
    return list(config.SEND_OK_MAIL_DOMAIN)


def domain_ok(address: str, allowed: list | None = None) -> bool:
    """このアドレスを許可リストに登録してよいか。

    許可ドメインが空なら制限なし（それでも許可リストへの登録自体は必要）。
    サブドメイン（sales.example.co.jp）も対象に含める。
    """
    if allowed is None:
        allowed = send_ok_domains()
    if not allowed:
        return True
    dom = str(address).strip().lower().rsplit("@", 1)[-1]
    return any(dom == d or dom.endswith("." + d) for d in allowed)


def allowed_domains_label(allowed: list | None = None) -> str:
    """画面に出す「登録できるドメイン」の表記。"""
    if allowed is None:
        allowed = send_ok_domains()
    return "、".join("@" + d for d in allowed) or "すべてのドメイン"


_HOSTNAME_RE = re.compile(r"^[A-Za-z0-9]([A-Za-z0-9.-]*[A-Za-z0-9])?$")


def _validate_server(data: dict) -> list[str]:
    """送信サーバ（接続情報）の点検。"""
    errors = []
    host = str(data.get("host") or "").strip()
    # 未入力は「まだメールを設定していない」状態として保存を許す。
    # ここで弾くと、SMTPを設定する前に警告メールのON/OFFすら保存できない
    # （足りない設定は problems() が画面上部に別途知らせる）。
    if host and not _HOSTNAME_RE.match(host):
        errors.append(f"ホスト名の形式が正しくありません: {host}")

    try:
        port = int(data.get("port") or 0)
    except (TypeError, ValueError):
        errors.append("ポート番号は数値で指定してください。")
    else:
        if not (1 <= port <= 65535):
            errors.append("ポート番号は 1〜65535 で指定してください。")

    try:
        t = int(data.get("timeout") or 0)
    except (TypeError, ValueError):
        errors.append("タイムアウトは数値で指定してください。")
    else:
        if not (1 <= t <= 300):
            errors.append("タイムアウトは 1〜300 秒で指定してください。")
    return errors


def validate_settings(data: dict) -> list[str]:
    """画面から来た設定の点検。1つでも返ったら保存しない。"""
    errors = _validate_server(data)

    # 許可ドメイン。宛先・通知先の検査は「これから保存する値」で行う
    doms = (_parse_domains(data.get("ok_domains")) if "ok_domains" in data
            else send_ok_domains())
    for d in doms:
        if not _HOSTNAME_RE.match(d) or "." not in d:
            errors.append(f"ドメインの形式が正しくありません: {d}（例: example.co.jp）")
    senders = [str(s).strip() for s in (data.get("senders") or []) if str(s).strip()]
    for s in senders:
        if not EMAIL_RE.match(s):
            errors.append(f"差出人アドレスの形式が正しくありません: {s}")
    sender = str(data.get("sender") or "").strip()
    # 差出人も未入力なら未設定として通す（ホスト名と同じ理由）
    if not sender:
        pass
    elif not EMAIL_RE.match(sender):
        errors.append(f"差出人アドレスの形式が正しくありません: {sender}")
    elif senders and sender not in senders:
        errors.append(f"{sender} は差出人の候補に入っていません。先に候補へ追加してください。")

    for a in (data.get("allow_addresses") or []):
        addr = str(a).strip()
        if not EMAIL_RE.match(addr):
            errors.append(f"宛先として登録できない形式です: {a}")
        elif not domain_ok(addr, doms):
            errors.append(f"{addr} は登録できません。"
                          f"登録できるのは {allowed_domains_label(doms)} のアドレスだけです"
                          "（「登録してよいドメイン」で変更できます）。")
    # 通知先の管理者も同じ縛り（許可ドメイン）。社外へは飛ばさない
    for a in (data.get("alert_to") or []):
        addr = str(a).strip()
        if not EMAIL_RE.match(addr):
            errors.append(f"通知先として登録できない形式です: {a}")
        elif not domain_ok(addr, doms):
            errors.append(f"{addr} は通知先に登録できません。"
                          f"登録できるのは {allowed_domains_label(doms)} のアドレスだけです。")

    try:
        n = int(data.get("max_recipients") or 0)
    except (TypeError, ValueError):
        errors.append("一度に送れる宛先数は数値で指定してください。")
    else:
        if not (1 <= n <= 500):
            errors.append("一度に送れる宛先数は 1〜500 で指定してください。")
    return errors


def _with_current(data: dict) -> dict:
    """省略されたキーを現在の値で埋める。

    画面は毎回すべて送ってくるが、一部だけ変えたい呼び出し方もできるように
    しておく。埋めずに検証すると「送っていない項目」で弾かれてしまう。
    """
    s = settings()
    merged = {"host": s.host, "port": s.port, "timeout": s.timeout,
              "sender": s.sender, "sender_name": s.sender_name,
              "senders": s.senders, "allow_addresses": s.allow_addresses,
              "max_recipients": s.max_recipients, "dry_run": s.dry_run,
              "alert_to": s.alert_to, "alert_enabled": s.alert_enabled,
              "alert_kinds": s.alert_kinds, "ok_domains": send_ok_domains()}
    # None は「指定なし」。空文字や空リストは「消したい」なので通す。
    merged.update({k: v for k, v in (data or {}).items()
                   if k in merged and v is not None})
    return merged


def save_settings(data: dict, user: str | None = None) -> SmtpSettings:
    """画面からの保存。検証してから書く。"""
    merged = _with_current(data)
    errors = validate_settings(merged)
    if errors:
        raise MailError(" / ".join(errors))
    keep = _read_overrides()
    keep.update({
        "host": str(merged["host"] or "").strip(),
        "port": int(merged["port"] or 25),
        "timeout": int(merged["timeout"] or 20),
        "sender": str(merged["sender"] or "").strip(),
        "sender_name": str(merged["sender_name"] or "").strip(),
        "senders": [str(s).strip() for s in (merged["senders"] or []) if str(s).strip()],
        "allow_addresses": [str(a).strip() for a in (merged["allow_addresses"] or [])
                            if str(a).strip()],
        "max_recipients": int(merged["max_recipients"] or 20),
        "dry_run": bool(merged["dry_run"]),
        "alert_to": [str(a).strip() for a in (merged.get("alert_to") or [])
                     if str(a).strip()],
        "alert_enabled": bool(merged.get("alert_enabled", True)),
        "alert_kinds": [k for k in (merged.get("alert_kinds") or []) if k in ALERT_KINDS],
        "ok_domains": _parse_domains(merged.get("ok_domains")),
    })
    _write_overrides(keep)
    print(f"[mailer] 設定を更新しました（{user or '不明'}）: "
          f"サーバ={keep['host']}:{keep['port']} / "
          f"差出人={keep['sender']} / 宛先制限="
          f"{len(keep['allow_addresses'])}アドレス / "
          f"テスト送信={keep['dry_run']}")
    return settings()


def mail_status() -> dict:
    """画面に渡す現在の設定。認証情報は含めない。"""
    s = settings()
    return {"configured": s.configured, "host": s.host, "port": s.port,
            "sender": s.sender,
            "sender_name": s.sender_name, "dry_run": s.dry_run,
            "senders": s.senders,
            "allow_addresses": s.allow_addresses,
            "alert_to": s.alert_to,
            "alert_enabled": s.alert_enabled, "alert_kinds": s.alert_kinds,
            "alert_kind_labels": dict(ALERT_KIND_LABELS),
            "restricted": s.restricted, "timeout": s.timeout,
            "allowed_domains": send_ok_domains(),
            "allowed_domains_label": allowed_domains_label(),
            "ok_domains_source": ("screen" if "ok_domains" in _read_overrides()
                                  else "config"),
            "max_recipients": s.max_recipients, "problems": s.problems(),
            "settings_file": str(config.SMTP_SETTINGS_FILE)}


# =============================================================================
# 宛先を探す
# =============================================================================

def _hit(name: str, hints) -> bool:
    low = str(name).lower()
    return any(h.lower() in low for h in hints)


def _is_identifier(name: str) -> bool:
    """その列が「番号・コードの類」か。画面に出す名前としては後回しにする。"""
    low = str(name).lower()
    return (low in ("id", "code", "no")
            or low.endswith(("_id", "_no", "_code", "番号", "コード")))


def _display_first(cols: list) -> list:
    """表示に使う候補を並べ替える。番号・コードの列は後ろへ。

    user_id は NAME_HINTS の "user" に当たってしまうが、宛先一覧に出したいのは
    氏名の方。dept_code と dept_name が両方ある表でも同じことが起きる。
    候補から外さないのは、番号しか持たない表では番号でも出せた方がよいため。
    """
    return sorted(cols, key=lambda c: 1 if _is_identifier(c) else 0)


def _looks_like_email_column(conn, table: str, column: str) -> bool:
    """列名で分からないときは、実際の値に @ が入っているかで判断する。"""
    try:
        cur = conn.execute(f'SELECT "{column}" FROM "{table}" '
                           f'WHERE "{column}" IS NOT NULL LIMIT 20')
        vals = [str(r[0]) for r in cur.fetchall()]
    except Exception:
        return False
    if not vals:
        return False
    return sum(1 for v in vals if "@" in v and "." in v.split("@")[-1]) >= max(1, len(vals) // 2)


def address_tables(scope: list[dict]) -> list[dict]:
    """選択中のDBから「人とメールアドレスが載っていそうな表」を探す。"""
    found = []
    for s in scope:
        try:
            conn = db.connect_ro(s["path"])
        except Exception:
            continue
        try:
            tables = [r[0] for r in conn.execute(
                "SELECT name FROM sqlite_master WHERE type='table' "
                "AND name NOT LIKE 'sqlite_%' ORDER BY name")]
            for t in tables:
                if s.get("tables") and t not in s["tables"]:
                    continue
                try:
                    cols = [r[1] for r in conn.execute(f'PRAGMA table_info("{t}")')]
                except Exception:
                    continue
                mail_cols = [c for c in cols if _hit(c, MAIL_HINTS)]
                if not mail_cols:
                    mail_cols = [c for c in cols if _looks_like_email_column(conn, t, c)]
                if not mail_cols:
                    continue
                found.append({
                    "alias": s["alias"], "table": t, "columns": cols,
                    "mail_columns": mail_cols,
                    "name_columns": _display_first([c for c in cols if _hit(c, NAME_HINTS)]),
                    "dept_columns": _display_first([c for c in cols if _hit(c, DEPT_HINTS)]),
                })
        finally:
            conn.close()
    return found


def find_recipients(scope: list[dict], query: str = "", *, limit: int = 50,
                    table: str | None = None) -> dict:
    """名前・部署・アドレスの断片から宛先候補を探す。

    どの表を見ればよいかは address_tables() が推測する。
    query が空なら、その表の先頭から候補を出す（一覧確認のため）。
    """
    sources = address_tables(scope)
    if table:
        sources = [s for s in sources
                   if s["table"] == table or f"{s['alias']}.{s['table']}" == table]
    if not sources:
        return {"ok": False, "candidates": [], "sources": [],
                "message": "メールアドレスが入っていそうな表が見つかりません。"
                           "宛先のアドレスを直接指定してください。"}

    conf = settings()
    q = str(query or "").strip()
    out, seen = [], set()
    for src in sources:
        cols = src["columns"]
        search_cols = list(dict.fromkeys(
            src["mail_columns"] + src["name_columns"] + src["dept_columns"]
            + [c for c in cols if c not in src["mail_columns"]]))[:12]
        where, params = "", {}
        if q:
            # 値はプレースホルダで渡す。列名は実在するものだけを使うので識別子として安全。
            where = " WHERE " + " OR ".join(f'CAST("{c}" AS TEXT) LIKE :q'
                                            for c in search_cols)
            params["q"] = f"%{q}%"
        sql = (f'SELECT {", ".join(chr(34) + c + chr(34) for c in cols)} '
               f'FROM "{src["table"]}"{where} LIMIT {int(limit)}')
        try:
            conn = db.connect_ro(next(s["path"] for s in scope
                                      if s["alias"] == src["alias"]))
        except Exception:
            continue
        try:
            rows = conn.execute(sql, params).fetchall()
        except Exception:
            rows = []
        finally:
            conn.close()

        for r in rows:
            rec = dict(zip(cols, r))
            mail = next((str(rec[c]).strip() for c in src["mail_columns"]
                         if rec.get(c) and "@" in str(rec[c])), "")
            if not mail or mail.lower() in seen:
                continue
            seen.add(mail.lower())
            out.append({
                "email": mail,
                "name": next((str(rec[c]) for c in src["name_columns"] if rec.get(c)), ""),
                "dept": next((str(rec[c]) for c in src["dept_columns"] if rec.get(c)), ""),
                "source": f"{src['alias']}.{src['table']}",
                "valid": bool(EMAIL_RE.match(mail)),
                # 許可リストの外なら、下書きを作る前に分かるようにしておく
                "allowed": conf.allows(mail),
                "row": {k: (str(v) if v is not None else "") for k, v in list(rec.items())[:8]},
            })
            if len(out) >= limit:
                break
        if len(out) >= limit:
            break

    msg = (f"{len(out)}件の宛先候補が見つかりました。" if out
           else f"「{q}」に一致する宛先が見つかりませんでした。"
                f"探した表: {', '.join(s['alias'] + '.' + s['table'] for s in sources)}")
    return {"ok": bool(out), "candidates": out, "message": msg,
            "sources": [{"table": f"{s['alias']}.{s['table']}",
                         "mail_columns": s["mail_columns"],
                         "name_columns": s["name_columns"],
                         "dept_columns": s["dept_columns"]} for s in sources]}


# =============================================================================
# 下書きの検証
# =============================================================================

def _norm_addresses(value) -> list[str]:
    if value is None:
        return []
    if isinstance(value, str):
        parts = re.split(r"[,;\n]", value)
    else:
        parts = list(value)
    out = []
    for p in parts:
        p = str(p).strip()
        if not p:
            continue
        name, addr = parseaddr(p)
        out.append(addr or p)
    return out


def validate_draft(draft: dict, *, system: bool = False) -> list[str]:
    """送る前の点検。1つでも返ったら送信できない。

    system=True は、アプリ自身が管理者へ送る通知（定期取り込みの失敗など）。
    宛先は「メール設定」の通知先（alert_to）そのものなので、利用者向けの
    許可リスト（allow_addresses）とは独立に通す。サーバ・差出人の設定は同じく必要。
    """
    s = settings()
    errors = [e for e in s.problems()
              if not (system and "送信できる宛先" in e)]
    to = _norm_addresses(draft.get("to"))
    cc = _norm_addresses(draft.get("cc"))
    bcc = _norm_addresses(draft.get("bcc"))
    if not to:
        errors.append("宛先(To)が空です。")
    for addr in to + cc + bcc:
        if not EMAIL_RE.match(addr):
            errors.append(f"アドレスの形式が正しくありません: {addr}")
    total = len(to) + len(cc) + len(bcc)
    if total > s.max_recipients:
        errors.append(f"宛先が {total} 件あります。一度に送れるのは "
                      f"{s.max_recipients} 件までです（SMTP_MAX_RECIPIENTS）。")
    if system:
        # 通知先として登録したアドレスにだけ送る
        ok = {a.lower() for a in s.alert_to}
        for addr in to + cc + bcc:
            if addr.lower() not in ok:
                errors.append(f"{addr} は通知先に登録されていません。")
    elif not s.allow_addresses:
        errors.append("送信できる宛先が1件も登録されていません。"
                      "「メール設定」画面で登録するまで、どこにも送信できません。")
    else:
        for addr in to + cc + bcc:
            if not s.allows(addr):
                errors.append(f"{addr} は送信が許可されていません"
                              f"（許可済みは {len(s.allow_addresses)}件）。"
                              "メール設定で追加してください。")
    if not str(draft.get("subject") or "").strip():
        errors.append("件名が空です。")
    if not str(draft.get("body") or "").strip():
        errors.append("本文が空です。")
    return errors


def build_message(draft: dict, attachments: list[dict] | None = None) -> EmailMessage:
    """EmailMessage を組み立てる（送信せずに中身を確認するのにも使う）。"""
    s = settings()
    msg = EmailMessage()
    msg["From"] = formataddr((str(Header(s.sender_name, "utf-8")), s.sender)) \
        if s.sender_name else s.sender
    msg["To"] = ", ".join(_norm_addresses(draft.get("to")))
    if _norm_addresses(draft.get("cc")):
        msg["Cc"] = ", ".join(_norm_addresses(draft.get("cc")))
    if draft.get("reply_to"):
        msg["Reply-To"] = str(draft["reply_to"])
    msg["Subject"] = str(draft.get("subject") or "")
    msg["Date"] = formatdate(localtime=True)
    msg["Message-ID"] = make_msgid()
    msg.set_content(str(draft.get("body") or ""))

    for a in (attachments or []):
        data, name = a.get("data"), a.get("filename") or "attachment"
        if not data:
            continue
        mime = a.get("mime") or mimetypes.guess_type(name)[0] or "application/octet-stream"
        maintype, _, subtype = mime.partition("/")
        msg.add_attachment(data, maintype=maintype, subtype=subtype or "octet-stream",
                           filename=name)
    return msg


def preview(draft: dict, attachments: list[dict] | None = None) -> dict:
    """送信せずに、送られる内容をそのまま見せる。"""
    s = settings()
    to = _norm_addresses(draft.get("to"))
    cc = _norm_addresses(draft.get("cc"))
    bcc = _norm_addresses(draft.get("bcc"))
    body = str(draft.get("body") or "")
    return {
        "from": (f"{s.sender_name} <{s.sender}>" if s.sender_name else s.sender),
        "to": to, "cc": cc, "bcc": bcc,
        "subject": str(draft.get("subject") or ""),
        "body": body,
        "body_lines": len(body.splitlines()),
        "attachments": [{"filename": a.get("filename"),
                         "size": len(a.get("data") or b"")}
                        for a in (attachments or [])],
        "errors": validate_draft(draft),
        "dry_run": s.dry_run,
        "smtp": f"{s.host}:{s.port} ({s.security})",
    }


# =============================================================================
# 送信
# =============================================================================

def _connect(s: SmtpSettings):
    if s.security == "ssl":
        server = smtplib.SMTP_SSL(s.host, s.port, timeout=s.timeout,
                                  context=ssl.create_default_context())
    else:
        server = smtplib.SMTP(s.host, s.port, timeout=s.timeout)
        if s.security == "starttls":
            server.starttls(context=ssl.create_default_context())
    if s.user:
        server.login(s.user, s.password)
    return server


def send(draft: dict, attachments: list[dict] | None = None,
         user: str | None = None, *, system: bool = False) -> dict:
    """実際に送る。呼ぶ前に必ずユーザーの承認を取ること。

    SMTP_DRY_RUN=true のあいだは接続せず、組み立てた内容だけ返す
    （本番のSMTPを教えてもらう前に画面を試せるようにするため）。
    system=True はアプリ自身からの管理者通知（validate_draft 参照）。
    """
    errors = validate_draft(draft, system=system)
    if errors:
        raise MailError(" / ".join(errors))
    s = settings()
    msg = build_message(draft, attachments)
    recipients = (_norm_addresses(draft.get("to")) + _norm_addresses(draft.get("cc"))
                  + _norm_addresses(draft.get("bcc")))

    record = {"at": datetime.now().isoformat(timespec="seconds"),
              "to": _norm_addresses(draft.get("to")),
              "cc": _norm_addresses(draft.get("cc")),
              "bcc_count": len(_norm_addresses(draft.get("bcc"))),
              "subject": msg["Subject"], "user": user,
              "attachments": [a.get("filename") for a in (attachments or [])],
              "dry_run": s.dry_run, "ok": False, "message": ""}

    if s.dry_run:
        record.update(ok=True, message="下書きの確認のみ（SMTP_DRY_RUN=true のため送信していません）")
        # 送らない代わりに、文面をサーバのログに残す（通知の中身を SMTP 無しで確かめるため）
        body_lines = "\n".join("  | " + l for l in str(draft.get("body") or "").splitlines())
        print(f"[mailer] DRY RUN → {', '.join(record['to'])}\n件名: {msg['Subject']}\n{body_lines}",
              flush=True)
    else:
        try:
            server = _connect(s)
        except Exception as e:
            record["message"] = f"SMTPサーバに接続できません（{s.host}:{s.port}）: {e}"
            _mailer_log(record)
            raise MailError(record["message"]) from e
        try:
            server.send_message(msg, from_addr=s.sender, to_addrs=recipients)
            record.update(ok=True, message=f"{len(recipients)}件の宛先に送信しました。")
        except Exception as e:
            record["message"] = f"送信に失敗しました: {e}"
            _mailer_log(record)
            raise MailError(record["message"]) from e
        finally:
            try:
                server.quit()
            except Exception:
                pass
    _mailer_log(record)
    return record


def test_connection() -> dict:
    """設定の疎通確認だけ行う（メールは送らない）。"""
    s = settings()
    problems = s.problems()
    if problems:
        return {"ok": False, "message": " / ".join(problems)}
    try:
        server = _connect(s)
    except Exception as e:
        return {"ok": False, "message": f"接続できませんでした（{s.host}:{s.port}）: {e}"}
    try:
        server.noop()
        return {"ok": True, "message": f"{s.host}:{s.port} に接続できました"
                                       f"（{s.security}{'・認証あり' if s.user else ''}）。"}
    finally:
        try:
            server.quit()
        except Exception:
            pass


def _mailer_log(record: dict) -> None:
    with _mailer_lock:
        _sent_log.append(record)
        while len(_sent_log) > _MAX_LOG:
            _sent_log.pop(0)


def sent_log(limit: int = 50) -> list[dict]:
    with _mailer_lock:
        return list(reversed(_sent_log[-limit:]))


# =============================================================================
# 定期取り込みの失敗を管理者に知らせる
#
# 状態が「健全 → 失敗」に変わった瞬間に1回だけ送る。失敗が続くあいだ毎周期
# 送ると（15分ごとの設定なら1日96通）読まれなくなるので、直るまで黙る。
# 直ったら「復旧しました」を1回送る。宛先は「メール設定」の通知先（管理者）。
# =============================================================================

def alert_import_problems(current: list[dict], previous: list[dict]) -> dict | None:
    """定期取り込みの状態変化を管理者に送る。送らなかったときは None。

    current / previous は jobs.problems() の結果（今回と前回）。
    """
    s = settings()
    if not s.alert_to or not s.alert_enabled:
        return None
    # 知らせる種類だけを見る。選ばれていない種類は、起きても復旧しても黙っている
    # （「値の型ずれは要らないが失敗は知りたい」のような使い分けのため）。
    kinds = set(s.alert_kinds or ())
    if not kinds:
        return None
    current = [p for p in current if p.get("kind") in kinds]
    previous = [p for p in previous if p.get("kind") in kinds]
    now_ids = {p["id"]: p for p in current}
    prev_ids = {p["id"]: p for p in previous}
    newly = [now_ids[i] for i in now_ids if i not in prev_ids]
    fixed = [prev_ids[i] for i in prev_ids if i not in now_ids]
    if not newly and not fixed:
        return None

    lines = []
    if newly:
        lines.append("■ 設定どおりに更新できなくなった定期取り込み")
        for p in newly:
            tag = {"failed": "失敗", "degraded": "値の型がずれた", "overdue": "動いていない"}.get(p.get("kind"), "")
            lines.append(f"  ・{p['name']}（{p['db_file']} / {p['table']}）{'［' + tag + '］' if tag else ''}")
            lines.append(f"     {p['message']}")
        lines.append("")
        kinds = {p.get("kind") for p in newly}
        if "failed" in kinds:
            lines.append("  失敗: 取り込み元のファイル・シート名・列構成を確認してください。"
                         "失敗している間、そのテーブルは前回の内容のまま変わりません。")
        if "degraded" in kinds:
            lines.append("  値の型がずれた: 取り込みはできていますが、数値の列に文字が混ざったため"
                         "文字として保存しました。元ファイルの値を直して次回の実行を待ってください。")
        if "overdue" in kinds:
            lines.append("  動いていない: アプリが起動しているか、自動実行が止まっていないか確認してください。")
        lines.append("")
    if fixed:
        lines.append("■ 復旧した定期取り込み")
        for p in fixed:
            lines.append(f"  ・{p['name']}（{p['db_file']} / {p['table']}）")
        lines.append("")
    lines.append("確認: データカタログ > DB・テーブル > 各テーブルの「管理」")
    subject = ("[DB分析アシスタント] 定期取り込みが失敗しています"
               if newly else "[DB分析アシスタント] 定期取り込みが復旧しました")
    draft = {"to": list(s.alert_to), "subject": subject, "body": "\n".join(lines)}
    try:
        return send(draft, [], user="scheduler", system=True)
    except Exception as e:
        print(f"[mailer] 通知を送れませんでした: {e}")
        return None


# ==========================================================================
# ===== 元 custom_tools.py
# ユーザーがUIから定義するツール（SQLテンプレート型）。
#
# ツール1つは「名前 + 説明 + パラメータ定義 + SQLテンプレート + 出力形式」で表す。
# Pythonコードは書かせない。SQLは既存の SELECT専用ガード（db.run_select）を通し、
# パラメータは SQLite のバインド変数として渡すのでSQLインジェクションは起こらない。
#
# 保存先は各DBの .meta.yaml の `tools:`。どのDBに置かれていても全DB共通で使える
# （collect_everywhere が全DBから集める。置き場はSQLが主に見ているDBに自動で決まる）。
#
#   tools:
#     - name: monthly_sales
#       description: 指定年の月別売上を返す。「今年の売上推移」などで使う。
#       parameters:
#         - name: year
#           type: string
#           description: "対象年 'YYYY'"
#           required: true
#       sql: |
#         SELECT strftime('%Y-%m', o.order_date) AS 月, ... WHERE ... = :year ...
#       render: chart          # table | chart | chart_dual | none
#       chart: {chart_type: line, x: 月, y: 売上, title: 月別売上}
#       enabled: true
#
# 組み込みツールの有効/無効と説明文の上書きは .meta.yaml の `builtin_tools:` に持つ。
#
#   builtin_tools:
#     plot_dual_axis: {enabled: false}
#     run_sql_query: {description: "…独自の言い回しに差し替え…"}
# ==========================================================================
import re

RENDER_KINDS = ("table", "chart", "chart_dual", "excel", "csv", "none")
PARAM_TYPES = ("string", "integer", "number", "boolean")

# ツール名はOpenAIのfunction名の制約に合わせる（英数字とアンダースコア）
_NAME_RE = re.compile(r"^[A-Za-z][A-Za-z0-9_]{0,47}$")
# SQL中のバインド変数 :name を拾う（:: は型キャストなので除外）
_BIND_RE = re.compile(r"(?<!:):([A-Za-z_][A-Za-z0-9_]*)")

# 組み込みツール名（ユーザー定義ツールと衝突させない）。
# 以前はここに4つだけ列挙していたが、実際の組み込みは38個ある。
# 漏れた名前（forecast など）でツールを作れてしまい、AIに同じ名前の関数が
# 2つ渡って、しかも実行されるのは組み込み側だけ、という不整合が起きていた。
# 一覧は tools.schemas が持っているので、そちらから引く（循環importを避けて遅延）。
def builtin_names() -> set:
    from tools import BUILTIN_TOOLS
    return {t["function"]["name"] for t in BUILTIN_TOOLS}


def custom_tool_safe_name(candidate: str, taken=()) -> str:
    """どんな文字列からでも、使えるツール名を作る。

    AIが起こした名前や日本語の説明が元でも、function名の制約
    （英字始まり・英数字と_・48文字以内）に収め、既存とも組み込みとも
    衝突しない名前にする。ユーザーに名前で悩ませないための道具。
    """
    ascii_ = re.sub(r"[^a-z0-9_]+", "_", str(candidate or "").lower()).strip("_")
    base = ascii_[:40] if re.match(r"^[a-z]", ascii_) else ""
    if not base:
        base = "tool"
    used = {str(t).lower() for t in taken} | {n.lower() for n in builtin_names()}
    name, n = base, 2
    while name.lower() in used:
        name = f"{base}_{n}"
        n += 1
    return name


def bind_names(sql: str) -> list[str]:
    """SQLテンプレートに現れるバインド変数名（重複なし・出現順）。"""
    seen, out = set(), []
    for m in _BIND_RE.finditer(sql or ""):
        if m.group(1) not in seen:
            seen.add(m.group(1))
            out.append(m.group(1))
    return out


def validate_custom_tool(tool: dict, existing_names: set = frozenset()) -> list[str]:
    """ツール定義を検証し、問題点のリストを返す（空なら妥当）。"""
    errs: list[str] = []
    name = str(tool.get("name") or "").strip()
    if not name:
        errs.append("ツール名は必須です。")
    elif not _NAME_RE.match(name):
        errs.append("ツール名は英字で始まる英数字とアンダースコアのみ（48文字以内）にしてください。")
    elif name in builtin_names():
        errs.append(f"'{name}' は組み込みツールと同じ名前です。別の名前にしてください。")
    elif name in existing_names:
        errs.append(f"'{name}' は既に存在します。")

    if not str(tool.get("description") or "").strip():
        errs.append("説明は必須です（AIがこのツールを使うかどうかの判断材料になります）。")

    sql = str(tool.get("sql") or "").strip()
    if not sql:
        errs.append("SQLは必須です。")

    params = tool.get("parameters") or []
    pnames = []
    for i, p in enumerate(params, start=1):
        pn = str((p or {}).get("name") or "").strip()
        if not pn:
            errs.append(f"パラメータ{i}: 名前が空です。")
            continue
        if not re.match(r"^[A-Za-z_][A-Za-z0-9_]*$", pn):
            errs.append(f"パラメータ '{pn}': 英数字とアンダースコアのみ使えます。")
        if pn in pnames:
            errs.append(f"パラメータ '{pn}' が重複しています。")
        pnames.append(pn)
        if (p or {}).get("type") not in PARAM_TYPES:
            errs.append(f"パラメータ '{pn}': 型は {', '.join(PARAM_TYPES)} のいずれかにしてください。")

    # SQL中の :name と パラメータ定義の対応
    if sql:
        binds = set(bind_names(sql))
        for miss in sorted(binds - set(pnames)):
            errs.append(f"SQLに :{miss} がありますが、パラメータが定義されていません。")
        for unused in sorted(set(pnames) - binds):
            errs.append(f"パラメータ '{unused}' がSQL中で使われていません（:{unused} と書きます）。")

    render = tool.get("render") or "table"
    if render not in RENDER_KINDS:
        errs.append(f"出力形式は {', '.join(RENDER_KINDS)} のいずれかにしてください。")
    if render == "chart":
        import charts
        c = tool.get("chart") or {}
        ct = str(c.get("chart_type") or "").strip()
        if not ct:
            errs.append("グラフ出力には chart.chart_type が必要です。")
        elif ct not in charts.CHART_TYPES:
            errs.append(f"未対応のグラフ種別です: {ct}")
        else:
            for k in charts.required_fields(ct):
                v = c.get(k)
                if (not v) if k != "path" else (not list(v or [])):
                    errs.append(f"{ct} には chart.{k} が必要です。")
    if render == "chart_dual":
        c = tool.get("chart") or {}
        if not str(c.get("x") or "").strip():
            errs.append("2軸グラフには chart.x が必要です。")
        if not (c.get("bar_y") or []):
            errs.append("2軸グラフには chart.bar_y（棒にする列）が1つ以上必要です。")
        if not (c.get("line_y") or []):
            errs.append("2軸グラフには chart.line_y（折れ線にする列）が1つ以上必要です。")
    return errs


def to_schema(tool: dict) -> dict:
    """ユーザー定義ツール → OpenAI function calling のJSON Schema。"""
    props, required = {}, []
    for p in (tool.get("parameters") or []):
        pn = str(p.get("name") or "").strip()
        if not pn:
            continue
        props[pn] = {"type": p.get("type") or "string",
                     "description": str(p.get("description") or "")}
        if p.get("required", True):
            required.append(pn)
    return {
        "type": "function",
        "function": {
            "name": tool["name"],
            "description": str(tool.get("description") or ""),
            "parameters": {"type": "object", "properties": props, "required": required},
        },
    }


def coerce_params(tool: dict, args: dict) -> dict:
    """LLMが渡してきた引数を、定義された型に寄せてバインド用の辞書にする。"""
    out = {}
    for p in (tool.get("parameters") or []):
        pn = str(p.get("name") or "").strip()
        if not pn:
            continue
        v = args.get(pn)
        if v is None:
            out[pn] = None
            continue
        t = p.get("type") or "string"
        try:
            if t == "integer":
                out[pn] = int(v)
            elif t == "number":
                out[pn] = float(v)
            elif t == "boolean":
                out[pn] = 1 if (v is True or str(v).lower() in ("true", "1", "yes")) else 0
            else:
                out[pn] = str(v)
        except (TypeError, ValueError):
            raise ValueError(f"パラメータ '{pn}' を {t} として解釈できません: {v!r}")
    return out


def collect_everywhere(selected: list[dict] | None = None) -> list[dict]:
    """全DBのユーザー定義ツールを集める。置き場のDBを選んでいなくても拾う。

    ツールは作るときにDBを意識させない（SQLがどのDBに入るかはAIが決める）ので、
    「置き場のDBを選んでいないと存在しないことになる」のは作った人の意図と食い違う。
    組み込みツールと同じで、DBの選択に関係なく在ることにする。

    selected を渡すと、そのSQLが名指ししているDBが1つも選ばれていないツールは外す。
    いま見ている範囲と関係のないツールまで並べると、AIの選び分けが鈍るため。
    """
    import catalog                       # 循環importを避けるため、使うときに読む
    import db as dbmod

    picked = {str(s.get("name") or "") for s in (selected or [])}
    out, seen = [], set()
    for p in dbmod.list_db_files():
        alias = dbmod.alias_for(p)
        for t in (catalog.load_meta(p).get("tools") or []):
            if not isinstance(t, dict):
                continue
            name = str(t.get("name") or "").strip()
            if not name or name in seen or t.get("enabled") is False:
                continue
            if picked:
                needs = set(dbmod.dbs_named_in(str(t.get("sql") or "")))
                # どのDBも名指ししていないSQLは、置き場のDBのものとして扱う
                if not (needs or {p.name}) & picked:
                    continue
            seen.add(name)
            # owner_file は編集画面が保存先を知るためのもの（aliasはファイル名と別物）
            out.append({**t, "owner": alias, "owner_file": p.name})
    return out


def builtin_overrides(entries: list[dict]) -> dict:
    """組み込みツールの有効/無効・説明上書きを合成する。

    無効化はどれか1つのDBで無効なら無効（安全側）。説明は最初に見つかったものを採用。
    """
    merged: dict[str, dict] = {}
    for e in entries:
        for name, ov in (e.get("meta", {}).get("builtin_tools") or {}).items():
            if not isinstance(ov, dict):
                continue
            cur = merged.setdefault(name, {})
            if ov.get("enabled") is False:
                cur["enabled"] = False
            desc = str(ov.get("description") or "").strip()
            if desc and not cur.get("description"):
                cur["description"] = desc
    return merged


# ==========================================================================
# ===== 元 tools/results.py
# ツールが取ったデータを短いあいだ覚えておく置き場。
#
# 同じSQLを何度も流し直さないための仕組み。「集計 → グラフ → レポート」と
# 進むとき、以前は各ツールが自分でSQLを実行していたので、1つの問いに対して
# 同じSQLが3回走っていた。往復の上限（config.MAX_AGENT_STEPS）も、そのぶん
# 無駄に消える。結果に名前（result_id）を付けて返し、後続のツールはSQLの
# 代わりにその名前を指せるようにする。
#
# 副産物として、表とグラフが必ず同じデータを見ることになる。
# 実行し直す方式では、その間にデータが入れ替わると数字がずれ得た。
#
# 置き方の約束:
#   * プロセス内の辞書に持つ。ワーカーは1つで運用する前提（末尾の起動コード参照）。
#   * 古いものから捨てる。件数と総セル数の両方に上限を設ける。
#   * 取り出すときは、預けたときと同じDBの組み合わせかを確かめる。
#     IDを当てずっぽうで指されても、選んでいないDBの中身は出さない。
#   * 「同じ質問の中で、同じSQLを2回実行しない」ためにも使う。
#     LLMには result_id を渡すよう指示しているが、指示に従わず同じSQLを
#     書き直してくることがある（gpt-4o-mini で頻発）。指示は強制力を持たないので、
#     ここで機械的に止める。find_by_sql を参照。
# ==========================================================================
import uuid
from collections import OrderedDict

#: 覚えておく結果の数。会話1本で使う量に対して十分な余裕を見た数。
MAX_ENTRIES = 40
#: 総セル数の上限（行×列の合計）。これを超えたら古いものから捨てる。
MAX_CELLS = 400_000

_store: "OrderedDict[str, dict]" = OrderedDict()

#: いま処理している質問の識別子。SQLの使い回しは「同じ質問の中」だけに限る。
#: 質問をまたいで使い回すと、その間にデータが入れ替わっていても古い結果を返してしまう
#: （このアプリには質問直前のリアルタイム取り込みがある）。
#: リクエストを処理しているスレッドに置き、web 側が質問のたびに入れ直す。
_turn_local = threading.local()


def new_turn() -> str:
    """質問の開始。ここから先に預けた結果を「同じ質問のもの」として扱う。"""
    tid = "t_" + uuid.uuid4().hex[:8]
    _turn_local.turn = tid
    return tid


def current_turn() -> str:
    return getattr(_turn_local, "turn", "")


def set_turn(turn_id: str) -> None:
    """別のスレッドで続きを処理するときに、同じ質問だと伝える。

    ストリーミングは応答を流す側が別スレッドになる構成があり、
    そのままだとターンが引き継がれず、同じSQLを2回実行してしまう。
    """
    _turn_local.turn = turn_id


def normalize_sql(sql: str) -> str:
    """SQLを見比べるための形。空白と改行の違いだけは同じものとして扱う。

    大文字小文字は揃えない。'A装置' と 'a装置' のように、
    文字列リテラルの違いが意味を持つことがあるため。
    """
    return " ".join(str(sql or "").split())


def scope_key(scope: list[dict]) -> str:
    """どのDBの組み合わせで取ったデータかを表す文字列。"""
    return "|".join(sorted(str((s or {}).get("path") or "") for s in (scope or [])))


def _cells(entry: dict) -> int:
    return len(entry["rows"]) * max(1, len(entry["columns"]))


def _evict() -> None:
    """上限を超えたぶんを、古い順に捨てる。"""
    while len(_store) > MAX_ENTRIES:
        _store.popitem(last=False)
    total = sum(_cells(e) for e in _store.values())
    while total > MAX_CELLS and len(_store) > 1:
        _, old = _store.popitem(last=False)
        total -= _cells(old)


def put(scope: list[dict], columns: list, rows: list, truncated: bool = False,
        sql: str | None = None, label: str | None = None) -> str:
    """結果を預けて result_id を返す。"""
    rid = "r_" + uuid.uuid4().hex[:8]
    _store[rid] = {
        "scope": scope_key(scope),
        "columns": list(columns),
        "rows": [tuple(r) for r in rows],
        "truncated": bool(truncated),
        "sql": sql,
        "norm_sql": normalize_sql(sql) if sql else "",
        "turn": current_turn(),
        "label": label,
    }
    _evict()
    return rid


def get(scope: list[dict], rid: str) -> dict | None:
    """預けた結果を取り出す。無い・別のDBの組み合わせ、のときは None。"""
    entry = _store.get(str(rid or ""))
    if entry is None or entry["scope"] != scope_key(scope):
        return None
    _store.move_to_end(rid)          # 使ったものは新しい扱いにして残す
    return entry


def find_by_sql(scope: list[dict], sql: str) -> str | None:
    """同じ質問の中で、同じSQLを既に実行していれば、その result_id を返す。

    LLMが result_id を使わずSQLを書き直してきたときに、実行を1回で済ませるためのもの。
    同じ数字を2回取りに行かないので、表とグラフで値がずれる事故も起きない。

    条件は3つとも一致すること: 同じ質問・同じDBの組み合わせ・同じSQL。
    1つでも違えば None を返して、普通に実行させる。
    """
    turn = current_turn()
    if not turn:
        return None                       # 質問の外（起動時の点検など）では使い回さない
    want = normalize_sql(sql)
    if not want:
        return None
    key = scope_key(scope)
    for rid, entry in reversed(_store.items()):
        if (entry.get("turn") == turn and entry.get("scope") == key
                and entry.get("norm_sql") == want):
            return rid
    return None


def describe(rid: str) -> str:
    """LLMに返す一言。何のデータなのかを思い出せるようにする。"""
    entry = _store.get(str(rid or ""))
    if entry is None:
        return ""
    return entry.get("label") or (entry.get("sql") or "")[:80]


def clear() -> None:
    """テスト用。"""
    _store.clear()


# ==========================================================================
# ===== 元 tools/common.py
# ツールの実処理が共通で使う小道具。
#
# LLMへ返すJSONの組み立てと、データを用意して advanced.py に渡す定型。
# データは sql から取ることも、前のツールが返した result_id を指すこともできる
# （results.py 参照）。
# ==========================================================================
import json

import advanced
import config
import db


def _json(obj) -> str:
    return json.dumps(obj, ensure_ascii=False, default=str)


def _err(message: str) -> dict:
    return {
        "ok": False,
        "llm_content": _json({"error": message}),
        "render": {"role": "assistant", "kind": "error", "message": message},
    }


def _total_rows(sql: str, scope: list[dict]) -> int | None:
    """上限で切り詰められたとき、本当は何行あるのかを数える。

    「全部見た」と誤解したまま結論を書かせないための添え物。
    重いSQLだと数え直しも失敗し得るので、そのときは黙って諦める。
    """
    try:
        _, rows, _ = db.run_select(f"SELECT COUNT(*) FROM ({sql.strip().rstrip(';')})",
                                   scope, max_rows=1)
        return int(rows[0][0])
    except Exception:
        return None


def _inline_table(spec: dict, scope: list[dict], *, label: str | None = None,
                  max_rows: int | None = None):
    """引数で直接渡された表を、fetch と同じ形に整える。

    受け取れる形は2つ。どちらでもよいのは、LLMがどちらでも書いてくるため。
      rows: [["分類A", 47], ["分類B", 18]]        + columns: ["分類", "値"]
      rows: [{"分類": "分類A", "値": 47}, ...]    （columns は省略可）

    列名を省かれたら、辞書なら鍵から、配列なら「列1, 列2…」を当てる。
    ここで例外にすると、AIは何を直せばよいか分からないまま同じ呼び出しを繰り返す。
    """
    raw = (spec or {}).get("rows")
    if not isinstance(raw, list):
        raise advanced.AnalysisError(
            "rows は表の行を並べた配列で渡してください"
            "（例: [[\"分類A\", 47], [\"分類B\", 18]]）。")
    # columns が別の意味（pivot_table の「列に展開する1列名」＝文字列）の
    # ツールもある。文字列を1文字ずつ列名に散らさないよう、配列だけを使う。
    raw_cols = (spec or {}).get("columns")
    columns = ([str(c) for c in raw_cols]
               if isinstance(raw_cols, (list, tuple)) else [])

    # 辞書の配列なら、出てきた順に列名を集める（行ごとに鍵が欠けていてもよい）
    if raw and isinstance(raw[0], dict):
        if not columns:
            for r in raw:
                if isinstance(r, dict):
                    for k in r:
                        if str(k) not in columns:
                            columns.append(str(k))
        rows = [tuple((r or {}).get(c) for c in columns)
                for r in raw if isinstance(r, dict)]
    else:
        rows = []
        for r in raw:
            rows.append(tuple(r) if isinstance(r, (list, tuple)) else (r,))
        width = max((len(r) for r in rows), default=len(columns))
        if not columns:
            columns = [f"列{i + 1}" for i in range(width)]
        # 行の長さが揃っていなくても落とさない。足りない分は空で埋める
        rows = [r + (None,) * (len(columns) - len(r)) if len(r) < len(columns)
                else r[: len(columns)] for r in rows]

    if not columns:
        raise advanced.AnalysisError("rows が空です。中身のある表を渡してください。")

    cap = max_rows or config.MAX_RESULT_ROWS
    truncated = len(rows) > cap
    rows = rows[:cap]
    rid = results.put(scope, columns, rows[: config.MAX_RESULT_ROWS], truncated,
                      sql=None, label=label or "AIが組み立てた表")
    return columns, rows, truncated, rid, (len(raw) if truncated else None)


def fetch(spec: dict, scope: list[dict], *, label: str | None = None,
          max_rows: int | None = None):
    """ツールが使うデータを用意する。

    spec に result_id があれば前の結果を使い、無ければ sql を実行する。
    どちらの道でも result_id を返すので、呼び出し側はそれをLLMに伝えて
    次のツールで使い回せるようにする。

    max_rows は取得の上限。省略時は画面向けの MAX_RESULT_ROWS(2,000)。
    ファイル出力は EXPORT_MAX_ROWS を渡してくる。「表とグラフは2,000行で足りるが、
    CSVは全行欲しい」ので、道具ごとに上限が違う。

    戻り値: (columns, rows, truncated, result_id, total_rows)
      total_rows は切り詰めが起きたときだけ入る（本当の総件数）。
    """
    # 表そのものを渡された場合。DBを介さないデータをファイルやグラフにする道。
    # これが無いと、ナレッジ検索で調べた内容や、AIが自分で組み立てた表を
    # Excel/CSVにできない（出力系がすべてSQL前提になってしまう）。
    # ただし「空の rows」と sql / result_id が同時に来たら、SQLの方を採る。
    # LLMは引数の雛形ごと rows:[] を付けてくることがあり、それを0行の表として
    # 返すと、データがあるのに「無い」と答えてしまう（実測: COUNTが常に0行）。
    if (spec or {}).get("rows") is not None:
        has_query = (str((spec or {}).get("sql") or "").strip()
                     or str((spec or {}).get("result_id") or "").strip())
        if (spec or {}).get("rows") or not has_query:
            return _inline_table(spec, scope, label=label, max_rows=max_rows)

    rid = str((spec or {}).get("result_id") or "").strip()
    if not rid:
        # result_id を渡さず、同じSQLを書き直してきた場合。
        # 既に実行した結果があるなら、それを指したものとして扱う。
        # LLMへの指示（system prompt）だけでは守られないので、ここで機械的に止める。
        # 以降は result_id を渡されたときとまったく同じ道を通るので、
        # 「行数の上限が足りなければ取り直す」といった扱いもそのまま効く。
        rid = results.find_by_sql(scope, (spec or {}).get("sql") or "") or ""
        if rid:
            print(f"[results] 同じSQLなので実行し直しません（{rid} を使い回します）")

    if rid:
        entry = results.get(scope, rid)
        if entry is None:
            raise advanced.AnalysisError(
                f"result_id '{rid}' のデータが見つかりません。"
                "古くなって捨てられたか、別の会話の結果です。"
                "sql を指定して取り直してください。")
        # 預かっているのは2,000行に切り詰めた結果のことがある。ファイル出力のように
        # もっと大きな上限で全行欲しい呼び出しなら、預けたときのSQLで取り直す。
        # （「集計→CSVに」の流れで、CSVだけ2,000行で欠ける事故を防ぐ）
        if (max_rows and entry["truncated"] and entry.get("sql")
                and max_rows > len(entry["rows"])):
            sql = entry["sql"]
            wide = db.widen_scope(sql, scope)
            columns, rows, truncated = db.run_select(sql, wide, max_rows=max_rows)
            return columns, rows, truncated, rid, (
                _total_rows(sql, wide) if truncated else None)
        return entry["columns"], entry["rows"], entry["truncated"], rid, None

    sql = str((spec or {}).get("sql") or "").strip()
    if not sql:
        raise advanced.AnalysisError(
            "sql と result_id のどちらも指定されていません。"
            "新しくデータを取るなら sql を、前のツールの結果を使うなら result_id を指定してください。")
    # スコープは質問ごとの自動判定なので、例文由来のSQLなどが範囲外のDBを
    # 名指しすることがある。必要なぶんは繋いで実行する（読み取り専用のまま）。
    scope = db.widen_scope(sql, scope)
    columns, rows, truncated = db.run_select(sql, scope, max_rows=max_rows)
    # 使い回し用の預かりは、後続の表・グラフに足りる行数まで。
    # 100万行をそのまま預けると、これ1つで置き場(MAX_CELLS)を食い潰すため。
    keep = rows[: config.MAX_RESULT_ROWS]
    rid = results.put(scope, columns, keep, truncated or len(rows) > len(keep),
                      sql=sql, label=label)
    return columns, rows, truncated, rid, (_total_rows(sql, scope) if truncated else None)


def source_note(row_count: int, truncated: bool, total: int | None,
                cap: int | None = None) -> dict:
    """LLMに渡す「元データの規模」。切り詰めのときは実際の件数も添える。"""
    out = {"source_row_count": row_count, "source_truncated": bool(truncated)}
    if truncated:
        out["source_total_rows"] = total
        out["warning"] = (
            f"上限 {cap or config.MAX_RESULT_ROWS:,} 行で切り詰めました。"
            + (f"実際は {total:,} 行あります。" if total else "")
            + "この結果は全体の一部です。全体を語るなら、SQL側で"
              "GROUP BY で集計するか、条件を絞って取り直してください。")
    return out


def render_source_note(truncated: bool, total: int | None) -> dict:
    """画面用の「元データを切り詰めた」印。

    切り詰めの警告は長らく llm_content にしか入っておらず、AIには見えるのに
    画面には出なかった。表だけは truncated を渡していたので「（上限で切り詰め）」が
    出たが、グラフとレポートには渡っておらず、2,000行を超える明細から作った
    散布図やピボットが、断り書きなしに一部だけを描いていた。
    """
    if not truncated:
        return {}
    return {"truncated": True, "source_total_rows": total}


def _select_for(args: dict, scope: list[dict]):
    """分析ツール共通の入口。データを用意して (columns, rows) を返す。"""
    columns, rows, truncated, rid, total = fetch(args, scope)
    if not rows:
        raise advanced.AnalysisError("データが0行でした。抽出条件を見直してください。")
    return columns, rows, truncated, rid, total


def _report_result(res: dict, *, source_rows: int | None = None,
                   truncated: bool = False, total: int | None = None,
                   result_id: str | None = None,
                   scope: list[dict] | None = None,
                   extra: dict | None = None) -> dict:
    """advanced.py の戻り値を、画面用アイテムとLLM用の要約に変換する。

    LLMには表を丸ごと渡さない。所見(notes)と各表の先頭数行があれば
    十分に説明でき、トークンも節約できる。
    表そのものを次のツールへ渡せるよう、1つ目の表は result_id を付けて預ける。
    """
    tables = res.get("tables") or []
    llm_tables = []
    for i, t in enumerate(tables):
        rows = t.get("rows") or []
        item = {
            "name": t.get("name"), "columns": t.get("columns"),
            "row_count": len(rows),
            "rows": [list(r) for r in rows[: config.SAMPLE_ROWS_FOR_LLM]],
        }
        # 分析結果の表もグラフやレポートの材料になる。指せるようにしておく。
        if scope is not None and rows:
            item["result_id"] = results.put(scope, t.get("columns") or [], rows,
                                            label=f"{res.get('title')} / {t.get('name')}")
        llm_tables.append(item)
    payload = {"status": "analysis_ready", "title": res.get("title"),
               "notes": res.get("notes") or [], "tables": llm_tables,
               "meta": res.get("meta") or {}}
    if source_rows is not None:
        payload.update(source_note(source_rows, truncated, total))
    if result_id:
        payload["source_result_id"] = result_id
    payload.update(extra or {})
    return {
        "ok": True,
        "llm_content": _json(payload),
        "render": {"role": "assistant", "kind": "report", "title": res.get("title"),
                   "tables": tables, "notes": res.get("notes") or [],
                   **render_source_note(truncated, total)},
    }


def _analysis_tool(fn):
    """データを用意して advanced.py の関数に渡す、共通のかたち。"""
    def run(args: dict, scope: list[dict]) -> dict:
        try:
            columns, rows, truncated, rid, total = _select_for(args, scope)
            res = fn(args, columns, rows)
        except advanced.AnalysisError as e:
            return _err(str(e))
        except Exception as e:
            return _err(f"分析に失敗しました: {e}")
        if args.get("title"):
            res["title"] = args["title"]
        return _report_result(res, source_rows=len(rows), truncated=truncated,
                              total=total, result_id=rid, scope=scope)
    return run


# ==========================================================================
# ===== 元 tools/schemas.py
# LLMに渡すツールの定義（JSON Schema）。
#
# ここは「何ができるか」の宣言だけを書く場所で、処理は置かない。
# 実処理は query / stats / reports / mail にある。
# ==========================================================================
import advanced
import analysis
import charts
import excel
import exports
import pptx_report
import usage

#: 前のツールが返したデータを、SQLを書き直さずに使い回すための指定。
#: 「集計 → グラフ → レポート」で同じSQLが何度も走るのを避ける。
_RESULT_ID = {
    "type": "string",
    "description": "前のツールが返した result_id。これを指定すると sql は不要で、"
                   "同じデータをそのまま使う（同じSQLを書き直さないこと）。",
}

#: DBを介さないデータの渡し方。ナレッジ検索で調べた内容や、会話の中で
#: 分かったことを、あなた自身が表に組み立ててファイルやグラフにするための指定。
#: これが無いと、DBに無いデータは何も出力できない。
_INLINE_ROWS = {
    "type": "array",
    "description": "表の中身そのもの。sql も result_id も無いデータ"
                   "（社内文書を調べて分かったこと、会話で決まったこと など）を"
                   "ファイルやグラフにするときに使う。"
                   "配列の配列 [[\"分類A\", 47], [\"分類B\", 18]] か、"
                   "オブジェクトの配列 [{\"分類\": \"分類A\", \"値\": 47}] のどちらでもよい。"
                   "配列の配列で渡すときは columns も一緒に指定すること。",
    "items": {},
}
_INLINE_COLUMNS = {
    "type": "array",
    "items": {"type": "string"},
    "description": "rows の列名。rows を配列の配列で渡すときに指定する"
                   "（オブジェクトの配列なら省略してよい）。",
}


# 指定の名前 -> スキーマ（同じ説明を何度も書かないためのまとめ）
_CHART_ARGS = {
    "x": {"type": "string", "description": "横軸／カテゴリにする列名。"},
    "y": {"type": "string", "description": "値にする列名（数値）。"},
    "y2": {"type": "string", "description": "もう一方の値の列（dumbbell の比較先）。"},
    "z": {"type": "string", "description": "3つ目の数値の列（scatter3d の高さ）。"},
    "color": {"type": "string",
              "description": "色分けに使う列名。積み上げや群分けにも使う。"},
    "size": {"type": "string", "description": "大きさ／幅に使う数値列。"},
    "text": {"type": "string", "description": "点や棒に添えるラベルの列名。"},
    "facet": {"type": "string", "description": "この列の値ごとに小さく分割して並べる。"},
    "path": {"type": "array", "items": {"type": "string"},
             "description": "階層。大きい分類から順に列名を並べる。"},
    "dimensions": {"type": "array", "items": {"type": "string"},
                   "description": "対象にする列名のリスト（3〜6列が読みやすい）。"},
    "lower": {"type": "string", "description": "下限の列（信頼区間や予測の幅）。"},
    "upper": {"type": "string", "description": "上限の列。"},
    "source": {"type": "string", "description": "流れの起点になる列。"},
    "target": {"type": ["string", "number"],
               "description": "流れの終点の列。指標では目標値（数値そのもの、"
                              "または目標が入っている列名）。"},
    "start": {"type": "string", "description": "開始日時の列。"},
    "end": {"type": "string", "description": "終了日時の列。"},
    "open": {"type": "string", "description": "始値の列。"},
    "high": {"type": "string", "description": "高値の列。"},
    "low": {"type": "string", "description": "安値の列。"},
    "close": {"type": "string", "description": "終値の列。"},
    "value": {"type": "string", "description": "指標にする数値の列。"},
    "agg": {"type": "string", "enum": ["sum", "mean", "max", "min", "last"],
            "description": "value をどうまとめるか。既定は sum。"},
    "max": {"type": "number", "description": "ゲージの上限値。省略すると自動。"},
    "suffix": {"type": "string", "description": "数値の後ろに付ける単位（円・%など）。"},
    "nbins": {"type": "integer", "description": "階級の数。既定は自動。"},
    "orientation": {"type": "string", "enum": ["v", "h"],
                    "description": "棒の向き。横棒は h。"},
    "barmode": {"type": "string", "enum": ["group", "stack", "relative"],
                "description": "棒の積み方。既定は group。"},
    "marginal": {"type": "string", "enum": ["box", "violin", "rug"],
                 "description": "ヒストグラムの上に添える分布。任意。"},
    "trendline": {"type": "boolean", "description": "散布図に回帰直線を重ねる。"},
    "colorscale": {"type": "string",
                   "description": "色の濃淡（Blues / Reds / Greens など）。"},
}


# ツール名 -> (分類, 説明, 使う指定, 必須)
_CHART_TOOLS = {
    "plot_comparison": (
        "比較",
        "項目どうしを比べるグラフ。「部署別」「商品別」「順位」「ランキング」"
        "「前年と比べて」「重点管理」を見せたいときに使う。",
        ("x", "y", "y2", "color", "size", "text", "facet", "orientation", "barmode"),
        ("sql", "chart_type", "x", "y", "title")),
    "plot_trend": (
        "推移",
        "時間とともにどう変わったかを見せるグラフ。「推移」「時系列」「予測の幅」"
        "「工程の期間」「日ごとの多寡」「異常な回」を扱うときに使う。",
        ("x", "y", "color", "text", "lower", "upper", "start", "end",
         "open", "high", "low", "close", "facet"),
        ("sql", "chart_type", "title")),
    "plot_composition": (
        "構成",
        "全体が何でできているかを見せるグラフ。「内訳」「構成比」「シェア」"
        "「階層」「増減の要因」「どこからどこへ流れたか」を扱うときに使う。",
        ("x", "y", "color", "path", "source", "target", "text"),
        ("sql", "chart_type", "title")),
    "plot_distribution": (
        "分布",
        "ばらつきの形を見せるグラフ。「分布」「ヒストグラム」「箱ひげ」"
        "「偏り」「正規分布か」「群ごとの散らばり」を扱うときに使う。",
        ("x", "y", "color", "nbins", "facet", "marginal"),
        ("sql", "chart_type", "title")),
    "plot_relationship": (
        "関係",
        "2つ以上の項目の関係を見せるグラフ。「相関」「散布図」「密度」"
        "「多変量」「総当たり」「つながり」を扱うときに使う。",
        ("x", "y", "z", "color", "size", "text", "dimensions", "nbins",
         "source", "target", "colorscale", "trendline", "facet"),
        ("sql", "chart_type", "title")),
    "plot_kpi": (
        "指標",
        "数字を1つ大きく見せるグラフ。「KPI」「達成率」「目標に対して」"
        "「今いくら」を見せたいときに使う。",
        ("value", "target", "agg", "max", "suffix", "colorscale"),
        ("sql", "chart_type", "value", "title")),
}


def _chart_tools() -> list[dict]:
    """用途ごとのグラフツール定義を作る。"""
    out = []
    for name, (cat, desc, fields, required) in _CHART_TOOLS.items():
        props = {
            "sql": {"type": "string",
                    "description": "グラフに使うデータを取る SELECT 文。"
                                   "集計が要るものは GROUP BY 済みにすること。"},
            "chart_type": {"type": "string", "enum": charts.types_in(cat),
                           "description": f"グラフ種別。{charts.type_help(cat)}"},
            "title": {"type": "string", "description": "グラフのタイトル。"},
            "purpose": {"type": "string",
                        "description": "このグラフで示したいことの短い説明。"},
        }
        props.update({f: _CHART_ARGS[f] for f in fields})
        out.append({"type": "function", "function": {
            "name": name,
            "description": (f"{desc}使える種別: {charts.type_help(cat)}"),
            "parameters": {"type": "object", "properties": props,
                           "required": list(required)},
        }})
    return out


BUILTIN_TOOLS = [
    {
        "type": "function",
        "function": {
            "name": "run_sql_query",
            "description": (
                "対象のテーブルに対して読み取り専用の SELECT 文を実行し、結果テーブルを取得する。"
                "SELECT(または WITH ... SELECT)以外は実行不可。"
                ""
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "sql": {
                        "type": "string",
                        "description": "実行する SQLite 用 SELECT 文。SELECT または WITH で始めること。",
                    },
                    "purpose": {
                        "type": "string",
                        "description": "このクエリで何を確認したいかの短い説明(日本語)。",
                    },
                },
                "required": ["sql", "purpose"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "plot_chart",
            "description": (
                "SELECT 文の結果をグラフ化する。時系列の推移や分布を可視化したいときに使用。"
                "内部で SELECT を実行し、指定の x / y / color 列でグラフを描画する。"
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "sql": {"type": "string", "description": "グラフ用データを取得する SELECT 文(GROUP BYで集計済みにする)。"},
                    "chart_type": {
                        "type": "string",
                        "enum": list(charts.CHART_TYPES),
                        "description": "グラフ種別。" + charts.type_help(),
                    },
                    "x": {"type": "string",
                          "description": "x軸に使う列名。pie/donut はカテゴリ、funnel は段階、"
                                         "radar は指標名、histogram は対象の数値列。"},
                    "y": {"type": "string",
                          "description": "y軸に使う列名。pie/donut/treemap/sunburst/funnel は値、"
                                         "radar は値。histogram では不要。"},
                    "color": {"type": "string",
                              "description": "系列(色分け)に使う列名。任意。heatmap ではマス目の値になる。"},
                    "size": {"type": "string", "description": "bubble の大きさに使う数値列名。"},
                    "text": {"type": "string", "description": "点や棒に表示するラベルの列名。任意。"},
                    "path": {
                        "type": "array", "items": {"type": "string"},
                        "description": "treemap/sunburst の階層。大きい分類から順に列名を並べる。",
                    },
                    "nbins": {"type": "integer", "description": "histogram の階級数。任意。"},
                    "orientation": {
                        "type": "string", "enum": ["v", "h"],
                        "description": "棒の向き。横棒にしたいときは h。bar以外では無視。",
                    },
                    "barmode": {
                        "type": "string",
                        "enum": ["group", "stack", "relative"],
                        "description": "棒グラフの積み方。積み上げ=stack、横並び比較=group(既定)。bar以外では無視。",
                    },
                    "title": {"type": "string", "description": "グラフのタイトル。"},
                    "purpose": {"type": "string", "description": "このグラフで示したいことの短い説明。"},
                },
                "required": ["sql", "chart_type", "title"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "plot_dual_axis",
            "description": (
                "棒グラフ(左軸)と折れ線グラフ(右軸)を組み合わせた2軸グラフを描く。"
                "件数(棒)と比率など単位の異なる指標(折れ線)を同時に見せたいときに使用。"
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "sql": {"type": "string", "description": "x列と数値の指標列を返す SELECT 文(GROUP BYで集計)。"},
                    "x": {"type": "string", "description": "x軸に使う列名。"},
                    "bar_y": {
                        "type": "array", "items": {"type": "string"},
                        "description": "左軸に棒で表示する数値列名のリスト(1つ以上)。例: 件数。",
                    },
                    "line_y": {
                        "type": "array", "items": {"type": "string"},
                        "description": "右軸に折れ線で表示する数値列名のリスト(1つ以上)。例: 比率(%)。",
                    },
                    "left_title": {"type": "string", "description": "左軸のラベル(任意)。"},
                    "right_title": {"type": "string", "description": "右軸のラベル(任意)。"},
                    "title": {"type": "string", "description": "グラフのタイトル。"},
                    "purpose": {"type": "string", "description": "このグラフで示したいことの短い説明。"},
                },
                "required": ["sql", "x", "bar_y", "line_y", "title"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "pivot_table",
            "description": (
                "クロス集計表（ピボットテーブル）を作る。"
                "「AとBのマトリクスで」「行に○○、列に△△」「クロス集計」「表形式で比較」"
                "などと言われたら使う。"
                "SQLite には PIVOT 構文が無く CASE WHEN を列の数だけ手書きする必要があるため、"
                "列に展開したい集計はSQLで書かずにこのツールを使うこと。"
                "sql では集計せず、明細または index/columns/values の3列を返すだけでよい。"
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "sql": {"type": "string",
                            "description": "元データを取得する SELECT 文。集計はこのツールが行う。"},
                    "index": {"type": "array", "items": {"type": "string"},
                              "description": "行にする列名（複数可）。"},
                    "columns": {"type": "string",
                                "description": "列に展開する列名。省略すると行ごとの集計表になる。"},
                    "values": {"type": "string", "description": "集計する値の列名。"},
                    "aggfunc": {"type": "string", "enum": list(analysis.AGG_FUNCS),
                                "description": "集計方法。既定は sum。"},
                    "margins": {"type": "boolean", "description": "総計の行と列を付けるか。既定は false。"},
                    "percent": {"type": "string", "enum": list(analysis.PERCENT_MODES),
                                "description": "実数の代わりに構成比(%)で出す。"
                                               + " / ".join(f"{k}={v}" for k, v
                                                            in analysis.PERCENT_MODES.items())},
                    "rank": {"type": "string",
                             "description": "大きい順に並べて順位を付ける。列名を書くとその列で、"
                                            "'total' と書くと行の合計で並べる。"},
                    "render": {"type": "string", "enum": ["table", "heatmap"],
                               "description": "表示方法。heatmap にすると色付きの行列で見せる。既定は table。"},
                    "title": {"type": "string", "description": "見出し。"},
                    "purpose": {"type": "string", "description": "この集計で確認したいことの短い説明。"},
                },
                "required": ["sql", "index", "values"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "analyze_stats",
            "description": (
                "統計的な分析を行う。"
                "SQLite には STDDEV / MEDIAN / CORR / PERCENTILE が無いため、"
                "「相関」「中央値」「ばらつき」「四分位」「外れ値」「異常値」を聞かれたら"
                "SQLで計算しようとせず必ずこのツールを使うこと。"
                "sql は集計せずに明細を返す（1行1件）ようにする。"
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "sql": {"type": "string",
                            "description": "分析対象の明細を取得する SELECT 文。集計はしない。"},
                    "method": {
                        "type": "string", "enum": ["describe", "correlation", "outliers"],
                        "description": "describe=基本統計量(件数/平均/標準偏差/最小/四分位/中央値/最大) / "
                                       "correlation=相関行列 / outliers=外れ値の抽出",
                    },
                    "columns": {"type": "array", "items": {"type": "string"},
                                "description": "対象の数値列。省略すると数値列を自動判定する。"},
                    "group_by": {"type": "string",
                                 "description": "describe のとき、この列ごとに分けて統計を出す。任意。"},
                    "target": {"type": "string",
                               "description": "outliers のとき、外れ値を調べる数値列。必須。"
                                              "mahalanobis のときは列名をカンマ区切りで複数。"},
                    "outlier_method": {"type": "string",
                                       "enum": list(advanced.OUTLIER_METHODS_EXT),
                                       "description": " / ".join(
                                           f"{k}={v}" for k, v
                                           in advanced.OUTLIER_METHODS_EXT.items())},
                    "threshold": {"type": "number",
                                  "description": "外れ値の閾値。iqr は既定1.5、zscore は既定3。"},
                    "corr_method": {"type": "string", "enum": list(analysis.CORR_METHODS),
                                    "description": "pearson=直線的な関係(既定) / spearman=順位の関係"},
                    "lag": {"type": "integer",
                            "description": "correlation のとき、何期先までずらして相関を見るか。"
                                           "「先月の投入量が今月の結果に効くか」のような遅れて出る効果を"
                                           "調べたいときに指定する。sql は時点の昇順で1行1期にすること。"},
                    "partial": {"type": "boolean",
                                "description": "correlation のとき true にすると偏相関にする。"
                                               "control で指定した列の影響を取り除いてから相関を見るので、"
                                               "「第3の変数のせいで関係して見えるだけ」を切り分けられる。"},
                    "control": {"type": "array", "items": {"type": "string"},
                                "description": "partial=true のとき、影響を取り除きたい列。"},
                    "title": {"type": "string", "description": "見出し。"},
                    "purpose": {"type": "string", "description": "この分析で確認したいことの短い説明。"},
                },
                "required": ["sql", "method"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "export_excel",
            "description": (
                "SELECT の結果を Excel ファイル(.xlsx)にまとめ、ユーザーがダウンロードできる状態にする。"
                "「エクセルで」「xlsxで」「ファイルにして」「ダウンロードしたい」"
                "などと言われたら使う。sheets に複数の SELECT を渡すと複数シートのブックになる。"
                "chart を書くと、そのシートのデータからExcelのグラフを作って貼る"
                "（画像ではないので、受け取った側が範囲や種類を変えられる）。"
                "数字を並べるだけのシートより、グラフを1つ付けた方が伝わる。"
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "sheets": {
                        "type": "array",
                        "description": "ブックに入れるシート。1要素につき1シート。",
                        "items": {
                            "type": "object",
                            "properties": {
                                "name": {"type": "string", "description": "シート名(31文字以内)。"},
                                "sql": {"type": "string", "description": "このシートに書き出す SELECT 文。"},
                                "note": {"type": "string", "description": "シート先頭に入れる補足(任意)。"},
                                "chart": {
                                    "type": "object",
                                    "description": "このシートのデータから作るグラフ（任意）。",
                                    "properties": {
                                        "type": {"type": "string",
                                                 "enum": list(excel.EXCEL_CHART_TYPES),
                                                 "description": "グラフの種類。"},
                                        "category_column": {"type": "string",
                                                            "description": "横軸にする列名。"},
                                        "value_columns": {"type": "array",
                                                          "items": {"type": "string"},
                                                          "description": "系列にする数値列。"},
                                        "title": {"type": "string", "description": "グラフの題名。"},
                                        "y_title": {"type": "string", "description": "縦軸の名前。"},
                                        "x_title": {"type": "string", "description": "横軸の名前。"},
                                        "data_labels": {"type": "boolean",
                                                        "description": "値ラベルを出すか。"},
                                    },
                                    "required": ["type"],
                                },
                                "charts": {
                                    "type": "array",
                                    "items": {"type": "object", "additionalProperties": True},
                                    "description": "グラフを複数貼るときはこちらに並べる。",
                                },
                            },
                            "required": ["name", "sql"],
                        },
                    },
                    "filename": {"type": "string", "description": "ファイル名(拡張子不要)。内容が分かる名前を付ける。"},
                    "purpose": {"type": "string", "description": "何のためのファイルかの短い説明。"},
                },
                "required": ["sheets", "filename"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "export_csv",
            "description": (
                "SELECT の結果を CSV ファイルとして書き出し、ユーザーがダウンロードできる状態にする。"
                "「CSVで」「csvにして」「取り込み用のファイル」などと言われたら使う。"
                "files に複数指定すると、まとめてZIPで渡す。"
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "files": {
                        "type": "array",
                        "description": "書き出すファイル。1要素につき1CSV。",
                        "items": {
                            "type": "object",
                            "properties": {
                                "name": {"type": "string", "description": "ファイル名(拡張子不要)。"},
                                "sql": {"type": "string", "description": "書き出す SELECT 文。"},
                            },
                            "required": ["name", "sql"],
                        },
                    },
                    "encoding": {
                        "type": "string", "enum": list(exports.ENCODINGS),
                        "description": "文字コード。既定は utf-8-sig（Excelで開いても文字化けしない）。"
                                       "Shift_JIS が要るときだけ cp932。",
                    },
                    "delimiter": {
                        "type": "string", "enum": list(exports.EXPORT_DELIMITERS),
                        "description": "区切り文字。既定は comma。TSVにしたいときは tab。",
                    },
                    "purpose": {"type": "string", "description": "何のためのファイルかの短い説明。"},
                },
                "required": ["files"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "export_text",
            "description": (
                "文章（レポート・要約・メモ）をテキストファイルとして書き出し、"
                "ユーザーがダウンロードできる状態にする。"
                "「テキストで」「レポートにして」「議事録に」「まとめを文書で」などと言われたら使う。"
                "body に本文を自分で書き、必要なら sections に SELECT を指定して集計表を差し込む。"
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "filename": {"type": "string", "description": "ファイル名(拡張子不要)。"},
                    "body": {
                        "type": "string",
                        "description": "本文。あなたが書いた文章をそのまま入れる。"
                                       "sections を使う場合、差し込みたい位置に {{見出し}} と書く。",
                    },
                    "sections": {
                        "type": "array",
                        "description": "本文に差し込む集計表。省略可。",
                        "items": {
                            "type": "object",
                            "properties": {
                                "heading": {"type": "string",
                                            "description": "見出し。本文の {{この文字列}} が表に置き換わる。"
                                                           "本文に無ければ末尾に追記される。"},
                                "sql": {"type": "string", "description": "表にする SELECT 文。"},
                            },
                            "required": ["heading", "sql"],
                        },
                    },
                    "format": {
                        "type": "string", "enum": ["md", "txt"],
                        "description": "md=Markdown（表が罫線付き）/ txt=プレーンテキスト。既定は md。",
                    },
                    "encoding": {
                        "type": "string", "enum": list(exports.ENCODINGS),
                        "description": "文字コード。既定は utf-8-sig。",
                    },
                },
                "required": ["filename", "body"],
            },
        },
    },
    # ---- 統計 ---------------------------------------------------------------
    {
        "type": "function",
        "function": {
            "name": "hypothesis_test",
            "description": (
                "統計的仮説検定。「差があると言えるか」「偶然ではないか」「有意か」"
                "「A/Bどちらが良いか」「効果があったか」を判断したいときに使う。"
                "平均の差・比率の差・分布の偏り・相関の有無を、p値と効果量つきで判定する。"
                "sql は集計せず明細（1行1件）を返すこと。"
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "sql": {"type": "string", "description": "検定対象の明細を取る SELECT 文。"},
                    "method": {"type": "string", "enum": list(advanced.TEST_METHODS),
                               "description": " / ".join(f"{k}={v}" for k, v
                                                         in advanced.TEST_METHODS.items())},
                    "value_col": {"type": "string", "description": "測定値の列（数値）。"},
                    "group_col": {"type": "string",
                                  "description": "群を表す列。2群比較・分散分析・カイ二乗で使う。"},
                    "value_col2": {"type": "string",
                                   "description": "対応のある検定や相関で、もう一方の列。"},
                    "popmean": {"type": "number", "description": "1標本t検定で比較する基準値。"},
                    "expected": {"type": "array", "items": {"type": "number"},
                                 "description": "適合度検定で期待する比率や度数。"},
                    "alternative": {"type": "string",
                                    "enum": ["two-sided", "less", "greater"],
                                    "description": "対立仮説。既定は両側(two-sided)。"},
                    "title": {"type": "string", "description": "見出し。"},
                },
                "required": ["sql", "method"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "regression",
            "description": (
                "回帰分析。「何が効いているか」「要因分析」「どの変数が影響するか」"
                "「予測式を作りたい」ときに使う。係数・p値・寄与の大きさ・"
                "あてはまり(R²)・多重共線性まで返す。"
                "目的変数が0/1なら logistic、件数なら poisson を選ぶ。"
                "文字列の説明変数（区分など）は自動でダミー変数にする。"
                "sql は集計せず明細（1行1件）を返すこと。"
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "sql": {"type": "string", "description": "分析対象の明細を取る SELECT 文。"},
                    "target": {"type": "string", "description": "目的変数（説明したい列）。"},
                    "features": {"type": "array", "items": {"type": "string"},
                                 "description": "説明変数の列名。"},
                    "method": {"type": "string", "enum": list(advanced.REGRESSION_METHODS),
                               "description": " / ".join(f"{k}={v}" for k, v
                                                         in advanced.REGRESSION_METHODS.items())},
                    "predict": {"type": "array", "items": {"type": "object"},
                                "description": "予測したい入力の一覧。例 [{\"説明変数の列名\":100}]",
                                "additionalProperties": True},
                    "title": {"type": "string", "description": "見出し。"},
                },
                "required": ["sql", "target", "features"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "distribution_analysis",
            "description": (
                "分布の形を調べる。「ばらつき」「ヒストグラム」「偏り」「どんな分布か」"
                "「上位下位の広がり」を見たいときに使う。度数分布・要約統計・"
                "正規分布などへの当てはめ判定を返す。"
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "sql": {"type": "string", "description": "明細を取る SELECT 文。"},
                    "target": {"type": "string", "description": "調べる数値列。"},
                    "bins": {"type": "integer", "description": "階級の数。既定20。"},
                    "group_col": {"type": "string", "description": "群ごとに比べるときの列。"},
                    "fit": {"type": "array", "items": {"type": "string",
                                                       "enum": list(advanced.DISTRIBUTIONS)},
                            "description": "当てはめを試す分布。既定は norm と lognorm。"},
                    "title": {"type": "string", "description": "見出し。"},
                },
                "required": ["sql", "target"],
            },
        },
    },
    # ---- 時系列 -------------------------------------------------------------
    {
        "type": "function",
        "function": {
            "name": "forecast",
            "description": (
                "将来の値を予測する。「来月はいくら」「このままいくと」「着地見込み」"
                "「予測」を聞かれたら使う。予測値と95%の幅、"
                "過去データで試した誤差率(MAPE)を返す。"
                "sql は時点ごとに1行（例: 月ごとの合計）にすること。"
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "sql": {"type": "string",
                            "description": "時点と値を返す SELECT 文（時点の昇順で1行1期）。"},
                    "time_col": {"type": "string", "description": "時点の列（年月など）。"},
                    "value_col": {"type": "string", "description": "予測する数値の列。"},
                    "periods": {"type": "integer", "description": "何期先まで予測するか。既定6。"},
                    "method": {"type": "string",
                               "enum": ["auto"] + list(advanced.FORECAST_METHODS),
                               "description": "auto=データ量から自動選択。" + " / ".join(
                                   f"{k}={v}" for k, v in advanced.FORECAST_METHODS.items())},
                    "season_length": {"type": "integer",
                                      "description": "季節の周期。月次で1年なら12、曜日なら7。"},
                    "exog": {
                        "type": "object", "additionalProperties": True,
                        "description": "説明変数つきで予測する場合の指定。"
                                       "{\"columns\": [\"説明変数の列名\"], \"future\": [[120],[130]]} の形で、"
                                       "future には予測する期数と同じ数だけ将来の値を並べる。"
                                       "「その変数をこう置いたら対象はどうなるか」に答えられる。",
                    },
                    "title": {"type": "string", "description": "見出し。"},
                },
                "required": ["sql", "time_col", "value_col"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "timeseries_analysis",
            "description": (
                "時系列の見方をまとめる。「推移」「トレンド」「季節性」「前年同月比」"
                "「移動平均」「周期」を聞かれたときに使う。"
                "sql は時点ごとに1行にすること。"
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "sql": {"type": "string", "description": "時点と値を返す SELECT 文。"},
                    "time_col": {"type": "string", "description": "時点の列。"},
                    "value_col": {"type": "string", "description": "値の列。"},
                    "window": {"type": "integer", "description": "移動平均の期間。既定3。"},
                    "season_length": {"type": "integer",
                                      "description": "季節の周期（月次なら12）。指定すると季節分解する。"},
                    "title": {"type": "string", "description": "見出し。"},
                },
                "required": ["sql", "time_col", "value_col"],
            },
        },
    },
    # ---- 試算・シミュレーション ---------------------------------------------
    {
        "type": "function",
        "function": {
            "name": "monte_carlo_simulation",
            "description": (
                "モンテカルロ・シミュレーション。「もし〜だったら」「リスク」"
                "「確率」「見込みの幅」「何%の確率で」を扱うときに使う。"
                "不確実な入力を分布で与え、式を何万回も試して結果の分布を出す。"
                "実データのばらつきをそのまま使いたい変数は dist=empirical と column を指定し、"
                "その列を返す sql も併せて渡すこと。"
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "formula": {"type": "string",
                                "description": "変数名を使った計算式。例 (単価 - 原価) * 数量 - 固定費"},
                    "variables": {
                        "type": "object", "additionalProperties": True,
                        "description": "{変数名: {dist, ...}}。dist は normal(mean,std) / "
                                       "uniform(min,max) / triangular(min,mode,max) / "
                                       "lognormal(mean,std) / poisson(lam) / binomial(n,p) / "
                                       "empirical(column) / fixed(value)。",
                    },
                    "trials": {"type": "integer", "description": "試行回数。既定10000。"},
                    "sql": {"type": "string",
                            "description": "empirical を使うときに、その列を含む明細を取る SELECT 文。"},
                    "targets": {"type": "array", "items": {"type": "number"},
                                "description": "「この値を超える確率」を知りたいしきい値。"},
                    "title": {"type": "string", "description": "見出し。"},
                },
                "required": ["formula", "variables"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "scenario_analysis",
            "description": (
                "シナリオ比較（楽観・標準・悲観など）。前提を数パターン置いて"
                "結果を並べ、どの変数の影響が大きいかも出す。"
                "確率分布まで置く必要がないときは、こちらの方が説明しやすい。"
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "formula": {"type": "string", "description": "変数名を使った計算式。"},
                    "scenarios": {"type": "object", "additionalProperties": True,
                                  "description": "{シナリオ名: {変数名: 値}}"},
                    "base": {"type": "object", "additionalProperties": True,
                             "description": "共通の前提値。感度分析の基準にもなる。"},
                    "title": {"type": "string", "description": "見出し。"},
                },
                "required": ["formula", "scenarios"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "bootstrap_estimate",
            "description": (
                "ブートストラップ法で平均などの信頼区間を出す。"
                "「この差は誤差の範囲か」「どのくらい確からしいか」を、"
                "分布の形を仮定せずに示せる。件数が少ないときにも使える。"
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "sql": {"type": "string", "description": "明細を取る SELECT 文。"},
                    "target": {"type": "string", "description": "対象の数値列。"},
                    "statistic": {"type": "string",
                                  "enum": ["mean", "median", "std", "sum", "p90"],
                                  "description": "推定する統計量。既定 mean。"},
                    "group_col": {"type": "string", "description": "群ごとに出すときの列。"},
                    "trials": {"type": "integer", "description": "再抽出の回数。既定5000。"},
                    "title": {"type": "string", "description": "見出し。"},
                },
                "required": ["sql", "target"],
            },
        },
    },
    # ---- 分ける -------------------------------------------------------------
    {
        "type": "function",
        "function": {
            "name": "clustering",
            "description": (
                "k-meansでグループ分けする。「セグメント」「タイプ分け」「似ている順に分類」"
                "「顧客を分けたい」ときに使う。列ごとの尺度差は標準化して吸収する。"
                "分け方が分からないときは k に \"auto\" を指定すると、"
                "最も素直に分かれる数を自動で選び、各グループの特徴も言葉で返す。"
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "sql": {"type": "string", "description": "明細を取る SELECT 文。"},
                    "features": {"type": "array", "items": {"type": "string"},
                                 "description": "分類に使う数値列。"},
                    "k": {"type": ["integer", "string"],
                          "description": "グループ数。既定3。\"auto\" にすると"
                                         "シルエット係数が最も高い数（2〜8）を自動で選ぶ。"},
                    "categorical": {"type": "array", "items": {"type": "string"},
                                    "description": "分類に使いたい区分の列（拠点・種別などの区分）。"
                                                   "0/1に開いてから一緒に分ける。"},
                    "label_col": {"type": "string", "description": "行の名前になる列（対象の名称）。"},
                    "title": {"type": "string", "description": "見出し。"},
                },
                "required": ["sql", "features"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "abc_analysis",
            "description": (
                "ABC分析（パレート分析）。「上位2割で全体の8割を占めるか」「重点管理」"
                "「上位集中度」を見るときに使う。累計構成比でA/B/Cに区分する。"
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "sql": {"type": "string", "description": "品目と値を返す SELECT 文。"},
                    "label_col": {"type": "string", "description": "対象の列（品目・拠点などの名称）。"},
                    "value_col": {"type": "string", "description": "金額や数量の列。"},
                    "thresholds": {"type": "array", "items": {"type": "number"},
                                   "description": "A/Bの境目。既定 [70, 90]（累計%）。"},
                    "title": {"type": "string", "description": "見出し。"},
                },
                "required": ["sql", "label_col", "value_col"],
            },
        },
    },
    # ---- レポート -----------------------------------------------------------
    {
        "type": "function",
        "function": {
            "name": "export_pptx",
            "description": (
                "PowerPointのレポートを作る。「パワポで」「スライドにして」「報告資料」"
                "「プレゼン」と言われたら使う。会議でそのまま映せる体裁で出力する。"
                "グラフは編集できるPowerPointのグラフとして入る。"
                "各スライドに sql を書くと、その場でSQLを実行して中身を埋める。"
                "\n"
                "重要: 各スライドに message（そのページで言いたいこと1行）を必ず書く。"
                "見出しの下に帯で表示され、聞き手はここだけ読めば分かる。"
                "「◯◯の推移」ではなく「3月の落ち込みは期ずれで、実勢は右肩上がり」と書く。"
                "\n"
                "構成の目安: title（表紙）→ kpi（数字の要約）→ "
                "section（章の区切り）→ chart / table（根拠）→ compare（案の比較）→ "
                "closing（まとめと次のアクション）。中扉が2つ以上あれば目次は自動で入る。"
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "filename": {"type": "string", "description": "ファイル名（拡張子不要）。"},
                    "title": {"type": "string", "description": "レポート全体の題名。"},
                    "subtitle": {"type": "string", "description": "副題（対象期間など）。"},
                    "footer": {"type": "string", "description": "各ページ下部に入れる文字。"},
                    "slides": {
                        "type": "array",
                        "description": "スライドの並び。",
                        "items": {
                            "type": "object",
                            "properties": {
                                "kind": {"type": "string", "enum": list(pptx_report.SLIDE_KINDS),
                                         "description": "title=表紙 / agenda=目次 / "
                                                        "section=中扉 / message=文字だけ / "
                                                        "table=表 / chart=グラフ / "
                                                        "kpi=数字を大きく / compare=2案の比較 / "
                                                        "closing=まとめと次のアクション"},
                                "title": {"type": "string"},
                                "subtitle": {"type": "string"},
                                "message": {"type": "string",
                                            "description": "そのページで言いたいこと1行。"
                                                           "見出しの下に帯で出る。必ず書く。"},
                                "sql": {"type": "string",
                                        "description": "table/chart のとき、中身を取る SELECT 文。"},
                                "chart": {"type": "string", "enum": list(pptx_report.PPTX_CHART_TYPES),
                                          "description": "グラフの種類。"},
                                "category_column": {"type": "string",
                                                    "description": "chart のとき横軸にする列。"},
                                "value_columns": {"type": "array", "items": {"type": "string"},
                                                  "description": "chart のとき系列にする数値列。"},
                                "bullets": {"type": "array",
                                            "items": {"type": "object",
                                                      "additionalProperties": True},
                                            "description": "箇条書き。文字列でも "
                                                           "{text, level, strong} でもよい。"},
                                "lead": {"type": "string",
                                         "description": "message のとき、箇条書きの前に置く導入文。"},
                                "body": {"type": "string", "description": "本文。"},
                                "items": {"type": "array", "items": {"type": "object",
                                                                     "additionalProperties": True},
                                          "description": "kpi のとき "
                                                         "[{label, value, unit, delta, "
                                                         "delta_unit, delta_label, "
                                                         "higher_is_better, note}]。"},
                                "panes": {"type": "array", "items": {"type": "object",
                                                                     "additionalProperties": True},
                                          "description": "compare のとき、左右2つ "
                                                         "[{title, value, unit, bullets}]。"},
                                "summary": {"type": "array", "items": {"type": "string"},
                                            "description": "closing のときのまとめ。"},
                                "actions": {"type": "array",
                                            "items": {"type": "object",
                                                      "additionalProperties": True},
                                            "description": "closing のとき "
                                                           "[{text, owner, due}]。"},
                                "callout": {"type": "string",
                                            "description": "下部の囲みで強調する一文。1枚に1つまで。"},
                                "comment": {"type": "string",
                                            "description": "図表の右に添える所見。"},
                                "source": {"type": "string",
                                           "description": "出所・集計条件。数字の資料には入れる。"},
                                "notes": {"type": "string", "description": "発表者ノート。"},
                                "data_labels": {"type": "boolean",
                                                "description": "グラフに数値ラベルを出す。"
                                                               "既定は自動判断。"},
                                "highlight_rows": {"type": "array",
                                                   "items": {"type": "integer"},
                                                   "description": "table で強調する行（0始まり）。"},
                                "max_rows": {"type": "integer",
                                             "description": "table で載せる最大行数。既定12。"},
                            },
                            "required": ["kind"],
                        },
                    },
                },
                "required": ["slides"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "build_report",
            "description": (
                "分析の結果を1つのレポートにまとめる。"
                "「レポートにして」「まとめて」「報告書」「分析結果を整理して」"
                "「結論と根拠を示して」と言われたら使う。"
                "画面に読みやすい形で出しつつ、ダウンロードできるファイルも作る。"
                "各セクションに sql を書けば表が入り、chart を書けばグラフも入る。"
                "要点(summary)と結論(conclusion)は必ず自分の言葉で書くこと。"
                "数字を並べるだけでなく『だから何か』を書く。"
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "title": {"type": "string", "description": "レポートの題名。"},
                    "subtitle": {"type": "string", "description": "対象期間や条件。"},
                    "summary": {
                        "type": "array", "items": {"type": "string"},
                        "description": "要点。最初に読む人が3行で分かるように書く。",
                    },
                    "sections": {
                        "type": "array",
                        "description": "本編。1セクション＝1つの論点。",
                        "items": {
                            "type": "object",
                            "properties": {
                                "heading": {"type": "string", "description": "見出し。"},
                                "body": {"type": "string",
                                         "description": "説明の文章。何が言えるかを書く。"},
                                "sql": {"type": "string",
                                        "description": "根拠として載せる表の SELECT 文。任意。"},
                                "chart": {
                                    "type": "object", "additionalProperties": True,
                                    "description": "グラフの指定。chart_type と x / y などを"
                                                   "入れる。sql の結果を使う。任意。",
                                },
                                "note": {"type": "string",
                                         "description": "この節の所見・注意点。任意。"},
                                "max_rows": {"type": "integer",
                                             "description": "表に載せる最大行数。既定20。"},
                            },
                            "required": ["heading"],
                        },
                    },
                    "conclusion": {"type": "string", "description": "結論。"},
                    "recommendations": {
                        "type": "array", "items": {"type": "string"},
                        "description": "推奨する打ち手。実行できる粒度で書く。",
                    },
                    "caveats": {
                        "type": "array", "items": {"type": "string"},
                        "description": "前提・制約・数字の読み方の注意。",
                    },
                    "format": {
                        "type": "string",
                        "enum": ["md", "docx", "pptx", "xlsx", "none"],
                        "description": "ダウンロード用ファイルの形式。"
                                       "md=軽い文書(既定) / docx=Word報告書（図表つき）/ "
                                       "pptx=スライド / xlsx=表ごとにシート / "
                                       "none=画面表示だけ",
                    },
                    "filename": {"type": "string", "description": "ファイル名（拡張子不要）。"},
                    "org": {"type": "string", "description": "表紙に入れる組織名（部署・拠点など）。"},
                    "footer": {"type": "string",
                               "description": "各ページ下部の文字（「社外秘」など）。"},
                },
                "required": ["title", "sections"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "export_docx",
            "description": (
                "Word文書（.docx）を作る。「Wordで」「docxで」「報告書にして」"
                "「配布資料」「回覧」と言われたら使う。"
                "表紙・目次・図表番号つきのキャプション・ページ番号が入り、"
                "そのまま配布できる体裁になる。"
                "各セクションに sql を書くと表が入り、chart も書くとグラフが図として入る。"
                "本文(body)は必ず自分の言葉で書くこと。表を貼っただけの文書は読まれない。"
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "title": {"type": "string", "description": "文書の題名。"},
                    "subtitle": {"type": "string", "description": "副題（対象期間など）。"},
                    "org": {"type": "string", "description": "表紙に入れる組織名（部署・拠点など）。"},
                    "author": {"type": "string", "description": "作成者名。"},
                    "footer": {"type": "string",
                               "description": "各ページ下部の文字（「社外秘」など）。"},
                    "toc": {"type": "boolean", "description": "目次を入れるか。既定は入れる。"},
                    "summary": {"type": "array", "items": {"type": "string"},
                                "description": "冒頭の要約。3点程度。"},
                    "sections": {
                        "type": "array",
                        "description": "本編。1セクション＝1つの見出し。",
                        "items": {
                            "type": "object",
                            "properties": {
                                "heading": {"type": "string", "description": "見出し。"},
                                "level": {"type": "integer",
                                          "description": "見出しの階層（1が大見出し）。"},
                                "body": {"type": "string", "description": "本文。"},
                                "bullets": {"type": "array", "items": {"type": "string"},
                                            "description": "箇条書き。"},
                                "sql": {"type": "string",
                                        "description": "表・グラフの元になる SELECT 文。"},
                                "chart": {"type": "object", "additionalProperties": True,
                                          "description": "グラフの指定（chart_type と x / y など）。"
                                                         "図として貼られる。"},
                                "table": {"type": "boolean",
                                          "description": "表も載せるか。既定は載せる。"},
                                "caption": {"type": "string", "description": "図のキャプション。"},
                                "table_caption": {"type": "string",
                                                  "description": "表のキャプション。"},
                                "note": {"type": "string", "description": "補足・注記。"},
                                "callout": {"type": "string",
                                            "description": "囲みで強調したい一文。"},
                                "max_rows": {"type": "integer",
                                             "description": "表に載せる最大行数。既定40。"},
                                "page_break": {"type": "boolean",
                                               "description": "このセクションの前で改ページする。"},
                            },
                            "required": ["heading"],
                        },
                    },
                    "conclusion": {"type": "string", "description": "結論。"},
                    "recommendations": {
                        "type": "array",
                        "items": {"type": "object", "additionalProperties": True},
                        "description": "推奨する打ち手。[{text, owner, due}] または文字列の並び。",
                    },
                    "caveats": {"type": "array", "items": {"type": "string"},
                                "description": "前提・注意。"},
                    "filename": {"type": "string", "description": "ファイル名（拡張子不要）。"},
                },
                "required": ["title", "sections"],
            },
        },
    },
    # ---- メール -------------------------------------------------------------
    {
        "type": "function",
        "function": {
            "name": "find_mail_recipients",
            "description": (
                "宛先を表から探す。名前・部署・アドレスの一部で検索できる。"
                "「〇〇部に送って」「田中さんに送って」と言われたら、"
                "まずこれで実在するアドレスを確認してから compose_email を呼ぶこと。"
                "アドレスを推測で作ってはいけない。"
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "query": {"type": "string",
                              "description": "検索語（氏名・部署名・アドレスの一部）。空なら一覧。"},
                    "table": {"type": "string", "description": "探す表を絞るとき。省略可。"},
                    "limit": {"type": "integer", "description": "最大件数。既定50。"},
                },
                "required": [],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "compose_email",
            "description": (
                "メールの下書きを作る。画面に確認カードが出て、"
                "ユーザーが「送信」を押したときだけ実際に送られる（自動送信はしない）。"
                "宛先は find_mail_recipients で確認した実在のアドレスを使うこと。"
                "直前に作った Excel / CSV / PowerPoint を添付できる。"
                "本文は挨拶・要点・詳細・結びの順で、日本語のビジネスメールとして書く。"
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "to": {"type": "array", "items": {"type": "string"},
                           "description": "宛先アドレス。"},
                    "to_query": {"type": "string",
                                 "description": "アドレスの代わりに検索語で指定する場合"
                                                "（例: 営業部）。DBから引いて宛先にする。"},
                    "cc": {"type": "array", "items": {"type": "string"}},
                    "bcc": {"type": "array", "items": {"type": "string"}},
                    "subject": {"type": "string", "description": "件名。"},
                    "body": {"type": "string", "description": "本文（プレーンテキスト）。"},
                    "attach_filenames": {
                        "type": "array", "items": {"type": "string"},
                        "description": "この会話で作ったファイル名。省略時は添付なし。"
                                       "'all' を入れると直近に作ったファイルを全部添付する。"},
                    "reply_to": {"type": "string", "description": "返信先アドレス。"},
                },
                "required": ["subject", "body"],
            },
        },
    },
    # ---- 業務でよく聞かれる分析 ---------------------------------------------
    {
        "type": "function",
        "function": {
            "name": "compare_periods",
            "description": (
                "2つの期間を比べ、差がどこから来たのかまで分解する。"
                "「先月と比べて」「前年同月比」「前期からどう変わったか」"
                "「なぜ落ちたのか」を聞かれたら使う。"
                "全体の増減だけでなく、どの区分が押し下げ／押し上げたか（寄与度）を出す。"
                "qty_col を渡すと、金額の変化を「数量が動いたぶん」と「単価が動いたぶん」に分ける。"
                "sql は期間の列・値の列（あれば区分の列）を含む形で、2期間ぶんまとめて取ること。"
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "sql": {"type": "string",
                            "description": "2期間ぶんのデータを取る SELECT 文。"
                                           "期間の列・値の列・（任意で）区分の列を返す。"},
                    "period_col": {"type": "string", "description": "期間を表す列（'2026-01' など）。"},
                    "value_col": {"type": "string", "description": "比べる数値の列（売上など）。"},
                    "dimension_col": {"type": "string",
                                      "description": "増減の内訳を見る区分の列（部門・商品など）。任意。"},
                    "qty_col": {"type": "string",
                                "description": "数量の列。指定すると数量要因と単価要因に分解する。"},
                    "current": {"type": "string", "description": "当期。省略すると最後の期。"},
                    "previous": {"type": "string", "description": "前期。省略すると最後から2番目の期。"},
                    "title": {"type": "string", "description": "見出し。"},
                },
                "required": ["sql", "period_col", "value_col"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "data_quality",
            "description": (
                "分析の前に、データそのものの異常を洗い出す。"
                "「数字が合わない」「件数がおかしい」「このデータは信用できるか」"
                "と言われたとき、また重要な集計を出す前の確認に使う。"
                "行数・主キーの重複・空の列・親に存在しない外部キー・日付の範囲を調べ、"
                "深刻な順に並べて返す。SQLは要らない（対象のテーブルを直接見る）。"
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "tables": {"type": "array", "items": {"type": "string"},
                               "description": "調べるテーブル名。省略すると対象のテーブルを順に見る。"},
                    "title": {"type": "string", "description": "見出し。"},
                },
                "required": [],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "detect_anomalies",
            "description": (
                "時系列から「いつもと違う時点」と「いつから変わったか」を見つける。"
                "「異常」「急に増えた」「おかしい日」「いつから悪化したか」を聞かれたら使う。"
                "前後の期間と比べるので、右肩上がりのデータでも直近を全部異常とは言わない。"
                "静的な外れ値（analyze_stats の outliers）とは用途が違う。"
                "sql は時点ごとに1行にすること。"
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "sql": {"type": "string", "description": "時点と値を返す SELECT 文。"},
                    "time_col": {"type": "string", "description": "時点の列。"},
                    "value_col": {"type": "string", "description": "監視する数値の列。"},
                    "window": {"type": "integer", "description": "比べる前後の期間。既定7。"},
                    "threshold": {"type": "number",
                                  "description": "何倍離れたら異常とするか。既定3。小さくすると多く拾う。"},
                    "season_length": {"type": "integer",
                                      "description": "曜日や月の周期。指定すると季節変動を除いてから判定する。"},
                    "changepoints": {"type": "boolean",
                                     "description": "水準が変わった時点も探すか。既定true。"},
                    "title": {"type": "string", "description": "見出し。"},
                },
                "required": ["sql", "time_col", "value_col"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "funnel_analysis",
            "description": (
                "段階ごとの通過・離脱・滞留を出す。"
                "「段階ごとの通過率」「どこで落ちているか」「滞留」"
                "「リードタイム」を聞かれたら使う。"
                "sql は 1行=1案件 にして、各段階の日付（または通過フラグ）の列を並べること。"
                "例: 受付日・着手日・完了日・確認日のように、進んだ順に日付列を並べたSELECT。"
                "値が入っていればその段階を通過した扱いになる。"
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "sql": {"type": "string",
                            "description": "1行=1案件で、各段階の日付列を並べた SELECT 文。"},
                    "steps": {"type": "array", "items": {"type": "string"},
                              "description": "段階を表す列名を、順番に並べる。"},
                    "labels": {"type": "array", "items": {"type": "string"},
                               "description": "画面に出す段階の名前。省略すると列名を使う。"},
                    "group_col": {"type": "string",
                                  "description": "区分ごとに通過率を比べるときの列（担当・地域など）。"},
                    "title": {"type": "string", "description": "見出し。"},
                },
                "required": ["sql", "steps"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "cohort_analysis",
            "description": (
                "いつ始めた人がどれだけ続いているかを見る。"
                "「継続率」「定着」「リピート」「離脱」「初回からの推移」を聞かれたら使う。"
                "初回の期でグループ分けし、経過期ごとの残存率をマトリクスで返す。"
                "sql は 1行=(対象, 期) の明細にすること（同じ人が複数期に出てよい）。"
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "sql": {"type": "string",
                            "description": "対象と期の明細を返す SELECT 文。"},
                    "id_col": {"type": "string", "description": "対象の列（顧客ID・社員IDなど）。"},
                    "period_col": {"type": "string",
                                   "description": "期の列（'2026-01' など、並べて正しい順になる形）。"},
                    "value_col": {"type": "string",
                                  "description": "金額なども見るときの数値列。任意。"},
                    "max_periods": {"type": "integer", "description": "何期先まで見るか。既定12。"},
                    "title": {"type": "string", "description": "見出し。"},
                },
                "required": ["sql", "id_col", "period_col"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "market_basket",
            "description": (
                "一緒に買われている（使われている）品目の組み合わせを見つける。"
                "「併売」「セット販売」「一緒に買われる」「関連商品」を聞かれたら使う。"
                "支持度・確信度・リフトを返す。SQLでは実質書けない分析。"
                "sql は 1行=(伝票, 品目) の明細にすること。"
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "sql": {"type": "string",
                            "description": "伝票と品目の明細を返す SELECT 文。"},
                    "transaction_col": {"type": "string",
                                        "description": "ひとまとまりを表す列（伝票ID・作業票IDなど）。"},
                    "item_col": {"type": "string", "description": "一緒に出てくる対象の列（品目名など）。"},
                    "min_support": {"type": "number",
                                    "description": "全伝票に占める最低の出現率(%)。既定1.0。"},
                    "top": {"type": "integer", "description": "返す組み合わせの数。既定25。"},
                    "title": {"type": "string", "description": "見出し。"},
                },
                "required": ["sql", "transaction_col", "item_col"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "survival_analysis",
            "description": (
                "「どれだけ持つか」「いつ辞めるか」を扱う。"
                "対象が何かの発生までどれだけ続いたかを見る分析で、"
                "在籍期間・契約の継続期間・機器の故障間隔(MTBF)・部品の寿命などに使える。"
                "まだ起きていない分（継続中・稼働中）を捨てずに計算するので、"
                "単純平均のように短く見積もることがない。"
                "Weibull分布の形から、経年で増える型か初期に集中する型か偶発型かも判定する。"
                "sql は 1行=1対象 にして、期間の列と（あれば）発生フラグの列を返すこと。"
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "sql": {"type": "string", "description": "1行=1対象の明細を返す SELECT 文。"},
                    "duration_col": {"type": "string",
                                     "description": "期間の列（経過日数・稼働時間など）。"},
                    "event_col": {"type": "string",
                                  "description": "起きたか(1)まだか(0)の列。省略すると全件で起きた扱い。"},
                    "group_col": {"type": "string",
                                  "description": "群ごとに比べるときの列（区分・種別など）。"},
                    "title": {"type": "string", "description": "見出し。"},
                },
                "required": ["sql", "duration_col"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "explore_import_files",
            "description": (
                "「データ取り込み」の取り込み元フォルダにあるファイルを調べる。読むだけで、"
                "取り込み・変更・削除はしない（取り込みは画面の操作でしか行われない）。"
                "「取り込み元に何がある？」「新しく届いたファイルは？」「このExcelは取り込める？」"
                "「まだ取り込んでいないファイルは？」と聞かれたら使う。"
                "まだ取り込んでいないファイルの話なので、SQLでは答えられない。"
                "\n"
                "file を指定しなければ一覧（拡張子を問わず、場所・サイズ・更新日時・"
                "取り込み済みか）、指定すればそのファイルの下見になる。"
                "一覧で得たパスをそのまま file に渡すこと（パスを推測して組み立てないこと）。"
                "\n"
                "下見では「そのまま取り込める / 手直しが要る / 取り込みに向かない」を判定し、"
                "理由と直し方を返す。セル結合・多段見出し・月が横に並んだクロス表・"
                "合計行の混入・見出しが1行目にない、といった"
                "「取り込めるが正しく使えない」形を見つけられる。"
                "取り込みを勧める前に、この判定を確認すること。"
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "path": {"type": "string",
                             "description": "見たいフォルダ。省略すると許可フォルダの直下を見る。"},
                    "recursive": {"type": "boolean",
                                  "description": "下の階層もまとめて見るか。既定 false。"},
                    "pattern": {"type": "string",
                                "description": "ファイル名に含まれる文字での絞り込み。"},
                    "only_not_imported": {"type": "boolean",
                                          "description": "まだ取り込んでいないファイルだけに絞る。"},
                    "check": {"type": "boolean",
                              "description": "一覧の各ファイルについて、表として使える形かどうかも"
                                             "判定する。1件ずつ開くので20件までにしている。"},
                    "file": {"type": "string",
                             "description": "中身を下見するファイルのパス。一覧で得たものを使う。"},
                    "sheet": {"type": "string",
                              "description": "Excelのとき、見たいシート名。省略すると先頭のシート。"},
                    "header_row": {"type": "integer",
                                   "description": "見出しの行（0始まり）。2行目が見出しなら1。"},
                    "rows": {"type": "integer",
                             "description": "下見で読む行数。既定5、最大20。"},
                    "title": {"type": "string", "description": "見出し。"},
                },
                "required": [],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "propose_glossary_term",
            "description": (
                "業務用語をデータカタログの用語集に登録する「登録カード」をチャットに出す。"
                "実際に登録するかはユーザーがカードのボタンで決める（勝手には登録されない）。"
                "\n使いどころ:"
                "\n- ユーザーが言葉の定義を教えてくれたとき"
                "（「有効な◯◯とは△△を除いたもの」など）"
                "\n- あいまいな用語をあなたが解釈し、その解釈をユーザーが認めたとき"
                "\n- 同じ言葉の意味を何度も聞き直していると気づいたとき"
                "\nまず「この定義で用語集に登録しますか？」と一言確認し、"
                "前向きな返事があったらこのツールでカードを出す。"
                "会話のたびに毎回は出さない（うるさくなる）。"
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "term": {"type": "string",
                             "description": "用語（現場で使われている言葉をそのまま）。"},
                    "description": {"type": "string",
                                    "description": "日本語の定義。ユーザーの言い回しを活かす。"},
                    "sql": {"type": "string",
                            "description": "SQLの条件式・計算式（任意）。"
                                           "例:「列名 != '除外コード'」のような条件式、"
                                           "または「SUM(列A) / SUM(列B)」のような計算式。"
                                           "会話で実際に使って正しかった式を入れる。"},
                    "how": {"type": "string",
                            "description": "どのデータをどこから取り、どう絞る/計算するのかを、"
                                           "SQLを知らない人に伝わる日本語で書く。テーブルや列は"
                                           "業務の言葉で呼ぶ。例:「◯◯データから、△△が□□の行を"
                                           "数える」の形で、対象・絞り込み条件・集計方法の3つが"
                                           "分かるように書く。必須。"},
                    "table": {"type": "string",
                              "description": "用語を置くテーブル。その用語が主に関わるテーブル名。"
                                             "複数テーブルにまたがる用語のときは省略し db を指定。"},
                    "db": {"type": "string",
                           "description": "DB全体の用語にするときのDBファイル名（拡張子込み。"
                                          "サイドバーに出ている名前をそのまま書く）。"},
                },
                "required": ["term", "description", "how"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "propose_example",
            "description": (
                "いまの質問と実行済みSQLを「例文」としてカタログに登録する登録カードを出す。"
                "例文はAIのお手本になり、似た質問への精度が上がる。"
                "実際に登録するかはユーザーがカードのボタンで決める。"
                "\n使いどころ:"
                "\n- ユーザーが「これを例文にして」「この答えを覚えて」と言ったとき"
                "\n- ユーザーが回答を「合っている」と認め、その質問が今後もよく出そうなとき"
                "\nsql には、この会話で実際に実行して正しかったSQLをそのまま入れる（書き直さない）。"
                "毎回は提案しない。"
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "question": {"type": "string",
                                 "description": "ユーザーの質問文（言い回しを変えない）。"},
                    "sql": {"type": "string",
                            "description": "実行して正しかったSELECT文そのまま。"},
                    "summary": {"type": "string",
                                "description": "どのデータをどこから取り、どう集計したかを、"
                                               "SQLを知らない人に伝わる日本語で。"
                                               "例:「◯◯データと△△マスタをつなぎ、"
                                               "□□ごとに件数を合計した」のように、"
                                               "使った表・つなぎ方・集計軸が分かるように書く。"},
                },
                "required": ["question", "sql", "summary"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "show_er_diagram",
            "description": (
                "ER図（テーブル同士の関係図）をチャット画面に表示する。"
                "「ER図を見せて」「テーブルの関係を図で」「データ構造を見たい」"
                "と言われたら使う。図は読み取り専用で、利用者が拡大縮小・全画面表示できる。"
                "表示と同時に結合の一覧も返るので、それを踏まえて補足してよい。"
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "purpose": {"type": "string",
                                "description": "何を確かめたくて表示するかの短い説明。"},
                },
                "required": [],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "open_table",
            "description": (
                "テーブルの中身（全行）を見る画面へのリンクをチャットに出す"
                "（利用者が「テーブル全体を開く」を押すと別タブで開く。勝手には開かない）。"
                "「〇〇テーブルを見せて」「中身を全部見たい」「データそのものを確認したい」"
                "と言われたら、SELECT * を書くのではなくこれを使う。"
                "画面では列ごとのフィルター・並べ替え・ページ送りができ、全行を辿れる。"
                "出したあとは、そのテーブルが何かを1〜2文で補足するだけでよい"
                "（中身を表で貼り直さない）。"
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "table": {"type": "string", "description": "開くテーブル名だけ。「DB名.テーブル名」の形にせず、DB名は db に分けて渡す。"},
                    "db": {"type": "string",
                           "description": "そのテーブルがあるDBファイル名。"
                                          "同じ表名が複数のDBにあるときは必須"
                                          "（サイドバーに出ている名前をそのまま書く）。"},
                    "purpose": {"type": "string",
                                "description": "何を確かめたくて開くかの短い説明。"},
                },
                "required": ["table"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "analyze_usage",
            "description": (
                "このアプリ自身の使われ方（利用状況）を調べる。分析対象のデータの中身ではなく、"
                "チャット履歴と取り込みの記録が材料なので、SQLでは答えられない。"
                "「このアプリはどれくらい使われている？」「誰が使っている？」"
                "「よく使われる機能は？」「どんな質問が多い？」「どこで失敗している？」"
                "「カタログのどこを直せばいい？」と聞かれたら使う。"
                "\n"
                "method で見る角度を選ぶ。まず summary で全体像を出し、"
                "気になった点を errors や users で掘るとよい。"
                "特に errors は失敗を「カタログを直せば減るもの」と"
                "「モデル・API側の問題」に分けて返すので、改善の打ち手を答えるときに使う。"
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "method": {
                        "type": "string",
                        "enum": [*usage.METHODS, "imports"],
                        "description": " / ".join(
                            f"{k}={v}" for k, v in usage.METHODS.items())
                        + " / imports=取り込みの実行実績",
                    },
                    "days": {"type": "integer",
                             "description": "直近何日ぶんを見るか。省略すると全期間。"},
                    "user": {"type": "string",
                             "description": "特定の利用者だけに絞る（ユーザー名）。任意。"},
                    "title": {"type": "string", "description": "見出し。"},
                    "purpose": {"type": "string",
                                "description": "この集計で確認したいことの短い説明。"},
                },
                "required": ["method"],
            },
        },
    },
    # ---- グラフ（用途別） ---------------------------------------------------
    *_chart_tools(),
    {
        "type": "function",
        "function": {
            "name": "describe_table",
            "description": (
                "テーブルの詳細（列・型・説明・コード値の意味・実値の分布・サンプル行）を取得する。"
                "初めて使うテーブルでSQLを書く前に呼んで、列名や値の実体を確認する。"
            ),
            "parameters": {
                "type": "object",
                "properties": {
                    "table": {"type": "string", "description": "テーブル名。"},
                },
                "required": ["table"],
            },
        },
    },
]


def _allow_result_id(node) -> None:
    """sql を受け取る所すべてに、SQL以外のデータの渡し方を足す。

    足すのは2つ。
      result_id … 前のツールが取ったデータをそのまま使う
      rows      … 表そのものを渡す（DBを介さないデータ）

    rows があることで、ナレッジ検索で調べた内容やAIが組み立てた表を、
    そのままExcel・CSV・グラフ・レポートにできる。これが無いと出力系が
    すべてSQL前提になり、DBを登録していない環境では何も出せない。

    レポートの節やExcelのシートのように、入れ子の中にも sql がある。
    1つずつ手で書き足すと必ず抜けるので、木をたどって機械的に付ける。
    """
    if isinstance(node, list):
        for v in node:
            _allow_result_id(v)
        return
    if not isinstance(node, dict):
        return
    props = node.get("properties")
    if isinstance(props, dict) and "sql" in props and "result_id" not in props:
        props["result_id"] = dict(_RESULT_ID)
        props["rows"] = dict(_INLINE_ROWS)
        # columns を自前の意味で宣言しているツールは上書きしない。
        # pivot_table の columns は「列に展開する1列名」（文字列）で、
        # ここで配列の宣言に置き換えると _coerce_lists が文字列を配列にし、
        # pandas 側で unhashable type: 'list' になる。
        if "columns" not in props:
            props["columns"] = dict(_INLINE_COLUMNS)
        req = node.get("required")
        if isinstance(req, list) and "sql" in req:
            node["required"] = [r for r in req if r != "sql"]
    for v in node.values():
        _allow_result_id(v)


_allow_result_id(BUILTIN_TOOLS)

# plot_chart は用途別のグラフツール（plot_comparison など）で完全に置き換えられる。
# 同じことが2通りでできると、AIはどちらを使うか毎回迷い、定義の文字数も倍かかる。
# 実処理は残してあるので、過去の会話やユーザー定義の上書きが壊れることはない。
_RETIRED = {"plot_chart"}
BUILTIN_TOOLS = [t for t in BUILTIN_TOOLS if t["function"]["name"] not in _RETIRED]


# ==========================================================================
# ===== 元 tools/business.py
# 業務でよく聞かれる分析のツール。
#
# business.py（期間比較・ファネル・コホート・併売）と
# advanced.py（異常検知・生存時間）をツールとして公開する。
# データ品質チェックだけは、SQLを受け取るのではなくDBそのものを見に行くので
# ここに実処理を置く。
# ==========================================================================
import advanced
import business
import catalog
import db

_compare_periods = _analysis_tool(lambda a, c, r: business.compare_periods(
    c, r, a.get("period_col"), a.get("value_col"),
    dimension_col=a.get("dimension_col"), current=a.get("current"),
    previous=a.get("previous"), qty_col=a.get("qty_col")))


_funnel_analysis = _analysis_tool(lambda a, c, r: business.funnel_analysis(
    c, r, a.get("steps") or [], labels=a.get("labels"),
    group_col=a.get("group_col")))


_cohort_analysis = _analysis_tool(lambda a, c, r: business.cohort_analysis(
    c, r, a.get("id_col"), a.get("period_col"), value_col=a.get("value_col"),
    max_periods=int(a.get("max_periods") or 12)))


_market_basket = _analysis_tool(lambda a, c, r: business.market_basket(
    c, r, a.get("transaction_col"), a.get("item_col"),
    min_support=float(a.get("min_support") or 1.0), top=int(a.get("top") or 25)))


_detect_anomalies = _analysis_tool(lambda a, c, r: advanced.detect_anomalies(
    c, r, a.get("time_col"), a.get("value_col"),
    window=int(a.get("window") or 7), threshold=float(a.get("threshold") or 3.0),
    season_length=int(a["season_length"]) if a.get("season_length") else None,
    changepoints=a.get("changepoints", True)))


_survival_analysis = _analysis_tool(lambda a, c, r: advanced.survival_analysis(
    c, r, a.get("duration_col"), event_col=a.get("event_col"),
    group_col=a.get("group_col")))


# =============================================================================
# データ品質チェック（DBを直接見る）
# =============================================================================

#: 1回のチェックで見るテーブルの上限。全部見ると時間がかかりすぎる。
_MAX_TABLES = 12
#: これより大きいテーブルでは、重い集計（種類数）を省く。
_HEAVY_ROWS = 200_000


def _q(alias: str, table: str) -> str:
    return f'{alias}."{table}"'


def _column_stats(scope: list[dict], alias: str, table: str, cols: list[dict],
                  rowcount: int | None) -> tuple:
    """1テーブルぶんの欠損・種類数を1本のSQLで数える。"""
    heavy = (rowcount or 0) <= _HEAVY_ROWS
    parts = ["COUNT(*) AS n"]
    for i, c in enumerate(cols):
        name = c["name"].replace('"', '""')
        parts.append(f'SUM(CASE WHEN "{name}" IS NULL THEN 1 ELSE 0 END) AS nul{i}')
        if heavy:
            parts.append(f'COUNT(DISTINCT "{name}") AS uni{i}')
        if str(c.get("type") or "").upper().startswith(("TEXT", "VARCHAR", "CHAR")):
            parts.append(f"SUM(CASE WHEN TRIM(\"{name}\") = '' THEN 1 ELSE 0 END) AS emp{i}")
        else:
            parts.append(f"0 AS emp{i}")
    sql = f"SELECT {', '.join(parts)} FROM {_q(alias, table)}"
    _, rows, _ = db.run_select(sql, scope, max_rows=1)
    got = rows[0]
    n = int(got[0] or 0)
    out, pos = [], 1
    for c in cols:
        nul = int(got[pos] or 0)
        pos += 1
        uni = int(got[pos] or 0) if heavy else None
        pos += 1 if heavy else 0
        emp = int(got[pos] or 0)
        pos += 1
        out.append({"column": c["name"], "type": c.get("type") or "",
                    "nulls": nul, "empty": emp, "unique": uni})
    return n, out


def _pk_duplicates(scope: list[dict], alias: str, table: str, pk: list) -> int | None:
    """主キーが重複している組み合わせの数。"""
    if not pk:
        return None
    keys = ", ".join(f'"{c}"' for c in pk)
    sql = (f"SELECT COUNT(*) FROM (SELECT {keys} FROM {_q(alias, table)} "
           f"GROUP BY {keys} HAVING COUNT(*) > 1)")
    try:
        _, rows, _ = db.run_select(sql, scope, max_rows=1)
        return int(rows[0][0] or 0)
    except Exception:
        return None


def _orphans(scope: list[dict], child: tuple, parent: tuple) -> int | None:
    """親に居ない子（孤立した外部キー）の件数。"""
    ca, ct, cc = child
    pa, pt, pc = parent
    sql = (f'SELECT COUNT(*) FROM {_q(ca, ct)} c WHERE c."{cc}" IS NOT NULL '
           f'AND NOT EXISTS (SELECT 1 FROM {_q(pa, pt)} p WHERE p."{pc}" = c."{cc}")')
    try:
        _, rows, _ = db.run_select(sql, scope, max_rows=1)
        return int(rows[0][0] or 0)
    except Exception:
        return None


def _date_range(scope: list[dict], alias: str, table: str, col: str) -> tuple | None:
    """日付列の最小・最大。データがいつまで入っているかを見る。"""
    try:
        _, rows, _ = db.run_select(
            f'SELECT MIN("{col}"), MAX("{col}") FROM {_q(alias, table)}',
            scope, max_rows=1)
        return rows[0][0], rows[0][1]
    except Exception:
        return None


def _looks_like_date(col: dict, sample: str = "") -> bool:
    name = str(col.get("name") or "").lower()
    return (str(col.get("type") or "").upper().startswith("DATE")
            or any(k in name for k in ("date", "日", "_at", "time", "月")))


def _data_quality(args: dict, scope: list[dict]) -> dict:
    """選択中のDBを見て、分析の前に気づいておくべき異常を洗い出す。"""
    if not scope:
        return _err("対象のDBがありません。")

    # テーブル名は '表名' でも 'DB名.表名' でもよい
    # （プロンプトが『DB名.テーブル名』で書くよう求めているので、後者で来ることが多い）
    want = [str(t) for t in (args.get("tables") or [])]
    def _wanted(alias, tname):
        return (not want) or tname in want or f"{alias}.{tname}" in want
    issues, tbl_rows, col_rows, ref_rows = [], [], [], []
    checked = 0

    # 結合定義はDBをまたぐので、スコープ全体で一度だけ組み立てる
    try:
        entries = [{"alias": e["alias"], "profile": catalog.profile_db(e["path"]),
                    "meta": e.get("meta") or catalog.load_meta(e["path"])} for e in scope]
        edges = catalog.collect_edges(entries)
    except Exception:
        edges = []

    for s in scope:
        alias = s["alias"]
        try:
            profile = catalog.profile_db(s["path"])
        except Exception as e:
            issues.append(("高", f"{alias}: プロファイルを読めませんでした（{e}）"))
            continue
        meta = s.get("meta") or catalog.load_meta(s["path"])
        allowed = set(s.get("tables") or profile.get("tables") or {})

        for tname, t in (profile.get("tables") or {}).items():
            if tname not in allowed:
                continue
            if not _wanted(alias, tname):
                continue
            if checked >= _MAX_TABLES:
                break
            checked += 1
            cols = t.get("columns") or []
            try:
                n, stats = _column_stats(scope, alias, tname, cols, t.get("row_count"))
            except Exception as e:
                issues.append(("中", f"{alias}.{tname}: 集計できませんでした（{e}）"))
                continue

            pk, pk_src = catalog.effective_pk(profile, meta, tname)
            dup = _pk_duplicates(scope, alias, tname, pk)
            tbl_rows.append([f"{alias}.{tname}", n, len(cols),
                             "、".join(pk) if pk else "（無し）",
                             dup if dup is not None else "—"])
            if n == 0:
                issues.append(("高", f"{alias}.{tname} は0行です。取り込みが済んでいない可能性があります。"))
                continue
            if dup:
                issues.append(("高", f"{alias}.{tname} は主キー（{'、'.join(pk)}）が "
                                     f"{dup} 組で重複しています。件数や金額が二重に数えられます。"))
            if not pk:
                issues.append(("中", f"{alias}.{tname} に主キーがありません。"
                                     "重複を検出できないので、カタログで指定してください。"))

            for st in stats:
                nul_pct = round(st["nulls"] / n * 100, 1) if n else 0.0
                emp_pct = round(st["empty"] / n * 100, 1) if n else 0.0
                col_rows.append([f"{alias}.{tname}", st["column"], st["type"],
                                 nul_pct, emp_pct,
                                 st["unique"] if st["unique"] is not None else "—"])
                if st["nulls"] == n:
                    issues.append(("中", f"{alias}.{tname}.{st['column']} は全て空です。"
                                         "この列は集計に使えません。"))
                elif nul_pct >= 30:
                    issues.append(("中", f"{alias}.{tname}.{st['column']} は "
                                         f"{nul_pct}% が空です。平均を取ると母数がずれます。"))
                elif emp_pct >= 10:
                    issues.append(("低", f"{alias}.{tname}.{st['column']} は "
                                         f"{emp_pct}% が空文字です。NULLと混在しています。"))
                if st["unique"] == 1 and n > 1:
                    issues.append(("低", f"{alias}.{tname}.{st['column']} は1種類の値しかありません。"))

            # データがいつまで入っているか（古いまま気づかないのを防ぐ）
            for c in cols:
                if not _looks_like_date(c):
                    continue
                rng = _date_range(scope, alias, tname, c["name"])
                if rng and rng[1]:
                    tbl_rows[-1].append(f"{c['name']}: {rng[0]} 〜 {rng[1]}")
                    break

        # 参照整合性。カタログの結合定義とDBのFK宣言の両方を見る。
        # 子と親は保存順ではなく主キーの位置から決める（手書きのYAMLが
        # 逆向きでも、「親に居ない子」を正しい向きで数えるため）
        for edge in edges:
            (ca, ct, cc), (pa, pt, pc) = catalog.child_parent(entries, edge)
            if ca != alias or ct not in allowed:
                continue
            miss = _orphans(scope, (ca, ct, cc), (pa, pt, pc))
            if miss is None:
                continue
            kind = "FK宣言" if edge.get("kind") == "fk" else "カタログの結合定義"
            ref_rows.append([f"{ca}.{ct}.{cc}", f"{pa}.{pt}.{pc}", miss, kind])
            if miss:
                issues.append(("高", f"{ca}.{ct}.{cc} の {miss} 件が "
                                     f"{pa}.{pt} に存在しません。"
                                     "内部結合すると、この件数ぶん落ちます。"))

    if not tbl_rows:
        # 実在する名前を例として出す。架空の例を出すと、AIがそれを真似て再び失敗する
        avail = "、".join(f"{s.get('alias')}.{t}" for s in scope
                          for t in (s.get("tables") or [])[:5])
        return _err("調べられるテーブルがありませんでした。"
                    "tables に指定した名前が合っているか確認してください"
                    + (f"（指定できる例: {avail}）。" if avail else "。"))

    head = ["テーブル", "行数", "列数", "主キー", "主キー重複"]
    if any(len(r) > 5 for r in tbl_rows):
        head.append("日付の範囲")
    tbl_rows = [r + [""] * (len(head) - len(r)) for r in tbl_rows]

    rank = {"高": 0, "中": 1, "低": 2}
    issues.sort(key=lambda x: rank.get(x[0], 3))
    tables = [_table_of("テーブル", head, tbl_rows),
              _table_of("見つかった問題", ["深刻度", "内容"],
                        [[lv, msg] for lv, msg in issues[:80]] or [["—", "問題は見つかりませんでした。"]]),
              _table_of("列ごとの状態", ["テーブル", "列", "型", "空の割合(%)",
                                        "空文字の割合(%)", "値の種類数"], col_rows)]
    if ref_rows:
        tables.append(_table_of("参照整合性", ["子", "親", "親に無い件数", "定義元"], ref_rows))

    high = [m for lv, m in issues if lv == "高"]
    notes = [f"{checked} テーブルを調べました。"
             + (f"深刻な問題が {len(high)} 件あります。" if high
                else "分析を止めるような問題は見つかりませんでした。")]
    notes += high[:5]
    if checked >= _MAX_TABLES:
        notes.append(f"テーブルが多いため {_MAX_TABLES} 件までにしています。"
                     "続きは tables で対象を指定してください。")
    notes.append("行数が0・主キーの重複・親に無い外部キーは、集計結果を直接ゆがめます。"
                 "先にここを直してから数字を読んでください。")

    return _report_result({"title": "データ品質チェック", "tables": tables, "notes": notes,
                           "meta": {"tables_checked": checked,
                                    "issues": len(issues), "critical": len(high)}},
                          scope=scope)


def _table_of(name: str, columns: list, rows: list) -> dict:
    return {"name": name, "columns": columns, "rows": [tuple(r) for r in rows]}


HANDLERS_business = {
    "compare_periods": _compare_periods,
    "funnel_analysis": _funnel_analysis,
    "cohort_analysis": _cohort_analysis,
    "market_basket": _market_basket,
    "detect_anomalies": _detect_anomalies,
    "survival_analysis": _survival_analysis,
    "data_quality": _data_quality,
}

# data_quality は自分でDBを見に行くので、SQLプレビューの対象外
SQL_TOOLS_business = {"compare_periods", "funnel_analysis", "cohort_analysis",
             "market_basket", "detect_anomalies", "survival_analysis"}


# ==========================================================================
# ===== 元 tools/files.py
# 取り込み元フォルダを調べるツール（読むだけ）。
#
# 「取り込み元に何が来ているか」「このCSVはどんな列か」「まだ取り込んでいない
# ファイルはあるか」に答えるためのもの。DBに入る前のファイルの話なので、
# SQLでは答えられない。
#
# 安全のうえで大事なところは、すべて importer.py の既存の仕組みに任せる。
#   allowed_dirs()   … 読んでよいフォルダ（env の IMPORT_DIRS ＋画面で追加した分）
#   is_allowed()     … .. やリンクで許可フォルダの外へ出ようとしても弾く
#   check_readable() … 読む直前にもう一度確かめる
# ここで新しくパスの判定を書かない（守りの仕組みを二重に持つと必ずズレる）。
#
# できるのは一覧と下見だけ。取り込み・作成・変更・削除は一切しない。
# 取り込みの実行は今までどおり「データ取り込み」画面の操作に限る。
# ==========================================================================
from datetime import datetime
from pathlib import Path

import config
import filecheck
import history
import importer

#: 一覧で返す最大件数。多すぎるとLLMに渡すだけで無駄になる。
_MAX_ROWS = 300
#: 下見で見せる行数の上限。
_MAX_PREVIEW_ROWS = 20
#: 一覧でまとめて形を判定するときの上限。1件ずつ開くので数を抑える。
_MAX_CHECK = 20


def _files_table(name: str, columns: list, rows: list) -> dict:
    return {"name": name, "columns": columns, "rows": [tuple(r) for r in rows]}


def _size(n: int | None) -> str:
    if n is None:
        return ""
    if n < 1024:
        return f"{n} B"
    if n < 1024 * 1024:
        return f"{n / 1024:.1f} KB"
    return f"{n / 1024 / 1024:.1f} MB"


def _mtime(p: Path) -> str:
    try:
        return datetime.fromtimestamp(p.stat().st_mtime).strftime("%Y-%m-%d %H:%M")
    except OSError:
        return ""


def _imported_index() -> dict:
    """取り込み済みファイルの索引。パスの表記ゆれに備えて小文字で引く。"""
    out = {}
    for src, rec in history.latest_by_source().items():
        out[src.strip().lower()] = rec
    return out


def _import_state(p: Path, index: dict) -> str:
    """そのファイルが取り込み済みかどうかの一言。"""
    rec = index.get(str(p).lower())
    if rec is None:
        # パスが変わっていても、ファイル名が一致すれば手がかりにはなる
        rec = next((r for key, r in index.items()
                    if Path(key).name == p.name.lower()), None)
        if rec is None:
            return "未取り込み"
        return (f"同名を取り込み済み（{rec.get('db_file')} / {rec.get('table')}"
                f"・{rec.get('at', '')[:16]}）")
    mark = "" if rec.get("ok") else "／前回は失敗"
    return f"{rec.get('db_file')} / {rec.get('table')}（{rec.get('at', '')[:16]}{mark}）"


def _roots_table() -> dict:
    """許可フォルダの状態。マウント切れや権限なしをここで切り分ける。"""
    rows = [[r["設定値"], r["状態"], r["source"]] for r in importer.dir_status()]
    return _files_table("取り込み元フォルダ", ["フォルダ", "状態", "設定元"], rows)


def _supported(p: Path) -> bool:
    return p.suffix.lower() in config.IMPORT_EXTENSIONS


def _listing(args: dict) -> dict:
    path = str(args.get("path") or "").strip()
    recursive = bool(args.get("recursive"))
    pattern = str(args.get("pattern") or "").strip().lower()
    only_new = bool(args.get("only_not_imported"))
    check = bool(args.get("check"))
    index = _imported_index()

    roots = importer.allowed_dirs()
    if not roots:
        return _err("取り込み元フォルダが設定されていません。"
                    "「データ取り込み」画面で追加するか、env の IMPORT_DIRS を設定してください。")

    if path and not importer.is_allowed_dir(Path(path)):
        return _err(f"そのフォルダは見られません（許可フォルダの外です）: {path}。"
                    f"見られるのは {'、'.join(str(d) for d in roots)} の中だけです。")

    here = Path(path).resolve() if path else None
    dirs: list = []
    files: list[Path] = []

    # 一覧は拡張子で絞らない。「何が置いてあるか」を知るのが目的なので、
    # 取り込めない形式（PDFなど）も見せて、可否は列で示す。
    if recursive:
        for p in importer.list_all_files():
            if here is None or here == p.parent or here in p.parents:
                files.append(p)
    else:
        for d in ([here] if here is not None else roots):
            try:
                entries = sorted(d.iterdir(), key=lambda x: (not x.is_dir(), x.name.lower()))
            except OSError:
                continue
            for p in entries:
                try:
                    if p.is_dir():
                        if not importer.is_noise(p.name):
                            dirs.append(str(p.relative_to(d)) if here is not None
                                        else f"{d.name}/{p.name}")
                    elif importer.is_within_allowed(p):
                        files.append(p.resolve())
                except OSError:
                    continue

    if pattern:
        files = [f for f in files if pattern in f.name.lower()]
        dirs = [d for d in dirs if pattern in str(d).lower()]

    rows, checked = [], 0
    for p in files:
        state = _import_state(p, index)
        if only_new and state != "未取り込み":
            continue
        if not _supported(p):
            shape = "取り込み対象外の形式"
        elif check and checked < _MAX_CHECK:
            checked += 1
            try:
                shape = filecheck.summary_line(filecheck.inspect_file(p))
            except Exception as e:
                shape = f"判定できず（{type(e).__name__}）"
        else:
            shape = "未判定"
        rows.append([importer.display_name(p), p.suffix.lower().lstrip(".") or "（なし）",
                     _size(p.stat().st_size if p.exists() else None), _mtime(p),
                     state, shape])
    rows.sort(key=lambda r: r[3], reverse=True)      # 新しいものから
    shown = rows[:_MAX_ROWS]

    tables = [_roots_table()]
    if dirs:
        tables.append(_files_table("フォルダ", ["名前"], [[d] for d in sorted(set(dirs))]))
    tables.append(_files_table("ファイル",
                         ["場所", "種類", "サイズ", "更新日時", "取り込み状況", "表として使えるか"],
                         shown))

    supported = [r for r in rows if r[5] != "取り込み対象外の形式"]
    notes = [f"{len(rows)} 件のファイルが見つかりました"
             + (f"（多いので新しい順に {len(shown)} 件だけ載せています）"
                if len(rows) > len(shown) else "")
             + f"。うち取り込みに対応した形式は {len(supported)} 件です"
             + f"（対応: {'、'.join(config.IMPORT_EXTENSIONS)}）。"]
    if not recursive and not path:
        notes.append("いまは許可フォルダの直下だけを見ています。"
                     "下の階層も見るなら recursive=true、"
                     "特定のフォルダを見るなら path を指定してください。")
    if not check:
        notes.append("「表として使えるか」は check=true を付けると調べます"
                     "（1件ずつ開くので、多いときは絞ってから）。")
    elif checked >= _MAX_CHECK:
        notes.append(f"判定は {_MAX_CHECK} 件までにしています。"
                     "pattern や path で絞ると続きを見られます。")
    bad = [r for r in rows if r[5].startswith(("取り込みに向かない", "手直しが要る"))]
    if bad:
        notes.append("そのままでは取り込めないものがあります: "
                     + "、".join(f"{r[0]}（{r[5]}）" for r in bad[:3]))
    fresh = [r for r in rows if r[4] == "未取り込み"]
    if fresh:
        notes.append(f"まだ取り込んでいないファイルが {len(fresh)} 件あります: "
                     + "、".join(r[0] for r in fresh[:5])
                     + ("（ほか）" if len(fresh) > 5 else ""))
    if not rows and not dirs:
        notes.append("ファイルはありませんでした。フォルダが空か、権限が無い可能性があります。")
    notes.append("中身と形を詳しく見るには file にパスを指定してください。"
                 "実際に取り込むのは「データ取り込み」画面の操作です。ここでは読むだけです。")

    return _report_result({"title": "取り込み元フォルダの中身",
                           "tables": tables, "notes": notes,
                           "meta": {"files": len(rows), "supported": len(supported),
                                    "not_imported": len(fresh)}})


def _issues_table(res: dict) -> dict:
    return _files_table("見つかった問題", ["深刻度", "内容", "直し方"],
                  [[i["level"], i["text"], i["fix"]] for i in res["issues"]]
                  or [["—", "気になる点はありませんでした。", ""]])


def _preview(args: dict) -> dict:
    raw = Path(str(args.get("file") or "").strip())
    # 場所の判定と「読める形式か」の判定を分ける。対象外の形式でも
    # 「なぜ取り込めないか」は答えられるようにする。
    if not importer.is_within_allowed(raw):
        return _err(f"そのファイルは見られません: {raw}。"
                    "許可フォルダの中のファイルだけを指定してください。"
                    "パスは一覧（file を指定しない呼び方）で得たものを使ってください。")
    target = raw.resolve()

    if not _supported(target):
        res = filecheck.inspect_file(target)
        return _report_result({
            "title": f"{target.name} は取り込みに対応していない形式です",
            "tables": [_issues_table(res)],
            "notes": [f"{importer.display_name(target)}"
                      f"（{_size(target.stat().st_size)}・更新 {_mtime(target)}）",
                      f"扱えるのは {'、'.join(config.IMPORT_EXTENSIONS)} です。"
                      "中身は読んでいません。"],
            "meta": {"verdict": res["verdict"]}})

    rows_want = max(1, min(int(args.get("rows") or 5), _MAX_PREVIEW_ROWS))

    try:
        sheets = importer.sheet_names(target)
    except importer.ImportError_ as e:
        return _err(str(e))
    sheet = args.get("sheet") or (sheets[0] if sheets else None)

    # まず形を見る。見出しが1行目に無ければ、その行で読み直す
    try:
        res = filecheck.inspect_file(target, sheet=sheet)
    except Exception as e:
        res = {"verdict": "判定できず", "issues": [
            {"level": "低", "text": f"形を調べられませんでした: {e}", "fix": ""}],
            "header_row": 0, "shape": {}}
    header_row = (int(args["header_row"]) if args.get("header_row") is not None
                  else int(res.get("header_row") or 0))

    try:
        df = importer.read_table(target, sheet=sheet, header_row=header_row,
                                 nrows=rows_want)
    except importer.ImportError_ as e:
        return _err(f"{target.name}: {e}")
    except Exception as e:
        return _err(f"{target.name} を読めませんでした: {e}")

    plan = importer.plan_columns(df)
    tables = [
        _files_table("判定", ["項目", "内容"],
               [["そのまま取り込めるか", res["verdict"]],
                ["見出しの行", f"{header_row + 1} 行目"],
                *[[k, v] for k, v in (res.get("shape") or {}).items() if k != "見出し行"]]),
        _issues_table(res),
        _files_table("列", ["元の列名", "取り込み後の列名", "推定される型"],
               [[c["元の列名"], c["列名"], c["型"]] for c in plan]),
        _files_table(f"先頭 {len(df)} 行", [str(c) for c in df.columns],
               [[("" if v is None else v) for v in r] for r in df.values.tolist()]),
    ]
    if sheets:
        tables.insert(0, _files_table("シート", ["シート名"], [[s] for s in sheets]))

    index = _imported_index()
    notes = [f"{importer.display_name(target)}（{_size(target.stat().st_size)}"
             f"・更新 {_mtime(target)}）",
             f"判定: {res['verdict']}",
             f"取り込み状況: {_import_state(target, index)}"]
    for i in res.get("issues", []):
        if i["level"] == "高":
            notes.append(f"{i['text']} → {i['fix']}")
    if sheets:
        notes.append(f"シートは {len(sheets)} 枚あります（いま見ているのは「{sheet}」）。"
                     "シートごとに形が違うので、使いたいシートを sheet で指定して確かめてください。")
    if res.get("encoding"):
        notes.append(f"文字コード: {res['encoding']} / 区切り: {res.get('delimiter')}")
    if header_row:
        notes.append(f"見出しが1行目ではないので、{header_row + 1} 行目を見出しとして読みました。"
                     f"取り込み画面でも「見出しの行」に {header_row + 1} を指定してください。")
    if df.empty:
        notes.append("中身の行が読めませんでした。header_row を変えて試してください。")
    notes.append("型は中身からの推定で、取り込み画面で直せます。"
                 "ここでは読むだけで、取り込みはしていません。")

    return _report_result({"title": f"{target.name} の下見（{res['verdict']}）",
                           "tables": tables, "notes": notes,
                           "meta": {"verdict": res["verdict"], "columns": len(plan),
                                    "header_row": header_row, "sheets": sheets}})


def _explore_import_files(args: dict, scope: list[dict]) -> dict:
    try:
        return _preview(args) if str(args.get("file") or "").strip() else _listing(args)
    except importer.ImportError_ as e:
        return _err(str(e))
    except PermissionError:
        return _err("読み取り権限がありません。共有フォルダの権限を確認してください。")
    except OSError as e:
        return _err(f"フォルダにアクセスできませんでした: {e.strerror or e}")


HANDLERS_files = {"explore_import_files": _explore_import_files}

# SQLは受け取らない
SQL_TOOLS_files: set = set()

# 管理者だけに渡すツール。「データ取り込み」画面が管理者専用なので、
# AI経由なら誰でも中身が見られる、という抜け道を作らない。
ADMIN_TOOLS_files = {"explore_import_files"}


# ==========================================================================
# ===== 元 tools/mail.py
# メールの宛先探しと下書き。送信そのものは画面のボタンからだけ。
# ==========================================================================
import mailer


def _find_mail_recipients(args: dict, scope: list[dict]) -> dict:
    res = mailer.find_recipients(scope, args.get("query") or "",
                                 limit=int(args.get("limit") or 50),
                                 table=args.get("table"))
    cands = res["candidates"]
    return {
        "ok": res["ok"],
        "llm_content": _json({
            "status": "recipients", "message": res["message"],
            "sources": res["sources"],
            "count": len(cands),
            "candidates": [{k: c[k] for k in ("email", "name", "dept", "source")}
                           for c in cands[:50]],
            "note": "ここに出たアドレスだけを宛先に使うこと。推測で作らない。",
        }),
        "render": {"role": "assistant", "kind": "table",
                   "columns": ["メールアドレス", "氏名", "部署", "出所"],
                   "rows": [[c["email"], c["name"], c["dept"], c["source"]]
                            for c in cands]} if cands else
                  {"role": "assistant", "kind": "text", "content": res["message"]},
    }


def _compose_email(args: dict, scope: list[dict]) -> dict:
    to = list(args.get("to") or [])
    matched = []
    if args.get("to_query"):
        res = mailer.find_recipients(scope, args["to_query"], limit=50)
        matched = res["candidates"]
        to += [c["email"] for c in matched if c["valid"]]
        if not matched:
            return _err(f"「{args['to_query']}」に一致する宛先が見つかりませんでした。"
                        "find_mail_recipients で候補を確認してください。")
    to = list(dict.fromkeys(a for a in to if a))
    draft = {"to": to, "cc": args.get("cc") or [], "bcc": args.get("bcc") or [],
             "subject": args.get("subject") or "", "body": args.get("body") or "",
             "reply_to": args.get("reply_to") or "",
             "attach_filenames": list(args.get("attach_filenames") or [])}
    view = mailer.preview(draft)
    # 添付は会話ログから web 側が解決する。ここでは名前だけ持たせる。
    return {
        "ok": not view["errors"],
        "llm_content": _json({
            "status": "mail_draft" if not view["errors"] else "mail_draft_invalid",
            "to": view["to"], "cc": view["cc"], "bcc_count": len(view["bcc"]),
            "subject": view["subject"], "body_lines": view["body_lines"],
            "attach_filenames": draft["attach_filenames"],
            "matched_from_db": [{"email": c["email"], "name": c["name"],
                                 "dept": c["dept"]} for c in matched[:20]],
            "problems": view["errors"],
            "note": ("下書きを画面に出した。送信するかはユーザーが画面のボタンで決める。"
                     "こちらから送信することはできないので、"
                     "『内容を確認して送信ボタンを押してください』と伝えること。"),
        }),
        "render": {"role": "assistant", "kind": "mail_draft", "draft": draft,
                   "preview": view},
    }

HANDLERS_mail = {
    "find_mail_recipients": _find_mail_recipients,
    "compose_email": _compose_email,
}

SQL_TOOLS_mail: set[str] = set()


# ==========================================================================
# ===== 元 tools/query.py
# 調べる・集計する・描く・出す。SQLの結果をそのまま扱うツール。
# ==========================================================================
from pathlib import Path

import advanced
import analysis
import catalog
import charts
import config
import excel
import exports


def _example_exists(sql: str) -> bool:
    """このSQLと同じ例文がどこかのDBに登録済みか。"""
    if not sql.strip():
        return False
    import db as dbmod
    for f in dbmod.list_db_files():
        if catalog.find_example(catalog.load_meta(f).get("examples") or [], sql):
            return True
    return False


def _run_sql_query(args: dict, scope: list[dict]) -> dict:
    try:
        columns, rows, truncated, rid, total = fetch(args, scope,
                                                     label=args.get("purpose"))
    except advanced.AnalysisError as e:
        return _err(str(e))
    except Exception as e:
        return _err(f"SQL実行エラー: {e}")

    sample = rows[: config.SAMPLE_ROWS_FOR_LLM]
    llm_content = _json({
        "columns": columns,
        "row_count": len(rows),
        "rows": [list(r) for r in sample],
        "result_id": rid,
        "note": (f"全{len(rows)}行中 先頭{len(sample)}行を表示。"
                 f"この結果は result_id '{rid}' で他のツールから使い回せます"
                 "（同じSQLを書き直さなくてよい）。"
                 if len(rows) > len(sample) else
                 f"この結果は result_id '{rid}' で他のツールから使い回せます。"),
        **source_note(len(rows), truncated, total),
        # AIが「例文に登録しますか？」と聞くべきかの手がかり。
        # 既に同じSQLの例文があれば聞かない（同じ提案の繰り返しは邪魔になる）。
        "example_registered": _example_exists(str(args.get("sql") or "")),
    })
    return {
        "ok": True,
        "llm_content": llm_content,
        "render": {
            "role": "assistant", "kind": "table",
            "columns": columns, "rows": rows, "truncated": truncated,
        },
    }


_CHART_FIELDS = ("chart_type", "x", "y", "color", "size", "text", "path",
                 "nbins", "orientation", "barmode", "title",
                 # 種別ごとに使う指定
                 "y2", "z", "lower", "upper", "facet", "dimensions",
                 "source", "target", "start", "end",
                 "open", "high", "low", "close",
                 "value", "agg", "max", "suffix", "valueformat",
                 "colorscale", "marginal", "trendline")


def _plot_chart(args: dict, scope: list[dict]) -> dict:
    try:
        columns, rows, truncated, rid, total = fetch(args, scope,
                                                     label=args.get("title"))
    except advanced.AnalysisError as e:
        return _err(str(e))
    except Exception as e:
        return _err(f"グラフ用SQLの実行エラー: {e}")

    item = {k: args.get(k) for k in _CHART_FIELDS}
    item["chart_type"] = item.get("chart_type") or "bar"
    errs = charts.validate(item, columns)
    if errs:
        return _err(" / ".join(errs))

    return {
        "ok": True,
        "llm_content": _json({
            "status": "chart_rendered",
            "chart_type": item["chart_type"],
            "columns": columns,
            "row_count": len(rows),
            "result_id": rid,
            **source_note(len(rows), truncated, total),
        }),
        "render": {
            "role": "assistant", "kind": "chart",
            "columns": columns, "rows": rows, **item,
            "title": args.get("title", ""),
            **render_source_note(truncated, total),
        },
    }


def _plot_dual_axis(args: dict, scope: list[dict]) -> dict:
    try:
        columns, rows, truncated, rid, total = fetch(args, scope,
                                                     label=args.get("title"))
    except advanced.AnalysisError as e:
        return _err(str(e))
    except Exception as e:
        return _err(f"2軸グラフ用SQLの実行エラー: {e}")

    x = args.get("x")
    bar_y = args.get("bar_y") or []
    line_y = args.get("line_y") or []
    needed = [x] + list(bar_y) + list(line_y)
    missing = [c for c in needed if c and c not in columns]
    if missing or not bar_y or not line_y:
        msg = (f"指定列が結果に存在しません: {missing} / 利用可能な列: {columns}"
               if missing else "bar_y と line_y にはそれぞれ1つ以上の数値列を指定してください。")
        return _err(msg)

    return {
        "ok": True,
        "llm_content": _json({
            "status": "dual_axis_chart_rendered",
            "columns": columns, "row_count": len(rows),
            "bar_y": bar_y, "line_y": line_y, "result_id": rid,
            **source_note(len(rows), truncated, total),
        }),
        "render": {
            "role": "assistant", "kind": "chart_dual",
            "columns": columns, "rows": rows,
            "x": x, "bar_y": bar_y, "line_y": line_y,
            "left_title": args.get("left_title"), "right_title": args.get("right_title"),
            "title": args.get("title", ""),
        },
    }


def _describe_table(args: dict, scope: list[dict]) -> dict:
    text = catalog.describe_table_text(scope, args.get("db", ""), args.get("table", ""))
    return {"ok": not text.startswith("エラー"), "llm_content": text, "render": None}


def _pivot_table(args: dict, scope: list[dict]) -> dict:
    try:
        columns, rows, truncated, rid, total = fetch(args, scope,
                                                     label=args.get("title"))
    except advanced.AnalysisError as e:
        return _err(str(e))
    except Exception as e:
        return _err(f"クロス集計用SQLの実行エラー: {e}")
    if not rows:
        return _err("データが0行でした。抽出条件を見直してください。")

    # columns は「列に展開する1列名」（文字列）。過去の宣言の名残やAIの癖で
    # 要素1つの配列で来ることがあるので、中身を取り出して受け付ける。
    pivot_col = args.get("columns") or None
    if isinstance(pivot_col, (list, tuple)):
        if len(pivot_col) > 1:
            return _err("columns には列に展開する列名を1つだけ指定してください"
                        f"（{list(pivot_col)} と複数指定されています）。")
        pivot_col = str(pivot_col[0]).strip() or None

    try:
        cols, prows = analysis.pivot(
            columns, rows,
            index=args.get("index") or [], cols=pivot_col,
            values=args.get("values"), aggfunc=args.get("aggfunc") or "sum",
            margins=bool(args.get("margins")),
            percent=args.get("percent"), rank_by=args.get("rank"),
        )
    except Exception as e:
        return _err(f"クロス集計に失敗しました: {e}")

    title = args.get("title") or "クロス集計"
    # 集計後の表もグラフやレポートの材料になるので、指せるようにして返す
    out_rid = results.put(scope, cols, prows, label=f"{title}（クロス集計の結果）")
    llm_content = _json({
        "status": "pivot_ready", "columns": cols, "row_count": len(prows),
        "rows": [list(r) for r in prows[: config.SAMPLE_ROWS_FOR_LLM]],
        "result_id": out_rid, "source_result_id": rid,
        "note": f"集計後の表は result_id '{out_rid}' でグラフやレポートに渡せます。",
        **source_note(len(rows), truncated, total),
    })
    if (args.get("render") or "table") == "heatmap":
        render = {"role": "assistant", "kind": "chart", "columns": cols, "rows": prows,
                  "chart_type": "matrix", "x": cols[0], "title": title}
    else:
        render = {"role": "assistant", "kind": "table", "columns": cols, "rows": prows,
                  "truncated": False}
    return {"ok": True, "llm_content": llm_content, "render": render}


def _analyze_stats(args: dict, scope: list[dict]) -> dict:
    method = args.get("method") or "describe"
    try:
        columns, rows, truncated, rid, total = fetch(args, scope,
                                                     label=args.get("title"))
    except advanced.AnalysisError as e:
        return _err(str(e))
    except Exception as e:
        return _err(f"分析用SQLの実行エラー: {e}")
    if not rows:
        return _err("データが0行でした。抽出条件を見直してください。")

    title = args.get("title") or {"describe": "基本統計量", "correlation": "相関",
                                  "outliers": "外れ値"}.get(method, "分析")
    note = f"（元データ {len(rows):,} 行"
    note += "／上限で切り詰め済み）" if truncated else "）"

    try:
        if method == "describe":
            cols, srows = _advanced_describe(columns, rows, args.get("columns"),
                                            args.get("group_by"))
            extra = {}
            render = {"role": "assistant", "kind": "table", "columns": cols, "rows": srows}
        elif method == "correlation":
            cm = args.get("corr_method") or "pearson"
            lag = int(args.get("lag") or 0)
            if args.get("partial"):
                # 交絡を取り除いた相関。「効いて見えるのは第3の変数のせい」を切り分ける
                res = advanced.partial_correlation(
                    columns, rows, args.get("columns"), args.get("control") or [],
                    method=cm)
                return _report_result(res, source_rows=len(rows), truncated=truncated,
                                      total=total, result_id=rid, scope=scope,
                                      extra={"method": "partial_correlation"})
            if lag:
                # 時差相関。「広告費は翌月の売上に効く」を見るための道具
                res = advanced.lag_correlation(
                    columns, rows, args.get("target"), args.get("columns"),
                    max_lag=lag, method=cm)
                return _report_result(res, source_rows=len(rows), truncated=truncated,
                                      total=total, result_id=rid, scope=scope,
                                      extra={"method": "lag_correlation"})
            cols, srows = analysis.correlation(columns, rows, args.get("columns"), cm)
            extra = {"method": cm, "strong_pairs": analysis.correlation_pairs(columns, rows, cm)[:8],
                     "caution": "相関は因果ではありません。効いている理由を確かめるには "
                                "partial=true で交絡を除くか、regression を使ってください。"}
            render = {"role": "assistant", "kind": "chart", "columns": cols, "rows": srows,
                      "chart_type": "matrix", "x": cols[0], "title": f"{title}{note}",
                      "colorscale": "RdBu"}
        else:
            target = args.get("target")
            if not target:
                return _err("outliers には target（外れ値を調べる数値列）が必要です。")
            om = args.get("outlier_method") or "iqr"
            # mahalanobis は複数列をまとめて見る。"売上, 客数" のような指定も許す。
            cols_in = [t.strip() for t in str(target).split(",")] if isinstance(target, str) \
                else list(target)
            res = advanced.outliers_ext(columns, rows,
                                        cols_in if len(cols_in) > 1 else cols_in[0],
                                        method=om, threshold=args.get("threshold"))
            return _report_result(res, source_rows=len(rows), truncated=truncated,
                                  total=total, result_id=rid, scope=scope,
                                  extra={"method": "outliers", "outlier_method": om})
    except Exception as e:
        return _err(f"{title}の計算に失敗しました: {e}")

    out_rid = results.put(scope, cols, srows, label=title)
    return {
        "ok": True,
        "llm_content": _json({
            "status": "stats_ready", "method": method, "columns": cols,
            "row_count": len(srows),
            "rows": [list(r) for r in srows[: config.SAMPLE_ROWS_FOR_LLM]],
            "result_id": out_rid, "source_result_id": rid,
            **source_note(len(rows), truncated, total), **extra,
        }),
        "render": render,
    }


def _propose_glossary_term(args: dict, scope: list[dict]) -> dict:
    """業務用語の登録カードをチャットに出す（提案まで。保存は人がボタンで確定）。

    メールと同じ型: 作るのはAI、確定するのは人。カタログは全員共通の土台なので、
    会話の「はい」だけで書き換えず、カードのボタンという明示の操作を挟む。
    SQL式があれば実データで確かめてから見せる（通らない式を提案しない）。
    """
    import db as dbmod

    term = str(args.get("term") or "").strip()
    desc = str(args.get("description") or "").strip()
    sql = str(args.get("sql") or "").strip()
    table = str(args.get("table") or "").strip()
    if not term or not desc:
        return _err("term（用語）と description（説明）は必須です。")

    # 置き場のDBを決める。table を持つDBを探す（複数あれば db で指定させる）
    name = str(args.get("db") or "").strip()
    files = dbmod.list_db_files()
    target = None
    if name:
        low = name.lower()
        target = next((f for f in files
                       if low in (f.name.lower(), f.stem.lower(),
                                  dbmod.alias_for(f).lower())), None)
        if target is None:
            return _err(f"DB '{name}' が見つかりません。")
    elif table:
        owners = [f for f in files
                  if table in (catalog.profile_db(f).get("tables") or {})]
        if len(owners) == 1:
            target = owners[0]
        elif not owners:
            return _err(f"テーブル '{table}' を持つDBが見つかりません。")
        else:
            return _err(f"テーブル '{table}' は複数のDBにあります。"
                        "db でどのDBか指定してください: "
                        + "、".join(f.name for f in owners))
    else:
        return _err("table（用語を置くテーブル）か db を指定してください。")

    alias = dbmod.alias_for(target)
    if table and table not in (catalog.profile_db(target).get("tables") or {}):
        return _err(f"{target.name} にテーブル '{table}' がありません。")

    # SQL式の検証。条件式→該当件数 / 計算式→計算例 / 通らない→エラーで差し戻し
    verdict, detail = "", ""
    if sql:
        ref = f"{alias}.{table}" if table else None
        wide = dbmod.widen_scope(f"{alias}. {sql}", [
            {"path": str(target), "alias": alias, "name": target.name, "tables": None}])
        wide = dbmod.widen_scope(sql, wide)
        try:
            if ref:
                _, rows, _ = dbmod.run_select(
                    f"SELECT COUNT(*) AS n, (SELECT COUNT(*) FROM {ref}) AS t "
                    f"FROM {ref} WHERE {sql}", wide, max_rows=1)
                n, t = rows[0]
                verdict = "条件式"
                detail = f"該当 {n:,} 行 / 全 {t:,} 行"
            else:
                _, rows, _ = dbmod.run_select(f"SELECT {sql} AS v", wide, max_rows=1)
                verdict, detail = "計算式", f"計算結果: {rows[0][0]}"
        except Exception:
            try:
                src = ref or f"{alias}.{(list(catalog.profile_db(target)['tables']) or [''])[0]}"
                _, rows, _ = dbmod.run_select(f"SELECT {sql} AS v FROM {src}",
                                              wide, max_rows=1)
                verdict, detail = "計算式", f"計算結果の例: {rows[0][0]}"
            except Exception as e:
                return _err(f"SQL式が実データで通りませんでした: {str(e).splitlines()[0][:120]} "
                            "式を直して提案し直すか、SQL式なし（説明だけ）で提案してください。")

    meta_now = catalog.load_meta(target)
    current = (catalog.table_glossary(meta_now, table) if table
               else catalog.db_glossary(meta_now)).get(term)
    return {
        "ok": True,
        "llm_content": _json({
            "status": "proposed", "db": target.name, "table": table or "(DB全体)",
            "term": term, "verdict": verdict or "説明のみ", "detail": detail,
            "already_exists": current is not None,
            "note": "登録カードをユーザーの画面に出しました。登録するかはユーザーが"
                    "カードのボタンで決めます。あなたはこれ以上の操作をしなくてよい。",
        }),
        "render": {"role": "assistant", "kind": "glossary_term",
                   "db": target.name, "table": table, "term": term,
                   "description": desc, "sql": sql,
                   "how": str(args.get("how") or "").strip(),
                   "verdict": verdict, "detail": detail,
                   "exists": current is not None,
                   "old": current or None},
    }


def _propose_example(args: dict, scope: list[dict]) -> dict:
    """例文（質問とSQLのペア）の登録カードをチャットに出す。保存は人がボタンで確定。

    カードにはSQLは出さない。代わりに「どのデータをどこから取って、どう集計したか」の
    日本語（summary。AIが書く）と、実際に動かした結果の先頭数行を載せる。
    SQLを読めない人でも、中身を見て正しさを判断できるようにするため。
    """
    import db as dbmod

    q = str(args.get("question") or "").strip()
    sql = str(args.get("sql") or "").strip()
    summary = str(args.get("summary") or "").strip()
    if not q or not sql or not summary:
        return _err("question（質問文）・sql・summary（何をどう集計したかの日本語）は必須です。")

    wide = dbmod.widen_scope(sql, scope)
    try:
        columns, rows, truncated = dbmod.run_select(sql, wide, max_rows=5)
    except Exception as e:
        return _err(f"SQLが実データで通りませんでした: {str(e).splitlines()[0][:120]} "
                    "実際に成功したSQLをそのまま渡してください。")
    total = _total_rows(sql, wide) if truncated else len(rows)

    # 置き場のDB = SQLが最初に名指ししているDB（例文はDBごとのファイルに残る）
    hits = []
    import re as _re
    for f in dbmod.list_db_files():
        m = _re.search(r'(?<![\w."])' + _re.escape(dbmod.alias_for(f)) + r"\s*\.",
                       sql, _re.IGNORECASE)
        if m:
            hits.append((m.start(), f))
    hits.sort(key=lambda t: t[0])
    home = hits[0][1] if hits else (Path(scope[0]["path"]) if len(scope) == 1 else None)
    if home is None:
        return _err("このSQLがどのDBのものか判断できませんでした。"
                    "テーブル名を『DB名.テーブル名』の形で書いたSQLを渡してください。")

    same = catalog.find_example(catalog.load_meta(home).get("examples") or [], sql)
    return {
        "ok": True,
        "llm_content": _json({
            "status": "proposed", "db": home.name, "question": q,
            "already_exists": same is not None,
            "note": "登録カードをユーザーの画面に出しました。登録するかはユーザーが"
                    "カードのボタンで決めます。あなたはこれ以上の操作をしなくてよい。",
        }),
        "render": {"role": "assistant", "kind": "example_proposal",
                   "db": home.name, "question": q, "sql": sql, "summary": summary,
                   "columns": columns, "rows": [list(r) for r in rows],
                   "total": total, "exists": same is not None,
                   "old_q": (same or {}).get("q", "")},
    }


def _show_er_diagram(args: dict, scope: list[dict]) -> dict:
    """ER図をチャットに出す（読み取り専用）。

    描くのはカタログ画面と同じキャンバス（er.js）。データも同じ er_payload なので、
    画面で見える図とAIの理解（結合定義）は必ず一致する。
    """
    import db as dbmod

    files = dbmod.list_db_files()
    if not files:
        return _err("data/ にDBがありません。")
    name = str(args.get("db") or "").strip()
    target = None
    if name:
        low = name.lower()
        target = next((f for f in files
                       if low in (f.name.lower(), f.stem.lower(),
                                  dbmod.alias_for(f).lower())), None)
        if target is None:
            return _err(f"DB '{name}' が見つかりません。指定できるのは: "
                        + "、".join(f.name for f in files))
    elif len(scope) == 1:
        target = Path(scope[0]["path"])
    else:
        return _err("どのDBのER図かを db で指定してください。候補: "
                    + "、".join(s_["name"] for s_ in scope))

    try:
        payload = catalog.er_payload(target)
    except Exception as e:
        return _err(f"ER図データの組み立てに失敗しました: {e}")

    own = [n for n in payload["nodes"] if not n.get("external")]
    rels = [{"from": e.get("from_ref") or ".".join(str(x) for x in e["from"]),
             "to": e.get("to_ref") or ".".join(str(x) for x in e["to"]),
             "cardinality": e.get("cardinality") or "",
             "kind": "FOREIGN KEY宣言" if e.get("kind") == "fk" else "カタログ登録",
             **({"note": f"複合キー（{len(e['pairs'])}列すべてを同時に結合条件にする）"}
                if len(e.get("pairs") or []) > 1 else {})}
            for e in payload["edges"]]
    return {
        "ok": True,
        "llm_content": _json({
            "status": "er_ready", "db": target.name,
            "tables": [n["table"] for n in own],
            "relationships": rels[:60],
            "note": "ER図はユーザーの画面に表示済み。関係を文章で説明し直す必要はない。"
                    "結合の一覧は上の relationships のとおり。",
        }),
        "render": {"role": "assistant", "kind": "er", "db": target.name,
                   "title": "ER図（テーブルの関係）", "er": payload},
    }


#: 画面のプレビューに載せる行数。画面が実際に出すのは先頭20行なので、
#: これ以上を会話に残しても誰も使わないまま重くなるだけ
#: （ファイルの中身は data に入っていて、そちらは全行ある）。
PREVIEW_ROWS = 40


def _preview_rows(columns, rows, name="", note=""):
    """会話に残すプレビュー。全行ではなく先頭だけ持つ。"""
    head = list(rows or [])[:PREVIEW_ROWS]
    cut = len(rows or []) - len(head)
    if cut > 0:
        note = (note + f"（表示は先頭{len(head)}行。ファイルには全{len(rows):,}行）").strip()
    out = {"name": name, "columns": columns, "rows": head,
           "total_rows": len(rows or [])}
    if note:
        out["note"] = note
    return out


def _export_excel(args: dict, scope: list[dict]) -> dict:
    sheets_in = args.get("sheets") or []
    if not sheets_in:
        return _err("sheets が空です。少なくとも1つ SELECT を指定してください。")

    built, summary, preview = [], [], []
    for i, sh in enumerate(sheets_in, start=1):
        name = (sh or {}).get("name") or f"Sheet{i}"
        # ファイルには全行入れる（画面向けの2,000行とは別枠）。
        # Excelのシートは仕様上 1,048,576 行までなので、見出しぶんを引いて丸める。
        cap = min(config.EXPORT_MAX_ROWS, 1_048_575)
        try:
            columns, rows, truncated, _, _ = fetch(sh or {}, scope, label=name,
                                                   max_rows=cap)
        except advanced.AnalysisError as e:
            return _err(f"シート '{name}': {e}")
        except Exception as e:
            return _err(f"シート '{name}' のSQL実行エラー: {e}")
        note = (sh or {}).get("note") or ""
        if truncated:
            note = (note + f"（{cap:,}行で切り詰め）").strip()
        charts = (sh or {}).get("charts") or (sh or {}).get("chart")
        if isinstance(charts, dict):
            charts = [charts]
        # Excelの中身は全行（build_excel に渡す）。会話に残すのは先頭だけ
        built.append({"name": name, "columns": columns, "rows": rows, "note": note,
                      "charts": charts or []})
        preview.append(_preview_rows(columns, rows, name, note))
        summary.append({"sheet": name, "columns": columns, "row_count": len(rows),
                        "truncated": truncated,
                        "charts": [c.get("type") for c in (charts or [])]})

    try:
        data = excel.build_excel(built, title=args.get("purpose") or args.get("filename"))
    except ValueError as e:
        return _err(f"Excelのグラフを作れませんでした: {e}")
    except Exception as e:
        return _err(f"Excelの作成に失敗しました: {e}")

    filename = exports.safe_filename(args.get("filename"), "xlsx")
    return {
        "ok": True,
        "llm_content": _json({
            "status": "file_ready",
            "filename": filename,
            "sheets": summary,
            "note": "ユーザーの画面に保存済み。ファイルの中身を再度説明する必要はない。",
        }),
        "render": {
            "role": "assistant", "kind": "file", "filename": filename,
            "mime": exports.XLSX_MIME, "data": data, "sheets": preview,
            "note": f"{len(built)}シート",
        },
    }


def _export_csv(args: dict, scope: list[dict]) -> dict:
    files_in = args.get("files") or []
    if not files_in:
        return _err("files が空です。少なくとも1つ SELECT を指定してください。")
    enc = args.get("encoding") or exports.DEFAULT_ENCODING
    delim = args.get("delimiter") or "comma"

    made, summary, preview = [], [], []
    for i, f in enumerate(files_in, start=1):
        name = (f or {}).get("name") or f"data{i}"
        try:
            columns, rows, truncated, _, _ = fetch(f or {}, scope, label=name,
                                                   max_rows=config.EXPORT_MAX_ROWS)
        except advanced.AnalysisError as e:
            return _err(f"'{name}': {e}")
        except Exception as e:
            return _err(f"'{name}' のSQL実行エラー: {e}")
        try:
            data = exports.build_csv(columns, rows, enc, delim)
        except Exception as e:
            return _err(f"'{name}' のCSV作成に失敗しました（文字コード {enc}）: {e}")
        made.append({"filename": exports.safe_filename(name, "csv"), "data": data})
        summary.append({"file": name, "columns": columns, "row_count": len(rows),
                        "truncated": truncated})
        preview.append(_preview_rows(columns, rows, name))

    if len(made) == 1:
        filename, data, mime = made[0]["filename"], made[0]["data"], exports.CSV_MIME
    else:
        filename = exports.safe_filename(args.get("purpose") or "csv_files", "zip")
        data, mime = exports.build_zip(made), exports.ZIP_MIME

    return {
        "ok": True,
        "llm_content": _json({
            "status": "file_ready", "filename": filename,
            "encoding": enc, "delimiter": delim, "files": summary,
            "note": "ユーザーの画面に保存済み。中身を再度全部説明する必要はない。",
        }),
        "render": {
            "role": "assistant", "kind": "file", "filename": filename,
            "mime": mime, "data": data, "sheets": preview,
            "note": f"文字コード {enc} / 区切り {delim}",
        },
    }


def _export_text(args: dict, scope: list[dict]) -> dict:
    body = str(args.get("body") or "")
    fmt = args.get("format") or "md"
    enc = args.get("encoding") or exports.DEFAULT_ENCODING
    style = "markdown" if fmt == "md" else "plain"

    summary, preview = [], []
    for sec in (args.get("sections") or []):
        heading = str((sec or {}).get("heading") or "")
        try:
            columns, rows, truncated, _, _ = fetch(sec or {}, scope, label=heading,
                                                   max_rows=config.EXPORT_MAX_ROWS)
        except advanced.AnalysisError as e:
            return _err(f"セクション '{heading}': {e}")
        except Exception as e:
            return _err(f"セクション '{heading}' のSQL実行エラー: {e}")
        table = exports.table_to_text(columns, rows, style)
        block = (f"## {heading}\n\n{table}\n" if fmt == "md"
                 else f"■ {heading}\n\n{table}\n")
        placeholder = "{{" + heading + "}}"
        if placeholder in body:
            body = body.replace(placeholder, block)
        else:
            body = body.rstrip() + "\n\n" + block
        summary.append({"heading": heading, "columns": columns, "row_count": len(rows),
                        "truncated": truncated})
        preview.append({"name": heading, "columns": columns, "rows": rows})

    try:
        data = exports.build_text(body, enc)
    except Exception as e:
        return _err(f"テキストの書き出しに失敗しました（文字コード {enc}）: {e}")

    filename = exports.safe_filename(args.get("filename"), fmt)
    return {
        "ok": True,
        "llm_content": _json({
            "status": "file_ready", "filename": filename,
            "format": fmt, "encoding": enc, "sections": summary,
            "chars": len(body),
            "note": "ユーザーの画面に保存済み。本文を再度全部繰り返す必要はない。",
        }),
        "render": {
            "role": "assistant", "kind": "file", "filename": filename,
            "mime": exports.MD_MIME if fmt == "md" else exports.TEXT_MIME,
            "data": data, "text": body, "sheets": preview,
            "note": f"{fmt} / 文字コード {enc} / {len(body):,} 文字",
        },
    }

def _open_table(args: dict, scope: list[dict]) -> dict:
    """テーブルの中身（全行）を別タブのビューアで開く。

    「テーブルを見せて」に対して、AIが SELECT * を打って先頭数行を貼るのは
    (1) 行数の上限で切れる (2) 会話が表で埋まる、の2つで具合が悪い。
    全行を辿れる画面（/table）を開いて、AIは要約に専念する。
    """
    import db as dbmod

    files = dbmod.list_db_files()
    if not files:
        return _err("data/ にDBがありません。")
    name = str(args.get("db") or "").strip()
    table = str(args.get("table") or "").strip()
    if not table:
        return _err("どのテーブルを開くかを table で指定してください。")
    # "alias.table" で来たら剥がす（describe_table と同じ扱い）。
    # 先頭がいずれかのDBの名前に一致するときだけ。テーブル名自体の "." は壊さない。
    if "." in table:
        head, rest = table.split(".", 1)
        low = head.lower()
        if rest and any(low in (f.name.lower(), f.stem.lower(), dbmod.alias_for(f).lower())
                        for f in files):
            if not name or name.lower() in (low, head.lower()):
                name, table = head, rest

    # DB名は 'sales.db' でも 'sales' でもよい。省略時は表名から探す
    target = None
    if name:
        low = name.lower()
        target = next((f for f in files
                       if low in (f.name.lower(), f.stem.lower(), dbmod.alias_for(f).lower())), None)
        if target is None:
            return _err(f"DB '{name}' が見つかりません。指定できるのは: "
                        + "、".join(f.name for f in files))
    else:
        hits = [f for f in files if table in (catalog.profile_db(f).get("tables") or {})]
        if not hits:
            return _err(f"テーブル '{table}' が見つかりません。db も指定してください。")
        if len(hits) > 1:
            return _err(f"テーブル '{table}' が複数のDBにあります（"
                        + "、".join(f.name for f in hits) + "）。db で指定してください。")
        target = hits[0]

    profile = catalog.profile_db(target)
    info = (profile.get("tables") or {}).get(table)
    if info is None:
        return _err(f"テーブル '{table}' が {target.name} にありません。"
                    "このDBのテーブル: " + "、".join(list((profile.get("tables") or {}).keys())[:20]))

    meta = catalog.load_meta(target)
    tmeta = (meta.get("tables") or {}).get(table) or {}
    cols = [c["name"] for c in (info.get("columns") or [])]
    rows = info.get("row_count")
    return {
        "ok": True,
        "llm_content": _json({
            "status": "table_view_opened", "db": target.name, "table": table,
            "rows": rows, "columns": cols,
            "note": "テーブルの全行を見る画面へのリンク（「テーブル全体を開く」ボタン）を"
                    "利用者のチャットに出した。利用者がそれを押すと別タブで開く。"
                    "中身を SELECT * で貼り直す必要はない。"
                    "何のテーブルか・何に使えるかを1〜2文で補足し、"
                    "「テーブル全体を開く」から見られることを伝えるだけでよい。",
        }),
        "render": {
            "role": "assistant", "kind": "table_link",
            "db": target.name, "table": table,
            "title": table,
            "rows": rows, "columns": cols,
            "description": tmeta.get("description") or "",
        },
    }


# このモジュールが受け持つツール
HANDLERS_query = {
    "run_sql_query": _run_sql_query,
    "describe_table": _describe_table,
    "show_er_diagram": _show_er_diagram,
    "open_table": _open_table,
    "propose_glossary_term": _propose_glossary_term,
    "propose_example": _propose_example,
    "pivot_table": _pivot_table,
    "analyze_stats": _analyze_stats,
    "plot_chart": _plot_chart,
    "plot_dual_axis": _plot_dual_axis,
    # 用途別のグラフツールは、中身はどれも同じ組み立てを通る
    **{name: _plot_chart for name in _CHART_TOOLS},
    "export_excel": _export_excel,
    "export_csv": _export_csv,
    "export_text": _export_text,
}

# SQLを受け取るツール（実行前プレビュー表示の対象）
SQL_TOOLS_query = {"run_sql_query", "plot_chart", "plot_dual_axis", "pivot_table",
             "analyze_stats", *_CHART_TOOLS}


# ==========================================================================
# ===== 元 tools/reports.py
# レポート出力（PowerPoint / Word / 画面用のレポート）。
# ==========================================================================
import re
from datetime import datetime

import advanced
import charts
import docx_report
import excel
import exports
import figures
import pptx_report


def _slide_from_sql(spec: dict, scope: list[dict], index: int) -> dict:
    """slides の1枚ぶん。sql か result_id があればここで中身に変える。"""
    out = dict(spec)
    kind = str(spec.get("kind") or "message").lower()
    if not (spec.get("sql") or spec.get("result_id")) or kind not in ("table", "chart"):
        return out

    columns, rows, truncated, _, _ = fetch(spec, scope, label=spec.get("title"))
    if not rows:
        raise pptx_report.PptxReportError(f"{index}枚目「{spec.get('title', '')}」の"
                                      "SQLが0行でした。抽出条件を見直してください。")
    if kind == "table":
        out["columns"], out["rows"] = columns, [list(r) for r in rows]
        if truncated:
            out["comment"] = (out.get("comment", "") + "　※ 上限で切り詰め済み").strip()
        return out

    cat = spec.get("category_column") or columns[0]
    if cat not in columns:
        raise pptx_report.PptxReportError(
            f"{index}枚目: 横軸の列 '{cat}' がSQLの結果にありません"
            f"（ある列: {', '.join(columns)}）。")
    vals = spec.get("value_columns") or [c for c in columns if c != cat]
    missing = [v for v in vals if v not in columns]
    if missing:
        raise pptx_report.PptxReportError(
            f"{index}枚目: 系列の列 {', '.join(missing)} がSQLの結果にありません"
            f"（ある列: {', '.join(columns)}）。")
    ci = columns.index(cat)
    out["categories"] = [r[ci] for r in rows]
    if str(spec.get("chart") or "bar").lower() == "scatter":
        xi = columns.index(vals[0])
        out["series"] = [{"name": vals[1] if len(vals) > 1 else vals[0],
                          "x": [r[xi] for r in rows],
                          "values": [r[columns.index(vals[1] if len(vals) > 1 else vals[0])]
                                     for r in rows]}]
    else:
        out["series"] = [{"name": v, "values": [r[columns.index(v)] for r in rows]}
                         for v in vals]
    return out


def _export_pptx(args: dict, scope: list[dict]) -> dict:
    slides_in = args.get("slides") or []
    if not slides_in:
        return _err("slides が空です。少なくとも1枚は指定してください。")
    built = []
    for i, spec in enumerate(slides_in, start=1):
        try:
            built.append(_slide_from_sql(spec or {}, scope, i))
        except (pptx_report.PptxReportError, advanced.AnalysisError) as e:
            return _err(f"{i}枚目: {e}" if isinstance(e, advanced.AnalysisError) else str(e))
        except Exception as e:
            return _err(f"{i}枚目のSQL実行エラー: {e}")
    try:
        data = pptx_report.build_pptx(built, title=args.get("title"),
                                 subtitle=args.get("subtitle"),
                                 footer=args.get("footer"))
    except pptx_report.PptxReportError as e:
        return _err(str(e))
    except Exception as e:
        return _err(f"PowerPointの作成に失敗しました: {e}")

    filename = pptx_report.pptx_safe_filename(args.get("filename") or args.get("title"))
    outline = pptx_report.outline_pptx(built)
    return {
        "ok": True,
        "llm_content": _json({
            "status": "file_ready", "filename": filename,
            "slides": outline,
            "note": "ユーザーの画面に保存済み。中身を再度説明する必要はない。",
        }),
        "render": {"role": "assistant", "kind": "file", "filename": filename,
                   "mime": PPTX_MIME, "data": data,
                   "note": f"{len(built)}スライド", "outline": outline},
    }


PPTX_MIME = ("application/vnd.openxmlformats-officedocument."
             "presentationml.presentation")


DOCX_MIME = ("application/vnd.openxmlformats-officedocument."
             "wordprocessingml.document")


def _export_docx(args: dict, scope: list[dict]) -> dict:
    """Word文書。図表つきで、そのまま配布できる体裁にする。"""
    secs_in = args.get("sections") or []
    if not secs_in:
        return _err("sections が空です。少なくとも1つの見出しを入れてください。")

    sections, figs, tbls = [], 0, 0
    for i, s in enumerate(secs_in, start=1):
        s = s or {}
        if not s.get("heading"):
            return _err(f"{i}番目のセクションに heading がありません。")
        sec = {k: s.get(k) for k in ("heading", "body", "bullets", "note",
                                     "callout", "level", "page_break")}
        # 図表のキャプションは、見出し頭の「1. 」を落として重複を避ける
        label = re.sub(r"^\s*\d+[.．)、]\s*", "", s["heading"])
        if s.get("sql") or s.get("result_id"):
            try:
                columns, rows, truncated, _, _ = fetch(s, scope, label=s["heading"])
            except advanced.AnalysisError as e:
                return _err(f"「{s['heading']}」: {e}")
            except Exception as e:
                return _err(f"「{s['heading']}」のSQL実行エラー: {e}")
            if not rows:
                return _err(f"「{s['heading']}」のデータが0行でした。")
            limit = int(s.get("max_rows") or 40)
            if s.get("chart"):
                chart = {**(s["chart"] or {}), "columns": columns,
                         "rows": [list(r) for r in rows],
                         "title": (s["chart"] or {}).get("title") or s["heading"]}
                chart.setdefault("chart_type", "bar")
                errs = charts.validate(chart, columns)
                if errs:
                    return _err(f"「{s['heading']}」のグラフ指定: {' / '.join(errs)}")
                img = _chart_image(chart)
                if img:
                    sec["image"] = img
                    sec["caption"] = s.get("caption") or label
                    figs += 1
            if s.get("table", True):        # 既定で表も載せる（根拠として残す）
                sec["table"] = {"columns": columns,
                                "rows": [list(r) for r in rows[:limit]]}
                if len(rows) > limit or truncated:
                    sec["table"]["note"] = f"全 {len(rows):,} 行から抜粋"
                sec["table_caption"] = s.get("table_caption") or label
                tbls += 1
        sections.append(sec)

    try:
        data = docx_report.build_docx(
            sections, title=args.get("title", "レポート"),
            subtitle=args.get("subtitle", ""),
            summary=args.get("summary") or [],
            conclusion=args.get("conclusion", ""),
            recommendations=args.get("recommendations") or [],
            caveats=args.get("caveats") or [],
            footer=args.get("footer", ""), org=args.get("org", ""),
            author=args.get("author", ""), toc=bool(args.get("toc", True)))
    except docx_report.DocxReportError as e:
        return _err(str(e))
    except Exception as e:
        return _err(f"Word文書の作成に失敗しました: {e}")

    filename = docx_report.docx_safe_filename(args.get("filename") or args.get("title"))
    note = ""
    if figs == 0 and any(s.get("chart") for s in secs_in):
        note = figures.why_unavailable()
    return {
        "ok": True,
        "llm_content": _json({
            "status": "file_ready", "filename": filename,
            "sections": docx_report.outline_docx(sections),
            "figures": figs, "tables": tbls,
            "warning": note or None,
            "note": "ユーザーの画面に保存済み。中身を再度説明する必要はない。",
        }),
        "render": {"role": "assistant", "kind": "file", "filename": filename,
                   "mime": DOCX_MIME, "data": data,
                   "note": f"{len(sections)}セクション / 図{figs} 表{tbls}"
                           + (f" ／ {note}" if note else ""),
                   "outline": docx_report.outline_docx(sections)},
    }


def _report_markdown(args: dict, sections: list[dict]) -> str:
    """レポートを1本の Markdown にする（ダウンロード用と、表示の下敷き）。"""
    out = [f"# {args.get('title', 'レポート')}"]
    if args.get("subtitle"):
        out.append(f"*{args['subtitle']}*")
    out.append(f"*作成: {datetime.now():%Y-%m-%d %H:%M}*")
    if args.get("summary"):
        out += ["", "## 要点"] + [f"- {s}" for s in args["summary"]]
    for i, s in enumerate(sections, 1):
        out += ["", f"## {i}. {s['heading']}"]
        if s.get("body"):
            out += ["", s["body"]]
        if s.get("table"):
            t = s["table"]
            out += ["", exports.table_to_text(t["columns"], t["rows"], "markdown")]
            if t.get("truncated"):
                out.append(f"（全 {t['total']:,} 行のうち上位 {len(t['rows'])} 行）")
        if s.get("chart"):
            out.append(f"（グラフ: {s['chart'].get('chart_type')} — 画面で確認できます）")
        if s.get("note"):
            out += ["", f"> {s['note']}"]
    if args.get("conclusion"):
        out += ["", "## 結論", "", args["conclusion"]]
    if args.get("recommendations"):
        out += ["", "## 推奨する打ち手"] + [f"{i}. {r}" for i, r
                                            in enumerate(args["recommendations"], 1)]
    if args.get("caveats"):
        out += ["", "## 前提・注意"] + [f"- {c}" for c in args["caveats"]]
    return "\n".join(out) + "\n"


# PowerPointのネイティブグラフは種類が限られる。近いもので描けるならそれを使い、
# 描けない種類（サンキー・箱ひげ等）は画像として貼る。
_PPTX_CHART_MAP = {
    "bar": "bar", "hbar": "hbar", "stacked_bar": "bar_stacked",
    "percent_bar": "bar_percent", "lollipop": "bar", "pareto": "bar",
    "line": "line", "step": "line", "bump": "line", "slope": "line",
    "area": "area", "area_percent": "area_stacked",
    "pie": "pie", "donut": "doughnut", "funnel": "hbar", "radar": "radar",
    "scatter": "scatter", "bubble": "scatter", "polar_bar": "radar",
}


# Excelのグラフも種類が限られる。近いものに寄せ、無理なものは表だけにする。
_XLSX_CHART_MAP = {
    "bar": "bar", "hbar": "hbar", "stacked_bar": "bar_stacked",
    "percent_bar": "bar_percent", "lollipop": "bar", "pareto": "bar",
    "line": "line", "step": "line", "bump": "line", "slope": "line",
    "area": "area", "area_percent": "area_stacked",
    "pie": "pie", "donut": "pie", "funnel": "hbar",
    "scatter": "scatter", "bubble": "scatter",
}


def _chart_image(chart: dict):
    """グラフを印刷向けの画像にする。できなければ None。"""
    try:
        return figures.for_print(charts.build_figure(chart))
    except Exception as e:
        print(f"[report] 画像化に失敗: {e}")
        return None


def _series_from(chart: dict, table: dict):
    """表からPowerPointのネイティブグラフ用の系列を組み立てる。"""
    cols = table["columns"]
    cat = chart.get("x") if chart.get("x") in cols else cols[0]
    ci = cols.index(cat)
    if chart.get("chart_type") in ("scatter", "bubble"):
        xs = [r[cols.index(chart["x"])] for r in table["rows"]]
        ys = [r[cols.index(chart["y"])] for r in table["rows"]]
        return [], [{"name": chart.get("y", "値"), "x": xs, "values": ys}]
    wanted = [chart["y"]] if chart.get("y") in cols else \
        [c for c in cols if c != cat][:3]
    return ([r[ci] for r in table["rows"]],
            [{"name": v, "values": [r[cols.index(v)] for r in table["rows"]]}
             for v in wanted])


def _split_message(s: dict) -> tuple[str, str]:
    """1ページの「言いたいこと1行」と、横に添える残りの文章に分ける。

    所見(note)があればそれを1行目にする。無ければ本文の最初の一文を使い、
    その場合は本文の残りだけを横に置く（同じ文を2回出さない）。
    """
    body = (s.get("body") or "").strip()
    if s.get("note"):
        return s["note"], body
    if not body:
        return "", ""
    head, sep, rest = body.partition("。")
    if not sep:
        return body[:90], ""
    return head + "。", rest.strip()


def _report_slides(args: dict, sections: list[dict]) -> list[dict]:
    """レポートの内容を、会議で映せるスライドの並びに翻訳する。"""
    summary = args.get("summary") or []
    slides = [{"kind": "title", "title": args.get("title", "レポート"),
               "subtitle": args.get("subtitle", ""),
               "lines": summary[:4], "org": args.get("org", "")}]
    if len(sections) >= 2:
        slides.append({"kind": "agenda", "title": "本日の内容",
                       "items": [s["heading"] for s in sections]})
    if summary:
        # 帯に出した1点目は繰り返さない（同じ文が2回出ると雑に見える）
        slides.append({"kind": "message", "title": "要約",
                       "message": summary[0],
                       "bullets": summary[1:] or summary,
                       "callout": args.get("conclusion", "")})

    for s in sections:
        message, comment = _split_message(s)
        base = {"title": s["heading"],
                # 見出しの下に置く1行。結論を先に言う。
                "message": message, "comment": comment,
                "notes": s.get("body") or "",
                "source": args.get("source", "")}
        chart, table = s.get("chart"), s.get("table")
        if chart and table:
            kind = chart.get("chart_type")
            native = _PPTX_CHART_MAP.get(kind)
            if native:
                cats, series = _series_from(chart, table)
                slides.append({**base, "kind": "chart", "chart": native,
                               "categories": cats, "series": series})
            else:
                img = _chart_image(chart)
                if img:
                    slides.append({**base, "kind": "chart", "image": img})
                else:
                    slides.append({**base, "kind": "table", **table})
        elif table:
            slides.append({**base, "kind": "table", **table})
        else:
            slides.append({**base, "kind": "message", "comment": None,
                           "lead": base.pop("comment", "") or "",
                           "bullets": s.get("bullets") or []})

    if args.get("conclusion") or args.get("recommendations"):
        slides.append({"kind": "closing", "title": "まとめと次のアクション",
                       "message": args.get("conclusion", ""),
                       "summary": (args.get("summary") or [])[:3],
                       "actions": args.get("recommendations") or []})
    return slides


def _report_docx_sections(args: dict, sections: list[dict]) -> list[dict]:
    """レポートの内容を、Wordのセクションに翻訳する（図は画像で貼る）。"""
    out = []
    for s in sections:
        label = re.sub(r"^\s*\d+[.．)、]\s*", "", s["heading"])
        sec = {"heading": s["heading"], "body": s.get("body", ""),
               "note": s.get("note", ""), "caption": label}
        if s.get("chart"):
            img = _chart_image(s["chart"])
            if img:
                sec["image"] = img
        if s.get("table"):
            sec["table"] = {"columns": s["table"]["columns"],
                            "rows": s["table"]["rows"]}
            sec["table_caption"] = label
            if s["table"].get("truncated"):
                sec["table"]["note"] = f"全 {s['table']['total']:,} 行から抜粋"
        out.append(sec)
    return out


def _build_report(args: dict, scope: list[dict]) -> dict:
    sections_in = args.get("sections") or []
    if not sections_in:
        return _err("sections が空です。少なくとも1つの論点を入れてください。")

    sections, dropped = [], []
    for i, s in enumerate(sections_in, 1):
        s = s or {}
        if not s.get("heading"):
            return _err(f"{i}番目のセクションに heading がありません。")
        out = {"heading": s["heading"], "body": s.get("body", ""),
               "note": s.get("note", "")}
        if s.get("sql") or s.get("result_id"):
            try:
                columns, rows, truncated, _, _ = fetch(s, scope, label=s["heading"])
            except advanced.AnalysisError as e:
                return _err(f"「{s['heading']}」: {e}")
            except Exception as e:
                return _err(f"「{s['heading']}」のSQL実行エラー: {e}")
            limit = int(s.get("max_rows") or 20)
            out["table"] = {"columns": columns, "rows": [list(r) for r in rows[:limit]],
                            "total": len(rows), "truncated": len(rows) > limit or truncated}
            out["sql"] = s.get("sql")
            if s.get("chart"):
                chart = {k: v for k, v in (s["chart"] or {}).items()}
                chart.setdefault("chart_type", "bar")
                errs = charts.validate(chart, columns)
                if errs:
                    # グラフ1つのためにレポート全体を捨てない。
                    # 表は根拠として残るので、図を落として作り切る方が役に立つ。
                    dropped.append(f"「{s['heading']}」のグラフ: {' / '.join(errs)}")
                else:
                    # グラフは全行を使う（表は読みやすさのために切っている）
                    out["chart"] = {**chart, "columns": columns,
                                    "rows": [list(r) for r in rows],
                                    "title": chart.get("title") or s["heading"]}
        elif s.get("chart"):
            dropped.append(f"「{s['heading']}」のグラフ: 元になる sql / result_id がありません。")
        sections.append(out)

    md = _report_markdown(args, sections)
    fmt = (args.get("format") or "md").lower()
    name = args.get("filename") or args.get("title") or "report"
    data = filename = mime = None
    try:
        if fmt == "md":
            data = exports.build_text(md)
            filename = exports.safe_filename(name, "md")
            mime = exports.TEXT_MIME
        elif fmt == "xlsx":
            sheets = []
            for s in sections:
                if not s.get("table"):
                    continue
                sheet = {"name": s["heading"], "columns": s["table"]["columns"],
                         "rows": s["table"]["rows"], "note": s.get("note", "")}
                # 画面のグラフ指定を、そのままExcelのグラフに読み替える
                if s.get("chart"):
                    ch = s["chart"]
                    kind = _XLSX_CHART_MAP.get(ch.get("chart_type"))
                    if kind:
                        cat = ch.get("x") if ch.get("x") in sheet["columns"] \
                            else sheet["columns"][0]
                        vals = ([ch["y"]] if ch.get("y") in sheet["columns"]
                                else [c for c in sheet["columns"] if c != cat][:3])
                        sheet["charts"] = [{"type": kind, "category_column": cat,
                                            "value_columns": vals,
                                            "title": s["heading"]}]
                sheets.append(sheet)
            if not sheets:
                return _err("xlsx にするには、表（sql）のあるセクションが1つ以上必要です。")
            data = excel.build_excel(sheets, title=args.get("title"))
            filename = exports.safe_filename(name, "xlsx")
            mime = exports.XLSX_MIME
        elif fmt == "pptx":
            data = pptx_report.build_pptx(_report_slides(args, sections),
                                     title=args.get("title"),
                                     subtitle=args.get("subtitle"),
                                     footer=args.get("footer", ""))
            filename = pptx_report.pptx_safe_filename(name)
            mime = PPTX_MIME
        elif fmt == "docx":
            data = docx_report.build_docx(
                _report_docx_sections(args, sections),
                title=args.get("title", "レポート"),
                subtitle=args.get("subtitle", ""),
                summary=args.get("summary") or [],
                conclusion=args.get("conclusion", ""),
                recommendations=args.get("recommendations") or [],
                caveats=args.get("caveats") or [],
                footer=args.get("footer", ""), org=args.get("org", ""))
            filename = docx_report.docx_safe_filename(name)
            mime = DOCX_MIME
    except Exception as e:
        return _err(f"ファイルの作成に失敗しました: {e}")

    render = {"role": "assistant", "kind": "report_doc",
              "title": args.get("title", "レポート"),
              "subtitle": args.get("subtitle", ""),
              "summary": args.get("summary") or [],
              "sections": sections,
              "conclusion": args.get("conclusion", ""),
              "recommendations": args.get("recommendations") or [],
              "caveats": args.get("caveats") or [],
              "markdown": md}
    if data:
        render.update(data=data, filename=filename, mime=mime)

    return {
        "ok": True,
        "llm_content": _json({
            "status": "report_ready", "title": args.get("title"),
            "sections": [{"heading": s["heading"],
                          "rows": (s.get("table") or {}).get("total"),
                          "chart": (s.get("chart") or {}).get("chart_type")}
                         for s in sections],
            "filename": filename,
            "dropped_charts": dropped or None,
            "note": "レポートは画面に表示済み。内容をもう一度書き出す必要はない。"
                    "次に何をするか（送付・追加分析など）だけ短く伝えること。"
                    + ("　※ 一部のグラフは指定が合わず省いた。作り直すなら、"
                       "dropped_charts の指摘どおりに列名を直して呼ぶこと"
                       "（同じ引数で呼び直さない）。" if dropped else ""),
        }),
        "render": render,
    }

HANDLERS_reports = {
    "export_pptx": _export_pptx,
    "export_docx": _export_docx,
    "build_report": _build_report,
}

SQL_TOOLS_reports: set[str] = set()


# ==========================================================================
# ===== 元 tools/stats.py
# 統計と試算。advanced.py の分析関数をツールとして公開する。
# ==========================================================================
import advanced


_hypothesis_test = _analysis_tool(lambda a, c, r: advanced.hypothesis_test(
    c, r, a.get("method"), value_col=a.get("value_col"), group_col=a.get("group_col"),
    value_col2=a.get("value_col2"), popmean=float(a.get("popmean") or 0),
    expected=a.get("expected"), alternative=a.get("alternative") or "two-sided"))


_regression = _analysis_tool(lambda a, c, r: advanced.regression(
    c, r, a.get("target"), a.get("features") or [], method=a.get("method") or "ols",
    predict=a.get("predict")))


_distribution_analysis = _analysis_tool(lambda a, c, r: advanced.distribution(
    c, r, a.get("target"), bins=int(a.get("bins") or 20), fit=a.get("fit"),
    group_col=a.get("group_col")))


_forecast = _analysis_tool(lambda a, c, r: advanced.forecast(
    c, r, a.get("time_col"), a.get("value_col"), periods=int(a.get("periods") or 6),
    method=a.get("method") or "auto",
    season_length=int(a["season_length"]) if a.get("season_length") else None,
    exog=a.get("exog")))


_timeseries_analysis = _analysis_tool(lambda a, c, r: advanced.timeseries(
    c, r, a.get("time_col"), a.get("value_col"), window=int(a.get("window") or 3),
    season_length=int(a["season_length"]) if a.get("season_length") else None))


_bootstrap_estimate = _analysis_tool(lambda a, c, r: advanced.bootstrap(
    c, r, a.get("target"), statistic=a.get("statistic") or "mean",
    trials=int(a.get("trials") or 5000), group_col=a.get("group_col")))


# k は "auto" も受けるので、ここで数値に変換しない
_clustering = _analysis_tool(lambda a, c, r: advanced.clustering(
    c, r, a.get("features") or [], k=a.get("k") or 3,
    label_col=a.get("label_col"), categorical=a.get("categorical")))


_abc_analysis = _analysis_tool(lambda a, c, r: advanced.abc_analysis(
    c, r, a.get("label_col"), a.get("value_col"), thresholds=a.get("thresholds")))


def _monte_carlo_simulation(args: dict, scope: list[dict]) -> dict:
    columns = rows = None
    if args.get("sql"):
        try:
            columns, rows, _ = _select_for(args, scope)
        except advanced.AnalysisError as e:
            return _err(str(e))
        except Exception as e:
            return _err(f"実データの取得に失敗しました: {e}")
    try:
        res = advanced.monte_carlo(
            args.get("formula", ""), args.get("variables") or {},
            trials=int(args.get("trials") or 10000), columns=columns, rows=rows,
            targets=args.get("targets"))
    except advanced.AnalysisError as e:
        return _err(str(e))
    except Exception as e:
        return _err(f"シミュレーションに失敗しました: {e}")
    if args.get("title"):
        res["title"] = args["title"]
    return _report_result(res)


def _scenario_analysis(args: dict, scope: list[dict]) -> dict:
    try:
        res = advanced.scenario(args.get("formula", ""), args.get("scenarios") or {},
                                base=args.get("base"))
    except advanced.AnalysisError as e:
        return _err(str(e))
    except Exception as e:
        return _err(f"シナリオ分析に失敗しました: {e}")
    if args.get("title"):
        res["title"] = args["title"]
    return _report_result(res)

HANDLERS_stats = {
    "hypothesis_test": _hypothesis_test,
    "regression": _regression,
    "distribution_analysis": _distribution_analysis,
    "forecast": _forecast,
    "timeseries_analysis": _timeseries_analysis,
    "monte_carlo_simulation": _monte_carlo_simulation,
    "scenario_analysis": _scenario_analysis,
    "bootstrap_estimate": _bootstrap_estimate,
    "clustering": _clustering,
    "abc_analysis": _abc_analysis,
}

# scenario_analysis と monte_carlo_simulation は SQL が任意なので含めない
SQL_TOOLS_stats = {"hypothesis_test", "regression", "distribution_analysis", "forecast",
             "timeseries_analysis", "bootstrap_estimate", "clustering", "abc_analysis"}


# ==========================================================================
# ===== 元 tools/usage.py
# このアプリ自身の使われ方を調べるツール（読むだけ）。
#
# usage.py の集計をLLMに公開する。材料はチャット履歴と取り込みの記録なので、
# 分析対象のDBを選んでいなくても答えられる（SQLでは答えられない話でもある）。
#
# 他人の質問文まで見えるため、管理者にだけ渡す。tools/files.py と同じ扱い。
# ==========================================================================
import usage


def _analyze_usage(args: dict, scope: list[dict]) -> dict:
    days = args.get("days")
    try:
        res = usage.analyze(str(args.get("method") or "summary"),
                            days=int(days) if days else None,
                            user=(args.get("user") or "").strip() or None)
    except ValueError as e:
        return _err(str(e))
    except Exception as e:
        return _err(f"利用状況の集計に失敗しました: {e}")

    if args.get("title"):
        res["title"] = args["title"]
    # scope を渡して表に result_id を付ける。グラフ化やExcel出力にそのまま繋げられる。
    return _report_result(res, scope=scope)


HANDLERS_usage = {"analyze_usage": _analyze_usage}

# SQLは受け取らない（材料はDBではなく履歴ファイル）
SQL_TOOLS_usage: set = set()

# 他の利用者の質問・失敗まで見えるので管理者だけに渡す。
ADMIN_TOOLS_usage = {"analyze_usage"}


# ==========================================================================
# ===== 元 rag/client.py
# LightRAG サーバー1台へのHTTPクライアント。
#
# LightRAG には必ずHTTP API経由で触る。ライブラリとして抱き込まない理由は2つ。
#   - LightRAG の索引作成は重く、このアプリとは別のサーバ・別のGPUで動かしたい
#   - 環境（ナレッジベース）はIT部門から「URLとAPIキー」の形で払い出される
#
# 呼ぶのは2つだけ。
#   health()   … 管理画面の接続テスト（APIキーの有効性確認を含む）
#   retrieve() … /query/data で「検索だけ」。回答は書かせない。
#
# 統合前の LightRAG アプリには generate()（/query の bypass モード）もあった。
# こちらでは使わない。回答を書くのはこのアプリのAgent（llm.py）であり、
# 検索結果はツールの戻り値としてAgentの手元に返る。LightRAG側にもう一度
# 生成させると、Agentが持っているSQLの集計結果やExcelの中身を知らないまま
# 文章を書くことになり、「複数のツールの結果を突き合わせて答える」という
# このアプリの前提が崩れる。生成は1問につき1回のまま（Agentが1回書く）。
#
# 注意: LightRAG の /health は認証不要で、APIキーが誤っていても 200 を返す
# （構成の詳細が伏せられるだけ）。接続テストをこれだけで済ませると、
# 「接続成功」と表示されたのに検索は全て403、という状態を見逃す。
# health() が保護されたエンドポイントも叩いているのはこのため。
# ==========================================================================
import threading

# requests が入っていなくてもアプリ自体は起動させる。ナレッジ検索を使うときだけ
# 困るようにしておく。ここで落とすと、RAGを使っていない利用者の分析まで
# 巻き添えで止まる（この依存は統合で新しく増えたもので、既存の環境には無い）。
try:
    import requests as _requests
except ImportError:                                  # pragma: no cover
    _requests = None

_REQUESTS_MISSING = ("ナレッジ検索には requests が必要です。"
                     "`pip install -r requirements.txt` を実行してください。")


class RagError(Exception):
    """ナレッジベースへの問い合わせが失敗したことを表す。"""


class RagClient:
    def __init__(self, base_url: str, api_key: str = "", *, timeout: int = 60) -> None:
        self.base_url = str(base_url or "").rstrip("/")
        self.api_key = api_key or ""
        self.timeout = timeout

    def _headers(self) -> dict:
        headers = {"Content-Type": "application/json"}
        if self.api_key:
            headers["X-API-Key"] = self.api_key      # LightRAG の認証ヘッダ
        return headers

    def _post(self, path: str, payload: dict, timeout: int) -> dict:
        if _requests is None:
            raise RagError(_REQUESTS_MISSING)
        url = f"{self.base_url}{path}"
        try:
            resp = _requests.post(url, json=payload, headers=self._headers(),
                                  timeout=timeout)
        except _requests.Timeout as exc:
            raise RagError(f"タイムアウトしました（{timeout}秒）: {url}") from exc
        except _requests.RequestException as exc:
            raise RagError(_rag_describe(exc, url)) from exc

        if resp.status_code in (401, 403):
            raise RagError("APIキーが拒否されました。「ナレッジベース」画面でキーを確認してください。")
        if resp.status_code >= 400:
            raise RagError(f"HTTP {resp.status_code}: {_rag_error_message(resp)}")
        try:
            return resp.json()
        except ValueError as exc:
            raise RagError("応答をJSONとして解釈できませんでした。") from exc

    def health(self) -> dict:
        if _requests is None:
            raise RagError(_REQUESTS_MISSING)
        url = f"{self.base_url}/health"
        try:
            resp = _requests.get(url, headers=self._headers(), timeout=15)
        except _requests.RequestException as exc:
            raise RagError(_rag_describe(exc, url)) from exc
        if resp.status_code in (401, 403):
            raise RagError("APIキーが拒否されました。")
        if resp.status_code >= 400:
            raise RagError(f"HTTP {resp.status_code}: {_rag_safe_body(resp)}")
        try:
            body = resp.json()
        except ValueError as exc:
            raise RagError("応答をJSONとして解釈できませんでした。") from exc
        self._verify_credentials()
        return body if isinstance(body, dict) else {}

    def _verify_credentials(self) -> None:
        """保護されたエンドポイントを叩いてAPIキーの有効性を確かめる。"""
        url = f"{self.base_url}/documents/pipeline_status"
        try:
            resp = _requests.get(url, headers=self._headers(), timeout=15)
        except _requests.RequestException as exc:
            raise RagError(_rag_describe(exc, url)) from exc
        if resp.status_code in (401, 403):
            raise RagError(
                "APIキーが拒否されました。「ナレッジベース」画面でキーを確認してください。"
                if self.api_key else
                "この環境はAPIキーを要求します。「ナレッジベース」画面でキーを設定してください。")
        if resp.status_code >= 400:
            raise RagError(f"HTTP {resp.status_code}: {_rag_safe_body(resp)}")

    def retrieve(self, question: str, *, mode: str = "mix", chunk_top_k: int = 10,
                 top_k: int | None = None) -> dict:
        """検索のみ。LightRAG側の回答生成（QUERYロール）は呼ばれない。"""
        payload = {
            "query": question,
            "mode": mode,
            "chunk_top_k": chunk_top_k,
            "include_chunk_content": True,
        }
        if top_k is not None:
            payload["top_k"] = top_k
        body = self._post("/query/data", payload, self.timeout)
        if body.get("status") not in (None, "success"):
            raise RagError(body.get("message") or "検索に失敗しました。")
        return body


def _rag_describe(exc, url: str) -> str:
    """接続失敗を「何を直せばよいか」が分かる日本語にする。

    requests の例外文字列はそのまま出すと読めないうえ、原因（綴り違い・FW・
    プロトコル取り違え）がどれも同じ見た目になる。管理画面で環境を登録するのは
    この機能の主要な操作なので、ここで切り分けておく。

    判定の順番が重要: SSLError と ConnectTimeout は ConnectionError の
    サブクラスなので、先に見ないと握り潰される。
    """
    if _requests is None:
        return _REQUESTS_MISSING
    if isinstance(exc, _requests.exceptions.SSLError):
        return (f"HTTPS接続に失敗しました: {url}\n"
                "URLの http/https が実体と合っているか、"
                "社内CAの証明書が信頼済みかを確認してください。")
    if isinstance(exc, _requests.exceptions.ProxyError):
        return (f"プロキシ経由の接続に失敗しました: {url}\n"
                "社内プロキシの設定（HTTP_PROXY / NO_PROXY）を確認してください。")
    if isinstance(exc, _requests.exceptions.Timeout):
        return (f"応答がありません（タイムアウト）: {url}\n"
                "ネットワーク経路とファイアウォールを確認してください。")
    if isinstance(exc, _requests.exceptions.ConnectionError):
        text = str(exc)
        if any(sign in text for sign in
               ("NameResolutionError", "getaddrinfo failed",
                "Name or service not known", "nodename nor servname")):
            return (f"ホスト名を解決できません: {url}\n"
                    "URLの綴りと、社内DNSで引ける名前かを確認してください。")
        return (f"接続を拒否されました: {url}\n"
                "ポート番号と、サーバーが起動しているかを確認してください。")
    return f"接続に失敗しました: {url}（{exc.__class__.__name__}）"


def _rag_error_message(resp, limit: int = 300) -> str:
    """エラー応答から人が読む部分だけを取り出す。

    LightRAG は失敗時に {"status","message","data","metadata"} を返す。
    生のJSONをそのまま画面へ出すと、環境の数だけ同じ塊が並んで読めない。
    """
    try:
        body = resp.json()
    except ValueError:
        return _rag_safe_body(resp, limit)
    if isinstance(body, dict):
        for key in ("message", "detail", "error"):
            value = body.get(key)
            if isinstance(value, str) and value.strip():
                return value.strip()[:limit]
    return _rag_safe_body(resp, limit)


def _rag_safe_body(resp, limit: int = 300) -> str:
    try:
        text = resp.text or ""
    except Exception:
        return "<本文を読めませんでした>"
    text = text.strip().replace("\n", " ")
    return text[:limit] + ("…" if len(text) > limit else "")


# ==========================================================================
# ===== 元 rag/registry.py
# ナレッジベース（LightRAG環境）の登録簿。
#
# 環境はIT部門から「URLとAPIキー」の形で払い出される前提なので、コードにも
# env にも書かず、管理画面から足せるようにする。件数の上限は設けない
# （設備マニュアル・トラブル事例・改善事例…と増えていく想定のため）。
#
# ファイルにAPIキーが平文で入る。可能な環境では所有者だけに絞っている。
# ==========================================================================
import os
import uuid
from datetime import datetime, timezone


class RegistryError(Exception):
    """入力が不正なときに送出する（呼び出し側が400にして返す）。"""


_kb_lock = threading.Lock()


def _kb_path():
    return config.KNOWLEDGE_BASES_FILE


def _kb_read() -> list[dict]:
    p = _kb_path()
    if not p.exists():
        return []
    try:
        data = json.loads(p.read_text(encoding="utf-8"))
    except (OSError, ValueError) as e:
        print(f"[rag] 登録簿を読めませんでした: {p}（{e}）")
        return []
    return data if isinstance(data, list) else []


def _kb_write(items: list[dict]) -> None:
    p = _kb_path()
    p.parent.mkdir(parents=True, exist_ok=True)
    tmp = p.with_suffix(".json.tmp")
    tmp.write_text(json.dumps(items, ensure_ascii=False, indent=2), encoding="utf-8")
    os.replace(tmp, p)
    try:
        os.chmod(p, 0o600)          # APIキーを含むので所有者だけに絞る
    except OSError:
        pass


def _kb_redact(item: dict) -> dict:
    """APIキーを画面やAPIレスポンスに出さないための写し。"""
    out = dict(item)
    key = out.pop("api_key", "")
    out["has_api_key"] = bool(key)
    return out


def _kb_normalize_url(url: str) -> str:
    url = str(url or "").strip().rstrip("/")
    if not url:
        return ""
    if not url.startswith(("http://", "https://")):
        url = "http://" + url
    return url


def kb_list(include_secrets: bool = False) -> list[dict]:
    items = _kb_read()
    return items if include_secrets else [_kb_redact(i) for i in items]


def kb_enabled() -> list[dict]:
    """検索対象として有効なナレッジベース（APIキー付き）。"""
    return [i for i in _kb_read() if i.get("enabled", True)]


def kb_get(kb_id: str, include_secrets: bool = False) -> dict | None:
    for item in _kb_read():
        if item.get("id") == kb_id:
            return item if include_secrets else _kb_redact(item)
    return None


def kb_add(*, name: str, base_url: str, api_key: str = "", description: str = "",
           enabled: bool = True) -> dict:
    name = str(name or "").strip()
    base_url = _kb_normalize_url(base_url)
    if not name:
        raise RegistryError("名前を入力してください。")
    if not base_url:
        raise RegistryError("URLを入力してください。")
    with _kb_lock:
        items = _kb_read()
        if any(i.get("base_url") == base_url for i in items):
            raise RegistryError(f"このURLは既に登録されています: {base_url}")
        if any(i.get("name") == name for i in items):
            # 名前はAIに見せるツールの選択肢そのもの。重なると、AIも人も
            # どちらを指しているか決められない。
            raise RegistryError(f"この名前は既に使われています: {name}")
        item = {
            "id": uuid.uuid4().hex[:12],
            "name": name,
            "base_url": base_url,
            "api_key": str(api_key or "").strip(),
            "description": str(description or "").strip(),
            "enabled": bool(enabled),
            "created_at": datetime.now(timezone.utc).isoformat(),
        }
        items.append(item)
        _kb_write(items)
    return _kb_redact(item)


def kb_update(kb_id: str, **fields) -> dict:
    with _kb_lock:
        items = _kb_read()
        for item in items:
            if item.get("id") != kb_id:
                continue
            if "name" in fields:
                name = str(fields["name"] or "").strip()
                if not name:
                    raise RegistryError("名前を入力してください。")
                if any(i.get("name") == name and i.get("id") != kb_id for i in items):
                    raise RegistryError(f"この名前は既に使われています: {name}")
                item["name"] = name
            if "base_url" in fields:
                url = _kb_normalize_url(fields["base_url"])
                if not url:
                    raise RegistryError("URLを入力してください。")
                if any(i.get("base_url") == url and i.get("id") != kb_id for i in items):
                    raise RegistryError(f"このURLは既に登録されています: {url}")
                item["base_url"] = url
            if "description" in fields:
                item["description"] = str(fields["description"] or "").strip()
            if "enabled" in fields:
                item["enabled"] = bool(fields["enabled"])
            # 空文字のAPIキーは「変更なし」として扱う。画面はキーを伏せて表示するので、
            # 未入力＝据え置きが自然なため。
            if fields.get("api_key"):
                item["api_key"] = str(fields["api_key"]).strip()
            _kb_write(items)
            return _kb_redact(item)
    raise RegistryError("対象のナレッジベースが見つかりません。")


def kb_delete(kb_id: str) -> bool:
    with _kb_lock:
        items = _kb_read()
        remaining = [i for i in items if i.get("id") != kb_id]
        if len(remaining) == len(items):
            return False
        _kb_write(remaining)
    return True


def kb_test(kb_id: str) -> dict:
    """接続テスト。画面の「接続テスト」ボタンから呼ぶ。"""
    env = kb_get(kb_id, include_secrets=True)
    if env is None:
        raise RegistryError("対象のナレッジベースが見つかりません。")
    client = RagClient(env["base_url"], env.get("api_key", ""),
                       timeout=config.RAG_RETRIEVE_TIMEOUT)
    body = client.health()
    return {"ok": True, "status": body.get("status") or "healthy",
            "detail": {k: v for k, v in body.items()
                       if k in ("status", "core_version", "api_version", "pipeline_busy")}}


# ==========================================================================
# ===== 元 rag/settings.py
# 検索の効き方（利用者ごと）。
#
# ここが唯一の正本で、画面のフォーム・入力検証・初期値・保存がすべてこの表から
# 作られる。項目を増やすときは RAG_SPECS に1行足すだけでよい。
#
# 統合前の LightRAG アプリにあった「回答の形式」「引き継ぐ会話の数」はここに無い。
# どちらも LightRAG 側に回答を書かせるための設定で、このアプリでは回答を書くのは
# Agent（llm.py）だから。会話の履歴も chats.py が持っている。
# ==========================================================================

_RAG_MODES = (
    ("mix", "mix（グラフ＋ベクトル・推奨）"),
    ("hybrid", "hybrid（local＋global）"),
    ("local", "local（関連エンティティ中心）"),
    ("global", "global（関係性・全体像）"),
    ("naive", "naive（ベクトル検索のみ）"),
)

#: (キー, 表示名, 種類, 説明, 最小, 最大, 選択肢)
#: 上限は LightRAG 側の制約（MAX_QUERY_TOP_K=1000）より狭く取ってある。
#: 実用外の値を入れられると、検索が返らないだけで理由が分からないため。
RAG_SPECS = (
    ("retrieve_mode", "検索モード", "choice",
     "mix はナレッジグラフとベクトル検索の併用。迷ったら mix のままで構いません。",
     None, None, _RAG_MODES),
    ("chunk_top_k", "1つのナレッジベースから取る文章数", "int",
     "増やすと拾える情報は増えますが、検索が遅くなり、AIに渡す量も増えます。",
     1, 100, ()),
    ("top_k", "エンティティ／関係の取得数", "int",
     "ナレッジグラフから取り出す要素数。naive モードでは使われません。",
     1, 200, ()),
    ("max_context_chars", "参考情報の上限（文字数）", "int",
     "複数のナレッジベースの結果を統合したあと、AIに渡す量の上限です。",
     1000, 20000, ()),
)

RAG_SPEC_BY_KEY = {s[0]: s for s in RAG_SPECS}


def _rag_default(key: str):
    """初期値は config（env）から取る。

    値を二重に持たないための遅延評価。管理者が env の既定を変えれば、
    「初期値に戻す」の戻り先も一緒に変わる。
    """
    return {
        "retrieve_mode": config.RAG_RETRIEVE_MODE,
        "chunk_top_k": config.RAG_CHUNK_TOP_K,
        "top_k": config.RAG_TOP_K,
        "max_context_chars": config.RAG_MAX_CONTEXT_CHARS,
    }[key]


def rag_defaults() -> dict:
    return {s[0]: _rag_default(s[0]) for s in RAG_SPECS}


def _rag_coerce(spec: tuple, raw):
    """入力値を正規化する。不正なら ValueError。"""
    key, label, kind, _help, lo, hi, choices = spec
    if kind == "int":
        try:
            value = int(raw)
        except (TypeError, ValueError):
            raise ValueError(f"{label}: 数値を入力してください。")
        if lo is not None and value < lo:
            raise ValueError(f"{label}: {lo} 以上にしてください。")
        if hi is not None and value > hi:
            raise ValueError(f"{label}: {hi} 以下にしてください。")
        return value
    value = str(raw)
    if value not in {c[0] for c in choices}:
        raise ValueError(f"{label}: 選択できない値です。")
    return value


def rag_merge_settings(stored: dict) -> dict:
    """保存値を初期値へ重ねる。未知のキーと壊れた値は捨てる。

    項目を削除したあとも古い保存値がファイルに残るので、RAG_SPECS に無いキーは
    常に無視する。
    """
    result = rag_defaults()
    for key, raw in (stored or {}).items():
        spec = RAG_SPEC_BY_KEY.get(key)
        if spec is None:
            continue
        try:
            result[key] = _rag_coerce(spec, raw)
        except ValueError:
            continue                          # 壊れた保存値は初期値のまま
    return result


def rag_validate_settings(payload: dict) -> dict:
    """画面から来た値を検証する。1つでも不正なら ValueError。"""
    # 辞書でないものが来ると、下の in 判定が「文字列の部分一致」に化けて
    # 別の失敗の仕方をする。入口で1つの型に揃える
    if payload is not None and not isinstance(payload, dict):
        raise ValueError("設定の形式が正しくありません。")
    cleaned = {}
    for spec in RAG_SPECS:
        if spec[0] in (payload or {}):
            cleaned[spec[0]] = _rag_coerce(spec, payload[spec[0]])
    return cleaned


def rag_form_fields() -> list[dict]:
    """画面へ渡すフォーム定義。"""
    return [{"key": k, "label": label, "kind": kind, "help": help_,
             "min": lo, "max": hi,
             "choices": [{"value": v, "label": t} for v, t in choices]}
            for k, label, kind, help_, lo, hi, choices in RAG_SPECS]


# ==========================================================================
# ===== 元 rag/retriever.py
# 全ナレッジベースへのfan-out → 統合。
#
# 統合はスコア順の単純な足切りではなく「ナレッジベースごとのラウンドロビン」。
# スコアで切ると1つのKBが枠を埋め尽くすことがあり、質問に関係するのに
# 別のKBが締め出される。どのKBにも必ず枠が回るようにしておくと、
# 振り分けを多少誤っても効く。
# ==========================================================================
import hashlib
from concurrent.futures import ThreadPoolExecutor

#: いまの利用者（ツールの実処理から prefs を引くために置く）。
#: dispatch は引数に利用者を持たない（35個のツールすべての形が変わるため）。
#: リクエストを処理しているスレッドに置いて、web 側が質問のたびに入れ直す。
_rag_local = threading.local()


def set_current_user(user) -> None:
    _rag_local.user = user
    # 質問の切り替わり。出典番号の通し番号を 0 に戻す
    # （results.new_turn() と同じ場所から呼ばれる）
    _rag_local.sources = 0


def _current_user():
    return getattr(_rag_local, "user", None)


def rag_user_settings(user=None) -> dict:
    """利用者ごとの検索設定（未設定なら env の初期値）。"""
    import prefs
    return rag_merge_settings(prefs.load(user or _current_user()).get("rag_settings") or {})


def rag_excluded_ids(user=None) -> list[str]:
    """利用者が検索対象から外したナレッジベースのid。

    「選んだもの」ではなく「外したもの」を保存している。選択リスト方式にすると、
    管理者が新しいナレッジベースを足したとき、既存の利用者全員にそれが見えない
    ままになるため。除外方式なら新しいものは既定で検索対象に入る。
    """
    import prefs
    off = prefs.load(user or _current_user()).get("rag_off")
    return [str(i) for i in off] if isinstance(off, list) else []


def excluded_tables(user=None) -> list[str]:
    """利用者が分析の対象から外したテーブル名。

    ナレッジベースと同じく「選んだもの」ではなく「外したもの」を保存する。
    選択リスト方式にすると、新しく取り込んだ表が既存の利用者全員に
    見えないままになるため。除外方式なら新しい表は既定で対象に入る。
    """
    import prefs
    off = prefs.load(user or _current_user()).get("tables_off")
    return [str(t) for t in off] if isinstance(off, list) else []


def rag_targets(user=None) -> list[dict]:
    """いまの利用者が検索できるナレッジベース（APIキー付き）。"""
    off = set(rag_excluded_ids(user))
    return [e for e in kb_enabled() if e.get("id") not in off]


def rag_available(user=None) -> bool:
    """AIにナレッジ検索ツールを渡すかどうか。

    利用者が1件も選んでいないなら渡さない。渡すと、AIは探せない情報源を
    探しに行って毎回空振りし、そのぶん往復と費用が増える。
    """
    return bool(rag_targets(user))


def rag_retrieve_all(question: str, environments: list[dict],
                     settings: dict) -> list[dict]:
    """全ナレッジベースに並列で検索をかける。1つ落ちても残りで続行する。

    戻り値は環境ごとの {name, chunks, error}。
    """
    if not environments:
        return []

    def _one(env: dict) -> dict:
        client = RagClient(env["base_url"], env.get("api_key", ""),
                           timeout=config.RAG_RETRIEVE_TIMEOUT)
        try:
            body = client.retrieve(question,
                                   mode=settings["retrieve_mode"],
                                   chunk_top_k=settings["chunk_top_k"],
                                   top_k=settings["top_k"])
        except RagError as exc:
            print(f"[rag] 検索失敗 kb={env.get('name')}: {exc}")
            return {"id": env.get("id"), "name": env.get("name"),
                    "chunks": [], "error": str(exc)}
        except Exception as exc:              # 想定外でも他のKBは活かす
            print(f"[rag] 検索でエラー kb={env.get('name')}: {exc}")
            return {"id": env.get("id"), "name": env.get("name"),
                    "chunks": [], "error": f"検索でエラー: {exc}"}
        data = body.get("data") or {}
        return {"id": env.get("id"), "name": env.get("name"),
                "chunks": _rag_normalize_chunks(data.get("chunks") or []),
                "error": None}

    workers = max(1, min(config.RAG_FANOUT_WORKERS, len(environments)))
    with ThreadPoolExecutor(max_workers=workers) as pool:
        return list(pool.map(_one, environments))


def _rag_normalize_chunks(raw: list) -> list[dict]:
    """chunks の形が実装差で揺れても壊れないように正規化する。"""
    out = []
    for item in raw:
        if isinstance(item, str):
            content, file_path, score = item, "", None
        elif isinstance(item, dict):
            content = (item.get("content") or item.get("text")
                       or item.get("chunk") or "")
            file_path = item.get("file_path") or item.get("source") or ""
            score = item.get("score")
        else:
            continue
        content = str(content).strip()
        if not content:
            continue
        out.append({"content": content, "file_path": str(file_path), "score": score})
    return out


#: 抜粋として画面に出す長さ。ここを変えたら画面の説明文も直すこと。
RAG_EXCERPT_CHARS = 200


def _rag_used_sources(add: int = 0) -> int:
    """この質問の中で、いままでに振った出典番号の数。

    dispatch は引数に「いまのターン」を持たないので、処理中のスレッドに置く。
    results.new_turn() と同じ考え方で、質問の切り替わりで 0 に戻す。
    """
    cur = getattr(_rag_local, "sources", 0)
    if add:
        _rag_local.sources = cur + add
    return cur


def rag_merge_context(results: list[dict], *, char_budget: int,
                      start_index: int = 0) -> tuple:
    """ナレッジベースをまたいでチャンクをラウンドロビンで拾い、予算内に収める。

    戻り値は (参考情報テキスト, 出典リスト, 予算で落とした数)。
    予算を文字数で見るのは、この先の渡し先がAgentのLLMだから（統合前は
    LightRAG に投げ直していたので 64KiB のバイト制限で見ていた）。

    start_index は出典番号の開始位置。1回の回答の中で search_knowledge_base を
    何度も呼ぶことがあり、毎回1から振り直すと [出典1] が別々の文書を指してしまう。
    """
    live = [r for r in results if not r.get("error") and r.get("chunks")]
    if not live:
        return "", [], 0

    seen: set = set()
    blocks: list = []
    sources: list = []
    state = {"used": 0, "index": int(start_index or 0), "dropped": 0}

    def take(result: dict, chunk: dict) -> bool:
        """1チャンクを採る。予算に入らなければ False（その場では捨てるだけ）。"""
        digest = hashlib.sha256(chunk["content"].encode("utf-8")).hexdigest()
        if digest in seen:
            return False                      # 同じ文章が複数のKBに入っていることがある
        index = state["index"] + 1
        label = f"[出典{index}] {result['name']}"
        if chunk["file_path"]:
            label += f" / {chunk['file_path']}"
        body, cut = chunk["content"], False
        block = f"{label}\n{body}\n"
        if state["used"] + len(block) > char_budget:
            # 1件も採れていないなら、切り詰めてでも採る。
            # 全部捨てると、検索は当たっているのにAIには「無かった」と伝わり、
            # 利用者にも「ありません」と答えてしまう
            room = char_budget - state["used"] - len(label) - 2
            if sources or room < 200:
                state["dropped"] += 1
                return False
            body, cut = body[:room] + "…（長いため以下省略）", True
            block = f"{label}\n{body}\n"
        seen.add(digest)
        blocks.append(block)
        state["used"] += len(block)
        state["index"] = index
        sources.append({
            "index": index,
            "knowledge_base": result["name"],
            "kb_id": result.get("id"),
            "file_path": chunk["file_path"],
            "score": chunk["score"],
            # 画面に出すのは先頭だけ。切れているときは印を付ける
            # （「全文」と書くと嘘になるので、画面の文言もそれに合わせてある）
            "excerpt": chunk["content"][:RAG_EXCERPT_CHARS],
            "excerpt_cut": len(chunk["content"]) > RAG_EXCERPT_CHARS,
            "excerpt_chars": RAG_EXCERPT_CHARS,
            "truncated": cut,
        })
        return True

    # 第1周: どのKBからも最低1件は拾う。
    # 短いチャンクから順に入れるのは、入る「KBの数」を最大にするため。
    # ここを素通りさせると、先に登録されたKBが予算を食い切り、後ろのKBが
    # 正常に応答しているのに1件も回答に出ない（無言の脱落）ことになる。
    first = [(r, r["chunks"][0]) for r in live if r["chunks"]]
    for result, chunk in sorted(first, key=lambda rc: len(rc[1]["content"])):
        take(result, chunk)

    # 第2周以降: 残りをラウンドロビンで配る。
    # 入らないチャンクがあっても打ち切らない（次のKBには入るかもしれない）。
    max_depth = max(len(r["chunks"]) for r in live)
    for depth in range(1, max_depth):
        for result in live:
            if depth < len(result["chunks"]):
                take(result, result["chunks"][depth])
    return "\n".join(blocks), sources, state["dropped"]


def rag_all_failed_message(failures: list[dict]) -> str:
    """全ナレッジベースが落ちたときの文面を作る。

    原因が全部同じことは多い（キーがまとめて期限切れ、サーバが止まっている等）。
    素直に連結すると同じ文章がKBの数だけ並んで読めなくなるので、
    同一のエラーはまとめて1回だけ出す。
    """
    grouped: dict = {}
    for f in failures:
        grouped.setdefault(f["error"], []).append(f["knowledge_base"])
    if len(grouped) == 1:
        (error,) = grouped
        return f"すべてのナレッジベースへの問い合わせに失敗しました。{error}"
    lines = [f"・{'、'.join(names)}: {error}" for error, names in grouped.items()]
    return "すべてのナレッジベースへの問い合わせに失敗しました。\n" + "\n".join(lines)


# ==========================================================================
# ===== 元 tools/knowledge.py
# 社内文書を調べるツール（情報検索系）。
#
# SQLが「数字を取る」のに対して、こちらは「文章を取る」。AIから見れば同じ
# ツールの1つで、どちらを使うか・両方使うかは依頼の内容から自分で決める。
#   「復旧手順を教えて」            → ナレッジ検索だけ
#   「8月のMTTRを教えて」           → SQLだけ
#   「MTTR悪化の原因を過去事例から」 → SQL → 分析 → ナレッジ検索 → Excel
#
# このツールは回答を書かない。取ってきた文章と出典を返すだけで、文章を書くのは
# 呼び出し元のAgent。そのため他のツール（SQLの集計結果・Excelの中身）と
# 突き合わせた回答が書ける。
# ==========================================================================


def _search_knowledge_base(args: dict, scope: list[dict]) -> dict:
    query = str(args.get("query") or "").strip()
    if not query:
        return _err("検索したい内容（query）を指定してください。")

    targets = rag_targets()
    if not targets:
        return _err("検索できるナレッジベースがありません。"
                    "「ナレッジベース」画面で登録されているか、"
                    "サイドバーの検索対象から全て外していないかを確認してください。")

    # 名前で絞る。AIが知らない名前を書いてきたら、黙って全件にせず教える
    # （そのKBが無いのに「調べたが無かった」と結論されるのを防ぐ）。
    wanted = args.get("knowledge_bases") or []
    if isinstance(wanted, str):
        wanted = [wanted]
    wanted = [str(w).strip() for w in wanted if str(w).strip()]
    if wanted:
        by_name = {t["name"]: t for t in targets}
        unknown = [w for w in wanted if w not in by_name]
        if unknown:
            return _err(f"「{'」「'.join(unknown)}」は、いまの検索対象に入っていません。"
                        + (f"検索できるのは「{'」「'.join(by_name)}」です。"
                           if by_name else "検索できる社内文書がありません。")
                        + "（サイドバーの LightRAG でチェックを入れると増やせます）")
        targets = [by_name[w] for w in wanted]

    settings = rag_user_settings()
    # AIが件数を指定してきたときはそれを優先する（「もっと広く探して」に応えるため）。
    # 上限は利用者の設定の2倍まで。青天井にすると1回の検索でAIの読める量を使い切る。
    want_k = args.get("chunk_top_k")
    if want_k:
        try:
            settings = {**settings,
                        "chunk_top_k": max(1, min(int(want_k),
                                                  settings["chunk_top_k"] * 2, 100))}
        except (TypeError, ValueError):
            pass

    results = rag_retrieve_all(query, targets, settings)
    failures = [{"knowledge_base": r["name"], "error": r["error"]}
                for r in results if r.get("error")]
    if failures and len(failures) == len(results):
        return _err(rag_all_failed_message(failures))

    # 出典番号は1回の質問の中で通しにする。検索のたびに1へ戻すと、
    # 同じ回答の中の [出典1] が別々の文書を指してしまう
    start = _rag_used_sources()
    context, sources, dropped = rag_merge_context(
        results, char_budget=int(settings["max_context_chars"]), start_index=start)
    _rag_used_sources(len(sources))

    if not sources:
        # 「予算に入らなくて渡せなかった」と「本当に無かった」は別のこと。
        # 一緒にすると、当たっているのに「ありません」と答えてしまう
        note = ("該当する文章は見つかりませんでした。言い回しを変えて"
                "もう一度だけ試すか、見つからなかったことをそのまま利用者に伝えること。"
                "推測で答えてはいけない。")
        if dropped:
            note = (f"文章は {dropped} 件見つかりましたが、1件も参考情報に入りませんでした"
                    "（1件あたりが長く、渡せる文字数の上限を超えたため）。"
                    "「見つからなかった」とは言わないこと。"
                    "利用者には、チャット画面の検索設定で"
                    "「参考情報の文字数上限」を上げるよう案内すること。")
        return {"ok": True, "llm_content": _json({
            "tool": "search_knowledge_base",
            "query": query,
            "searched": [r["name"] for r in results if not r.get("error")],
            "found": 0,
            "dropped_for_budget": dropped,
            "failures": failures,
            "note": note,
        }), "render": {"role": "assistant", "kind": "sources", "query": query,
                       "sources": [], "failures": failures, "dropped": dropped,
                       "searched": [r["name"] for r in results if not r.get("error")]}}

    return {"ok": True, "llm_content": _json({
        "tool": "search_knowledge_base",
        "query": query,
        "searched": [r["name"] for r in results if not r.get("error")],
        "found": len(sources),
        "context": context,
        "sources": [{k: v for k, v in s.items() if k != "excerpt"} for s in sources],
        "failures": failures,
        "note": "回答は context に書かれていることだけを根拠にすること。"
                "根拠にした箇所には [出典1] のように出典番号を必ず添える。"
                "context に無いことは推測せず、分からないと述べること。"
                "複数のナレッジベースで内容が食い違う場合は、両方を出典付きで併記する。",
    }), "render": {"role": "assistant", "kind": "sources", "query": query,
                   "sources": sources, "failures": failures,
                   "searched": [r["name"] for r in results if not r.get("error")]}}


HANDLERS_knowledge = {"search_knowledge_base": _search_knowledge_base}

# SQLは受け取らない（材料はDBではなくナレッジベース）
SQL_TOOLS_knowledge: set = set()

# 全員に渡す。社内文書を調べるのは、このアプリの一般利用者の主目的そのもの。
ADMIN_TOOLS_knowledge: set = set()

#: 宣言。BUILTIN_TOOLS に足しておくと、必須引数の検査（_missing_required）と
#: 配列引数の直し（_coerce_lists）がそのまま効く。
#: 実際にAIへ渡すときは build_tools が、登録中のナレッジベースの名前と説明を
#: 差し込んだものに置き換える（登録は運用中に増えるため）。
KNOWLEDGE_TOOLS = [{
    "type": "function",
    "function": {
        "name": "search_knowledge_base",
        "description": (
            "社内文書（マニュアル・トラブル事例・改善事例・規程など）を検索して、"
            "根拠になる本文と出典を取り出す。手順・原因・対処・規則・過去の事例など、"
            "表の数値では答えられない『書かれていること』を調べるときに使う。"
            "数値の集計はこのツールではなく run_sql_query を使うこと。"
            "回答では取得した [出典n] を必ず示すこと。"),
        "parameters": {
            "type": "object",
            "properties": {
                "query": {
                    "type": "string",
                    "description": "検索したい内容。利用者の言葉のままではなく、"
                                   "探したい事柄が伝わる語を並べるとよい"
                                   "（対象・部位・行いたいこと、のように語を並べる）。",
                },
                "knowledge_bases": {
                    "type": "array",
                    "items": {"type": "string"},
                    "description": "探しに行くナレッジベースの名前。"
                                   "省略すると全体から探す。迷うときは省略してよい。",
                },
                "chunk_top_k": {
                    "type": "integer",
                    "description": "1つのナレッジベースから取る文章の数。"
                                   "省略時は利用者の設定に従う。"
                                   "見つからなかったときだけ増やして呼び直す。",
                },
            },
            "required": ["query"],
        },
    },
}]


def knowledge_tool_schemas() -> list[dict]:
    """いま登録されているナレッジベースを差し込んだ宣言。

    名前と説明を選択肢としてAIに見せる。AIがツールを選ぶ材料は名前と説明しか
    ないので（要件定義書 6章）、ここが検索の当たり外れをいちばん左右する。
    説明は「ナレッジベース」画面で管理者が書いたものがそのまま入る。

    見せるのは、その利用者がサイドバーで選んでいるものだけ。外したものまで
    並べると、AIはあると思って呼び、そのたびに断られて往復を1回損する
    （利用者から見ると「選んでいないものの話をした末に失敗する」ことになる）。
    """
    envs = rag_targets()
    if not envs:
        return []
    names = [e["name"] for e in envs]
    lines = []
    for e in envs:
        desc = (e.get("description") or "").strip()
        lines.append(f"- {e['name']}" + (f": {desc}" if desc else ""))
    schema = json.loads(json.dumps(KNOWLEDGE_TOOLS[0]))     # 原本を壊さない
    fn = schema["function"]
    fn["description"] += "\n\n登録されているナレッジベース:\n" + "\n".join(lines)
    fn["parameters"]["properties"]["knowledge_bases"]["items"]["enum"] = names
    return [schema]


# ==========================================================================
# ===== 元 tools/__init__.py
# LLM(OpenAI互換 function calling)に渡すツール定義と実行ロジック。
#
# 組み込みツールの系統:
#   調べる    run_sql_query / describe_table / search_knowledge_base
#   集計する  pivot_table / analyze_stats
#   描く      plot_comparison / plot_trend / plot_composition / plot_distribution /
#             plot_relationship / plot_kpi / plot_dual_axis
#   統計      hypothesis_test / regression / distribution_analysis
#   時系列    forecast / timeseries_analysis / detect_anomalies
#   試算      monte_carlo_simulation / scenario_analysis / bootstrap_estimate
#   分ける    clustering / abc_analysis
#   業務分析  compare_periods / funnel_analysis / cohort_analysis / market_basket /
#             survival_analysis / data_quality
#   ファイル  explore_import_files（取り込み元フォルダを読むだけ。管理者のみ）
#   文書検索  search_knowledge_base（LightRAGのナレッジベース。登録が1件も無ければ渡さない）
#   自己分析  analyze_usage（このアプリ自身の使われ方。管理者のみ）
#   出す      export_excel / export_csv / export_text / export_pptx / export_docx /
#             build_report
#   送る      find_mail_recipients / compose_email
#             ※ 実際の送信はユーザーが画面のボタンを押したときだけ。
#               誤送信は取り消せないので、LLMには下書きまでしかさせない。
#
# データを取るツールは sql の代わりに result_id を受け取れる。前のツールが返した
# データをそのまま使えるので、同じSQLを何度も流さずに済む（results.py 参照）。
#
# これに加えて、ユーザーが画面から定義したSQLテンプレート型ツール
# （各DBの .meta.yaml の tools:）を実行時に合成する。custom_tools.py 参照。
# 組み込みツールは .meta.yaml の builtin_tools: で無効化・説明の上書きができる。
#
# dispatch(name, arguments_json, scope, entries, admin) の戻り値:
#   {
#     "ok": bool,
#     "llm_content": str,        # LLMへ返すテキスト(JSON)。トークン節約のため要約。
#     "render": dict | None,     # UI描画用アイテム(app側が kind を見て解釈)
#   }
#
# scope は質問ごとに自動判定された対象DB群（web/chat_bp.py の _auto_scope）:
#   [{"path": str, "alias": str, "tables": [...]}, ...]
#
# ファイルの分かれ方:
#   schemas.py  … LLMに見せるツールの宣言（JSON Schema）。処理は書かない
#   common.py   … 実処理が共通で使う小道具（データの取り出しもここ）
#   results.py  … 取ったデータの置き場（result_id で使い回す）
#   query.py    … 調べる・集計する・描く・出す
#   stats.py    … 統計と試算
#   business.py … 期間比較・ファネル・コホート・併売・異常検知・生存時間・品質
#   reports.py  … PowerPoint / Word / 画面用レポート
#   mail.py     … 宛先探しと下書き
#   files.py    … 取り込み元フォルダの調査（管理者のみ）
#   usage.py    … このアプリ自身の利用状況（管理者のみ）
#   knowledge.py… 社内文書の検索（rag/ の連携層を使う）
# ツールを1つ足すときは、宣言(schemas)と実処理(各モジュール)の2箇所を触る。
# 実処理を置いたモジュールの HANDLERS に名前を登録すれば dispatch から引ける。
# 管理者だけに渡したいものは、そのモジュールの ADMIN_TOOLS にも名前を入れる。
# ==========================================================================
import json

import config
import custom_tools
import db
import excel
import exports
import verify

# ナレッジ検索の宣言もここで合流させる。こうしておくと _missing_required と
# _coerce_lists（どちらも BUILTIN_TOOLS から表を作る）がそのまま効く。
# ただしAIへ実際に渡す宣言は build_tools が組み立て直す（_DYNAMIC_TOOLS を参照）。
BUILTIN_TOOLS = BUILTIN_TOOLS + KNOWLEDGE_TOOLS

#: 宣言を実行時に組み立て直すツール。BUILTIN_TOOLS の固定の宣言は使わない。
#: ナレッジベースは運用中に増減するので、選択肢（名前と説明）を起動時には決められない。
_DYNAMIC_TOOLS = {"search_knowledge_base"}

# (実処理, SQLを受け取るもの, 管理者専用) をモジュールごとに並べる。
# 並び順は統合前の (query, stats, reports, mail, business, files, usage) のまま。
# 名前が重なったときにどちらが残るかを変えないため、順序は動かさないこと。
_MODULES = (
    (HANDLERS_query, SQL_TOOLS_query, ()),
    (HANDLERS_stats, SQL_TOOLS_stats, ()),
    (HANDLERS_reports, SQL_TOOLS_reports, ()),
    (HANDLERS_mail, SQL_TOOLS_mail, ()),
    (HANDLERS_business, SQL_TOOLS_business, ()),
    (HANDLERS_files, SQL_TOOLS_files, ADMIN_TOOLS_files),
    (HANDLERS_usage, SQL_TOOLS_usage, ADMIN_TOOLS_usage),
    (HANDLERS_knowledge, SQL_TOOLS_knowledge, ADMIN_TOOLS_knowledge),
)

# ツール名 -> 実処理。各モジュールが自分のぶんを申告する。
_HANDLERS = {name: fn for m in _MODULES for name, fn in m[0].items()}

# SQLを受け取る組み込みツール（実行前プレビュー表示の対象）
SQL_TOOLS = {name for m in _MODULES for name in m[1]}

# 管理者にだけ渡すツール。画面側で管理者専用になっているものは、
# AI経由でも同じ制限にしないと抜け道になる（取り込み元フォルダの中身など）。
ADMIN_TOOLS = {name for m in _MODULES for name in m[2]}


def render_sql(tool: dict) -> str:
    """UIのプレビュー用。実行時は :name のままバインドするので置換はしない。"""
    return str(tool.get("sql") or "").strip()


def _run_custom(tool: dict, args: dict, scope: list[dict]) -> dict:
    sql = render_sql(tool)
    try:
        params = custom_tools.coerce_params(tool, args)
    except ValueError as e:
        return _err(str(e))
    # ツールは作るときにDBを意識させないので、SQLが選択外のDBに入ることがある。
    # 必要なぶんは繋いでから実行する（結果を預ける先も同じ範囲にする）。
    scope = db.widen_scope(sql, scope)
    # ファイルに出すツールは全行（Excelはシート上限で丸める）。画面用は2,000行。
    kind = tool.get("render") or "table"
    cap = (min(config.EXPORT_MAX_ROWS, 1_048_575) if kind in ("excel", "csv")
           else None)
    try:
        columns, rows, truncated = db.run_select(sql, scope, params=params,
                                                 max_rows=cap)
    except Exception as e:
        return _err(f"ツール '{tool.get('name')}' のSQL実行エラー: {e}")

    chart = tool.get("chart") or {}
    sample = rows[: config.SAMPLE_ROWS_FOR_LLM]

    # 取った表を預けて result_id を返す。組み込みツールは前からこうしているのに
    # ユーザー定義ツールだけ返しておらず、「このツールの結果をグラフにして」と
    # 言われてもAIには渡す手段が無かった（SQLはAIに見せていないので取り直せない）。
    # これがあれば、表で作ったツールでも後からグラフ・Excel・統計に回せる。
    keep = rows[: config.MAX_RESULT_ROWS]
    rid = _results.put(scope, columns, keep, truncated or len(rows) > len(keep),
                       sql=sql, label=f"{tool.get('name')}（ユーザー定義ツール）")
    llm_content = _json({
        "tool": tool.get("name"),
        "columns": columns,
        "row_count": len(rows),
        "truncated": truncated,
        "rows": [list(r) for r in sample],
        "result_id": rid,
        "note": ((f"全{len(rows)}行中 先頭{len(sample)}行を表示。"
                  if len(rows) > len(sample) else "")
                 + f"この結果は result_id '{rid}' で他のツールに渡せます"
                   "（グラフを描く・集計する・統計をかける・"
                   "Excel/CSV/PowerPoint/Word にする、など）。"),
    })

    if kind == "none":
        return {"ok": True, "llm_content": llm_content, "render": None}
    if kind in ("excel", "csv"):
        sheet = {"name": chart.get("title") or tool.get("name") or "Sheet1",
                 "columns": columns, "rows": rows,
                 "note": f"{cap:,}行で切り詰め" if truncated else ""}
        # 会話に残すのは先頭だけ（ファイルの中身は data にある）
        sheet_preview = _preview_rows(columns, rows, sheet["name"], sheet["note"])
        base = chart.get("filename") or tool.get("name")
        try:
            if kind == "excel":
                data = excel.build_excel([sheet], title=tool.get("description"))
                filename, mime = exports.safe_filename(base, "xlsx"), exports.XLSX_MIME
            else:
                enc = chart.get("encoding") or exports.DEFAULT_ENCODING
                data = exports.build_csv(columns, rows, enc)
                filename, mime = exports.safe_filename(base, "csv"), exports.CSV_MIME
        except Exception as e:
            return _err(f"ファイルの作成に失敗しました: {e}")
        return {"ok": True, "llm_content": _json({
            "status": "file_ready", "tool": tool.get("name"),
            "filename": filename, "columns": columns, "row_count": len(rows),
            "note": "ユーザーの画面に保存済み。",
        }), "render": {
            "role": "assistant", "kind": "file", "filename": filename,
            "mime": mime, "data": data, "sheets": [sheet_preview],
        }}
    if kind == "chart":
        missing = [c for c in (chart.get("x"), chart.get("y")) if c and c not in columns]
        if missing:
            return _err(f"グラフ用の列が結果にありません: {missing} / 利用可能: {columns}")
        return {"ok": True, "llm_content": llm_content, "render": {
            "role": "assistant", "kind": "chart", "columns": columns, "rows": rows,
            "chart_type": chart.get("chart_type", "bar"),
            "x": chart.get("x"), "y": chart.get("y"), "color": chart.get("color"),
            "barmode": chart.get("barmode"), "title": chart.get("title", "") or tool.get("name"),
        }}
    if kind == "chart_dual":
        bar_y = chart.get("bar_y") or []
        line_y = chart.get("line_y") or []
        needed = [chart.get("x")] + list(bar_y) + list(line_y)
        missing = [c for c in needed if c and c not in columns]
        if missing:
            return _err(f"2軸グラフ用の列が結果にありません: {missing} / 利用可能: {columns}")
        return {"ok": True, "llm_content": llm_content, "render": {
            "role": "assistant", "kind": "chart_dual", "columns": columns, "rows": rows,
            "x": chart.get("x"), "bar_y": bar_y, "line_y": line_y,
            "left_title": chart.get("left_title"), "right_title": chart.get("right_title"),
            "title": chart.get("title", "") or tool.get("name"),
        }}
    return {"ok": True, "llm_content": llm_content, "render": {
        "role": "assistant", "kind": "table",
        "columns": columns, "rows": rows, "truncated": truncated,
    }}


#: SQLを組み立てるツールに共通で足す「日本語の解説」の引数。
#: 画面ではSQLの下に出る。SQLを読めない人が、根拠を読み解けるようにするためのもの。
#: 定義を1か所に置いて build_tools で配るので、SQLツールを足しても自動で付く。
_EXPLANATION_PARAM = {
    "type": "string",
    "description": (
        "組み立てたSQLの日本語の解説（3〜5行）。SQLを読めない人にも分かる言葉で、"
        "「どの表を使うか」「どうつないだか（結合の条件と、なぜその条件か）」"
        "「どう絞ったか・集計したか」「1行が何を表すか」の順に書く。"
        "INNER JOIN のような構文用語を並べるのではなく、何をしているかを説明すること。"
    ),
}


def _with_explanation(t: dict) -> dict:
    """SQLツールの宣言に explanation 引数を足したものを返す。"""
    fn = t.get("function") or {}
    params = fn.get("parameters") or {}
    props = params.get("properties") or {}
    if "explanation" in props:
        return t
    return {**t, "function": {**fn, "parameters": {
        **params,
        "properties": {**props, "explanation": _EXPLANATION_PARAM},
        # 必須にはしない。解説が無くてもSQLは実行できるべきで、
        # 必須にすると解説を書き損ねただけで質問全体が止まる
        "required": list(params.get("required") or ()),
    }}}


def build_tools(entries: list[dict], admin: bool = False) -> list[dict]:
    """組み込み（無効化・説明上書きを反映）＋ナレッジ検索＋ユーザー定義 のツール定義一覧。

    admin=False のときは管理者専用のツールを渡さない。渡さなければ
    AIはその存在を知らないので、呼ばれること自体が起きない。
    """
    ov = custom_tools.builtin_overrides(entries)
    out = []
    for t in BUILTIN_TOOLS:
        name = t["function"]["name"]
        if name in _DYNAMIC_TOOLS:
            continue                       # 宣言は下で組み立て直す
        if name in ADMIN_TOOLS and not admin:
            continue
        o = ov.get(name) or {}
        if o.get("enabled") is False:
            continue
        if o.get("description"):
            t = {**t, "function": {**t["function"], "description": o["description"]}}
        out.append(_with_explanation(t) if name in SQL_TOOLS else t)
    # ナレッジ検索。登録が1件も無ければ渡さない（存在しない情報源を
    # 探しに行かせても、往復と費用が増えるだけで何も出てこない）。
    for t in knowledge_tool_schemas():
        o = ov.get(t["function"]["name"]) or {}
        if o.get("enabled") is False:
            continue
        if o.get("description"):
            t = {**t, "function": {**t["function"], "description": o["description"]}}
        out.append(t)
    for tool in custom_tools.collect_everywhere(entries):
        if not custom_tools.validate_custom_tool(tool, set()):     # 壊れた定義はAIに渡さない
            out.append(custom_tools.to_schema(tool))
    return out


def _required_params() -> dict:
    """スキーマで必須になっている引数。{ツール名: (引数名, ...)}"""
    return {t["function"]["name"]: tuple((t["function"].get("parameters") or {})
                                         .get("required") or ())
            for t in BUILTIN_TOOLS}


_REQUIRED = _required_params()


#: スキーマ上は必須でも、実処理が既定値を持っている引数。
#: ここまで止めると、これまで通っていた呼び出しが弾かれてしまう
#: （表題やファイル名は無ければツール側が付ける）。
_HAS_DEFAULT = {"title", "filename", "chart_type", "purpose"}


def _missing_required(name: str, args: dict) -> list[str]:
    """必須なのに渡ってこなかった引数（既定値を持つものは除く）。

    LLMは required を落とすことがある。そのまま実処理へ渡すと、
    pandas の "'[None] not in index'" のような内部エラーになって返る。
    これでは何を直せばよいか分からず、同じ呼び出しを繰り返して打ち切られる。
    ここで止めて、足りない引数の名前をそのまま返す。
    """
    if not isinstance(args, dict):
        return []
    out = []
    for k in _REQUIRED.get(name) or ():
        if k in _HAS_DEFAULT:
            continue
        v = args.get(k)
        # 0 や False は正しい値なので、空とみなすのは None と空の入れ物だけ
        if v is None or (isinstance(v, (str, list, dict, tuple)) and len(v) == 0):
            out.append(k)
    return out


def _string_list_params() -> dict:
    """スキーマ上「文字列の配列」になっている引数。{ツール名: {引数名, ...}}"""
    out: dict = {}
    for t in BUILTIN_TOOLS:
        fn = t["function"]
        props = (fn.get("parameters") or {}).get("properties", {})
        names = {k for k, v in props.items()
                 if v.get("type") == "array"
                 and (v.get("items") or {}).get("type") == "string"}
        if names:
            out[fn["name"]] = names
    return out


_LIST_PARAMS = _string_list_params()


def _coerce_lists(name: str, args: dict) -> dict:
    """配列で受ける引数に文字列が1つ来たら、要素1つの配列として扱う。

    LLMは列名が1つのとき index="地域" のように素の文字列で渡してくることがある。
    そのまま渡すと文字列が1文字ずつに散り、「'地' という列がありません」という
    人には意味の分からないエラーになる（日本語の列名だと必ずこうなる）。
    ここで直せば、13個ある同じ形の引数すべてに効く。
    """
    wanted = _LIST_PARAMS.get(name)
    if not wanted or not isinstance(args, dict):
        return args
    for k in wanted:
        v = args.get(k)
        if isinstance(v, str):
            args[k] = [v.strip()] if v.strip() else []
    return args


def _gather_sqls(node, scope: list[dict], acc: list) -> None:
    """呼び出しの引数から、実行されるSQLを全部拾う。

    レポートの節・Excelのシートのように入れ子の中にも sql がある。
    result_id で前の結果を使い回している場合は、その元のSQLを引く。
    """
    if isinstance(node, dict):
        for k, v in node.items():
            if k == "sql" and isinstance(v, str) and v.strip():
                acc.append(v)
            elif k == "result_id" and isinstance(v, str) and v.strip():
                entry = _results.get(scope, v)
                if entry and entry.get("sql"):
                    acc.append(entry["sql"])
            else:
                _gather_sqls(v, scope, acc)
    elif isinstance(node, list):
        for v in node:
            _gather_sqls(v, scope, acc)


def _attach_verification(res: dict, sqls: list[str], scope: list[dict]) -> dict:
    """実行後の相互検証。触れたテーブルに関係する検算を突き合わせる。

    不一致があれば res["verify_alerts"] に積む。画面とLLMへの出し方は
    呼び出し側（chat側）が決める（同じ警告を会話の中で繰り返さないため）。
    検証自体の失敗で回答を止めない。
    """
    if not res.get("ok") or not sqls:
        return res
    try:
        alerts = verify.alerts_for(sqls, scope)
    except Exception as e:
        print(f"[verify] 検算でエラー（回答は続行）: {e}")
        return res
    if alerts:
        res["verify_alerts"] = alerts
    return res


def dispatch(name: str, arguments_json: str | None, scope: list[dict],
             entries: list[dict] | None = None, admin: bool = False) -> dict:
    try:
        args = json.loads(arguments_json) if arguments_json else {}
    except json.JSONDecodeError as e:
        return _err(f"ツール引数のJSON解析に失敗しました: {e}")

    # 渡していないツールを名指しで呼ばれても実行しない（守りは2箇所で持つ）
    if name in ADMIN_TOOLS and not admin:
        return _err(f"'{name}' は管理者だけが使えます。")

    args = _coerce_lists(name, args)
    missing = _missing_required(name, args)
    if missing:
        return _err(f"'{name}' の必須の引数が指定されていません: {'、'.join(missing)}。"
                    f"（{name} の必須引数は {'、'.join(_REQUIRED[name])}）"
                    "この引数を入れて呼び直してください。列名が分からないときは、"
                    "先に describe_table か run_sql_query で列を確認すること。")

    sqls: list = []
    _gather_sqls(args, scope, sqls)

    handler = _HANDLERS.get(name)
    if handler:
        try:
            return _attach_verification(handler(args, scope), sqls, scope)
        except Exception as e:  # ツールの例外でアプリを落とさない
            return _err(f"ツール '{name}' の実行でエラー: {e}")

    tool = next((t for t in custom_tools.collect_everywhere(entries or []) if t.get("name") == name), None)
    if tool is None:
        return {"ok": False, "llm_content": _json({"error": f"未知のツール: {name}"}), "render": None}
    try:
        sqls.append(render_sql(tool))
        return _attach_verification(_run_custom(tool, args, scope), sqls, scope)
    except Exception as e:
        return _err(f"ツール '{name}' の実行でエラー: {e}")


# ==========================================================================
# ===== 元 llm.py
# OpenAI / OpenAI互換API クライアント / system prompt 生成 / AI下書き。
# ==========================================================================
import json
import re
import time
from datetime import datetime
from pathlib import Path

from openai import OpenAI

import catalog
import config
import custom_tools
import tools

_client: OpenAI | None = None
_models_client: OpenAI | None = None


def _derived_base(full_url: str, suffix: str) -> str:
    """フルパスのURLから、SDKに渡す base_url を導出する（SDKが末尾を自動付与するため）。"""
    u = str(full_url or "").strip().rstrip("/")
    return u[: -len(suffix)] if u.endswith(suffix) else ""


def is_configured() -> bool:
    # 接続先・キーとも「モデル設定」画面で保存した値 > env の順
    return bool(llm_chat_url() and llm_api_key())


def client() -> OpenAI:
    """チャット（AI呼び出し）用クライアント。"""
    global _client
    if _client is None:
        _client = OpenAI(
            base_url=_derived_base(llm_chat_url(), "/chat/completions") or None,
            api_key=llm_api_key() or "not-set",
        )
    return _client


def models_client() -> OpenAI:
    """モデル一覧（/models）用クライアント。チャットと接続先を分けられるよう別に持つ。
    2本のURLが同じ場所を指しているあいだは、チャット用と同じものを使い回す。"""
    global _models_client
    if _models_client is None:
        base = _derived_base(llm_models_url(), "/models")
        if base == _derived_base(llm_chat_url(), "/chat/completions"):
            _models_client = client()
        else:
            _models_client = OpenAI(base_url=base or None,
                                    api_key=llm_api_key() or "not-set")
    return _models_client


def reset_llm_client() -> None:
    """画面から接続先URLやAPIキーが変更されたときに呼ぶ。次回の呼び出しで作り直す。"""
    global _client, _models_client
    _client = None
    _models_client = None


# --- モデルごとの作法の違いを吸収する ---------------------------------------------
#
# 同じ OpenAI互換API でも、モデルによって受け付ける引数が違う。実測した例:
#   gpt-5.6-sol : ツールを使うなら reasoning_effort='none' が必須。temperature も既定値のみ
#   gpt-4o 系   : reasoning_effort を送ると「Unrecognized request argument」で拒否
# モデル名の一覧を持って場合分けすると、ゲートウェイや新モデルのたびに保守が要る。
# そこで「1回投げて、断られた理由を読んで直して、覚える」方式にする。
# 400 は推論前に弾かれるので、やり直しても費用はかからない。
#
# 覚えた内容はプロセスの寿命だけ持つ。再起動後の最初の1回だけ余計に往復する。

#: モデル名 -> 学習した調整 {"set": {引数: 値}, "drop": {引数, ...}}
_QUIRKS: dict[str, dict] = {}
#: 1回の呼び出しで引数を直しにいく上限。無限に投げ続けないための歯止め。
_MAX_FIX = 4


def _apply_quirks(kwargs: dict) -> dict:
    q = _QUIRKS.get(str(kwargs.get("model") or ""))
    if not q:
        return kwargs
    out = {k: v for k, v in kwargs.items() if k not in q.get("drop", set())}
    out.update(q.get("set", {}))
    return out


def _learn(model: str, *, set_: dict | None = None, drop: str | None = None) -> None:
    q = _QUIRKS.setdefault(model, {"set": {}, "drop": set()})
    if set_:
        q["set"].update(set_)
    if drop:
        q["drop"].add(drop)
        q["set"].pop(drop, None)
    print(f"[llm] {model} の呼び出し方を調整しました: "
          f"set={q['set']} drop={sorted(q['drop'])}")


def _fix_for(message: str, kwargs: dict) -> tuple | None:
    """エラー文から、次に試す直し方を決める。戻り値: (set_, drop) か None。"""
    low = message.lower()

    # ツールと推論モードを同時に使えない → 推論を切れば chat/completions で通る
    if "reasoning_effort" in low and "function tools" in low:
        return ({"reasoning_effort": "none"}, None)
    # 値が受け付けられない（'minimal' など）→ 'none' に寄せる
    if "reasoning_effort" in low and "does not support" in low:
        return ({"reasoning_effort": "none"}, None)
    # そもそもこの引数を知らないモデル → 落とす
    if "reasoning_effort" in low and "unrecognized" in low:
        return (None, "reasoning_effort")
    # temperature / top_p が既定値しか許されないモデル → 落として既定に任せる
    for name in ("temperature", "top_p"):
        if f"'{name}'" in low and ("does not support" in low
                                   or "only the default" in low):
            return (None, name)
    # 新しいモデルは max_tokens ではなく max_completion_tokens を使う
    if "max_tokens" in low and "max_completion_tokens" in low:
        v = kwargs.get("max_tokens")
        if v is not None:
            return ({"max_completion_tokens": v}, "max_tokens")
    return None


# --- レート制限（429）------------------------------------------------------------
#
# 短時間に問い合わせが集中すると 429 が返る。SDKも2回までは自動で投げ直すが、
# それを使い切ってもなお混んでいることがある。そこで諦めずに、
# サーバが言ってきた待ち時間だけ待って投げ直す。
#
# 待ち時間の手がかりは2つある。どちらも無ければ 2秒・4秒・8秒 と延ばす。
#   retry-after ヘッダ          … 秒数
#   エラー文の "try again in X" … "300ms" / "1.5s" / "20s" のような書き方

_RETRY_IN = re.compile(r"try again in\s+([\d.]+)\s*(ms|s|m)\b", re.IGNORECASE)


def _is_rate_limit(e: Exception) -> bool:
    return getattr(e, "status_code", None) == 429 or "rate_limit" in str(e).lower()


def _rate_limit_wait(e: Exception, attempt: int) -> float:
    """次に待つ秒数を決める。attempt は今回が何回目の待ちか（1から）。"""
    # 1) サーバが retry-after で指示してきた場合
    try:
        raw = e.response.headers.get("retry-after")
        if raw:
            return min(float(raw), config.LLM_RATE_LIMIT_MAX_WAIT)
    except Exception:
        pass
    # 2) エラー文に「◯◯後に再試行」と書いてある場合
    m = _RETRY_IN.search(str(e))
    if m:
        value = float(m.group(1))
        unit = m.group(2).lower()
        sec = value / 1000 if unit == "ms" else (value * 60 if unit == "m" else value)
        # 指示どおりだと早すぎて再び弾かれることがあるので、少しだけ余裕を持たせる
        return min(sec + 0.5, config.LLM_RATE_LIMIT_MAX_WAIT)
    # 3) 手がかりが無ければ 2秒・4秒・8秒…
    return min(2.0 ** attempt, config.LLM_RATE_LIMIT_MAX_WAIT)


class RateLimited(Exception):
    """待って投げ直しても解消しなかったレート制限。

    生のJSONをそのまま画面に出すと何をすればよいか分からないので、
    人が読んで動ける文面にして投げ直す。
    """


def _create(**kwargs):
    """chat.completions.create の呼び出し口。

    モデルが受け付けない引数を、エラーの内容を見て直しながら投げ直す。
    直し方が分からないエラーはそのまま投げる（画面にそのまま出す）。
    """
    model = str(kwargs.get("model") or "")
    attempt = _apply_quirks(kwargs)
    fixes = waits = 0
    waited_total = 0.0
    while fixes < _MAX_FIX:
        try:
            return client().chat.completions.create(**attempt)
        except Exception as e:
            # レート制限は「引数の直し方」の問題ではないので、先に片付ける。
            # 待てば通ることがほとんどなので、その場で終わらせない。
            if _is_rate_limit(e):
                if waits >= config.LLM_RATE_LIMIT_RETRIES:
                    raise RateLimited(
                        f"混み合っています（レート制限）。{waits}回・"
                        f"合計{waited_total:.1f}秒待ちましたが解消しませんでした。"
                        "少し時間をおいてから「続けて」と送ってください。"
                        "ここまでに取得したデータは残っています。") from e
                waits += 1
                sec = _rate_limit_wait(e, waits)
                waited_total += sec
                print(f"[llm] レート制限。{sec:.1f}秒待って投げ直します"
                      f"（{waits}/{config.LLM_RATE_LIMIT_RETRIES}回目）")
                time.sleep(sec)
                continue
            fixes += 1
            fix = _fix_for(str(e), attempt)
            if fix is None:
                raise
            set_, drop = fix
            if set_ and all(attempt.get(k) == v for k, v in set_.items()):
                raise                      # 同じ直しを繰り返している
            if drop and drop not in attempt:
                raise
            _learn(model, set_=set_, drop=drop)
            attempt = _apply_quirks(kwargs)
    return client().chat.completions.create(**attempt)


# --- 表ルーター（質問に関係するテーブルだけを選ぶ） --------------------------------

_ROUTE_TABLE_SYSTEM = """あなたはデータ分析アプリの振り分け係です。
質問に答えるために必要なテーブルを、下の一覧から選んでください。

出力はJSONの配列だけ（説明文は書かない）:
  ["alias.table", "alias.other_table"]

守ること:
- 説明・行数・業務用語を手がかりに選ぶ。集計対象だけでなく、名前や区分を引くための
  マスタ（結合相手）も一緒に選ぶ。
- 迷ったら含める。外しすぎて答えられないより、多めの方がよい。
- 同じ領域が拠点・区分ごとに複数の表に分かれていることがある。「全社」「会社全体」
  「合計」のように範囲を限定しない質問なら、その領域の表をすべて選ぶ。
- 全部要る・判断できないときは ["*"] と書く。
- 名前は一覧に出てくる「alias.table」の形で書く。"""


def route_tables(question: str, scope: list[dict],
                 history: list[str] | None = None) -> dict | None:
    """質問に関係するテーブルを選ぶ。判断できなければ None（=絞らない）。

    カタログ全体がモデルの文脈に収まらないときの絞り込み。
    世に言うスキーマ・リンキングで、渡す情報の濃度を保つための仕組み。
    失敗したら None を返して従来どおり（要約モードへのフォールバック）に任せる。
    戻り値は {DBファイル名: [テーブル名, ...]}。
    """
    # 対応表: 「alias.table」「table単独（一意なら）」のどちらでも引けるようにする。
    # ルーターLLMは形式を崩しがちで、厳格一致だと正しい選択まで捨ててしまう。
    known: dict = {}
    ambiguous: set = set()
    total = 0
    for s in scope:
        prof = profile_db(s["path"])
        for t in prof["tables"].keys():
            total += 1
            for key in (f"{s['alias']}.{t}", t):
                k = key.strip().lower()
                if known.get(k, (s["name"], t)) != (s["name"], t):
                    ambiguous.add(k)
                known[k] = (s["name"], t)
    for k in ambiguous:
        known.pop(k, None)
    for s in scope:                    # alias.table だけは必ず引けるようにする
        prof = profile_db(s["path"])
        for t in prof["tables"].keys():
            known[f"{s['alias']}.{t}".lower()] = (s["name"], t)

    if total <= 8:
        return None                    # この数なら絞る意味がない（全部渡した方が確実）

    cards = "\n".join(
        db_text_cached(s["alias"], s["path"], s.get("tables"), full=False)
        for s in scope)
    ask = []
    for h in (history or [])[-3:]:
        ask.append(f"（直前の質問: {h}）")
    ask.append(f"今回の質問: {question}")
    try:
        resp = _create(
            model=config.OPENAI_MODEL,
            messages=[{"role": "system", "content": _ROUTE_TABLE_SYSTEM},
                      {"role": "user", "content": f"{cards}\n\n{chr(10).join(ask)}"}],
            temperature=0, max_tokens=300,
        )
        m = re.search(r"\[.*?\]", resp.choices[0].message.content or "", re.DOTALL)
        picked = json.loads(m.group(0)) if m else []
    except Exception as e:
        print(f"[router] 表の振り分けに失敗したため絞りません: {e}")
        return None
    if "*" in picked:
        return None
    out: dict = {}
    for p in picked:
        hit = known.get(str(p).strip().lower())
        if hit:
            out.setdefault(hit[0], set()).add(hit[1])
    if not out:
        return None
    return {name: sorted(ts) for name, ts in out.items()}


def expand_tables_by_relations(picked: dict, scope: list[dict]) -> dict:
    """選ばれた表に、カタログの関係線でつながる相手を1ホップぶん足す。

    ルーターLLMが結合相手（マスタ等）を選び忘れても、宣言済みの関係から
    機械的に補える。「必要な表の取りこぼしがSQL精度の上限を決める」ため、
    ここはAIの判断に任せず決め打ちで行う。
    """
    by_alias = {s["alias"]: s for s in scope}
    by_name = {s["name"]: s for s in scope}
    # 判定は「元の選択」のスナップショットに対して行う。育っていく集合に対して
    # 判定すると、追加された表が次の判定を呼んで連鎖し（1ホップのつもりが2ホップ）、
    # 設備マスタのようなハブを経由して無関係な表まで雪だるま式に増える。
    base = {k: set(v) for k, v in picked.items()}
    out = {k: set(v) for k, v in picked.items()}

    def add(db_name: str, table: str):
        s = by_name.get(db_name)
        if s is None:
            return
        if table in profile_db(s["path"])["tables"]:
            out.setdefault(db_name, set()).add(table)

    for s in scope:
        meta = load_meta(s["path"])
        for rel in (meta.get("relationships") or []):
            try:
                fa, ft, _fc = parse_endpoint_cols(rel.get("from", ""), s["alias"])
                ta, tt, _tc = parse_endpoint_cols(rel.get("to", ""), s["alias"])
            except Exception:
                continue
            f_scope = by_alias.get(fa)
            t_scope = by_alias.get(ta)
            # 方向は「子（from）→ 参照先（to）」だけ。選んだ表が参照している
            # マスタを足すのが目的で、逆（選んだマスタを参照している表を全部
            # 足す）をやると、ハブ表を選んだ瞬間にDBの大半が付いてくる。
            if (f_scope is not None and t_scope is not None
                    and ft in base.get(f_scope["name"], set())):
                add(t_scope["name"], tt)
    return {name: sorted(ts) for name, ts in out.items()}


# --- system prompt -----------------------------------------------------------

def build_system_prompt(scope: list[dict], admin: bool = False,
                        model: str | None = None) -> str:
    """選択スコープのデータカタログを埋め込んだ system prompt を組み立てる。

    admin は「管理者だけに渡すツール」を一覧に載せるかどうか。
    渡していないツールを説明に書くと、AIが呼ぼうとして失敗するだけになる。
    model を渡すと、カタログをインラインするかの判定を「そのモデルが読める量」で行う
    （渡さなければ管理者設定/envの上限）。
    """
    inline_cap = None
    if model:
        import models as models_mod        # 循環importを避ける
        inline_cap = models_mod.inline_limit_for(model)
    aliases = [s["alias"] for s in scope]

    # 設定どおりに更新できていないテーブルがあれば、AIに教えておく。
    # そのテーブルを使った回答に「データが古い可能性」を添えさせるため。
    import jobs as jobs_mod
    stale_lines = []
    in_scope = {s["name"] for s in scope}
    for (db_file, table), ps in jobs_mod.problems_by_table().items():
        if db_file in in_scope:
            alias = next((s["alias"] for s in scope if s["name"] == db_file), db_file)
            since = (ps[0].get("since") or "")[:16].replace("T", " ")
            stale_lines.append(f"- {alias}.{table}: {ps[0]['message']}"
                               + (f"（{since} 以降）" if since else ""))
    stale_note = ""
    if stale_lines:
        # 「古い」だけではない。取り込めたが数値列が文字に落ちた場合（degraded）は
        # データは新しく、集計の方が信用できない。どちらかを断定せず両方を伝える。
        stale_note = ("\n# 状態に問題があるデータ（重要）\n"
                      "次のテーブルは定期取り込みが設定どおりに動いていない。"
                      "中身が古いか、値の型が想定と違う（数値の列が文字で入っている）可能性がある。\n"
                      + "\n".join(stale_lines) + "\n"
                      "これらのテーブルを使って答えるときは、回答の冒頭に"
                      "「※ このデータは○月○日以降、正しく更新できていない可能性があります（理由）」"
                      "と必ず一言添える。数値の列が文字になっている場合は、合計・平均がずれうることも書く。"
                      "使わない質問では触れなくてよい。\n")
    # 社内文書（ナレッジベース）。1件も無ければ、この節ごと出さない
    # （使えないものを説明すると、AIは呼べないツールを呼ぼうとするだけになる）。
    # 見せるのは、その利用者がサイドバーで選んでいるものだけ。外したものまで
    # 並べると「選んでいないものの中身」を語り出す。
    kbs = rag_targets()
    kb_note = ""
    if kbs:
        kb_lines = "\n".join(
            f"- {e['name']}" + (f": {e['description']}" if e.get("description") else "")
            for e in kbs)
        kb_note = f"""# 社内文書（ナレッジベース）
手順・原因・対処・規則・過去の事例など「文書に書かれていること」は
search_knowledge_base で調べる。数値の集計は表（run_sql_query）、
書かれていることはナレッジベース、と使い分ける。両方使ってよい
（例: 表の集計で悪化した装置を特定し、その原因を文書から探す）。

検索できるナレッジベース:
{kb_lines}

守ること:
- 取得した本文(context)に書かれていることだけを根拠にする。書かれていないことは推測しない。
- 根拠にした箇所には [出典1] のように出典番号を必ず添える。番号は検索結果のものをそのまま使う。
- 文書を使って答えたときは、回答の最後に「根拠」としてその出典を並べる
  （出典番号・ナレッジベース名・ファイル名）。人が原本に当たれるようにするため。
- 見つからなければ、語を変えて1〜2回まで試す。それでも無ければ「文書には見当たらない」と述べる。
  見つからなかったことを、自分の知識で埋めてはいけない。
- 複数のナレッジベースで内容が食い違うときは、どちらの出典かを明示して両論を併記する。

"""

    if len(aliases) > 1:
        naming = (f"複数のDBが対象です（{', '.join(aliases)}）。"
                  "テーブル名は必ず『エイリアス.テーブル名』で修飾すること"
                  f"（例: {aliases[0]}.xxx）。DBをまたぐ JOIN も可能。")
    elif aliases:
        # DBは1つだけ。その名前をAIに教えると、回答やSQLにまで出てきてしまう
        # （画面ではDBという概念を見せていないので、利用者には意味が通じない）。
        naming = "テーブル名はそのまま書く（前に何かを付けて修飾しない）。"
    elif kbs:
        # 数値のデータは無いがナレッジベースはある構成。SQLの話に持っていかず文書で答える。
        naming = ("いま対象にできるデータがありません。数値の集計はできないので、"
                  "ナレッジベースの検索で答えられる範囲で答えること。"
                  "数値が要る質問には、対象のデータが選ばれていないことを伝える。")
    else:
        naming = "対象にできるDBがありません。「データ取り込み」でDBを作るよう案内すること。"

    # 実際に渡すツール一覧（無効化・説明の上書き・ユーザー定義ツールを反映）
    lines = []
    for t in tools.build_tools(scope, admin=admin):
        fn = t["function"]
        args = ", ".join((fn.get("parameters") or {}).get("properties", {}).keys())
        lines.append(f"- {fn['name']}({args}) : {fn['description']}")
    custom = custom_tools.collect_everywhere(scope)
    if custom:
        lines.append("※ 上記のうち次はこの環境専用に用意されたツールです。"
                     "目的が合致するときは自分でSQLを書かずにこちらを優先して使ってください: "
                     + ", ".join(t["name"] for t in custom))
    tool_list = "\n".join(lines)

    intro = ("あなたは社内の業務アシスタントです。\n"
             "読み取り専用(SELECTのみ)でデータベースにアクセスでき、"
             "社内文書のナレッジベースも検索できます。"
             if kbs else
             "あなたはSQLiteデータベースの分析アシスタントです。\n"
             "読み取り専用(SELECTのみ)でデータベースにアクセスできます。")

    return f"""{intro}

# 振る舞い
- ユーザーの質問に答えるため、必要に応じてツールを呼び出し、必ず実データに基づいて回答する。
- 推測で数値を答えてはいけない。データが必要なら run_sql_query を使う。
- 列の意味や値の実体が不確かなら、SQLを書く前に describe_table で確認する。
- 範囲や対象の取り方が複数あって答えが変わるとき（似た表が複数ある、期間が示されていない等）は、
  勝手にどれかへ決めて答えず、一言確認するか、選び方を明記して答える。
- ツールを使うか・どのSQLを書くかはあなたが判断する(挨拶や一般的な雑談ならツール不要)。
- 回答は日本語。まず結論、次に根拠(表やグラフの要点)を簡潔に述べる。
- 後述の「業務用語」に載っている言葉が質問に出たら、必ずその定義に従う。自分の常識で解釈し直さない。
- カタログを育てる提案は、ユーザーが頼まなくても、あなたから聞く。
  ただし押しつけない: 聞くのは1回の回答につき最大1つ、同じものを何度も聞かない。
  - 用語: ユーザーが業務用語の意味を教えてくれたら（「有効な受注とはキャンセル以外のこと」等）、
    回答の最後に「この定義を用語集に登録しますか？」と一言添える。
  - 例文: run_sql_query の結果に example_registered: false が入っていて、
    回答が問題なく出せたら、回答の最後に
    「この質問と答え方を例文として登録しますか？（似た質問に強くなります）」と一言添える。
    example_registered: true なら聞かない（すでに登録済み）。
    エラーで言い直した回答や、雑談・確認だけのやり取りでは聞かない。
  - 前向きな返事（「はい」「お願い」「登録して」等）が来たら、用語は propose_glossary_term、
    例文は propose_example でカードを出す。登録するかはユーザーがカードのボタンで決める。
    例文の sql は実行して正しかったものをそのまま使う。
  - 「SQL式:」が書かれている用語は、その式をそのまま WHERE や SELECT に埋め込む。
  - SQL式が無く説明文だけの用語は、その説明と列情報から自分でSQLを組み立てる。
    どの列をどう使ったかを回答の中で一言添える（人が誤りに気づけるようにするため）。
- ツールの結果に verification_warnings（検算の不一致）が入っていることがある。
  これは「同じ数字を別の経路で数えたら食い違った」という自動検算の結果で、
  あなたのSQLの誤りとは限らない。入っていたら、回答の末尾で
  「どの数字を使ったか」と「別の経路では値が異なること」を必ず1〜2文で注記する。
  差異の原因は、検算結果に示された内訳の範囲でだけ述べ、推測で断定しない。

# 可視化の方針（チャットにグラフを描く）
- ユーザーが「グラフ」「可視化」「チャート」「推移」「トレンド」「割合」「内訳」「分布」等を求めたら、必ず plot_* のツールを呼んでグラフを描く。
- 明示が無くても、結果が次に該当するなら積極的にグラフにする。目的でツールを選ぶ：
  - 項目どうしの比較・順位 → plot_comparison（bar / hbar / stacked_bar / pareto / radar など）
  - 時系列・推移 → plot_trend（line / step / area / calendar など）
  - 構成比・内訳・増減の要因 → plot_composition（pie / donut / treemap / funnel / waterfall / sankey）
  - ばらつき・分布 → plot_distribution（histogram / box / violin など）
  - 2つ以上の項目の関係 → plot_relationship（scatter / bubble / heatmap など）
  - 1つの数字を大きく・目標との対比 → plot_kpi（indicator / gauge / bullet）
- sql は集計済み(GROUP BY)にし、x/y にする列を AS で明示する。色分けは color に列名を渡す。
- 棒グラフの積み方は種別で指定する：「積み上げ」なら chart_type="stacked_bar"、「横並び」「比較」なら "bar"。
- 「2軸」「二軸」「棒と折れ線」「件数と比率を一緒に」など、単位の異なる2指標を重ねたい時は plot_dual_axis を使う。
  bar_y(左軸=棒, 件数など) と line_y(右軸=折れ線, 比率など) に列名を渡す。
- 必要なら run_sql_query で数値を確認しつつ、可視化は plot_* で別途描く（両方呼んでよい）。
- 前のツールと同じデータを使うときは、sql を書き直さず、そのツールが返した result_id を渡す。
  同じSQLを2回実行すると、その間にデータが変わって表とグラフの数字がずれることがある。

# SQLで書けないこと（必ず専用ツールを使う）
このSQLiteには STDDEV / VARIANCE / MEDIAN / CORR / PERCENTILE / SQRT / POWER が無く、
PIVOT構文も無い。次はSQLで計算しようとせず、必ずツールを使うこと。
- 「クロス集計」「行に○○・列に△△」「マトリクスで」 → pivot_table
  （sql では集計せず、必要な列を返すだけにする。集計はツールが行う）
- 「相関」「中央値」「ばらつき」「標準偏差」「四分位」「分布の要約」 → analyze_stats
- 「外れ値」「異常値」「突出しているもの」 → analyze_stats の method="outliers"
  （sql は集計せず明細を返す。1行1件の状態にしてから渡すこと）

# ファイル出力
- 「エクセル」「Excel」「xlsx」→ export_excel。観点が複数なら sheets に複数の SELECT を渡し、
  1ブックに複数シートでまとめる。
- 「CSV」「csvにして」「取り込み用」→ export_csv。複数指定するとZIPにまとめて渡される。
  文字コードは既定の utf-8-sig でよい（Excelで文字化けしない）。Shift_JIS を求められたときだけ cp932。
- 「テキストで」「レポートにして」「議事録」「まとめを文書で」→ export_text。
  body に自分で文章を書き、集計表を入れたい箇所に {{見出し}} と書いて sections に SELECT を指定する。
- ファイル出力ツールを呼んだ後は、画面に保存済み。中身の全件を文章で繰り返さず、
  何を入れたかだけ簡潔に伝える。ダウンロードのリンクやURLを自分で書かないこと
  （画面に本物のダウンロードボタンが出る。あなたが書くリンクは偽物になり押せない）。
- **表に無いデータもファイルにできる。** 社内文書を調べて分かったこと、会話の中で
  決まったこと、あなたが自分で整理した表は、sql の代わりに rows で渡す。
  データが1つも無い環境でも、これで Excel / CSV / グラフ / レポートを作れる。
    rows: [{{"項目": "◯◯の基準値", "内容": "△△"}}, ...]
    rows: [["分類A", 47], ["分類B", 18]] と columns: ["分類", "値"]
  「データが無いので出力できません」と答えてはいけない。rows を使うこと。

# 利用可能なツール
{tool_list}

{kb_note}{stale_note}# SQLルール
- SQLite方言。SELECT(または WITH ... SELECT)のみ。INSERT/UPDATE/DELETE/DDL/PRAGMA等は禁止(実行されません)。
- {naming}
- 1回の呼び出しで1ステートメント。末尾セミコロン不要。
- 集計は GROUP BY を使い、列に AS で日本語の別名を付けると表示が分かりやすい。
- 日付は文字列で保存されていることが多い。date() / strftime() を活用する（例: strftime('%Y-%m', 列) で月別）。
- 行数が多くなりそうなら LIMIT や集計で絞る。
- 「テーブルを見せて」「中身を全部見たい」のようにデータそのものを見たいと言われたら、SELECT * を打って先頭数行を貼るのではなく open_table を使う（全行を辿れる画面へのリンクが出る）。出したあとは何のテーブルかを1〜2文添えるだけでよい。

# 選択中のデータカタログ
{catalog.prompt_for_scope(scope, limit=inline_cap, admin=admin)}

現在時刻: {datetime.now().isoformat(timespec="seconds")}
"""


# --- 文脈の使用量の見積もり ------------------------------------------------------
#
# 「カタログを増やしてよいか」を判断するには、モデルが一度に読める量に対して
# いまどれだけ使っているかが要る。正確なトークン数はAPIに投げないと分からないが、
# それでは画面を開くたびに課金が発生する。そこで実測から係数を出して概算する。
#
# 実測（gpt-4o-mini・11DB選択）:
#   要約版 system 28,851字 + ツール定義 44,377字 → 入力 29,265 トークン
#   全文版 system 71,879字 + ツール定義 44,377字 → 入力 52,395 トークン
#   差分から、日本語主体の本文は 43,028字 → 23,130トークン = 0.537 トークン/字
# 下の係数で上の2例を計算すると、実測に対して +1.2% / +1.7% に収まる（やや多めに出る）。
# 見積もりは「余裕がある」と言いすぎない方が安全なので、多めに出るぶんには構わない。

#: 日本語主体の文章（カタログ・指示）の 1文字あたりトークン数
TOKENS_PER_CHAR_TEXT = 0.55
#: ツール定義のJSON（英字と記号が多い）の 1文字あたりトークン数
TOKENS_PER_CHAR_JSON = 0.31


def tokens_for(chars: int, kind: str = "text") -> int:
    """文字数からトークン数の概算を出す。"""
    ratio = TOKENS_PER_CHAR_JSON if kind == "json" else TOKENS_PER_CHAR_TEXT
    return int(max(0, chars) * ratio)


def budget(scope: list[dict], model: str | None = None, admin: bool = False) -> dict:
    """このスコープ・このモデルで、文脈をどれだけ使うかの概算。

    「いま」と「上限までカタログが育ったとき」の両方を返す。
    上限を決める画面で、変えた結果どうなるかを見せるため。
    """
    import models as models_mod

    model = model or config.OPENAI_MODEL
    context, known = models_mod.context_window(model)
    limit = models_mod.inline_limit_for(model)      # そのモデルの自動上限

    # 推定せず、実際に組み立てたものを測る（カタログはキャッシュ済みなので速い）
    system = build_system_prompt(scope, admin=admin, model=model)
    used_catalog = len(catalog.prompt_for_scope(scope))
    catalog_chars = catalog.inline_length(scope)       # 全文にした場合の長さ
    tool_chars = len(json.dumps(tools.build_tools(scope, admin=admin), ensure_ascii=False))

    # カタログ以外（SQLルール・ツールの使い分け・ツール名の一覧など）
    fixed_chars = max(0, len(system) - used_catalog)

    tool_tokens = tokens_for(tool_chars, "json")
    now = tokens_for(len(system)) + tool_tokens
    at_limit = tokens_for(fixed_chars + limit) + tool_tokens

    def pct(n: int) -> float:
        return round(n / context * 100, 1) if context else 0.0

    return {
        "model": model, "context": context, "context_known": known,
        "limit_chars": limit,
        "catalog_chars": catalog_chars,
        "catalog_inlined": catalog_chars <= limit,
        "tool_tokens": tool_tokens,
        # カタログ以外の固定ぶん。画面で上限を動かしたときに、
        # サーバと同じ式で計算し直せるように渡す。
        "base_tokens": tool_tokens + tokens_for(fixed_chars),
        "tokens_per_char": TOKENS_PER_CHAR_TEXT,
        "now_tokens": now, "now_pct": pct(now), "headroom_pct": round(100 - pct(now), 1),
        "at_limit_tokens": at_limit, "at_limit_pct": pct(at_limit),
        # 上限をここまで上げても文脈の半分に収まる、という目安
        "suggest_max_chars": max(0, int((context * 0.5 - tool_tokens
                                         - tokens_for(fixed_chars))
                                        / TOKENS_PER_CHAR_TEXT)),
    }


# --- 画像つきのメッセージ --------------------------------------------------------

# 受け付ける画像。ここに無い形式は送らない（APIが解釈できないため）
IMAGE_MIMES = {"image/png", "image/jpeg", "image/gif", "image/webp"}


def user_message(text: str, images: list[dict] | None = None) -> dict:
    """ユーザー発言を作る。画像があればマルチモーダル形式にする。

    images: [{"mime": "image/png", "b64": "..."}] （b64はデータ本体のみ）
    画像が無いときは、これまで通り content が文字列のメッセージを返す
    （画像非対応のモデルに配列を渡すと弾かれることがあるため）。
    """
    if not images:
        return {"role": "user", "content": text}
    parts: list[dict] = []
    if text:
        parts.append({"type": "text", "text": text})
    for img in images:
        mime = img.get("mime") or "image/png"
        if mime not in IMAGE_MIMES:
            continue
        parts.append({"type": "image_url",
                      "image_url": {"url": f"data:{mime};base64,{img['b64']}",
                                    "detail": img.get("detail") or "auto"}})
    return {"role": "user", "content": parts or text}


# --- チャット補完 --------------------------------------------------------------

def _msg_chars(m) -> int:
    """1メッセージがプロンプトで占める、おおよその文字数。"""
    if not isinstance(m, dict):
        return len(str(m))
    n = 0
    c = m.get("content")
    if isinstance(c, str):
        n += len(c)
    elif isinstance(c, list):                 # 画像つきのマルチモーダル形式
        for p in c:
            if not isinstance(p, dict):
                continue
            if p.get("type") == "text":
                n += len(p.get("text") or "")
            else:                             # data: URI がそのまま長さになる
                n += len(str((p.get("image_url") or {}).get("url") or ""))
    if m.get("tool_calls"):
        n += len(json.dumps(m["tool_calls"], ensure_ascii=False, default=str))
    return n


def _turn_blocks(messages: list) -> tuple:
    """履歴を「質問のかたまり」に割る。

    かたまり = role:"user" から次の role:"user" の手前まで。
    ツール呼び出しとその結果が離れると API に弾かれるので、
    落とすときは必ずこの単位で落とす。
    戻り値は (先頭に置くもの, かたまりの並び)。
    """
    head, blocks, cur = [], [], None
    for m in messages:
        role = m.get("role") if isinstance(m, dict) else None
        if role == "system":
            head.append(m)
        elif role == "user":
            if cur is not None:
                blocks.append(cur)
            cur = [m]
        elif cur is None:
            head.append(m)
        else:
            cur.append(m)
    if cur is not None:
        blocks.append(cur)
    return head, blocks


def drop_old_images(messages: list, keep_turns: int | None = None) -> int:
    """古い質問に付いた画像を、短い注記に差し替える。落とした枚数を返す。

    差し替えるだけで *メッセージは消さない*。消すと巻き戻しの
    「何回目の発言か」がずれる（画面側の render_log は触らないため）。
    """
    keep = config.CHAT_KEEP_IMAGE_TURNS if keep_turns is None else keep_turns
    _head, blocks = _turn_blocks(messages)
    if len(blocks) <= keep:
        return 0
    dropped = 0
    for b in blocks[:len(blocks) - keep]:
        for m in b:
            if not isinstance(m, dict) or m.get("role") != "user":
                continue
            parts = m.get("content")
            if not isinstance(parts, list):
                continue
            imgs = [p for p in parts if isinstance(p, dict) and p.get("type") == "image_url"]
            if not imgs:
                continue
            text = "".join(p.get("text") or "" for p in parts
                           if isinstance(p, dict) and p.get("type") == "text").strip()
            m["content"] = (text + f"\n（この質問には画像が{len(imgs)}枚ありましたが、"
                            "会話が長くなったため履歴からは外しました。"
                            "画像の中身を聞かれたら、貼り直してもらうよう伝えること）").strip()
            dropped += len(imgs)
    return dropped


def history_for_llm(messages: list, model: str | None = None,
                    tool_defs: list | None = None) -> list:
    """LLMへ渡す用の履歴。長すぎるときは古い質問から落とした控えを返す。

    保存側（chat["messages"]）は触らない。保存側から消すと、画面の
    render_log とユーザー発言の番号がずれ、巻き戻しが別の位置を切る。

    使える量は「モデルの文脈 −（システムプロンプト＋ツール定義＋回答の枠）」。
    落としたときはシステムプロンプトの末尾に一言添えて、AIが
    「前の話は残っていない」と分かるようにする。
    """
    import models as models_mod

    head, blocks = _turn_blocks(messages)
    if len(blocks) <= config.CHAT_KEEP_TURNS_MIN:
        return messages

    context, _known = models_mod.context_window(model or config.OPENAI_MODEL)
    used = sum(tokens_for(_msg_chars(m)) for m in head)
    if tool_defs:
        used += tokens_for(len(json.dumps(tool_defs, ensure_ascii=False, default=str)), "json")
    left = context - used - config.CHAT_ANSWER_RESERVE_TOKENS
    if left <= 0:
        left = 0
    room = int(left / TOKENS_PER_CHAR_TEXT)

    sizes = [sum(_msg_chars(m) for m in b) for b in blocks]
    total, cut = sum(sizes), 0
    while total > room and len(blocks) - cut > config.CHAT_KEEP_TURNS_MIN:
        total -= sizes[cut]
        cut += 1
    if not cut:
        return messages

    out = [dict(m) if isinstance(m, dict) else m for m in head]
    if out and isinstance(out[0], dict) and out[0].get("role") == "system":
        out[0]["content"] = (str(out[0].get("content") or "")
                             + f"\n\n※ この会話は長くなったため、古い方の"
                               f"{cut}回ぶんのやり取りは履歴から外れています。"
                               "そこで話した内容は覚えていないので、"
                               "必要なら聞き直してください。")
    for b in blocks[cut:]:
        out.extend(b)
    return out


def chat(messages: list[dict], tool_defs: list[dict] | None = None,
         model: str | None = None):
    """messages を渡して1回の補完を取得。tools 付き。message オブジェクトを返す。

    tool_defs の意味は3通り:
      None … 省略。組み込みツールのみ渡す
      []   … ツールを1つも渡さない（文章で答えるしかない状態にする）
      list … そのツールを渡す。通常はチャット側で tools.build_tools(entries)

    空リストのときに tools を空配列で送ると弾くAPIがあるので、
    キーごと外して送る（None との違いはここ）。
    model は画面で選ばれたモデル。省略時は env の既定。
    """
    kwargs = dict(
        model=model or config.OPENAI_MODEL,
        messages=messages,
        temperature=config.OPENAI_TEMPERATURE,
    )
    defs = tools.BUILTIN_TOOLS if tool_defs is None else tool_defs
    if defs:
        kwargs["tools"] = defs
        kwargs["tool_choice"] = "auto"
    if config.OPENAI_TOP_P is not None:
        kwargs["top_p"] = config.OPENAI_TOP_P
    if config.OPENAI_MAX_TOKENS is not None:
        kwargs["max_tokens"] = config.OPENAI_MAX_TOKENS
    resp = _create(**kwargs)
    msg = resp.choices[0].message
    # 出力が上限で切れたか（"length"）を呼び出し側が見られるようにしておく。
    # SDKのモデルは属性の後付けを許すが、万一拒まれても本体の動作は変えない。
    try:
        msg.finish_reason = resp.choices[0].finish_reason
    except (AttributeError, ValueError, TypeError):
        pass
    return msg


class StreamedMessage:
    """ストリーミングで組み立てた1回ぶんの応答。

    chat() が返す message オブジェクトと同じ形（content / tool_calls）に
    見えるようにしておく。呼び出し側はどちらでも同じ扱いができる。
    """

    class _Fn:
        def __init__(self, name="", arguments=""):
            self.name, self.arguments = name, arguments

    class _Call:
        def __init__(self, id="", name="", arguments=""):
            self.id, self.type = id, "function"
            self.function = StreamedMessage._Fn(name, arguments)

    def __init__(self):
        self.content = ""
        self.tool_calls = None
        self.finish_reason = None      # "length" なら出力が上限で切れている
        self._parts: dict[int, dict] = {}

    def _finish(self):
        if not self._parts:
            self.tool_calls = None
            return
        self.tool_calls = [
            StreamedMessage._Call(p["id"], p["name"], p["arguments"])
            for _, p in sorted(self._parts.items())
        ]


def chat_stream(messages: list[dict], tool_defs: list[dict] | None = None,
                model: str | None = None):
    """1回の補完をストリーミングで受け取る。

    文字が届くたびに ("text", 差分) を yield し、
    最後に ("done", StreamedMessage) を1回だけ yield する。
    ツール呼び出しは途中経過を出さない（引数のJSONは途中では読めないため）。
    """
    kwargs = dict(
        model=model or config.OPENAI_MODEL,
        messages=messages,
        tools=tool_defs if tool_defs is not None else tools.BUILTIN_TOOLS,
        tool_choice="auto",
        temperature=config.OPENAI_TEMPERATURE,
        stream=True,
    )
    if config.OPENAI_TOP_P is not None:
        kwargs["top_p"] = config.OPENAI_TOP_P
    if config.OPENAI_MAX_TOKENS is not None:
        kwargs["max_tokens"] = config.OPENAI_MAX_TOKENS

    out = StreamedMessage()
    for chunk in _create(**kwargs):
        if not chunk.choices:
            continue
        if chunk.choices[0].finish_reason:
            out.finish_reason = chunk.choices[0].finish_reason
        delta = chunk.choices[0].delta
        if getattr(delta, "content", None):
            out.content += delta.content
            yield ("text", delta.content)
        for tc in (getattr(delta, "tool_calls", None) or []):
            # 同じツール呼び出しが複数のチャンクに分かれて届くので、index で束ねる
            part = out._parts.setdefault(tc.index, {"id": "", "name": "", "arguments": ""})
            if tc.id:
                part["id"] = tc.id
            fn = getattr(tc, "function", None)
            if fn is not None:
                if getattr(fn, "name", None):
                    part["name"] += fn.name
                if getattr(fn, "arguments", None):
                    part["arguments"] += fn.arguments
    out._finish()
    yield ("done", out)


# --- AI下書き（データカタログ用） ----------------------------------------------

def _ask_json(system: str, user: str, what: str = "AIの応答") -> dict:
    """AIに聞いて、応答からJSONオブジェクトを取り出す。

    下書き系（表の説明・用語のSQL式・ビュー・ツール）が全部この形なので、
    「温度0で聞く → 最初の { から最後の } までを取り出す → dict か確かめる」
    をここ1本にまとめてある。``` で囲まれて返ってきても取り出せる。
    what はエラー文に出す呼び名（どの下書きで失敗したかが分かるように）。
    """
    resp = _create(
        model=config.OPENAI_MODEL,
        messages=[{"role": "system", "content": system},
                  {"role": "user", "content": user}],
        temperature=0,
    )
    content = resp.choices[0].message.content or ""
    m = re.search(r"\{.*\}", content, re.DOTALL)
    if not m:
        raise ValueError(f"{what}をJSONとして解析できませんでした: {content[:200]}")
    data = json.loads(m.group(0))
    if not isinstance(data, dict):
        raise ValueError(f"{what}が想定した形式ではありません。")
    return data


_DRAFT_SYSTEM = """あなたはデータカタログ作成の専門家です。
与えられたテーブルのプロファイル（列名・型・実値の分布・サンプル行）から、
テーブルと各列の業務的な説明文を日本語で推測し、JSONだけを出力してください。

出力形式（JSON以外の文字を含めないこと）:
{
  "description": "テーブルの説明。1行 = 何のレコードかを必ず含める。",
  "columns": {
    "列名": {
      "description": "列の説明",
      "values": {"コード値": "意味"}   // 値がコード(区分値)と思われる列のみ。それ以外は省略
    }
  }
}

注意:
- 確信が持てない場合は「〜と思われる」と書く。
- values は実値一覧にある値だけを対象にする。
- すべての列に説明を付ける。"""


#: まとまりのメモの点検。書き直させるのではなく、間違いを指摘させる。
#: この欄は「人しか知らない決めごと」を書く場所なので、AIに書かせると
#: 人が書いた決めごとと見分けが付かなくなる（この欄だけ由来の印も履歴も無い）。
#: 一方「書いてあることがデータと合っているか」はAIが確かめられる。
_MEMO_REVIEW_SYSTEM = """あなたはデータカタログの点検役です。
「まとまりのメモ」に書かれている文を1行ずつ確かめ、問題のある行だけを挙げてください。
文章を書き直してはいけません。指摘だけを返します。

メモは、複数の表にまたがる前提・決めごと（指標の定義、記録の範囲、社内規則など）を
書くための欄です。表1つで完結する説明や、列の意味は、別の欄からAIに渡るので不要です。

出力形式（JSON以外の文字を含めないこと）:
{
  "findings": [
    {"level": "err",
     "line": "問題のある行をそのまま引用",
     "problem": "何が問題か（1〜2文）",
     "suggestion": "どうするとよいか（1文）"}
  ]
}

level の使い分け:
- err  … 資料と食い違う。存在しない表・列・区分値を指している。数字が資料と合わない。
          同じ内容の行が2回書かれている（別のまとまりからの貼り付けなど）ときもこれ。
- warn … 出どころが1つも書かれていない決めごと。
          この欄は「人しか知らないこと」を書く場所なので、資料で裏が取れないこと自体は
          問題ではない。「就業規則第17条」「作業標準STD-102」「ナレッジベースのトラブル事例」
          のように、どこを見れば確かめられるかが書いてあれば指摘しないこと。
          対象は、しきい値・上限・判定基準・因果の主張のように、間違えると
          出てくる数字が変わるものだけ。単位・通貨・表記のように、決めれば済むこと
          （「金額はすべて円（税抜）」など）に出どころを求めてはいけない。
- info … 資料に既に載っていて、メモに書く必要がない（列の意味・主キー・結合のキー・
          値の範囲など、機械が毎回自動で書き出している事柄）。

守ること:
- 問題が無ければ findings を空の配列にする。無理に指摘を作らない。
  指摘は少ないほどよい。直す人は専門家ではないので、多いと全部読まれなくなる。
- メモの1行目は「このまとまりが何のデータか」を一言で書く決まりなので、指摘しない。
  同じ趣旨の要約が2行目以降にもう一度出てきたときだけ、重複として挙げる。
- line は必ずメモの原文をそのまま引用する。要約しない。
- 表名・列名は資料に実在するものだけを挙げる。
- メモに書かれた社内規則・指標の定義そのものの是非は論じない。
  「その定義でよいか」は人が決めることで、点検役の仕事ではない。"""


_GLOSSARY_SYSTEM = """あなたはSQLiteに詳しいデータカタログ作成の専門家です。
与えられたテーブル定義（列・型・実値の分布・サンプル行）をもとに、
業務用語の「自然言語の説明」をSQLの式に翻訳してください。

用語ごとに「式にできるか」を判断し、できる場合だけ式を書いてください。

出力形式（JSON以外の文字を含めないこと。用語をキーにしたオブジェクト）:
{
  "用語": {"ok": true, "sql": "SQL式", "explanation": "その式が何をしているかの日本語の説明"},
  "用語2": {"ok": false, "reason": "式にできない理由と、どう書き足せばよいか（日本語）"}
}

守ること:
- WHERE にそのまま入る条件式（例: status != '9' AND amount >= 1000000）か、
  SELECT にそのまま入る計算式（例: SUM(amount) * 1.0 / COUNT(*)）だけを書く。
- SELECT や FROM で始まる文全体は書かない。末尾にセミコロンを付けない。
- 列名は与えられたテーブルに実在するものだけを使う。値は実値一覧にあるものを使う。
- SQLiteに無い関数(STDDEV, MEDIAN, PERCENTILE_CONT, SQRT, POWER など)は使わない。

ok: false にする場合（推測で書かないこと。理由は具体的に）:
- 必要な列がテーブルに無いとき → どんな列が要るのに無いのかを書く。
- 説明があいまいで、しきい値や範囲が決まらないとき
  → 何がはっきりしないか（「最近」が何日か、「大口」がいくら以上か等）を尋ねる。
- SQLiteに無い関数が必要なとき → 何が必要で、代わりに何なら書けるかを書く。
- 1つの式では表せない（複数の表の結合が要る等）とき → その旨と代わりの手段を書く。

explanation（解説）の書き方:
- 日本語で1〜3行。どの列をどう見て、何を満たす行（または何を計算した値）かを書く。
- 列名はそのまま出してよいが、何を表す列かを添える。"""


def draft_glossary_sql(db_path, table_name: str | None, terms: list[dict]) -> dict:
    """業務用語の説明文からSQL式の下書きを作る。

    terms: [{"term": 用語, "description": 自然言語の説明}, ...]
    戻り値: {用語: {"ok": True, "sql": 式, "explanation": 解説}
             または {"ok": False, "reason": 書けない理由}}
    """
    if not terms:
        return {}
    profile = catalog.profile_db(Path(db_path))
    meta = catalog.load_meta(Path(db_path))
    if table_name:
        context = catalog.table_text("db", table_name, profile, meta, full=True)
    else:   # テーブルをまたぐ用語。DB全体を見せる
        context = catalog.db_text("db", Path(db_path), None, full=True)
    asked = "\n".join(f"- {t['term']}: {t.get('description') or ''}" for t in terms)

    data = _ask_json(_GLOSSARY_SYSTEM, f"{context}\n\n翻訳したい業務用語:\n{asked}", "用語のSQL式の下書き")
    wanted = {t["term"] for t in terms}
    out = {}
    for k, v in data.items():
        if k not in wanted:
            continue
        if isinstance(v, str):        # 旧形式（式だけ）の応答にも耐える
            if v.strip():
                out[k] = {"ok": True, "sql": v.strip(), "explanation": ""}
            continue
        if not isinstance(v, dict):
            continue
        sql = str(v.get("sql") or "").strip().rstrip(";")
        if v.get("ok") is False or not sql:
            out[k] = {"ok": False,
                      "reason": str(v.get("reason")
                                    or "この説明からはSQL式を決められませんでした。").strip()}
        else:
            out[k] = {"ok": True, "sql": sql,
                      "explanation": str(v.get("explanation") or "").strip()}
    # 応答に出てこなかった用語も「書けなかった」として理由を付ける
    for t in terms:
        out.setdefault(t["term"], {
            "ok": False,
            "reason": "この説明からはSQL式を決められませんでした。"
                      "対象の列や、しきい値（いくつ以上か・何日以内か）を書き足してみてください。"})
    return out


def draft_table_meta(db_path, table_name: str) -> dict:
    """テーブルのプロファイルからメタ情報の下書きを生成する。"""
    profile = catalog.profile_db(Path(db_path))
    meta = catalog.load_meta(Path(db_path))
    text = catalog.table_text("db", table_name, profile, meta, full=True)
    return _ask_json(_DRAFT_SYSTEM, f"ファイル名: {Path(db_path).name}\n\n{text}", "表の説明の下書き")


def review_group_memo(db_path, group: str, memo: str) -> list[dict]:
    """まとまりのメモを点検して、指摘の一覧を返す。メモは書き換えない。

    memo は画面のいまの文字（未保存でも点検できるようにするため）。
    資料には保存済みの版が混ざるので、どちらを見るかを本文で明示している。
    """
    path = Path(db_path)
    profile = catalog.profile_db(path)
    tables = [t for t in profile["tables"] if t.startswith(group + "__")]
    if not tables:
        raise ValueError(f"まとまり「{group}」の表がありません。")
    if not memo.strip():
        # 空欄を「点検済み・問題なし」と読ませない（未記入であることを見せる）
        return [{"level": "warn", "line": "",
                 "problem": "このまとまりのメモはまだ空です。",
                 "suggestion": "複数の表にまたがる前提や決めごと"
                               "（指標の定義、記録の範囲、社内規則）を書いてください。"}]
    # 表を絞って渡す。db_text は自動生成の注記（兄弟まとまりの合算指示など）も
    # 含むので、「メモに書く必要がないこと」をAIが判断できる
    context = catalog.db_text("db", path, tables, full=True)
    # まとまりの外の表は、名前と列だけ添える。メモは「複数の表にまたがる前提」を
    # 書く欄なので、他のまとまりを指す行が普通にある。絞った資料だけを見せると、
    # 実在する表を「存在しません」と誤って指摘する（実際に誤検知した）。
    others = "\n".join(
        f"- {t}: " + ", ".join(c["name"] for c in profile["tables"][t]["columns"])
        for t in sorted(profile["tables"]) if t not in tables)
    ask = (f"{context}\n\n"
           f"## このDBにある他の表（名前と列だけ）\n"
           f"メモが他のまとまりの表に触れていたら、実在するかどうかはこの一覧で"
           f"確かめること。ここに名前があるものを「存在しない」と書いてはいけない。\n"
           f"{others}\n\n"
           f"## 点検するメモ（まとまり「{group}」）\n"
           f"上の資料の中にも同じ見出しでメモが載っていることがありますが、"
           f"それは保存済みの古い版です。点検するのは下の文章です。\n\n{memo}")
    data = _ask_json(_MEMO_REVIEW_SYSTEM, ask, "メモの点検")
    out = []
    for f in (data.get("findings") or []):
        if not isinstance(f, dict):
            continue
        level = str(f.get("level") or "warn")
        out.append({"level": level if level in ("err", "warn", "info") else "warn",
                    "line": str(f.get("line") or "").strip(),
                    "problem": str(f.get("problem") or "").strip(),
                    "suggestion": str(f.get("suggestion") or "").strip()})
    return out


# --- ユーザー定義ツールの下書き ---------------------------------------------------
#
# SQLを書けない人でもツールを作れるようにするための入口。
# 「何をするツールか」を日本語で書いてもらい、SQLとパラメータはAIに起こさせる。
# 起こしたSQLは呼び出し側で必ず実データに当てて確かめる（推測のまま保存させない）。

_TOOL_SYSTEM = """あなたはSQLiteに詳しいデータ分析アプリの設定担当です。
利用者が日本語で書いた「やりたいこと」を、AIが呼び出せるツールの定義に変換してください。

まず「そもそもツールを作れるか」を判断し、作れる場合だけSQLを書いてください。

【作れる場合】出力形式（JSON以外の文字を含めないこと）:
{
  "ok": true,
  "name": "英小文字と_のみの短い名前（例: monthly_sales）",
  "description": "このツールが何を返すかの説明。AIがこれを読んで使うかどうかを決める",
  "sql": "SELECT ...（1文だけ。末尾のセミコロンは不要）",
  "parameters": [
    {"name": "year", "type": "string", "description": "対象年 YYYY",
     "required": true, "example": "2026"}
  ],
  "chart": {"chart_type": "line", "x": "月", "y": "売上", "title": "月別売上"},
  "explanation": "組み立てたSQLの日本語の解説（下記の書き方に従う）"
}

【作れない場合】理由だけを返す:
{
  "ok": false,
  "reason": "なぜ作れないかの日本語の説明。代わりにどうすればよいかも書く"
}

ok: false にする場合（無理にSQLを書かないこと）:
- カタログに、指示に合うテーブルや列が見当たらないとき
  → 何を探したが見つからなかったかを具体的に書く。近そうなテーブルがあれば挙げる。
- 指示があいまいで、どの表・どの列を使えばよいか決められないとき
  → 何がはっきりしないか、どう書き足せばよいかを伝える。
- 集計も絞り込みも要らず、1つの表をそのまま返すだけのとき
  → その表をそのまま使えばよいこと（ツールにする必要が薄いこと）を伝える。

explanation（解説）の書き方:
- 日本語で3〜6行。SQLを読めない人にも分かる言葉で書く。
- 「どの表を使うか」「どうつないだか（結合の条件と理由）」「どう絞ったか／集計したか」
  「1行が何を表すか」「毎回変える値は何か」の順に書く。
- SQLの構文用語（INNER JOIN など）をそのまま並べない。何をしているかを説明する。

守ること:
- SQLは SELECT（または WITH ... SELECT）だけ。書き込み・DDLは書かない。
- 列名・テーブル名は、与えられたカタログに実在するものだけを使う。推測で作らない。
- 「毎回変えたい値」は、利用者の日本語から自分で見極めて parameters にする。
  聞かれ方が変わるたびに差し替える値（「指定した年の」「ある部署の」「任意の期間で」など）は
  パラメータにし、SQLでは :名前 の形で参照する。
  一方「部署ごと」「月別」のような集計の切り口は、パラメータではなく GROUP BY で表す。
  迷ったらパラメータにしない。引数が増えるほどAIは呼びにくくなる。
- parameters の type は string / integer / number / boolean のいずれか。
- parameters には description（日本語）と example を必ず書く。
  example は「カタログの実値・期間に実在し、実際に行が返る値」にする。
  この値で試し実行して見せるので、0行になる値を書かないこと。
- 複数のDBにまたがるときは「DB名.テーブル名」で修飾する。
- 列には日本語の別名を AS で付ける（画面にそのまま出るため）。
- SQLiteに無い関数(STDDEV, MEDIAN, PERCENTILE_CONT, SQRT, POWER)やPIVOT構文は使わない。
- 見せ方が「グラフ」のときだけ chart を書く。x と y には SELECT の別名をそのまま使う。
  グラフでないときは chart を省略する。"""


_VIEW_SYSTEM = """あなたはSQLiteに詳しいデータ分析アプリの設定担当です。
利用者が日本語で書いた「欲しい一覧」を、ビュー（名前を付けたSELECT）の定義に変換してください。

まず「そもそもビューを作るべきか」を判断し、作れる場合だけSQLを書いてください。

【作れる場合】出力形式（JSON以外の文字を含めないこと）:
{
  "ok": true,
  "name": "まとまり__名前 の形。まとまりは使う表の接頭辞に合わせる（例: 部品_在庫__発注一覧）",
  "description": "この一覧が何かの説明。AIがこれを読んで使うかどうかを決める",
  "sql": "SELECT ...（1文だけ。末尾のセミコロンは不要）",
  "explanation": "組み立てたSQLの日本語の解説（下記の書き方に従う）"
}

【作らない方がよい場合】理由だけを返す:
{
  "ok": false,
  "reason": "なぜ作れない・作らない方がよいかの日本語の説明。代わりにどうすればよいかも書く"
}

ok: false にする場合（無理にSQLを書かないこと）:
- 求められている内容が、既にある1つのテーブルそのものであるとき
  → 「その内容は ○○ というテーブルがそのまま持っています。ビューを作らなくても
     そのテーブルを使えます」のように、そのテーブル名を挙げて伝える。
- カタログに、指示に合うテーブルや列が見当たらないとき
  → 何を探したが見つからなかったかを具体的に書く。近そうなテーブルがあれば挙げる。
- 指示があいまいで、どの表・どの列を使えばよいか決められないとき
  → 何がはっきりしないか、どう書き足せばよいかを伝える。
- 毎回変わる値（対象年・担当者など）を指定する必要があるとき
  → ビューは固定の一覧なので作れない旨と、「ツール」なら毎回値を変えられることを伝える。
- 集計や結合が要らず、単に1つの表を並べ替える・列を選ぶだけのとき
  → チャットでそのつど聞けばよく、ビューにする必要は薄いことを伝える。

explanation（解説）の書き方:
- 日本語で3〜6行。SQLを読めない人にも分かる言葉で書く。
- 「どの表を使うか」「どうつないだか（結合の条件と理由）」「どう絞ったか／集計したか」
  「1行が何を表すか」の順に書く。
- 複合キーで結合したときは、なぜ全列を条件にしたのかを一言添える。
- SQLの構文用語（INNER JOIN など）をそのまま並べない。何をしているかを説明する。

守ること:
- SQLは SELECT（または WITH ... SELECT）だけ。書き込み・DDLは書かない。
- 列名・テーブル名は、与えられたカタログに実在するものだけを使う。推測で作らない。
- カタログの「結合キー」に書かれた条件をそのまま使う。複合キー（※印つき）は
  必ず全列を AND でつなぐ。片方の列だけで結ばない。
- 列には日本語の別名を付けてよい（AS 部品名 など）。人が読む一覧になるように。
- パラメータ（毎回変える値）は使えない。ビューは固定の一覧。
- 名前は必ず「まとまり__名前」の形にする（__ は2つの下線）。
"""


def draft_view(db_path, purpose: str, previous: dict | None = None,
               error: str | None = None) -> dict:
    """日本語の「欲しい一覧」から、ビューの下書き（名前・説明・SQL）を起こす。"""
    import db                       # 循環importを避けるため、使うときに読む

    paths = [Path(db_path)] if db_path else db.list_db_files()
    context = catalog.prompt_for_scope(
        [{"path": str(p), "alias": db.alias_for(p), "tables": None} for p in paths])
    ask = [f"欲しい一覧: {purpose}"]
    if previous and error:
        ask.append("\n前回の下書きは実際のデータで失敗しました。原因を直して書き直してください。")
        ask.append(f"前回のSQL:\n{previous.get('sql', '')}")
        ask.append(f"エラー: {error}")

    data = _ask_json(_VIEW_SYSTEM, f"{context}\n\n{chr(10).join(ask)}", "ビューの下書き")
    # 「作らない方がよい」判断。SQLが無い応答も同じ扱いにする（無理に書かせない）
    sql = str(data.get("sql") or "").strip().rstrip(";")
    if data.get("ok") is False or not sql:
        return {"ok": False,
                "reason": str(data.get("reason")
                              or "この指示ではビューを作れませんでした。").strip()}
    return {
        "ok": True,
        "name": str(data.get("name") or "").strip(),
        "description": str(data.get("description") or purpose).strip(),
        "sql": sql,
        "explanation": str(data.get("explanation") or "").strip(),
    }


def draft_tool(db_path, purpose: str, params_wanted: list[str] | None = None,
               render: str = "table", previous: dict | None = None,
               error: str | None = None) -> dict:
    """日本語の「やりたいこと」から、ユーザー定義ツールの下書きを起こす。

    db_path        … None なら全DBのカタログを見せる（作る人にDBを選ばせない）。
                     特定のDBに限りたいときだけパスを渡す。
    purpose        … 何をするツールか（日本語）。毎回変えたい値もこの文から読み取らせる
                     ので、呼び出し側が指定を組み立てる必要はない。
    params_wanted  … 毎回変えたい項目を明示したいときだけ渡す（例: ["対象年", "部署"]）。
                     省略すれば purpose の書き方からAIが判断する。
    render         … 結果の見せ方（table / chart / chart_dual / excel / csv / none）
    previous/error … 前回の下書きが実データで失敗したときの、SQLとエラー文。
                     渡すと「どこが間違っていたか」を踏まえて書き直す。
    """
    import db                       # 循環importを避けるため、使うときに読む

    # db_path が None なら全DBを見せる。どのDBに書くかは、やりたいことを読んだAIが
    # 決める（作る人にDBを選ばせない）。量が上限を超えるときは要約に落ちる。
    if db_path is None:
        paths = db.list_db_files()
    else:
        paths = [Path(db_path)]
    context = catalog.prompt_for_scope(
        [{"path": str(p), "alias": db.alias_for(p), "tables": None} for p in paths])
    ask = [f"やりたいこと: {purpose}"]
    if params_wanted:
        ask.append("毎回変えたい項目: " + "、".join(params_wanted))
    ask.append(f"結果の見せ方: {render}")
    if previous and error:
        ask.append("\n前回の下書きは実際のデータで失敗しました。原因を直して書き直してください。")
        ask.append(f"前回のSQL:\n{previous.get('sql', '')}")
        ask.append(f"エラー: {error}")

    data = _ask_json(_TOOL_SYSTEM, f"{context}\n\n{chr(10).join(ask)}", "ツールの下書き")

    # 「作れない」判断。SQLが無い応答も同じ扱いにする（無理に書かせない）
    if data.get("ok") is False or not str(data.get("sql") or "").strip():
        return {"ok": False,
                "reason": str(data.get("reason")
                              or "この指示ではツールを作れませんでした。").strip()}
    out = {
        "ok": True,
        "name": str(data.get("name") or "").strip(),
        "description": str(data.get("description") or purpose).strip(),
        "sql": str(data.get("sql") or "").strip().rstrip(";"),
        "explanation": str(data.get("explanation") or "").strip(),
        "parameters": [],
        "render": render,
        "enabled": True,
    }
    for p in (data.get("parameters") or []):
        if not isinstance(p, dict) or not str(p.get("name") or "").strip():
            continue
        t = str(p.get("type") or "string")
        item = {
            "name": str(p["name"]).strip(),
            "type": t if t in custom_tools.PARAM_TYPES else "string",
            "description": str(p.get("description") or "").strip(),
            "required": p.get("required", True) is not False,
        }
        # 試し実行に使う値。空のまま流すと0行になり、動くかどうか確かめられない。
        if p.get("example") not in (None, ""):
            item["example"] = p["example"]
        out["parameters"].append(item)
    if render in ("chart", "chart_dual") and isinstance(data.get("chart"), dict):
        out["chart"] = data["chart"]
    return out
